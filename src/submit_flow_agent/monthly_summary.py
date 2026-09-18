"""Independent single-voucher monthly summary flow for a fixed non-school site."""

from __future__ import annotations

import calendar
import hashlib
import json
import os
import re
import shutil
import tempfile
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.properties import CalcProperties

from submit_flow_agent.ocr.base import OcrAdapter, OcrRawResult, write_ocr_raw_result
from submit_flow_agent.ocr.factory import create_ocr_adapter
from submit_flow_agent.pdf_renderer import RenderedPage, render_pdf_to_images


BUSINESS_PROFILE = "full_grid_sale_monthly_summary"
SITE_KEY = "tangxi_nursing_home"
SITE_NAME = "\u6c64\u897f\u656c\u8001\u9662"
SITE_ALIASES = ("\u656c\u8001\u9662", "\u656c\u8001\u9662\u5149\u4f0f\u9879\u76ee")
STATION_NAME = "\u656c\u8001\u9662\u5149\u4f0f\u9879\u76ee"
OUTPUT_NAME = "\u6c64\u897f\u656c\u8001\u9662\u6c47\u603b\u8868.xlsx"
CAPACITY_KW = Decimal("86.14")
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".pdf"}
REQUIRED_FIELDS = ("month", "total_energy", "tax_inclusive_fee")
TITLE_ROW = 1
HEADER_ROW = 2
DATA_START_ROW = 3


class MonthlySummaryError(RuntimeError):
    """Raised when the independent monthly-summary contract cannot advance."""


@dataclass(frozen=True)
class MonthlySummaryValues:
    month: str
    total_energy: Decimal
    tax_inclusive_fee: Decimal
    tax_exclusive_fee: Decimal | None = None
    tax_amount: Decimal | None = None

    def to_dict(self) -> dict[str, str]:
        result = {
            "month": self.month,
            "total_energy": _decimal_text(self.total_energy),
            "tax_inclusive_fee": _decimal_text(self.tax_inclusive_fee),
        }
        if self.tax_exclusive_fee is not None:
            result["tax_exclusive_fee"] = _decimal_text(self.tax_exclusive_fee)
        if self.tax_amount is not None:
            result["tax_amount"] = _decimal_text(self.tax_amount)
        return result


def prepare_monthly_summary(runtime_root: Path | str, *, task_id: str, month: str) -> dict[str, Any]:
    clean_month = _validate_month(month)
    clean_task_id = _safe_task_id(task_id)
    root = _task_root(runtime_root, clean_task_id)
    if root.exists():
        task = _read_json(root / "task.json")
        _assert_task_identity(task, clean_task_id, clean_month)
        return task
    root.mkdir(parents=True, exist_ok=False)
    task = {
        "task_id": clean_task_id,
        "business_profile": BUSINESS_PROFILE,
        "site_key": SITE_KEY,
        "site_name": SITE_NAME,
        "month": clean_month,
        "revision": _next_revision(runtime_root, clean_month),
        "status": "collecting_file",
        "stage": "collecting_file",
        "source_file": None,
        "recognized_data": {},
        "confirmed_data": {},
        "outputs": [],
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
    }
    _write_json(root / "task.json", task)
    return task


def collect_monthly_summary_source(
    runtime_root: Path | str,
    task_id: str,
    source: Path | str,
    *,
    ocr_adapter: OcrAdapter | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_file", "awaiting_confirmation", "need_review"}:
        raise MonthlySummaryError("Current monthly-summary task does not accept a source file.")
    source_path = Path(source)
    _validate_source_file(source_path)
    inputs_dir = root / "inputs"
    inputs_dir.mkdir(parents=True, exist_ok=True)
    digest = _sha256_file(source_path)
    target = inputs_dir / f"voucher-{digest[:16]}{source_path.suffix.lower()}"
    if not target.exists():
        shutil.copy2(source_path, target)

    pages = _source_pages(target, root / "rendered")
    raw = (ocr_adapter or create_ocr_adapter()).recognize(target, pages)
    write_ocr_raw_result(raw, root / "raw_ocr.json", relative_to=root)
    recognized, confidence = recognize_monthly_summary_values(raw)
    source_month = recognized.get("month")
    if source_month and source_month != task["month"]:
        raise MonthlySummaryError(
            f"\u51ed\u8bc1\u8d2d\u7535\u6708\u4efd\u4e3a {source_month}\uff0c"
            f"\u4e0e\u5f53\u524d\u586b\u62a5\u6708\u4efd {task['month']} \u4e0d\u4e00\u81f4\uff1b\u672c\u6b21\u6587\u4ef6\u672a\u767b\u8bb0\u3002"
        )
    missing = [field for field in REQUIRED_FIELDS if not recognized.get(field)]
    warnings = _recognition_warnings(recognized)
    task.update(
        {
            "source_file": {
                "name": source_path.name,
                "stored_path": str(target.relative_to(root)).replace("\\", "/"),
                "sha256": digest,
            },
            "recognized_data": recognized,
            "recognition_confidence": confidence,
            "missing_fields": missing,
            "warnings": warnings,
            "status": "need_review" if missing else "awaiting_confirmation",
            "stage": "need_review" if missing else "awaiting_confirmation",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def update_monthly_summary(runtime_root: Path | str, task_id: str, values: dict[str, Any]) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"awaiting_confirmation", "need_review"}:
        raise MonthlySummaryError("Current monthly-summary task is not awaiting data review.")
    unknown = sorted(set(values) - set(REQUIRED_FIELDS))
    if unknown:
        raise MonthlySummaryError(f"Unsupported monthly-summary fields: {', '.join(unknown)}")
    normalized = _normalize_values({**task.get("recognized_data", {}), **values}, expected_month=task["month"])
    task.update(
        {
            "confirmed_data": normalized.to_dict(),
            "missing_fields": [],
            "status": "awaiting_confirmation",
            "stage": "awaiting_confirmation",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def confirm_monthly_summary(
    runtime_root: Path | str,
    task_id: str,
    values: dict[str, Any] | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"awaiting_confirmation", "need_review"}:
        raise MonthlySummaryError("Current monthly-summary task is not awaiting confirmation.")
    normalized = _normalize_values(
        {**task.get("recognized_data", {}), **task.get("confirmed_data", {}), **(values or {})},
        expected_month=task["month"],
    )
    task.update(
        {
            "confirmed_data": normalized.to_dict(),
            "missing_fields": [],
            "status": "ready_to_run",
            "stage": "ready_to_run",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def run_monthly_summary(
    runtime_root: Path | str,
    task_id: str,
    *,
    template_path: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] == "completed":
        return task
    if task["status"] != "ready_to_run":
        raise MonthlySummaryError("\u8bf7\u5148\u6838\u5bf9\u5e76\u786e\u8ba4\u8bc6\u522b\u6570\u636e\uff0c\u518d\u53d1\u9001\u201c\u5f00\u59cb\u8fd0\u884c\u201d\u3002")
    values = _normalize_values(task["confirmed_data"], expected_month=task["month"])
    template = Path(template_path)
    if not template.is_file():
        raise MonthlySummaryError(f"Controlled monthly-summary template does not exist: {template}")
    records = _effective_records(runtime_root)
    records[values.month] = values
    output = root / "outputs" / OUTPUT_NAME
    write_monthly_summary_workbook(
        template,
        output,
        records,
        materialize_calculations=_summary_materialization_enabled(runtime_root),
    )
    task.update(
        {
            "status": "completed",
            "stage": "completed",
            "outputs": [str(output.relative_to(root)).replace("\\", "/")],
            "completed_at": _utc_now(),
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    _promote_effective(runtime_root, task, values, output)
    return task


def get_monthly_summary(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    return _load_task(runtime_root, task_id)[1]


def list_monthly_summary_outputs(runtime_root: Path | str, task_id: str) -> list[str]:
    root, task = _load_task(runtime_root, task_id)
    return [str((root / item).resolve()) for item in task.get("outputs", []) if (root / item).is_file()]


def recognize_monthly_summary_values(raw: OcrRawResult) -> tuple[dict[str, str], dict[str, float]]:
    blocks = [block for page in raw.pages for block in page.blocks if block.text.strip()]
    text = "\n".join(block.text.strip() for block in blocks)
    fields: dict[str, str] = {}
    confidence: dict[str, float] = {}
    patterns = {
        "month": r"(?:\u8d2d\u7535\u6708\u4efd|\u7535\u8d39\u5e74\u6708|\u7ed3\u7b97\u6708\u4efd)\s*[:\uff1a]?\s*(20\d{4}|20\d{2}[-./\u5e74]\d{1,2}\u6708?)",
        "total_energy": r"(?:\u603b\u7535\u91cf|\u4e0a\u7f51\u7535\u91cf)\s*[:\uff1a]?\s*([\d,]+(?:\.\d+)?)",
        "tax_inclusive_fee": r"(?:\u542b\u7a0e\u7535\u8d39|\u4ef7\u7a0e\u5408\u8ba1)\s*[:\uff1a]?\s*([\d,]+(?:\.\d+)?)",
        "tax_exclusive_fee": r"\u4e0d\u542b\u7a0e\u7535\u8d39\s*[:\uff1a]?\s*([\d,]+(?:\.\d+)?)",
        "tax_amount": r"(?:\u7a0e\u91d1|\u7a0e\u989d)\s*[:\uff1a]?\s*([\d,]+(?:\.\d+)?)",
    }
    for key, pattern in patterns.items():
        match = re.search(pattern, text, flags=re.IGNORECASE) or _match_adjacent_blocks(blocks, pattern)
        if not match:
            continue
        raw_value = match.group(1).replace(",", "")
        fields[key] = _normalize_month(raw_value) if key == "month" else _decimal_text(_decimal(raw_value, key))
        confidence[key] = _field_confidence(blocks, match.group(0))
    coordinate_values = _coordinate_table_values(blocks)
    for key, (raw_value, score) in coordinate_values.items():
        fields[key] = _normalize_month(raw_value) if key == "month" else _decimal_text(_decimal(raw_value, key))
        confidence[key] = score
    return fields, confidence


def _coordinate_table_values(blocks: list[Any]) -> dict[str, tuple[str, float]]:
    labels = {
        "month": {"\u8d2d\u7535\u6708\u4efd", "\u7535\u8d39\u5e74\u6708", "\u7ed3\u7b97\u6708\u4efd"},
        "total_energy": {"\u603b\u7535\u91cf", "\u4e0a\u7f51\u7535\u91cf"},
        "tax_inclusive_fee": {"\u542b\u7a0e\u7535\u8d39", "\u4ef7\u7a0e\u5408\u8ba1"},
        "tax_exclusive_fee": {"\u4e0d\u542b\u7a0e\u7535\u8d39"},
        "tax_amount": {"\u7a0e\u91d1", "\u7a0e\u989d"},
    }
    result: dict[str, tuple[str, float]] = {}
    for key, aliases in labels.items():
        header = next((block for block in blocks if _compact_text(block.text) in aliases), None)
        if header is None:
            continue
        header_x = (header.bbox[0] + header.bbox[2]) / 2
        header_y = (header.bbox[1] + header.bbox[3]) / 2
        candidates: list[tuple[float, float, Any, str]] = []
        for block in blocks:
            value = block.text.replace(",", "").strip()
            if not re.fullmatch(r"\d+(?:\.\d+)?", value):
                continue
            value_x = (block.bbox[0] + block.bbox[2]) / 2
            value_y = (block.bbox[1] + block.bbox[3]) / 2
            y_gap = value_y - header_y
            x_gap = abs(value_x - header_x)
            if 8 <= y_gap <= 180 and x_gap <= 90:
                candidates.append((y_gap, x_gap, block, value))
        if not candidates:
            continue
        _, _, selected, value = min(candidates, key=lambda item: (item[0], item[1]))
        result[key] = (value, round(min(float(header.confidence), float(selected.confidence)), 4))
    return result


def _compact_text(value: str) -> str:
    return re.sub(r"[\s:\uff1a]", "", value)


def write_monthly_summary_workbook(
    template_path: Path | str,
    output_path: Path | str,
    records: dict[str, MonthlySummaryValues],
    *,
    materialize_calculations: bool = False,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path)
    sheet = workbook["Sheet1"]
    template_row = DATA_START_ROW if sheet.max_row >= DATA_START_ROW else HEADER_ROW
    style = _row_style_snapshot(sheet, template_row)
    row_height = sheet.row_dimensions[template_row].height
    existing: dict[str, MonthlySummaryValues] = {}
    for row in range(DATA_START_ROW, sheet.max_row + 1):
        month = _month_from_excel_value(sheet.cell(row, 4).value)
        if month:
            existing[month] = MonthlySummaryValues(
                month,
                _decimal(sheet.cell(row, 7).value, "total_energy"),
                _decimal(sheet.cell(row, 11).value, "tax_inclusive_fee"),
            )
    existing.update(records)
    if sheet.max_row >= DATA_START_ROW:
        sheet.delete_rows(DATA_START_ROW, sheet.max_row - DATA_START_ROW + 1)
    for row, month in enumerate(sorted(existing), start=DATA_START_ROW):
        _apply_row_style(sheet, row, style, row_height)
        values = existing[month]
        year, month_number = map(int, month.split("-"))
        days = calendar.monthrange(year, month_number)[1]
        row_values: list[Any] = [
            STATION_NAME,
            "\u5168\u989d\u4e0a\u7f51",
            "\u56fa\u5b9a\u7535\u4ef7",
            date(year, month_number, 1),
            0,
            float(values.total_energy),
            float(values.total_energy),
            0,
            float(values.tax_inclusive_fee / values.total_energy) if materialize_calculations else f"=IFERROR(J{row}/G{row},0)",
            float(values.tax_inclusive_fee),
            float(values.tax_inclusive_fee),
            float(values.total_energy / Decimal(str(CAPACITY_KW)) / Decimal(days))
            if materialize_calculations
            else f"=IFERROR(G{row}/{CAPACITY_KW}/{days},0)",
        ]
        for column, value in enumerate(row_values, start=1):
            sheet.cell(row, column, value)
        sheet.cell(row, 4).number_format = "yyyy-mm"
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output)
    return output


def _row_style_snapshot(sheet: Any, row: int) -> list[dict[str, Any]]:
    return [
        {
            "style": copy(sheet.cell(row, column)._style),
            "font": copy(sheet.cell(row, column).font),
            "fill": copy(sheet.cell(row, column).fill),
            "border": copy(sheet.cell(row, column).border),
            "alignment": copy(sheet.cell(row, column).alignment),
            "protection": copy(sheet.cell(row, column).protection),
            "number_format": sheet.cell(row, column).number_format,
        }
        for column in range(1, 13)
    ]


def _apply_row_style(sheet: Any, row: int, style: list[dict[str, Any]], height: float | None) -> None:
    sheet.row_dimensions[row].height = height
    for column, source in enumerate(style, start=1):
        target = sheet.cell(row, column)
        target._style = copy(source["style"])
        target.font = copy(source["font"])
        target.fill = copy(source["fill"])
        target.border = copy(source["border"])
        target.alignment = copy(source["alignment"])
        target.protection = copy(source["protection"])
        target.number_format = source["number_format"]


def _source_pages(source: Path, output_dir: Path) -> list[RenderedPage]:
    if source.suffix.lower() == ".pdf":
        return render_pdf_to_images(source, output_dir)
    return [RenderedPage(source_file=source, page=1, image_path=source)]


def _normalize_values(values: dict[str, Any], *, expected_month: str) -> MonthlySummaryValues:
    month = _normalize_month(str(values.get("month") or expected_month))
    if month != expected_month:
        raise MonthlySummaryError(f"Voucher month {month} does not match task month {expected_month}.")
    total_energy = _decimal(values.get("total_energy"), "total_energy")
    fee = _decimal(values.get("tax_inclusive_fee"), "tax_inclusive_fee")
    if total_energy <= 0:
        raise MonthlySummaryError("\u603b\u7535\u91cf\u5fc5\u987b\u5927\u4e8e 0\u3002")
    if fee < 0:
        raise MonthlySummaryError("\u542b\u7a0e\u7535\u8d39\u4e0d\u80fd\u5c0f\u4e8e 0\u3002")
    return MonthlySummaryValues(
        month,
        total_energy,
        fee,
        _optional_decimal(values.get("tax_exclusive_fee")),
        _optional_decimal(values.get("tax_amount")),
    )


def _effective_records(runtime_root: Path | str) -> dict[str, MonthlySummaryValues]:
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    result: dict[str, MonthlySummaryValues] = {}
    for month, record in (series.get("months") or {}).items():
        if isinstance(record, dict) and isinstance(record.get("values"), dict):
            result[month] = _normalize_values(record["values"], expected_month=month)
    return result


def _summary_materialization_enabled(runtime_root: Path | str) -> bool:
    series = _read_json(_series_path(runtime_root), default={})
    return series.get("summary_materialize_calculations") is True


def _promote_effective(runtime_root: Path | str, task: dict[str, Any], values: MonthlySummaryValues, output: Path) -> None:
    target = _series_path(runtime_root)
    series = _read_json(target, default={"site_key": SITE_KEY, "months": {}})
    months = series.setdefault("months", {})
    current = months.get(values.month) if isinstance(months.get(values.month), dict) else {}
    history = list(current.get("history") or [])
    if current.get("task_id"):
        history.append({key: current.get(key) for key in ("task_id", "revision", "values", "output")})
    months[values.month] = {
        "task_id": task["task_id"],
        "revision": task["revision"],
        "values": values.to_dict(),
        "output": str(output),
        "history": history,
        "updated_at": _utc_now(),
    }
    _write_json(target, series)


def _next_revision(runtime_root: Path | str, month: str) -> int:
    current = (_read_json(_series_path(runtime_root), default={"months": {}}).get("months") or {}).get(month)
    return int(current.get("revision") or 0) + 1 if isinstance(current, dict) else 1


def _task_root(runtime_root: Path | str, task_id: str) -> Path:
    return Path(runtime_root) / "monthly_summary_tasks" / task_id


def _series_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "monthly_summaries" / SITE_KEY / "series.json"


def _load_task(runtime_root: Path | str, task_id: str) -> tuple[Path, dict[str, Any]]:
    clean = _safe_task_id(task_id)
    root = _task_root(runtime_root, clean)
    task = _read_json(root / "task.json")
    _assert_task_identity(task, clean, str(task.get("month") or ""))
    return root, task


def _assert_task_identity(task: dict[str, Any], task_id: str, month: str) -> None:
    if task.get("task_id") != task_id or task.get("business_profile") != BUSINESS_PROFILE:
        raise MonthlySummaryError("Monthly-summary task identity mismatch.")
    if task.get("site_key") != SITE_KEY or task.get("month") != month:
        raise MonthlySummaryError("Monthly-summary task crossed site or month boundaries.")


def _validate_source_file(path: Path) -> None:
    if not path.is_file() or path.is_symlink():
        raise MonthlySummaryError("Monthly-summary source must be a regular file.")
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise MonthlySummaryError("\u4ec5\u652f\u6301 JPG\u3001PNG \u6216 PDF \u51ed\u8bc1\u3002")


def _safe_task_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", value):
        raise MonthlySummaryError("Invalid monthly-summary task id.")
    return value


def _validate_month(value: str) -> str:
    if not re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", value):
        raise MonthlySummaryError("Month must use YYYY-MM format.")
    return value


def _normalize_month(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) == 6:
        return _validate_month(f"{digits[:4]}-{digits[4:]}")
    match = re.fullmatch(r"(20\d{2})[-./\u5e74](\d{1,2})(?:\u6708)?", value.strip())
    if match:
        return _validate_month(f"{match[1]}-{int(match[2]):02d}")
    return _validate_month(value)


def _month_from_excel_value(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, str):
        try:
            return _normalize_month(value)
        except MonthlySummaryError:
            return None
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
        except (TypeError, ValueError, OverflowError):
            return None
        if isinstance(converted, (date, datetime)):
            return f"{converted.year:04d}-{converted.month:02d}"
    return None


def _recognition_warnings(values: dict[str, Any]) -> list[str]:
    inclusive = _optional_decimal(values.get("tax_inclusive_fee"))
    exclusive = _optional_decimal(values.get("tax_exclusive_fee"))
    tax = _optional_decimal(values.get("tax_amount"))
    if inclusive is None or exclusive is None or tax is None:
        return []
    difference = abs((exclusive + tax) - inclusive)
    if difference <= Decimal("0.02"):
        return []
    return [
        "\u4e0d\u542b\u7a0e\u7535\u8d39\u4e0e\u7a0e\u91d1\u4e4b\u548c\u4e0e\u542b\u7a0e\u7535\u8d39\u4e0d\u4e00\u81f4\uff0c\u8bf7\u6838\u5bf9\u51ed\u8bc1\u3002"
    ]


def _decimal(value: Any, field: str) -> Decimal:
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise MonthlySummaryError(f"{field} must be numeric.") from exc


def _optional_decimal(value: Any) -> Decimal | None:
    return None if value in {None, ""} else _decimal(value, "optional amount")


def _decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _match_adjacent_blocks(blocks: list[Any], pattern: str) -> re.Match[str] | None:
    for size in (2, 3, 4):
        for index in range(0, max(0, len(blocks) - size + 1)):
            match = re.search(pattern, " ".join(block.text for block in blocks[index : index + size]), re.IGNORECASE)
            if match:
                return match
    return None


def _field_confidence(blocks: list[Any], matched_text: str) -> float:
    tokens = [token for token in re.split(r"\s+", matched_text) if token]
    scores = [block.confidence for block in blocks if any(token in block.text for token in tokens)]
    return round(min(scores), 4) if scores else 0.0


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_json(path: Path, default: dict[str, Any] | None = None) -> dict[str, Any]:
    if not path.exists() and default is not None:
        return default
    if not path.is_file() or path.is_symlink():
        raise MonthlySummaryError(f"Controlled JSON file is unavailable: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise MonthlySummaryError(f"Controlled JSON file is invalid: {path}") from exc
    if not isinstance(payload, dict):
        raise MonthlySummaryError("Controlled JSON payload must be an object.")
    return payload


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(payload, ensure_ascii=False, indent=2)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temp = Path(temp_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            handle.write(serialized)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def new_monthly_summary_task_id(month: str) -> str:
    clean = _validate_month(month)
    return f"monthly_{SITE_KEY}_{clean.replace('-', '_')}_{uuid.uuid4().hex[:12]}"
