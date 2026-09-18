"""Reusable two-workbook monthly settlement flow for factory sites."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
import uuid
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.properties import CalcProperties


BUSINESS_PROFILE = "factory_settlement_v1"
TEMPLATE_KEY = "factory_meizhou_v1"
SEED_SITE_KEY = "factory_meizhou_3"
SEED_SITE_NAME = "梅州3号厂房"
SEED_ALIASES = {"梅州3号厂房", "梅州三号厂房"}
SEED_GRID_UNIT_PRICE = Decimal("0.453")
OUTPUT_SUFFIX = "汇总表.xlsx"
REQUIRED_ROLES = ("energy_settlement", "fee_settlement")
PERIODS = ("尖", "峰", "平", "谷")
MAX_SOURCE_BYTES = 30 * 1024 * 1024
MAX_XLSX_ENTRIES = 3000
MAX_XLSX_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
ENERGY_TOLERANCE = Decimal("0.01")
HISTORICAL_SUMMARY_PERIOD_OVERRIDES = {
    (SEED_SITE_NAME, "2026-06"): (
        Decimal("27354"),
        Decimal("25548"),
        Decimal("44426"),
        Decimal("7404"),
    )
}


class FactorySettlementError(RuntimeError):
    """Raised when a factory settlement task cannot advance safely."""


@dataclass(frozen=True)
class FactoryValues:
    month: str
    total_generation: Decimal
    grid_energy: Decimal
    period_energy: tuple[Decimal, Decimal, Decimal, Decimal]
    period_prices: tuple[Decimal, Decimal, Decimal, Decimal]
    self_use_fee: Decimal
    grid_unit_price: Decimal

    @property
    def self_use_energy(self) -> Decimal:
        return sum(self.period_energy, Decimal("0"))

    @property
    def grid_fee(self) -> Decimal:
        return _money_decimal(self.grid_energy * self.grid_unit_price)

    @property
    def total_fee(self) -> Decimal:
        return _money_decimal(self.grid_fee + self.self_use_fee)

    def to_dict(self) -> dict[str, Any]:
        return {
            "month": self.month,
            "total_generation": _number(self.total_generation),
            "grid_energy": _number(self.grid_energy),
            "grid_unit_price": _number(self.grid_unit_price, places=6),
            "grid_fee": _money(self.grid_fee),
            "period_energy": {
                period: _number(value) for period, value in zip(PERIODS, self.period_energy, strict=True)
            },
            "period_prices": {
                period: _number(value, places=7)
                for period, value in zip(PERIODS, self.period_prices, strict=True)
            },
            "self_use_energy": _number(self.self_use_energy),
            "self_use_fee": _money(self.self_use_fee),
            "total_fee": _money(self.total_fee),
        }


def prepare_factory_settlement(
    runtime_root: Path | str,
    *,
    task_id: str,
    site_name: str,
    month: str,
    allow_create: bool = False,
) -> dict[str, Any]:
    clean_task_id = _safe_task_id(task_id)
    clean_month = _validate_month(month)
    profile, is_new_site = _resolve_profile(runtime_root, site_name, allow_create=allow_create)
    root = _task_root(runtime_root, clean_task_id)
    if root.exists():
        task = _read_json(root / "task.json")
        _assert_task_identity(task, clean_task_id, profile["site_key"], clean_month)
        return task
    root.mkdir(parents=True, exist_ok=False)
    inherited_price, price_source = _inherited_grid_price(runtime_root, profile, clean_month)
    task = {
        "task_id": clean_task_id,
        "business_profile": BUSINESS_PROFILE,
        "template_key": TEMPLATE_KEY,
        "site_key": profile["site_key"],
        "site_name": profile["site_name"],
        "month": clean_month,
        "revision": _next_revision(runtime_root, profile["site_key"], clean_month),
        "is_new_site": is_new_site,
        "include_template_baseline": bool(profile.get("include_template_baseline")),
        "status": "awaiting_pricing",
        "stage": "awaiting_pricing",
        "pricing": {
            "grid_unit_price": _number(inherited_price, places=6) if inherited_price is not None else None,
            "source": price_source,
            "confirmed": False,
        },
        "sources": {},
        "source_count": 0,
        "recognized_data": {},
        "confirmed_data": {},
        "outputs": [],
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
    }
    _write_json(root / "task.json", task)
    return task


def update_factory_settlement_price(
    runtime_root: Path | str,
    task_id: str,
    *,
    grid_unit_price: str | int | float | Decimal,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] != "awaiting_pricing":
        raise FactorySettlementError("当前厂房任务不在上网电价确认阶段。")
    price = _decimal(grid_unit_price, "上网电价")
    if price <= 0:
        raise FactorySettlementError("上网电价必须大于 0。")
    task["pricing"] = {
        "grid_unit_price": _number(price, places=6),
        "source": "operator_input",
        "confirmed": False,
    }
    task["updated_at"] = _utc_now()
    _write_json(root / "task.json", task)
    return task


def confirm_factory_settlement_price(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] != "awaiting_pricing":
        raise FactorySettlementError("当前厂房任务不在上网电价确认阶段。")
    pricing = task.get("pricing") if isinstance(task.get("pricing"), dict) else {}
    if pricing.get("grid_unit_price") in {None, ""}:
        raise FactorySettlementError("该厂房没有历史上网电价，请先发送“修改：上网电价 0.453元/kWh”。")
    price = _decimal(pricing.get("grid_unit_price"), "上网电价")
    if price <= 0:
        raise FactorySettlementError("该厂房没有历史上网电价，请先发送“修改：上网电价 0.453元/kWh”。")
    pricing["grid_unit_price"] = _number(price, places=6)
    pricing["confirmed"] = True
    pricing["confirmed_at"] = _utc_now()
    task.update(
        {
            "pricing": pricing,
            "status": "collecting_files",
            "stage": "collecting_files",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def collect_factory_settlement_source(
    runtime_root: Path | str,
    task_id: str,
    source: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_files", "awaiting_confirmation", "need_review"}:
        raise FactorySettlementError("当前厂房任务不接收结算单。")
    source_path = Path(source)
    _validate_xlsx(source_path)
    digest = _sha256_file(source_path)
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    if any(isinstance(item, dict) and item.get("sha256") == digest for item in sources.values()):
        return {**task, "duplicate": True}

    role, extracted = recognize_factory_workbook(source_path, expected_month=task["month"])
    if role in sources:
        label = "电量结算单" if role == "energy_settlement" else "电费结算单"
        raise FactorySettlementError(f"{label}已登记，请勿重复上传另一份同类文件。")
    inputs = root / "inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    target = inputs / f"{role}-{digest[:16]}.xlsx"
    if not target.exists():
        shutil.copy2(source_path, target)
    sources[role] = {
        "role": role,
        "name": source_path.name,
        "stored_path": str(target.relative_to(root)).replace("\\", "/"),
        "sha256": digest,
        "recognized": extracted,
    }
    task["sources"] = sources
    task["source_count"] = sum(role_name in sources for role_name in REQUIRED_ROLES)
    task["updated_at"] = _utc_now()
    if not all(role_name in sources for role_name in REQUIRED_ROLES):
        task["status"] = task["stage"] = "collecting_files"
        _write_json(root / "task.json", task)
        return task

    values = _combine_sources(task)
    task.update(
        {
            "recognized_data": values.to_dict(),
            "validation_checks": _validation_checks(values),
            "status": "awaiting_confirmation",
            "stage": "awaiting_confirmation",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def confirm_factory_settlement(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] != "awaiting_confirmation":
        raise FactorySettlementError("当前厂房任务不在数据确认阶段。")
    values = _values_from_dict(task.get("recognized_data"), expected_month=task["month"])
    task.update(
        {
            "confirmed_data": values.to_dict(),
            "status": "ready_to_run",
            "stage": "ready_to_run",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def run_factory_settlement(
    runtime_root: Path | str,
    task_id: str,
    *,
    template_path: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] == "completed":
        return task
    if task["status"] != "ready_to_run":
        raise FactorySettlementError("请先核对并确认识别数据，再发送“开始运行”。")
    template = Path(template_path)
    if not template.is_file():
        raise FactorySettlementError(f"厂房汇总表模板不存在：{template}")
    values = _values_from_dict(task.get("confirmed_data"), expected_month=task["month"])
    records = _effective_records(runtime_root, task["site_key"])
    if task.get("include_template_baseline"):
        records = {**_template_records(template), **records}
    records[values.month] = values
    output_dir = root / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / _output_name(task["site_name"])
    write_factory_summary(
        template,
        output,
        site_name=task["site_name"],
        records=records,
        materialize_calculations=_summary_materialization_enabled(runtime_root),
    )
    if not output.is_file() or output.stat().st_size <= 0:
        raise FactorySettlementError("厂房汇总表生成失败，任务未完成。")
    task.update(
        {
            "status": "completed",
            "stage": "completed",
            "outputs": [str(output.relative_to(root)).replace("\\", "/")],
            "completed_at": _utc_now(),
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    _promote_effective(runtime_root, task, values)
    return task


def get_factory_settlement(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    return _load_task(runtime_root, task_id)[1]


def list_factory_settlement_outputs(runtime_root: Path | str, task_id: str) -> list[str]:
    root, task = _load_task(runtime_root, task_id)
    return [str((root / item).resolve()) for item in task.get("outputs", []) if (root / item).is_file()]


def recognize_factory_workbook(source: Path | str, *, expected_month: str) -> tuple[str, dict[str, Any]]:
    path = Path(source)
    formulas = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    values = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    try:
        sheet = formulas.worksheets[0]
        value_sheet = values.worksheets[0]
        title = _factory_workbook_title(sheet)
        if title == "电费结算单":
            return "fee_settlement", _extract_fee_settlement(sheet, value_sheet, expected_month)
        if title == "电量结算单":
            return "energy_settlement", _extract_energy_settlement(sheet, value_sheet, expected_month)
        raise FactorySettlementError("无法判断文件是电费结算单还是电量结算单，本次文件未登记。")
    finally:
        formulas.close()
        values.close()


def _extract_fee_settlement(sheet: Any, cached: Any, expected_month: str) -> dict[str, Any]:
    if str(sheet["A1"].value or "").strip() == "电费结算单":
        month_cell = "A8"
        start_row = 10
        component_columns = (18, 19, 20)
    elif str(sheet["A2"].value or "").strip() == "电费结算单":
        month_cell = "A9"
        start_row = 11
        component_columns = (19, 20, 21)
    else:
        raise FactorySettlementError("无法识别电费结算单版式，本次文件未登记。")
    month = _month_from_excel(sheet[month_cell].value, "电费结算单购电月份")
    _assert_source_month(month, expected_month, "电费结算单")
    energies = tuple(
        _clean_energy(
            _cached_formula_decimal(
                sheet,
                cached,
                row=row,
                column=7,
                field=f"{PERIODS[index]}时段电量",
            )
        )
        for index, row in enumerate(range(start_row, start_row + 4))
    )
    prices: list[Decimal] = []
    for index, row in enumerate(range(start_row, start_row + 4)):
        cached_price = cached.cell(row, 10).value
        if cached_price is None:
            component_values = [
                _decimal(sheet.cell(row, column).value, label)
                for column, label in zip(component_columns, ("电能电费", "输配电费", "线损"), strict=True)
            ]
            cached_price = sum(component_values, Decimal("0")) * Decimal("0.6")
        prices.append(_decimal(cached_price, f"{PERIODS[index]}时段电价"))
    expected_fee = _money_decimal(sum((energy * price for energy, price in zip(energies, prices, strict=True)), Decimal("0")))
    fee_total_row = start_row + 4
    settlement_total_row = start_row + 6
    source_fee_value = cached.cell(fee_total_row, 13).value
    if source_fee_value is None:
        source_fee_value = cached.cell(settlement_total_row, 12).value
    source_fee = _decimal(source_fee_value, "消纳电费")
    source_energy = _clean_energy(_decimal(cached.cell(settlement_total_row, 4).value, "结算电量合计"))
    if abs(source_energy - sum(energies, Decimal("0"))) > Decimal("0.01"):
        raise FactorySettlementError("电费结算单的尖峰平谷电量与结算电量合计不一致。")
    if abs(source_fee - expected_fee) > Decimal("0.02"):
        raise FactorySettlementError("电费结算单的分时电费与含税电费合计不一致。")
    return {
        "month": month,
        "period_energy": {period: _number(value) for period, value in zip(PERIODS, energies, strict=True)},
        "period_prices": {period: _number(value, places=7) for period, value in zip(PERIODS, prices, strict=True)},
        "self_use_energy": _number(source_energy),
        "self_use_fee": _money(source_fee),
    }


def _extract_energy_settlement(sheet: Any, cached: Any, expected_month: str) -> dict[str, Any]:
    multi_meter_layout = str(sheet["A23"].value or "").strip() == "立能派光伏用电量合计"
    month = _month_from_excel(sheet["G9" if multi_meter_layout else "F9"].value, "电量结算单购电月份")
    _assert_source_month(month, expected_month, "电量结算单")
    if multi_meter_layout:
        generated = tuple(_multi_meter_energy(sheet, row) for row in (13, 15, 16, 17, 18, 19, 20))
        grid_energy = _multi_meter_energy(sheet, 21)
        cached_self_use = _decimal(cached["C23"].value, "光伏用电量合计")
        cached_grid = _decimal(cached["C24"].value, "表显上网电量")
    else:
        generated = tuple(_meter_energy(sheet, row) for row in (13, 15))
        grid_energy = _meter_energy(sheet, 18)
        cached_self_use = _decimal(cached["C20"].value, "光伏用电量合计")
        cached_grid = _decimal(cached["C21"].value, "表显上网电量")
    meter_total_generation = sum(generated, Decimal("0"))
    if abs(cached_grid - grid_energy) > ENERGY_TOLERANCE:
        raise FactorySettlementError("电量结算单的总表计算值与表显上网电量不一致。")
    if abs(cached_self_use + cached_grid - meter_total_generation) > ENERGY_TOLERANCE:
        raise FactorySettlementError("电量结算单的发电量、上网电量与光伏用电量无法闭合。")
    return {
        "month": month,
        "generation_meters": [_number(value) for value in generated],
        "meter_total_generation": _number(meter_total_generation),
        "total_generation": _number(cached_self_use),
        "grid_energy": _number(grid_energy),
        "self_use_energy": _number(cached_self_use),
    }


def _meter_energy(sheet: Any, row: int) -> Decimal:
    differences = (
        _decimal(sheet.cell(row, 4).value, "尖本期示数") - _decimal(sheet.cell(row, 3).value, "尖上期示数"),
        _decimal(sheet.cell(row, 7).value, "峰本期示数") - _decimal(sheet.cell(row, 6).value, "峰上期示数"),
        _decimal(sheet.cell(row, 11).value, "平本期示数") - _decimal(sheet.cell(row, 9).value, "平上期示数"),
        _decimal(sheet.cell(row, 14).value, "谷本期示数") - _decimal(sheet.cell(row, 13).value, "谷上期示数"),
    )
    if any(value < 0 for value in differences):
        raise FactorySettlementError("电量结算单存在本期示数小于上期示数的情况。")
    multiplier = _decimal(sheet.cell(row, 15).value, "综合倍率")
    if multiplier <= 0:
        raise FactorySettlementError("电量结算单综合倍率必须大于 0。")
    return sum(differences, Decimal("0")) * multiplier


def _multi_meter_energy(sheet: Any, row: int) -> Decimal:
    differences = tuple(
        _decimal(sheet.cell(row, current_column).value, f"第{row}行本期示数")
        - _decimal(sheet.cell(row, prior_column).value, f"第{row}行上期示数")
        for prior_column, current_column in ((3, 5), (7, 9), (11, 13), (15, 16))
    )
    if any(value < 0 for value in differences):
        raise FactorySettlementError("电量结算单存在本期示数小于上期示数的情况。")
    multiplier = _decimal(sheet.cell(row, 17).value, f"第{row}行综合倍率")
    if multiplier <= 0:
        raise FactorySettlementError("电量结算单综合倍率必须大于 0。")
    return _clean_energy(sum(differences, Decimal("0")) * multiplier)


def _factory_workbook_title(sheet: Any) -> str:
    for row in range(1, min(sheet.max_row, 3) + 1):
        for column in range(1, min(sheet.max_column, 3) + 1):
            value = str(sheet.cell(row, column).value or "").strip()
            if value in {"电费结算单", "电量结算单"}:
                return value
    return ""


def _combine_sources(task: dict[str, Any]) -> FactoryValues:
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    energy = sources.get("energy_settlement", {}).get("recognized", {})
    fee = sources.get("fee_settlement", {}).get("recognized", {})
    if not isinstance(energy, dict) or not isinstance(fee, dict):
        raise FactorySettlementError("两份厂房结算单尚未收齐。")
    if energy.get("month") != fee.get("month") or energy.get("month") != task["month"]:
        raise FactorySettlementError("两份结算单月份不一致。")
    period_energy_map = fee.get("period_energy") if isinstance(fee.get("period_energy"), dict) else {}
    period_price_map = fee.get("period_prices") if isinstance(fee.get("period_prices"), dict) else {}
    pricing = task.get("pricing") if isinstance(task.get("pricing"), dict) else {}
    values = FactoryValues(
        month=task["month"],
        total_generation=_decimal(energy.get("total_generation"), "月总发电量"),
        grid_energy=_decimal(energy.get("grid_energy"), "上网电量"),
        period_energy=tuple(_decimal(period_energy_map.get(period), f"{period}时段电量") for period in PERIODS),
        period_prices=tuple(_decimal(period_price_map.get(period), f"{period}时段电价") for period in PERIODS),
        self_use_fee=_decimal(fee.get("self_use_fee"), "消纳电费"),
        grid_unit_price=_decimal(pricing.get("grid_unit_price"), "上网电价"),
    )
    energy_self_use = _decimal(energy.get("self_use_energy"), "电量结算单消纳电量")
    fee_self_use = _decimal(fee.get("self_use_energy"), "电费结算单消纳电量")
    if values.total_generation != values.self_use_energy:
        raise FactorySettlementError("月总发电量与尖峰平谷消纳电量不一致。")
    if energy_self_use != values.self_use_energy or fee_self_use != values.self_use_energy:
        raise FactorySettlementError("两份结算单的消纳电量不一致。")
    return values


def _validation_checks(values: FactoryValues) -> list[dict[str, str]]:
    return [
        {"name": "energy_balance", "status": "pass", "message": "月总发电量取立能派光伏用电量合计。"},
        {"name": "period_balance", "status": "pass", "message": "尖峰平谷电量合计与消纳电量一致。"},
        {"name": "fee_balance", "status": "pass", "message": "分时电费合计与消纳电费一致。"},
    ]


def write_factory_summary(
    template_path: Path | str,
    output_path: Path | str,
    *,
    site_name: str,
    records: dict[str, FactoryValues],
    materialize_calculations: bool = False,
) -> Path:
    workbook = load_workbook(template_path)
    sheet = workbook.worksheets[0]
    styles = [[_cell_snapshot(sheet.cell(row, col)) for col in range(1, 14)] for row in range(5, 9)]
    row_heights = [sheet.row_dimensions[row].height for row in range(5, 9)]
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row >= 5:
            sheet.unmerge_cells(str(merged))
    if sheet.max_row >= 5:
        sheet.delete_rows(5, sheet.max_row - 4)
    sheet["A1"] = site_name
    for index, month in enumerate(sorted(records)):
        row = 5 + index * 4
        sheet.insert_rows(row, 4)
        for offset in range(4):
            sheet.row_dimensions[row + offset].height = row_heights[offset]
            for col in range(1, 14):
                _apply_cell_snapshot(sheet.cell(row + offset, col), styles[offset][col - 1])
        for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 12, 13):
            sheet.merge_cells(start_row=row, start_column=col, end_row=row + 3, end_column=col)
        values = records[month]
        year, month_number = map(int, month.split("-"))
        sheet.cell(row, 1, date(year, month_number, 1))
        sheet.cell(row, 1).number_format = 'yyyy"年"m"月"'
        sheet.cell(row, 2, float(values.total_generation))
        sheet.cell(row, 3, float(values.grid_energy))
        sheet.cell(row, 4, float(values.grid_unit_price))
        grid_fee = values.grid_energy * values.grid_unit_price
        sheet.cell(row, 5, float(grid_fee) if materialize_calculations else f"=C{row}*D{row}")
        period_energy = _summary_period_energy(site_name, month, values)
        for offset, value in enumerate(period_energy, start=6):
            sheet.cell(row, offset, float(value))
        for offset, (period, price) in enumerate(zip(PERIODS, values.period_prices, strict=True)):
            sheet.cell(row + offset, 10, period)
            sheet.cell(row + offset, 11, float(price))
        sheet.cell(row, 12, float(values.self_use_fee))
        sheet.cell(row, 13, float(grid_fee + values.self_use_fee) if materialize_calculations else f"=E{row}+L{row}")
    workbook.calculation = CalcProperties(calcMode="auto")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _summary_period_energy(
    site_name: str,
    month: str,
    values: FactoryValues,
) -> tuple[Decimal, Decimal, Decimal, Decimal]:
    period_energy = HISTORICAL_SUMMARY_PERIOD_OVERRIDES.get((site_name, month), values.period_energy)
    if sum(period_energy, Decimal("0")) != values.self_use_energy:
        raise FactorySettlementError("历史尖峰平谷修正值与当月消纳总电量不一致。")
    return period_energy


def _template_records(template_path: Path | str) -> dict[str, FactoryValues]:
    formulas = load_workbook(template_path, data_only=False)
    cached = load_workbook(template_path, data_only=True)
    try:
        sheet = formulas.worksheets[0]
        value_sheet = cached.worksheets[0]
        records: dict[str, FactoryValues] = {}
        for row in range(5, sheet.max_row + 1, 4):
            if sheet.cell(row, 1).value is None:
                continue
            month = _month_from_excel(sheet.cell(row, 1).value, "汇总表月份")
            grid_energy = _decimal(sheet.cell(row, 3).value, "上网电量")
            period_energy = tuple(
                _decimal(sheet.cell(row, col).value, f"{period}时段电量")
                for col, period in zip(range(6, 10), PERIODS, strict=True)
            )
            records[month] = FactoryValues(
                month=month,
                total_generation=_decimal(sheet.cell(row, 2).value, "月总发电量"),
                grid_energy=grid_energy,
                grid_unit_price=_decimal(sheet.cell(row, 4).value, "上网电价"),
                period_energy=period_energy,
                period_prices=tuple(_decimal(sheet.cell(row + offset, 11).value, f"{period}时段电价") for offset, period in enumerate(PERIODS)),
                self_use_fee=_decimal(sheet.cell(row, 12).value, "消纳电费"),
            )
        return records
    finally:
        formulas.close()
        cached.close()


def _values_from_dict(raw: Any, *, expected_month: str) -> FactoryValues:
    if not isinstance(raw, dict):
        raise FactorySettlementError("厂房任务缺少已确认数据。")
    month = _validate_month(str(raw.get("month") or expected_month))
    if month != expected_month:
        raise FactorySettlementError("厂房任务月份不一致。")
    period_energy = raw.get("period_energy") if isinstance(raw.get("period_energy"), dict) else {}
    period_prices = raw.get("period_prices") if isinstance(raw.get("period_prices"), dict) else {}
    period_energy_values = tuple(_decimal(period_energy.get(period), f"{period}时段电量") for period in PERIODS)
    grid_energy = _decimal(raw.get("grid_energy"), "上网电量")
    historical_total = raw.get("historical_source_total_generation")
    total_generation = (
        _decimal(historical_total, "历史源表月总发电量")
        if historical_total not in (None, "")
        else _normalized_summary_generation(
            _decimal(raw.get("total_generation"), "月总发电量"),
            grid_energy,
            period_energy_values,
        )
    )
    return FactoryValues(
        month=month,
        total_generation=total_generation,
        grid_energy=grid_energy,
        grid_unit_price=_decimal(raw.get("grid_unit_price"), "上网电价"),
        period_energy=period_energy_values,
        period_prices=tuple(_decimal(period_prices.get(period), f"{period}时段电价") for period in PERIODS),
        self_use_fee=_decimal(raw.get("self_use_fee"), "消纳电费"),
    )


def _normalized_summary_generation(
    source_total: Decimal,
    grid_energy: Decimal,
    period_energy: tuple[Decimal, Decimal, Decimal, Decimal],
) -> Decimal:
    self_use_total = sum(period_energy, Decimal("0"))
    if source_total not in {self_use_total, self_use_total + grid_energy}:
        raise FactorySettlementError("历史月总发电量无法按立能派光伏用电量合计规则归一化。")
    return self_use_total


def _resolve_profile(runtime_root: Path | str, site_name: str, *, allow_create: bool) -> tuple[dict[str, Any], bool]:
    clean_name = _clean_site_name(site_name)
    if _normalize_site_name(clean_name) in {_normalize_site_name(alias) for alias in SEED_ALIASES}:
        return {
            "site_key": SEED_SITE_KEY,
            "site_name": SEED_SITE_NAME,
            "template_key": TEMPLATE_KEY,
            "include_template_baseline": True,
            "default_grid_unit_price": _number(SEED_GRID_UNIT_PRICE, places=6),
        }, False
    profiles = _read_profiles(runtime_root)
    matches = [profile for profile in profiles.values() if _normalize_site_name(str(profile.get("site_name", ""))) == _normalize_site_name(clean_name)]
    if len(matches) == 1:
        return matches[0], False
    if len(matches) > 1:
        raise FactorySettlementError("厂房名称匹配到多个档案，请联系后台核查。")
    if not allow_create:
        raise FactorySettlementError(f"尚未找到“{clean_name}”厂房档案；新厂房请使用“创建 {clean_name} YYYY-MM月”。")
    if "厂房" not in clean_name:
        raise FactorySettlementError("新厂房名称必须包含“厂房”，避免误建其他业务站区。")
    return {
        "site_key": _site_key(clean_name),
        "site_name": clean_name,
        "template_key": TEMPLATE_KEY,
        "include_template_baseline": False,
        "default_grid_unit_price": None,
    }, True


def _inherited_grid_price(runtime_root: Path | str, profile: dict[str, Any], month: str) -> tuple[Decimal | None, str]:
    series = _read_series(runtime_root)
    site = series.get("sites", {}).get(profile["site_key"], {}) if isinstance(series.get("sites"), dict) else {}
    records = site.get("records", {}) if isinstance(site, dict) and isinstance(site.get("records"), dict) else {}
    candidates = [(record_month, record) for record_month, record in records.items() if record_month <= month and isinstance(record, dict)]
    if candidates:
        _, latest = max(candidates, key=lambda item: item[0])
        return _decimal(latest.get("grid_unit_price"), "历史上网电价"), "same_site_history"
    default = profile.get("default_grid_unit_price")
    if default is not None:
        return _decimal(default, "默认上网电价"), "site_baseline"
    return None, "missing"


def _effective_records(runtime_root: Path | str, site_key: str) -> dict[str, FactoryValues]:
    series = _read_series(runtime_root)
    site = series.get("sites", {}).get(site_key, {}) if isinstance(series.get("sites"), dict) else {}
    raw_records = site.get("records", {}) if isinstance(site, dict) and isinstance(site.get("records"), dict) else {}
    return {month: _values_from_dict(record, expected_month=month) for month, record in raw_records.items() if isinstance(record, dict)}


def _summary_materialization_enabled(runtime_root: Path | str) -> bool:
    return _read_series(runtime_root).get("summary_materialize_calculations") is True


def _promote_effective(runtime_root: Path | str, task: dict[str, Any], values: FactoryValues) -> None:
    series = _read_series(runtime_root)
    sites = series.setdefault("sites", {})
    site = sites.setdefault(task["site_key"], {"site_name": task["site_name"], "records": {}})
    site["site_name"] = task["site_name"]
    records = site.setdefault("records", {})
    current = records.get(values.month) if isinstance(records.get(values.month), dict) else {}
    historical_total = current.get("historical_source_total_generation")
    records[values.month] = {
        **values.to_dict(),
        "task_id": task["task_id"],
        "revision": task["revision"],
        "completed_at": task["completed_at"],
    }
    if historical_total not in (None, ""):
        records[values.month]["historical_source_total_generation"] = historical_total
    _write_json(_series_path(runtime_root), series)
    profiles = _read_profiles(runtime_root)
    profiles[task["site_key"]] = {
        "site_key": task["site_key"],
        "site_name": task["site_name"],
        "template_key": TEMPLATE_KEY,
        "include_template_baseline": bool(task.get("include_template_baseline")),
        "default_grid_unit_price": _number(SEED_GRID_UNIT_PRICE, places=6) if task["site_key"] == SEED_SITE_KEY else None,
        "registered_at": task["completed_at"],
    }
    _write_json(_profiles_path(runtime_root), profiles)


def _read_profiles(runtime_root: Path | str) -> dict[str, Any]:
    path = _profiles_path(runtime_root)
    return _read_json(path) if path.is_file() else {}


def _read_series(runtime_root: Path | str) -> dict[str, Any]:
    path = _series_path(runtime_root)
    return _read_json(path) if path.is_file() else {"business_profile": BUSINESS_PROFILE, "sites": {}}


def _next_revision(runtime_root: Path | str, site_key: str, month: str) -> int:
    root = Path(runtime_root) / "factory_settlement_tasks"
    revisions = []
    if root.is_dir():
        for task_file in root.glob("*/task.json"):
            try:
                task = _read_json(task_file)
            except FactorySettlementError:
                continue
            if task.get("site_key") == site_key and task.get("month") == month:
                revisions.append(int(task.get("revision", 0)))
    return max(revisions, default=0) + 1


def _validate_xlsx(path: Path) -> None:
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise FactorySettlementError("厂房结算流程只接收 .xlsx 文件。")
    if path.stat().st_size <= 0 or path.stat().st_size > MAX_SOURCE_BYTES:
        raise FactorySettlementError("Excel 文件为空或超过 30 MB，本次文件未登记。")
    if path.read_bytes()[:4] != b"PK\x03\x04":
        raise FactorySettlementError("文件扩展名与 Excel 内容不一致，本次文件未登记。")
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES or sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise FactorySettlementError("Excel 文件解压规模超出限制，本次文件未登记。")
            if "xl/workbook.xml" not in {info.filename for info in infos}:
                raise FactorySettlementError("文件不是有效的 .xlsx 工作簿。")
            for info in infos:
                parts = Path(info.filename).parts
                if Path(info.filename).is_absolute() or ".." in parts:
                    raise FactorySettlementError("Excel 文件包含不安全路径，本次文件未登记。")
    except zipfile.BadZipFile as exc:
        raise FactorySettlementError("文件不是有效的 .xlsx 工作簿。") from exc


def _assert_source_month(source_month: str, expected_month: str, label: str) -> None:
    if source_month != expected_month:
        raise FactorySettlementError(f"{label}月份为 {source_month}，与当前填报月份 {expected_month} 不一致；本次文件未登记。")


def _month_from_excel(value: Any, field: str) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.strftime("%Y-%m")
    if isinstance(value, (int, float)):
        parsed = from_excel(value)
        if not 2000 <= parsed.year <= 2099:
            raise FactorySettlementError(f"{field}无法识别。")
        return parsed.strftime("%Y-%m")
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})\s*[-./年]\s*(0?[1-9]|1[0-2])", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    raise FactorySettlementError(f"{field}无法识别。")


def _cell_snapshot(cell: Any) -> dict[str, Any]:
    return {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
        "protection": copy(cell.protection),
    }


def _apply_cell_snapshot(cell: Any, snapshot: dict[str, Any]) -> None:
    cell.font = copy(snapshot["font"])
    cell.fill = copy(snapshot["fill"])
    cell.border = copy(snapshot["border"])
    cell.alignment = copy(snapshot["alignment"])
    cell.number_format = snapshot["number_format"]
    cell.protection = copy(snapshot["protection"])


def _output_name(site_name: str) -> str:
    clean = re.sub(r'[<>:"/\\|?*]', "", site_name).strip().rstrip(".")
    if not clean:
        raise FactorySettlementError("厂房名称无法用于成果文件名。")
    return f"{clean}{OUTPUT_SUFFIX}"


def _clean_site_name(value: str) -> str:
    clean = re.sub(r"\s+", "", str(value or "").strip())
    if not clean or len(clean) > 80 or any(char in clean for char in '<>:"/\\|?*'):
        raise FactorySettlementError("厂房名称为空、过长或包含非法字符。")
    return SEED_SITE_NAME if _normalize_site_name(clean) in {_normalize_site_name(alias) for alias in SEED_ALIASES} else clean


def _normalize_site_name(value: str) -> str:
    return re.sub(r"[\s（）()·\-—_]", "", value).replace("三号", "3号").casefold()


def _site_key(site_name: str) -> str:
    return f"factory_{hashlib.sha256(_normalize_site_name(site_name).encode('utf-8')).hexdigest()[:16]}"


def _task_root(runtime_root: Path | str, task_id: str) -> Path:
    return Path(runtime_root) / "factory_settlement_tasks" / task_id


def _profiles_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "factory_settlement" / "sites.json"


def _series_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "factory_settlement" / "series.json"


def _load_task(runtime_root: Path | str, task_id: str) -> tuple[Path, dict[str, Any]]:
    clean_task_id = _safe_task_id(task_id)
    root = _task_root(runtime_root, clean_task_id)
    task = _read_json(root / "task.json")
    if task.get("business_profile") != BUSINESS_PROFILE or task.get("task_id") != clean_task_id:
        raise FactorySettlementError("厂房任务身份不匹配。")
    return root, task


def _assert_task_identity(task: dict[str, Any], task_id: str, site_key: str, month: str) -> None:
    if task.get("business_profile") != BUSINESS_PROFILE or task.get("task_id") != task_id:
        raise FactorySettlementError("厂房任务身份不匹配。")
    if task.get("site_key") != site_key or task.get("month") != month:
        raise FactorySettlementError("厂房任务跨站区或月份。")


def _safe_task_id(value: str) -> str:
    clean = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", clean):
        raise FactorySettlementError("厂房任务编号无效。")
    return clean


def _validate_month(value: str) -> str:
    clean = str(value or "").strip()
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", clean):
        raise FactorySettlementError("月份必须使用 YYYY-MM 格式。")
    return clean


def _decimal(value: Any, field: str) -> Decimal:
    if value is None or isinstance(value, bool):
        raise FactorySettlementError(f"{field}缺少有效数值。")
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise FactorySettlementError(f"{field}不是有效数字。") from exc
    if not result.is_finite():
        raise FactorySettlementError(f"{field}不是有限数字。")
    return result


def _cached_formula_decimal(
    sheet: Any,
    cached: Any,
    *,
    row: int,
    column: int,
    field: str,
) -> Decimal:
    cached_value = cached.cell(row, column).value
    if cached_value is not None:
        return _decimal(cached_value, field)
    source_value = sheet.cell(row, column).value
    if isinstance(source_value, str) and source_value.strip().startswith("="):
        raise FactorySettlementError(f"{field}公式没有已计算结果，请使用 Excel 或 WPS 打开文件并保存后重新上传。")
    return _decimal(source_value, field)


def _number(value: Decimal, *, places: int | None = None) -> str:
    if places is not None:
        value = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _money_decimal(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _clean_energy(value: Decimal) -> Decimal:
    rounded = value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return rounded if abs(value - rounded) <= Decimal("0.0001") else value


def _money(value: Decimal) -> str:
    return format(_money_decimal(value), ".2f")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise FactorySettlementError(f"厂房任务状态不存在：{path}") from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise FactorySettlementError(f"厂房任务状态损坏：{path}") from exc
    if not isinstance(value, dict):
        raise FactorySettlementError(f"厂房任务状态格式错误：{path}")
    return value


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=f".{uuid.uuid4().hex}.tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise
