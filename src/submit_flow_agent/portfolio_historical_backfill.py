"""Auditable portfolio-wide historical summary backfill."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from copy import copy
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from submit_flow_agent.config import load_site_config
from submit_flow_agent.excel_writer import write_summary_excel
from submit_flow_agent.factory_settlement import FactoryValues, write_factory_summary
from submit_flow_agent.historical_backfill import _next_history_revision, _publish_history_revision
from submit_flow_agent.huizhou_zexin import ZexinValues, write_summary_workbook
from submit_flow_agent.monthly_summary import MonthlySummaryValues, write_monthly_summary_workbook
from submit_flow_agent.nanke import NankeValues, write_nanke_summary
from submit_flow_agent.task_series import record_historical_revision, resolve_history_workbooks
from submit_flow_agent.yangjiang_hongneng import (
    PeriodReadings,
    YangjiangSummaryValues,
    effective_yangjiang_summary_records,
    write_yangjiang_summary,
)


class PortfolioHistoricalBackfillError(RuntimeError):
    """Raised when a portfolio migration cannot be proven safe."""


SCHOOL_SOURCES = (
    ("site_166839c0594a", "兴围小学汇总表.xlsx", "2026-07"),
    ("site_91cc194f1159", "塘头学校汇总表.xlsx", "2026-07"),
    ("site_ec9cde5061de", "新安二外汇总表.xlsx", "2026-07"),
    ("site_eed6ce312d0d", "松岗实验学校汇总表.xlsx", "2026-06"),
    ("site_fd4aa59269bb", "沙井中学汇总表.xlsx", "2026-07"),
    ("xixiang_middle_school", "西乡中学汇总表.xlsx", "2026-07"),
)

SPECIAL_SOURCES = {
    "huizhou_zexin": "惠州泽鑫汇总表.xlsx",
    "factory_meizhou_multi": "梅州2、4、5厂房汇总表.xlsx",
    "factory_meizhou_3": "梅州3号厂房汇总表.xlsx",
    "tangxi": "汤西敬老院汇总表.xlsx",
    "nanke": "南科大汇总表.xlsx",
    "yangjiang": "阳江宏能科技汇总表.xlsx",
}


def backfill_portfolio_history(
    source_dir: Path | str,
    *,
    runtime_root: Path | str = "runtime",
    config_path: Path | str = "config/sites.json",
    apply: bool = False,
    rehearsal_root: Path | str | None = None,
) -> dict[str, Any]:
    """Stage all configured historical summaries and optionally publish them."""

    source_root = Path(source_dir)
    runtime = Path(runtime_root)
    config = Path(config_path)
    if not source_root.is_dir() or source_root.is_symlink():
        raise PortfolioHistoricalBackfillError(f"Historical source directory is missing or unsafe: {source_root}")
    required = [name for _, name, _ in SCHOOL_SOURCES] + list(SPECIAL_SOURCES.values())
    missing = [name for name in required if not (source_root / name).is_file()]
    if missing:
        raise PortfolioHistoricalBackfillError(f"Historical source files are missing: {', '.join(missing)}")

    workspace = Path(rehearsal_root) if rehearsal_root else Path(tempfile.mkdtemp(prefix="portfolio-history-backfill-"))
    workspace.mkdir(parents=True, exist_ok=True)
    staged_runtime = workspace / "runtime"
    staged_config = workspace / "config" / "sites.json"
    _copy_file(config, staged_config)
    config_changes = _prepare_school_config(staged_config)

    school_results = [
        _stage_school(
            source_root / source_name,
            site_key=site_key,
            through_month=through_month,
            production_runtime=runtime,
            staged_runtime=staged_runtime,
            config_path=staged_config,
            source_config_dir=config.parent.resolve(),
        )
        for site_key, source_name, through_month in SCHOOL_SOURCES
    ]

    series_results = _stage_special_series(source_root, runtime, staged_runtime)
    rehearsal_outputs = _render_special_rehearsals(staged_runtime, workspace / "rehearsal_outputs")
    result: dict[str, Any] = {
        "status": "ready_to_apply" if not apply else "applied",
        "source_dir": str(source_root.resolve()),
        "workspace": str(workspace.resolve()),
        "production_write": False,
        "config_changes": config_changes,
        "schools": school_results,
        "series": series_results,
        "rehearsal_outputs": rehearsal_outputs,
    }
    _write_json_atomic(workspace / "backfill_plan.json", result)
    if not apply:
        return result

    backup_root = runtime / "backups" / f"portfolio_historical_backfill_{_timestamp()}"
    _backup_scope(backup_root, runtime=runtime, config_path=config)
    _copy_file(staged_config, config)
    for item in school_results:
        staged_history = Path(item["staged_history"])
        target_history = runtime / "site_history" / item["site_key"] / item["through_month"]
        _publish_history_revision(
            staged_history,
            target_history=target_history,
            backup_root=backup_root,
            replace_existing=target_history.exists(),
        )
        manifest = _read_json(target_history / "manifest.json")
        item["series_result"] = record_historical_revision(runtime, manifest)
        item["history_dir"] = str(target_history.resolve())
    for relative in _series_paths():
        _copy_file(staged_runtime / relative, runtime / relative)
    result.update({
        "production_write": True,
        "backup_root": str(backup_root.resolve()),
    })
    _write_json_atomic(backup_root / "apply_receipt.json", result)
    return result


def _prepare_school_config(path: Path) -> list[dict[str, Any]]:
    payload = _read_json(path)
    changes: list[dict[str, Any]] = []
    for site in payload.get("sites") or []:
        if not isinstance(site, dict) or site.get("site_key") not in {item[0] for item in SCHOOL_SOURCES}:
            continue
        for field in ("summary_materialize_calculations", "summary_retroactive_full_history"):
            if site.get(field) is not True:
                changes.append({"site_key": site["site_key"], "field": field, "before": site.get(field), "after": True})
                site[field] = True
        if site.get("site_key") == "site_eed6ce312d0d":
            requested_meter = "09001SF00000042508942832"
            if site.get("meter_no") != requested_meter:
                changes.append({
                    "site_key": site["site_key"],
                    "field": "meter_no",
                    "before": site.get("meter_no"),
                    "after": requested_meter,
                })
                site["meter_no"] = requested_meter
            meters = site.get("meters") if isinstance(site.get("meters"), list) else []
            if meters:
                meters[0]["meter_no"] = requested_meter
        if site.get("dynamic_site_signature"):
            site["dynamic_site_signature"] = _site_signature(site)
    _write_json_atomic(path, payload)
    return changes


def _stage_school(
    source: Path,
    *,
    site_key: str,
    through_month: str,
    production_runtime: Path,
    staged_runtime: Path,
    config_path: Path,
    source_config_dir: Path,
) -> dict[str, Any]:
    site = load_site_config(site_key, config_path)
    revision = _next_history_revision(production_runtime / "site_history" / site_key / through_month)
    revision_name = f"r{revision:04d}"
    history_dir = staged_runtime / "site_history" / site_key / through_month
    revision_dir = history_dir / "revisions" / revision_name
    revision_dir.mkdir(parents=True, exist_ok=True)

    project = revision_dir / ".project"
    project.mkdir(parents=True, exist_ok=True)
    summary_working = project / site.summary_workbook
    months = _materialize_school_summary(source, summary_working, site=site, through_month=through_month)
    verified = _overlay_verified_school_months(
        summary_working,
        site=site,
        months=months,
        production_runtime=production_runtime,
        config_path=config_path,
    )
    _rewrite_school_meter_labels(summary_working, site=site, through_month=through_month)

    summary_name = site.output_workbook_name("summary", through_month)
    detail_name = site.output_workbook_name("detail", through_month)
    summary_path = revision_dir / summary_name
    shutil.copy2(summary_working, summary_path)
    detail_source = _latest_school_detail(
        production_runtime,
        site=site,
        through_month=through_month,
    ) or _production_template_path(
        site.template_dir,
        site.detail_workbook,
        staged_config_dir=config_path.parent.resolve(),
        source_config_dir=source_config_dir,
    )
    if not detail_source.is_file():
        raise PortfolioHistoricalBackfillError(f"School detail baseline is missing: {detail_source}")
    detail_path = revision_dir / detail_name
    shutil.copy2(detail_source, detail_path)
    shutil.rmtree(project)

    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    task_id = f"portfolio_historical_backfill_{site_key}_{through_month.replace('-', '_')}_{revision_name}"
    manifest = {
        "site_key": site_key,
        "site_name": site.site_name,
        "month": through_month,
        "revision": revision,
        "revision_name": revision_name,
        "source_task_id": task_id,
        "effective_revision": revision,
        "effective_task_id": task_id,
        "created_at": now,
        "updated_at": now,
        "migration": {
            "kind": "portfolio_historical_summary_backfill",
            "source": str(source.resolve()),
            "source_sha256": _sha256(source),
            "months": months,
            "verified_runtime_months": verified,
        },
        "baseline_dependency": None,
        "staleness": {
            "is_stale": False,
            "upstream_changed_month": None,
            "upstream_changed_revision": None,
            "upstream_changed_task_id": None,
            "marked_at": None,
            "cleared_at": now,
        },
        "workbooks": {
            "detail_workbook": {"path": f"revisions/{revision_name}/{detail_name}", "sha256": _sha256(detail_path)},
            "summary_workbook": {"path": f"revisions/{revision_name}/{summary_name}", "sha256": _sha256(summary_path)},
        },
    }
    _write_json_atomic(revision_dir / "manifest.json", {k: v for k, v in manifest.items() if k not in {"effective_revision", "effective_task_id"}})
    _write_json_atomic(history_dir / "manifest.json", manifest)
    return {
        "site_key": site_key,
        "site_name": site.site_name,
        "through_month": through_month,
        "months": months,
        "revision": revision,
        "verified_runtime_months": verified,
        "staged_history": str(history_dir.resolve()),
    }


def _materialize_school_summary(source: Path, output: Path, *, site: Any, through_month: str) -> list[str]:
    formulas = load_workbook(source, data_only=False)
    values = load_workbook(source, data_only=True)
    try:
        sheet = formulas[site.summary_sheet]
        value_sheet = values[site.summary_sheet]
        rows_to_delete: list[int] = []
        months: list[str] = []
        for row in range(site.summary_start_row, sheet.max_row + 1):
            month = _month_value(value_sheet.cell(row, 2).value)
            if month and month > through_month:
                rows_to_delete.append(row)
                continue
            if month:
                months.append(month)
            end_column = site.summary_end_column
            if isinstance(end_column, str):
                end_column = column_index_from_string(end_column)
            for column in range(1, int(end_column) + 1):
                formula_value = sheet.cell(row, column).value
                if isinstance(formula_value, str) and formula_value.startswith("="):
                    cached = value_sheet.cell(row, column).value
                    if cached is None:
                        raise PortfolioHistoricalBackfillError(
                            f"Historical formula has no cached value: {source.name}!{sheet.cell(row, column).coordinate}"
                        )
                    sheet.cell(row, column).value = cached
        for row in reversed(rows_to_delete):
            sheet.delete_rows(row)
        unique_months = list(dict.fromkeys(months))
        if not unique_months or unique_months[-1] != through_month:
            raise PortfolioHistoricalBackfillError(f"{source.name} does not contain complete history through {through_month}.")
        output.parent.mkdir(parents=True, exist_ok=True)
        formulas.save(output)
        return unique_months
    finally:
        formulas.close()
        values.close()


def _overlay_verified_school_months(
    summary_path: Path,
    *,
    site: Any,
    months: list[str],
    production_runtime: Path,
    config_path: Path,
) -> list[str]:
    verified: list[str] = []
    for month in months:
        manifest_path = production_runtime / "site_history" / site.site_key / month / "manifest.json"
        if not manifest_path.is_file():
            continue
        manifest = _read_json(manifest_path)
        task_id = manifest.get("effective_task_id") or manifest.get("source_task_id")
        task_path = production_runtime / "tasks" / str(task_id) / "task.json"
        if not task_path.is_file():
            continue
        task = _read_json(task_path)
        report_rel = next((item for item in task.get("outputs") or [] if str(item).endswith("validation_report.json")), None)
        report_path = task_path.parent / str(report_rel) if report_rel else None
        if report_path is None or not report_path.is_file():
            continue
        output = summary_path.parent / f"verified-{month}.xlsx"
        write_summary_excel(
            report_path,
            summary_path.parent,
            output,
            site_key=site.site_key,
            config_path=config_path,
            pricing_snapshot=task.get("pricing_snapshot"),
        )
        os.replace(output, summary_path)
        verified.append(month)
    return verified


def _rewrite_school_meter_labels(path: Path, *, site: Any, through_month: str) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook[site.summary_sheet]
        month_offsets: dict[str, int] = {}
        for row in range(site.summary_start_row, sheet.max_row + 1):
            month = _month_value(sheet.cell(row, 2).value)
            if not month or month > through_month:
                continue
            offset = month_offsets.get(month, 0)
            meter = site.meters[min(offset, len(site.meters) - 1)]
            sheet.cell(row, 1).value = f"表号{meter.meter_no}"
            month_offsets[month] = offset + 1
        workbook.save(path)
    finally:
        workbook.close()


def _latest_school_detail(runtime: Path, *, site: Any, through_month: str) -> Path | None:
    history_root = runtime / "site_history" / site.site_key
    if not history_root.is_dir():
        return None
    for month_dir in sorted((p for p in history_root.iterdir() if p.is_dir() and p.name <= through_month), reverse=True):
        resolved = resolve_history_workbooks(
            runtime,
            site_key=site.site_key,
            month=month_dir.name,
            workbook_names=(site.detail_workbook, site.summary_workbook),
        )
        if resolved:
            return resolved["detail_workbook"]
    return None


def _production_template_path(
    staged_template_dir: Path | str,
    workbook_name: str,
    *,
    staged_config_dir: Path,
    source_config_dir: Path,
) -> Path:
    template_dir = Path(staged_template_dir).resolve()
    try:
        relative = template_dir.relative_to(staged_config_dir)
    except ValueError:
        return template_dir / workbook_name
    return source_config_dir / relative / workbook_name


def _stage_special_series(source_root: Path, production_runtime: Path, staged_runtime: Path) -> dict[str, Any]:
    for relative in _series_paths():
        source = production_runtime / relative
        target = staged_runtime / relative
        if source.is_file():
            _copy_file(source, target)

    counts: dict[str, Any] = {}
    counts["huizhou_zexin"] = _stage_huizhou(source_root / SPECIAL_SOURCES["huizhou_zexin"], staged_runtime)
    counts["factory_meizhou_3"] = _stage_factory(
        source_root / SPECIAL_SOURCES["factory_meizhou_3"], staged_runtime, "factory_meizhou_3", "梅州3号厂房"
    )
    counts["factory_meizhou_multi"] = _stage_factory(
        source_root / SPECIAL_SOURCES["factory_meizhou_multi"],
        staged_runtime,
        "factory_efd23018c0aeefc2",
        "梅州2、4、5厂房",
    )
    counts["tangxi"] = _stage_tangxi(source_root / SPECIAL_SOURCES["tangxi"], staged_runtime)
    counts["nanke"] = _stage_nanke(source_root / SPECIAL_SOURCES["nanke"], staged_runtime)
    counts["yangjiang"] = _stage_yangjiang(source_root / SPECIAL_SOURCES["yangjiang"], staged_runtime)
    return counts


def _stage_huizhou(source: Path, runtime: Path) -> dict[str, Any]:
    path = runtime / "huizhou_zexin" / "series.json"
    series = _read_json(path, default={"site_key": "huizhou_zexin", "months": {}})
    series["summary_materialize_calculations"] = True
    months = series.setdefault("months", {})
    records = _sheet_rows(source, start=4, month_column=1, max_column=5)
    added = []
    for month, row_values in records:
        if month in months:
            continue
        inclusive = _decimal(row_values[4])
        exclusive = (inclusive / Decimal("1.13")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        record = ZexinValues(month, _decimal(row_values[1]), _decimal(row_values[2]), exclusive, inclusive)
        months[month] = _migration_month_record("huizhou_zexin", month, record.to_dict())
        added.append(month)
    _write_json_atomic(path, series)
    return {"source_months": [m for m, _ in records], "added_months": added, "preserved_months": sorted(set(months) - set(added))}


def _stage_factory(source: Path, runtime: Path, site_key: str, site_name: str) -> dict[str, Any]:
    path = runtime / "factory_settlement" / "series.json"
    series = _read_json(path, default={"business_profile": "factory_settlement_v1", "sites": {}})
    series["summary_materialize_calculations"] = True
    site = series.setdefault("sites", {}).setdefault(site_key, {"site_name": site_name, "records": {}})
    records = site.setdefault("records", {})
    formulas = load_workbook(source, data_only=False)
    values = load_workbook(source, data_only=True)
    added = []
    try:
        sheet = formulas.active
        value_sheet = values.active
        for row in range(5, sheet.max_row + 1, 4):
            month = _month_value(value_sheet.cell(row, 1).value)
            if not month:
                continue
            historical_total = _decimal(value_sheet.cell(row, 2).value)
            if month in records:
                if isinstance(records[month], dict):
                    records[month]["historical_source_total_generation"] = _number(historical_total)
                continue
            period_energy = tuple(_decimal(value_sheet.cell(row, column).value) for column in range(6, 10))
            period_prices = tuple(_decimal(value_sheet.cell(row + offset, 11).value) for offset in range(4))
            value = FactoryValues(
                month=month,
                total_generation=historical_total,
                grid_energy=_decimal(value_sheet.cell(row, 3).value),
                period_energy=period_energy,
                period_prices=period_prices,
                self_use_fee=_decimal(value_sheet.cell(row, 12).value),
                grid_unit_price=_decimal(value_sheet.cell(row, 4).value),
            )
            records[month] = {
                **value.to_dict(),
                "historical_source_total_generation": _number(value.total_generation),
                "task_id": f"historical_backfill_{site_key}_{month.replace('-', '_')}",
                "revision": 1,
                "completed_at": _utc_now(),
            }
            added.append(month)
    finally:
        formulas.close()
        values.close()
    _write_json_atomic(path, series)
    return {"added_months": added, "preserved_months": sorted(set(records) - set(added))}


def _stage_tangxi(source: Path, runtime: Path) -> dict[str, Any]:
    path = runtime / "monthly_summaries" / "tangxi_nursing_home" / "series.json"
    series = _read_json(path, default={"site_key": "tangxi_nursing_home", "months": {}})
    series["summary_materialize_calculations"] = True
    months = series.setdefault("months", {})
    records = _sheet_rows(source, start=3, month_column=4, max_column=11)
    added = []
    for month, row_values in records:
        if month in months:
            continue
        fee = _decimal(row_values[10])
        exclusive = (fee / Decimal("1.13")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        record = MonthlySummaryValues(month, _decimal(row_values[6]), fee, exclusive, fee - exclusive)
        months[month] = _migration_month_record("tangxi_nursing_home", month, record.to_dict(), output_key="output")
        added.append(month)
    _write_json_atomic(path, series)
    return {"source_months": [m for m, _ in records], "added_months": added, "preserved_months": sorted(set(months) - set(added))}


def _stage_nanke(source: Path, runtime: Path) -> dict[str, Any]:
    path = runtime / "nanke" / "series.json"
    series = _read_json(path, default={"site_key": "nanke", "months": {}})
    series["summary_materialize_calculations"] = True
    months = series.setdefault("months", {})
    values = load_workbook(source, data_only=True)
    added = []
    source_months = []
    try:
        sheet = values.active
        for row in range(4, sheet.max_row + 1):
            label = str(sheet.cell(row, 2).value or "").strip()
            number = int(label.replace("月", "")) if label.endswith("月") and label[:-1].isdigit() else 0
            if not number:
                continue
            month = f"{2025 if number == 12 else 2026:04d}-{number:02d}"
            source_months.append(month)
            if month in months:
                continue
            base_energy = _decimal(sheet.cell(row, 8).value)
            base_price = _decimal(sheet.cell(row, 9).value)
            record = NankeValues(
                month=month,
                meter_number="09001SF00000042509389672",
                prior_reading=_decimal(sheet.cell(row, 3).value),
                current_reading=_decimal(sheet.cell(row, 4).value),
                multiplier=_decimal(sheet.cell(row, 5).value),
                total_generation=_decimal(sheet.cell(row, 6).value),
                grid_energy=_decimal(sheet.cell(row, 7).value),
                grid_price=_decimal(sheet.cell(row, 10).value),
                grid_fee=_decimal(sheet.cell(row, 11).value),
                base_station_energy=base_energy,
                base_station_price=base_price,
                base_station_fee=(base_energy * base_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
            )
            months[month] = _migration_month_record("nanke", month, record.to_dict())
            added.append(month)
    finally:
        values.close()
    _write_json_atomic(path, series)
    return {"source_months": source_months, "added_months": added, "preserved_months": sorted(set(months) - set(added))}


def _stage_yangjiang(source: Path, runtime: Path) -> dict[str, Any]:
    path = runtime / "yangjiang_hongneng" / "series.json"
    series = _read_json(path, default={"site_key": "yangjiang_hongneng", "months": {}})
    series["summary_materialize_calculations"] = True
    operational = series.setdefault("months", {})
    historical = series.setdefault("historical_months", {})
    values = load_workbook(source, data_only=True)
    added = []
    source_months = []
    try:
        sheet = values.active
        for row in range(4, sheet.max_row + 1, 4):
            month = _month_value(sheet.cell(row, 2).value)
            if not month:
                continue
            source_months.append(month)
            if month in operational or month in historical:
                continue
            summary = YangjiangSummaryValues(
                month=month,
                summary_total_generation=_decimal(sheet.cell(row, 3).value),
                grid_meter_period_total_energy=_decimal(sheet.cell(row, 4).value),
                grid_settlement_price=_decimal(sheet.cell(row, 5).value),
                self_use_energy=_decimal(sheet.cell(row, 6).value),
                period_prices=PeriodReadings(*(_decimal(sheet.cell(row + offset, 8).value) for offset in range(4))),
                self_use_fee_before_tax_adjustment=_decimal(sheet.cell(row, 9).value),
                note=str(sheet.cell(row, 14).value or ""),
            )
            historical[month] = {
                "task_id": f"historical_backfill_yangjiang_hongneng_{month.replace('-', '_')}",
                "revision": 1,
                "values": summary.to_dict(),
                "source_type": "historical_summary_only",
                "updated_at": _utc_now(),
            }
            added.append(month)
    finally:
        values.close()
    _write_json_atomic(path, series)
    return {
        "source_months": source_months,
        "added_summary_only_months": added,
        "preserved_operational_months": sorted(operational),
    }


def _render_special_rehearsals(runtime: Path, output_dir: Path) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, str] = {}

    from submit_flow_agent.factory_settlement import _effective_records as factory_records
    from submit_flow_agent.huizhou_zexin import _effective_records as huizhou_records
    from submit_flow_agent.monthly_summary import _effective_records as tangxi_records
    from submit_flow_agent.nanke import _effective_records as nanke_records

    targets = {
        "huizhou_zexin": (
            write_summary_workbook,
            Path("config/templates/huizhou_zexin/惠州泽鑫汇总表.xlsx"),
            output_dir / "惠州泽鑫汇总表.xlsx",
            huizhou_records(runtime),
        ),
        "tangxi": (
            write_monthly_summary_workbook,
            Path("config/templates/汤西敬老院汇总表.xlsx"),
            output_dir / "汤西敬老院汇总表.xlsx",
            tangxi_records(runtime),
        ),
        "nanke": (
            write_nanke_summary,
            Path("config/templates/nanke/南科大汇总表.xlsx"),
            output_dir / "南科大汇总表.xlsx",
            nanke_records(runtime),
        ),
        "yangjiang": (
            write_yangjiang_summary,
            Path("config/templates/yangjiang_hongneng/阳江宏能科技汇总表.xlsx"),
            output_dir / "阳江宏能科技汇总表.xlsx",
            effective_yangjiang_summary_records(runtime),
        ),
    }
    for key, (writer, template, output, records) in targets.items():
        writer(template, output, records, materialize_calculations=True)
        outputs[key] = str(output.resolve())

    for site_key, site_name in (("factory_meizhou_3", "梅州3号厂房"), ("factory_efd23018c0aeefc2", "梅州2、4、5厂房")):
        output = output_dir / f"{site_name}汇总表.xlsx"
        write_factory_summary(
            Path("config/templates/factory_meizhou_v1/梅州3号厂房汇总表.xlsx"),
            output,
            site_name=site_name,
            records=factory_records(runtime, site_key),
            materialize_calculations=True,
        )
        outputs[site_key] = str(output.resolve())
    return outputs


def _sheet_rows(
    source: Path,
    *,
    start: int,
    month_column: int,
    max_column: int,
) -> list[tuple[str, tuple[Any, ...]]]:
    values = load_workbook(source, data_only=True)
    try:
        result: list[tuple[str, tuple[Any, ...]]] = []
        sheet = values.active
        for row in range(start, sheet.max_row + 1):
            month = _month_value(sheet.cell(row, month_column).value)
            if month:
                result.append((month, tuple(sheet.cell(row, column).value for column in range(1, max_column + 1))))
        return result
    finally:
        values.close()


def _migration_month_record(site_key: str, month: str, values: dict[str, Any], *, output_key: str = "outputs") -> dict[str, Any]:
    return {
        "task_id": f"historical_backfill_{site_key}_{month.replace('-', '_')}",
        "revision": 1,
        "values": values,
        output_key: [] if output_key == "outputs" else "",
        "history": [],
        "updated_at": _utc_now(),
    }


def _series_paths() -> tuple[Path, ...]:
    return (
        Path("huizhou_zexin/series.json"),
        Path("factory_settlement/series.json"),
        Path("monthly_summaries/tangxi_nursing_home/series.json"),
        Path("nanke/series.json"),
        Path("yangjiang_hongneng/series.json"),
    )


def _backup_scope(backup: Path, *, runtime: Path, config_path: Path) -> None:
    backup.mkdir(parents=True, exist_ok=False)
    _copy_file(config_path, backup / "config" / "sites.json")
    for relative in _series_paths():
        source = runtime / relative
        if source.is_file():
            _copy_file(source, backup / "runtime" / relative)
    for site_key, _, through_month in SCHOOL_SOURCES:
        for scope in ("site_history", "site_series", "site_pricing"):
            source = runtime / scope / site_key
            if source.exists():
                shutil.copytree(source, backup / "runtime" / scope / site_key)


def _site_signature(site: dict[str, Any]) -> str:
    payload = {key: value for key, value in site.items() if key != "dynamic_site_signature"}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _month_value(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value or "").strip()
    if len(text) >= 7 and text[:4].isdigit():
        compact = text.replace("年", "-").replace("月", "")
        parts = compact.split("-")
        if len(parts) >= 2 and parts[1].isdigit():
            return f"{int(parts[0]):04d}-{int(parts[1]):02d}"
    return None


def _decimal(value: Any) -> Decimal:
    if value is None or value == "":
        raise PortfolioHistoricalBackfillError("Required historical value is blank.")
    return Decimal(str(value))


def _number(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S_%f")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path, default: dict[str, Any] | None = None) -> dict[str, Any]:
    if not path.is_file():
        if default is not None:
            return default
        raise PortfolioHistoricalBackfillError(f"JSON file is missing: {path}")
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict):
        raise PortfolioHistoricalBackfillError(f"JSON file must contain an object: {path}")
    return payload


def _copy_file(source: Path, target: Path) -> None:
    if not source.is_file() or source.is_symlink():
        raise PortfolioHistoricalBackfillError(f"Required file is missing or unsafe: {source}")
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    finally:
        Path(temporary_name).unlink(missing_ok=True)
