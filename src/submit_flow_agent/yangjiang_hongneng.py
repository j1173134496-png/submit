"""Independent monthly reporting flow for Yangjiang Hongneng Technology."""

from __future__ import annotations

import calendar
import hashlib
import json
import re
import shutil
import unicodedata
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

from submit_flow_agent.ocr.base import OcrAdapter, OcrBlock, OcrRawResult, write_ocr_raw_result
from submit_flow_agent.ocr.factory import create_ocr_adapter
from submit_flow_agent.pdf_renderer import RenderedPage, render_pdf_to_images


BUSINESS_PROFILE = "yangjiang_hongneng_monthly_reporting_v1"
SITE_KEY = "yangjiang_hongneng"
SITE_NAME = "阳江宏能科技"
PROJECT_TITLE = "阳江宏能科技有限公司分布式屋顶光伏发电项目"
OUTPUT_DETAIL = "阳江宏能科技用电统计表.xlsx"
OUTPUT_SUMMARY = "阳江宏能科技汇总表.xlsx"
REQUIRED_ROLES = (
    "generation_meter_2065",
    "generation_meter_2066",
    "electricity_bill",
    "grid_settlement",
    "grid_meter_export",
)
PERIODS = ("尖", "峰", "平", "谷")
PERIOD_KEYS = ("sharp", "peak", "flat", "valley")
METER_SUFFIXES = ("2065", "2066")
METER_NUMBERS = {
    "2065": "03591SF00000002504222065",
    "2066": "03591SF00000002504222066",
}
GENERATION_MULTIPLIER = Decimal("100")
SETTLEMENT_FACTOR = Decimal("0.67")
CAPACITY_KWP = Decimal("680.4")
EFFECTIVE_HOUR_CAPACITY_KW = Decimal("600")
TAX_DIVISOR = Decimal("1.06")
TAX_MULTIPLIER = Decimal("1.01")
GRID_METER_NUMBER = "03591SG00000002100073001"
GRID_METER_MULTIPLIER = Decimal("1000")
# Operations confirmed June's small sharp line is an adjustment, not an active sharp tariff.
SHARP_TARIFF_INACTIVE_MONTHS = frozenset({"2026-06"})
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".pdf", ".xlsx"}
MAX_SOURCE_BYTES = 30 * 1024 * 1024

# April is the controlled seed month. Later months must chain from the immediately
# preceding effective record so cumulative meter readings never span missing months.
SEED_MONTH = "2026-04"
SEED_GENERATION_READINGS = {
    "2065": {"sharp": "116.80", "peak": "105.40", "flat": "159.13", "valley": "9.90"},
    "2066": {"sharp": "117.33", "peak": "105.38", "flat": "158.92", "valley": "9.78"},
}
SEED_GRID_PERIOD_READINGS = {"sharp": "11.69", "peak": "10.24", "flat": "16.31", "valley": "0.65"}
SEED_GRID_TOTAL_READING = Decimal("38.90")


class YangjiangHongnengError(RuntimeError):
    """Raised when the fixed Yangjiang Hongneng contract cannot advance."""


@dataclass(frozen=True)
class PeriodReadings:
    sharp: Decimal
    peak: Decimal
    flat: Decimal
    valley: Decimal

    def values(self) -> tuple[Decimal, Decimal, Decimal, Decimal]:
        return self.sharp, self.peak, self.flat, self.valley

    def to_dict(self) -> dict[str, str]:
        return {key: _number(value) for key, value in zip(PERIOD_KEYS, self.values(), strict=True)}

    @classmethod
    def from_dict(cls, raw: Any, field: str) -> "PeriodReadings":
        if not isinstance(raw, dict):
            raise YangjiangHongnengError(f"{field}缺少尖峰平谷数据。")
        return cls(*(_decimal(raw.get(key), f"{field}{period}") for key, period in zip(PERIOD_KEYS, PERIODS, strict=True)))


@dataclass(frozen=True)
class YangjiangSummaryValues:
    """Historical values that are valid for the cumulative summary only."""

    month: str
    summary_total_generation: Decimal
    grid_meter_period_total_energy: Decimal
    grid_settlement_price: Decimal
    self_use_energy: Decimal
    period_prices: PeriodReadings
    self_use_fee_before_tax_adjustment: Decimal
    note: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "month": self.month,
            "summary_total_generation": _number(self.summary_total_generation),
            "grid_meter_period_total_energy": _number(self.grid_meter_period_total_energy),
            "grid_settlement_price": _number(self.grid_settlement_price, places=6),
            "self_use_energy": _number(self.self_use_energy),
            "period_prices": self.period_prices.to_dict(),
            "self_use_fee_before_tax_adjustment": _money(self.self_use_fee_before_tax_adjustment),
            "note": self.note,
        }

    @classmethod
    def from_dict(cls, raw: Any, *, expected_month: str) -> "YangjiangSummaryValues":
        if not isinstance(raw, dict):
            raise YangjiangHongnengError("阳江宏能科技历史汇总记录缺少数据。")
        month = _validate_month(str(raw.get("month") or expected_month))
        if month != expected_month:
            raise YangjiangHongnengError("阳江宏能科技历史汇总记录月份不一致。")
        return cls(
            month=month,
            summary_total_generation=_decimal(raw.get("summary_total_generation"), "汇总总发电量"),
            grid_meter_period_total_energy=_decimal(raw.get("grid_meter_period_total_energy"), "汇总上网电量"),
            grid_settlement_price=_decimal(raw.get("grid_settlement_price"), "汇总上网电价"),
            self_use_energy=_decimal(raw.get("self_use_energy"), "汇总消纳电量"),
            period_prices=PeriodReadings.from_dict(raw.get("period_prices"), "汇总分时电价"),
            self_use_fee_before_tax_adjustment=_decimal(
                raw.get("self_use_fee_before_tax_adjustment"), "汇总消纳电费"
            ),
            note=str(raw.get("note") or ""),
        )


@dataclass(frozen=True)
class YangjiangValues:
    month: str
    prior_generation_readings: dict[str, PeriodReadings]
    current_generation_readings: dict[str, PeriodReadings]
    grid_prior_readings: PeriodReadings
    grid_current_readings: PeriodReadings
    grid_prior_total_reading: Decimal
    grid_current_total_reading: Decimal
    grid_meter_multiplier: Decimal
    period_prices: PeriodReadings
    price_components: dict[str, dict[str, Decimal]]
    market_allocation_fee: Decimal
    billed_period_energy: PeriodReadings
    grid_settlement_energy: Decimal
    grid_settlement_price: Decimal
    grid_settlement_fee: Decimal

    @property
    def generation_period_energy(self) -> PeriodReadings:
        totals: list[Decimal] = []
        for index in range(4):
            energy = sum(
                (
                    self.current_generation_readings[suffix].values()[index]
                    - self.prior_generation_readings[suffix].values()[index]
                )
                * GENERATION_MULTIPLIER
                for suffix in METER_SUFFIXES
            )
            totals.append(_clean_energy(energy))
        return PeriodReadings(*totals)

    @property
    def grid_meter_period_energy(self) -> PeriodReadings:
        return PeriodReadings(
            *(
                _clean_energy((current - prior) * self.grid_meter_multiplier)
                for prior, current in zip(
                    self.grid_prior_readings.values(),
                    self.grid_current_readings.values(),
                    strict=True,
                )
            )
        )

    @property
    def grid_meter_total_energy(self) -> Decimal:
        return _clean_energy(
            (self.grid_current_total_reading - self.grid_prior_total_reading) * self.grid_meter_multiplier
        )

    @property
    def grid_meter_period_total_energy(self) -> Decimal:
        return sum(self.grid_meter_period_energy.values(), Decimal("0"))

    @property
    def self_use_period_energy(self) -> PeriodReadings:
        return PeriodReadings(
            *(
                generation - grid
                for generation, grid in zip(
                    self.generation_period_energy.values(),
                    self.grid_meter_period_energy.values(),
                    strict=True,
                )
            )
        )

    @property
    def self_use_energy(self) -> Decimal:
        return sum(self.self_use_period_energy.values(), Decimal("0"))

    @property
    def self_use_fee_before_tax_adjustment(self) -> Decimal:
        return sum(
            (energy * price for energy, price in zip(self.self_use_period_energy.values(), self.period_prices.values(), strict=True)),
            Decimal("0"),
        )

    @property
    def self_use_fee(self) -> Decimal:
        return self.self_use_fee_before_tax_adjustment / TAX_DIVISOR * TAX_MULTIPLIER

    @property
    def summary_total_generation(self) -> Decimal:
        return sum(self.generation_period_energy.values(), Decimal("0"))

    @property
    def grid_income(self) -> Decimal:
        return self.grid_meter_period_total_energy * self.grid_settlement_price

    @property
    def total_income(self) -> Decimal:
        return self.self_use_fee_before_tax_adjustment + self.grid_income

    @property
    def sharp_price_source(self) -> str:
        return (
            "bill_sharp"
            if _has_effective_sharp_price(self.month, self.billed_period_energy, self.price_components)
            else "peak_fallback"
        )

    @property
    def self_use_ratio(self) -> Decimal:
        return self.self_use_energy / self.summary_total_generation

    @property
    def effective_hours(self) -> Decimal:
        year, month_number = map(int, self.month.split("-"))
        days = Decimal(calendar.monthrange(year, month_number)[1])
        return self.summary_total_generation / EFFECTIVE_HOUR_CAPACITY_KW / days

    def to_dict(self) -> dict[str, Any]:
        return {
            "month": self.month,
            "site_name": SITE_NAME,
            "settlement_factor": _number(SETTLEMENT_FACTOR),
            "prior_generation_readings": {key: value.to_dict() for key, value in self.prior_generation_readings.items()},
            "current_generation_readings": {key: value.to_dict() for key, value in self.current_generation_readings.items()},
            "generation_period_energy": self.generation_period_energy.to_dict(),
            "grid_prior_readings": self.grid_prior_readings.to_dict(),
            "grid_current_readings": self.grid_current_readings.to_dict(),
            "grid_prior_total_reading": _number(self.grid_prior_total_reading),
            "grid_current_total_reading": _number(self.grid_current_total_reading),
            "grid_meter_multiplier": _number(self.grid_meter_multiplier),
            "grid_meter_period_energy": self.grid_meter_period_energy.to_dict(),
            "grid_meter_period_total_energy": _number(self.grid_meter_period_total_energy),
            "grid_meter_total_energy": _number(self.grid_meter_total_energy),
            "billed_period_energy": self.billed_period_energy.to_dict(),
            "self_use_period_energy": self.self_use_period_energy.to_dict(),
            "self_use_energy": _number(self.self_use_energy),
            "period_prices": self.period_prices.to_dict(),
            "sharp_price_source": self.sharp_price_source,
            "price_components": {
                period: {key: _number(value) for key, value in components.items()}
                for period, components in self.price_components.items()
            },
            "market_allocation_fee": _money(self.market_allocation_fee),
            "self_use_fee_before_tax_adjustment": _money(self.self_use_fee_before_tax_adjustment),
            "self_use_fee": _money(self.self_use_fee),
            "grid_settlement_energy": _number(self.grid_settlement_energy),
            "grid_settlement_price": _number(self.grid_settlement_price, places=6),
            "grid_settlement_fee": _money(self.grid_settlement_fee),
            "summary_total_generation": _number(self.summary_total_generation),
            "grid_income": _money(self.grid_income),
            "total_income": _money(self.total_income),
            "self_use_ratio": _number(self.self_use_ratio, places=8),
            "effective_hours": _number(self.effective_hours, places=8),
        }


def prepare_yangjiang_hongneng(runtime_root: Path | str, *, task_id: str, month: str) -> dict[str, Any]:
    clean_task_id = _safe_task_id(task_id)
    clean_month = _validate_month(month)
    root = _task_root(runtime_root, clean_task_id)
    if root.exists():
        task = _read_json(root / "task.json")
        _assert_task_identity(task, clean_task_id, clean_month)
        return task
    _prior_generation_readings(runtime_root, clean_month)
    _prior_grid_export(runtime_root, clean_month)
    root.mkdir(parents=True, exist_ok=False)
    task = {
        "task_id": clean_task_id,
        "business_profile": BUSINESS_PROFILE,
        "site_key": SITE_KEY,
        "site_name": SITE_NAME,
        "month": clean_month,
        "revision": _next_revision(runtime_root, clean_month),
        "settlement_factor": _number(SETTLEMENT_FACTOR),
        "status": "collecting_files",
        "stage": "collecting_files",
        "sources": {},
        "source_count": 0,
        "recognized_data": {},
        "confirmed_data": {},
        "missing_fields": [],
        "warnings": [],
        "outputs": [],
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
    }
    _write_json(root / "task.json", task)
    return task


def collect_yangjiang_hongneng_source(
    runtime_root: Path | str,
    task_id: str,
    source: Path | str,
    *,
    ocr_adapter: OcrAdapter | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_files", "awaiting_confirmation", "need_review"}:
        raise YangjiangHongnengError("当前阳江宏能科技任务不接收文件。")
    source_path = Path(source)
    _validate_source(source_path)
    digest = _sha256_file(source_path)
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    if any(isinstance(item, dict) and item.get("sha256") == digest for item in sources.values()):
        return {**task, "duplicate": True}

    staged = root / "inputs" / f"source-{digest[:16]}{source_path.suffix.lower()}"
    staged.parent.mkdir(parents=True, exist_ok=True)
    if not staged.exists():
        shutil.copy2(source_path, staged)
    if staged.suffix.lower() == ".xlsx":
        role, extracted = recognize_yangjiang_hongneng_meter_export(staged, expected_month=task["month"])
        _write_json(root / "ocr" / f"{role}.json", {"structured_source": True, "recognized": extracted})
    else:
        pages = _source_pages(staged, root / "rendered" / digest[:16])
        raw = (ocr_adapter or create_ocr_adapter()).recognize(staged, pages)
        role, extracted = recognize_yangjiang_hongneng_source(raw, expected_month=task["month"])
        write_ocr_raw_result(raw, root / "ocr" / f"{role}.json", relative_to=root)
    if role in sources:
        staged.unlink(missing_ok=True)
        raise YangjiangHongnengError(f"{_role_label(role)}已登记，请勿重复上传同类文件。")
    source_month = extracted.get("month")
    if source_month and source_month != task["month"]:
        staged.unlink(missing_ok=True)
        raise YangjiangHongnengError(
            f"文件账期为 {source_month}，与当前填报月份 {task['month']} 不一致；本次文件未登记。"
        )
    sources[role] = {
        "role": role,
        "name": source_path.name,
        "stored_path": str(staged.relative_to(root)).replace("\\", "/"),
        "sha256": digest,
        "recognized": extracted,
    }
    task["sources"] = sources
    task["source_count"] = sum(role_name in sources for role_name in REQUIRED_ROLES)
    task["updated_at"] = _utc_now()
    if not all(role_name in sources for role_name in REQUIRED_ROLES):
        task["status"] = task["stage"] = "collecting_files"
        _write_json(root / "task.json", task)
        return task

    values = _combine_sources(runtime_root, task)
    warnings = []
    negative_periods = [period for period, value in zip(PERIODS, values.self_use_period_energy.values(), strict=True) if value < 0]
    if negative_periods:
        warnings.append(f"按四月模板公式计算，以下时段消纳量为负数：{'、'.join(negative_periods)}。已保留原始业务口径。")
    settlement_difference = values.grid_meter_total_energy - values.grid_settlement_energy
    if settlement_difference:
        warnings.append(
            "抄表反向总电量与供电所结算电量相差"
            f" {_number(abs(settlement_difference))} kWh；请核对供电所结算材料。"
        )
    period_total_difference = values.grid_meter_period_total_energy - values.grid_meter_total_energy
    if period_total_difference:
        warnings.append(
            "上网表分时电量合计与总示数差相差"
            f" {_number(abs(period_total_difference))} kWh；用电统计表和汇总表按分时电量合计填报。"
        )
    task.update(
        {
            "recognized_data": values.to_dict(),
            "missing_fields": [],
            "warnings": warnings,
            "status": "awaiting_confirmation",
            "stage": "awaiting_confirmation",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def confirm_yangjiang_hongneng(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] != "awaiting_confirmation":
        raise YangjiangHongnengError("当前阳江宏能科技任务不在数据确认阶段。")
    values = _values_from_dict(task.get("recognized_data"), expected_month=task["month"])
    task.update(
        {
            "confirmed_data": values.to_dict(),
            "status": "ready_to_run",
            "stage": "ready_to_run",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def run_yangjiang_hongneng(
    runtime_root: Path | str,
    task_id: str,
    *,
    template_dir: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] == "completed":
        return task
    if task["status"] != "ready_to_run":
        raise YangjiangHongnengError("请先核对并确认识别数据，再发送“开始运行”。")
    values = _values_from_dict(task.get("confirmed_data"), expected_month=task["month"])
    templates = Path(template_dir)
    detail_template = templates / OUTPUT_DETAIL
    summary_template = templates / OUTPUT_SUMMARY
    for template in (detail_template, summary_template):
        if not template.is_file():
            raise YangjiangHongnengError(f"受控模板不存在：{template}")
    output_dir = root / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    detail_output = output_dir / OUTPUT_DETAIL
    summary_output = output_dir / OUTPUT_SUMMARY
    write_yangjiang_detail(detail_template, detail_output, values)
    records = effective_yangjiang_summary_records(runtime_root)
    records[values.month] = values
    write_yangjiang_summary(
        summary_template,
        summary_output,
        records,
        materialize_calculations=_summary_materialization_enabled(runtime_root),
    )
    outputs = [detail_output, summary_output]
    if any(not output.is_file() or output.stat().st_size <= 0 for output in outputs):
        raise YangjiangHongnengError("阳江宏能科技成果生成不完整，任务未完成。")
    task.update(
        {
            "status": "completed",
            "stage": "completed",
            "outputs": [str(output.relative_to(root)).replace("\\", "/") for output in outputs],
            "completed_at": _utc_now(),
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    _promote_effective(runtime_root, task, values, outputs)
    return task


def get_yangjiang_hongneng(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    return _load_task(runtime_root, task_id)[1]


def list_yangjiang_hongneng_outputs(runtime_root: Path | str, task_id: str) -> list[str]:
    root, task = _load_task(runtime_root, task_id)
    return [str((root / item).resolve()) for item in task.get("outputs", []) if (root / item).is_file()]


def recognize_yangjiang_hongneng_source(
    raw: OcrRawResult,
    *,
    expected_month: str,
) -> tuple[str, dict[str, Any]]:
    text = "\n".join(block.text for page in raw.pages for block in page.blocks)
    compact = _compact(text)
    if "2065" in compact and ("正向有功" in compact or "功率1电能" in compact):
        return "generation_meter_2065", {
            "meter_suffix": "2065",
            "current_readings": _extract_generation_meter_readings(raw).to_dict(),
        }
    if "2066" in compact and ("正向有功" in compact or "功率1电能" in compact):
        return "generation_meter_2066", {
            "meter_suffix": "2066",
            "current_readings": _extract_generation_meter_readings(raw).to_dict(),
        }
    if (
        _looks_like_grid_meter_number(compact)
        and "反向有" in compact
        and "综合" in compact
        and "倍率" in compact
    ):
        return "grid_meter_export", _extract_grid_meter_image(raw, expected_month)
    month = _month_from_text(text, expected_month)
    if "电费通知单" in compact or "0317100301492144" in compact:
        return "electricity_bill", _extract_electricity_bill(raw, month)
    if "电费结算单" in compact and ("市场化电费" in compact or "抄见电量" in compact):
        return "grid_settlement", _extract_grid_settlement(raw, month)
    raise YangjiangHongnengError(
        "无法识别文件角色，请上传两张2065/2066电表截图、电费账单、阳江宏能科技结算单和电表抄表 XLSX/图片/PDF。"
    )


def recognize_yangjiang_hongneng_meter_export(
    source: Path | str,
    *,
    expected_month: str,
) -> tuple[str, dict[str, Any]]:
    path = Path(source)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise YangjiangHongnengError("电表抄表 XLSX 无法读取，本次文件未登记。") from exc
    try:
        sheet = workbook.active
        meter_number = str(sheet["D2"].value or "").strip()
        if meter_number != GRID_METER_NUMBER:
            raise YangjiangHongnengError(
                f"电表抄表 XLSX 的表号不是 {GRID_METER_NUMBER}，本次文件未登记。"
            )
        reading_date = _excel_reading_date(sheet["G2"].value)
        source_month = _previous_month(f"{reading_date.year:04d}-{reading_date.month:02d}")
        if reading_date.day != 1 or source_month != expected_month:
            raise YangjiangHongnengError(
                f"电表抄表 XLSX 对应账期为 {source_month}，与当前填报月份 {expected_month} 不一致。"
            )

        time_of_use_values = [_optional_decimal(sheet.cell(2, column).value) for column in range(14, 19)]
        export_multiplier = _optional_decimal(sheet["W2"].value)
        has_time_of_use_columns = all(value is not None for value in time_of_use_values) and export_multiplier is not None
        if has_time_of_use_columns:
            current_total = time_of_use_values[0]
            current_period = PeriodReadings(*time_of_use_values[1:])
            multiplier = export_multiplier
            schema = "time_of_use_v2"
        elif expected_month == SEED_MONTH:
            current_total = _decimal(sheet["J2"].value, "反向有功总电能")
            current_period = PeriodReadings.from_dict(SEED_GRID_PERIOD_READINGS, "四月反向有功历史基准")
            multiplier = _decimal(sheet["O2"].value, "电表综合倍率")
            schema = "legacy_total_with_controlled_april_period_seed"
        else:
            raise YangjiangHongnengError(
                "电表抄表 XLSX 缺少反向有功尖、峰、平、谷列；仅四月历史基线允许旧版格式。"
            )
    finally:
        workbook.close()

    if multiplier <= 0:
        raise YangjiangHongnengError("电表抄表 XLSX 的综合倍率必须大于 0。")
    if current_total < 0 or any(value < 0 for value in current_period.values()):
        raise YangjiangHongnengError("电表抄表 XLSX 的反向有功累计示数不能为负数。")
    return "grid_meter_export", {
        "month": source_month,
        "grid_meter_number": meter_number,
        "reading_date": reading_date.isoformat(),
        "reverse_total_reading": _number(current_total),
        "reverse_period_readings": current_period.to_dict(),
        "grid_meter_multiplier": _number(multiplier),
        "schema": schema,
    }


def _extract_grid_meter_image(raw: OcrRawResult, expected_month: str) -> dict[str, Any]:
    blocks = [block for page in raw.pages for block in page.blocks]
    rows = _ocr_rows(raw, tolerance=12)
    data_rows = [
        row
        for row in rows
        if _looks_like_grid_meter_number(_compact("".join(block.text for block in row)))
    ]
    if len(data_rows) != 1:
        raise YangjiangHongnengError("电表抄表图片无法唯一定位目标电表数据行，本次文件未登记。")
    data_row = data_rows[0]
    data_y = sum((block.bbox[1] + block.bbox[3]) / 2 for block in data_row) / len(data_row)
    header_columns = _grid_meter_header_columns([block for block in blocks if (block.bbox[1] + block.bbox[3]) / 2 < data_y])

    def column_text(label: str) -> str:
        matches = [
            x_center
            for x_center, text in header_columns
            if label in text and not (label == "反向有功总" and "无功" in text)
        ]
        if len(matches) != 1:
            raise YangjiangHongnengError(f"电表抄表图片无法唯一识别{label}列，本次文件未登记。")
        target_x = matches[0]
        return min(
            data_row,
            key=lambda block: abs(((block.bbox[0] + block.bbox[2]) / 2) - target_x),
        ).text.strip()

    meter_number = GRID_METER_NUMBER
    reading_date = _excel_reading_date(column_text("数据时间"), source_label="电表抄表图片")
    source_month = _previous_month(f"{reading_date.year:04d}-{reading_date.month:02d}")
    if reading_date.day != 1 or source_month != expected_month:
        raise YangjiangHongnengError(
            f"电表抄表图片对应账期为 {source_month}，与当前填报月份 {expected_month} 不一致。"
        )
    current_total = _decimal(column_text("反向有功总"), "电表抄表图片反向有功总电能")
    current_period = PeriodReadings(
        *(
            _decimal(column_text(f"反向有功{period}"), f"电表抄表图片反向有功{period}")
            for period in PERIODS
        )
    )
    multiplier = _decimal(column_text("综合倍率"), "电表抄表图片综合倍率")
    if multiplier <= 0:
        raise YangjiangHongnengError("电表抄表图片的综合倍率必须大于 0。")
    if current_total < 0 or any(value < 0 for value in current_period.values()):
        raise YangjiangHongnengError("电表抄表图片的反向有功累计示数不能为负数。")
    return {
        "month": source_month,
        "grid_meter_number": meter_number,
        "reading_date": reading_date.isoformat(),
        "reverse_total_reading": _number(current_total),
        "reverse_period_readings": current_period.to_dict(),
        "grid_meter_multiplier": _number(multiplier),
        "schema": "time_of_use_image_v1",
    }


def _grid_meter_header_columns(blocks: list[OcrBlock]) -> list[tuple[float, str]]:
    clusters: list[list[OcrBlock]] = []
    for block in sorted(blocks, key=lambda item: (item.bbox[0] + item.bbox[2]) / 2):
        x_center = (block.bbox[0] + block.bbox[2]) / 2
        target = next(
            (
                cluster
                for cluster in clusters
                if abs(
                    x_center
                    - sum((item.bbox[0] + item.bbox[2]) / 2 for item in cluster) / len(cluster)
                ) <= 24
            ),
            None,
        )
        if target is None:
            clusters.append([block])
        else:
            target.append(block)
    return [
        (
            sum((block.bbox[0] + block.bbox[2]) / 2 for block in cluster) / len(cluster),
            _compact("".join(block.text for block in sorted(cluster, key=lambda item: item.bbox[1]))),
        )
        for cluster in clusters
    ]


def _looks_like_grid_meter_number(text: str) -> bool:
    expected = GRID_METER_NUMBER.upper()
    for candidate in re.findall(r"03591SG\d+", text.upper()):
        if _within_one_edit(candidate, expected):
            return True
    return False


def _within_one_edit(candidate: str, expected: str) -> bool:
    if abs(len(candidate) - len(expected)) > 1:
        return False
    if len(candidate) == len(expected):
        return sum(left != right for left, right in zip(candidate, expected, strict=True)) <= 1
    shorter, longer = (candidate, expected) if len(candidate) < len(expected) else (expected, candidate)
    short_index = long_index = differences = 0
    while short_index < len(shorter) and long_index < len(longer):
        if shorter[short_index] == longer[long_index]:
            short_index += 1
            long_index += 1
            continue
        differences += 1
        long_index += 1
        if differences > 1:
            return False
    return True


def write_yangjiang_detail(template_path: Path | str, output_path: Path | str, values: YangjiangValues) -> Path:
    workbook = load_workbook(template_path)
    sheet = workbook.active
    year, month_number = map(int, values.month.split("-"))
    month_label = _chinese_month(month_number)
    sheet["A1"] = f"{PROJECT_TITLE}{month_label}月份电量对账单"
    sheet["E2"] = date(year, month_number, 1)
    sheet["I2"] = date(year, month_number, calendar.monthrange(year, month_number)[1])
    sheet["E2"].number_format = sheet["I2"].number_format = "yyyy/m/d"

    for row, suffix in ((5, "2066"), (6, "2065")):
        prior = values.prior_generation_readings[suffix]
        current = values.current_generation_readings[suffix]
        sheet.cell(row, 1, f"发电量表：表号{METER_NUMBERS[suffix]}")
        _write_reading_row(sheet, row, prior, current, GENERATION_MULTIPLIER)
        sheet.cell(row, 11, f"=(C{row}-B{row})*J{row}+(E{row}-D{row})*J{row}+(G{row}-F{row})*J{row}+(I{row}-H{row})*J{row}")

    sheet["A7"] = f"上网电量表：表号{GRID_METER_NUMBER}"
    _write_reading_row(sheet, 7, values.grid_prior_readings, values.grid_current_readings, values.grid_meter_multiplier)
    sheet["K7"] = "=(C7-B7)*J7+(E7-D7)*J7+(G7-F7)*J7+(I7-H7)*J7"
    sheet["B8"] = "=(C5-B5)*100+(C6-B6)*100-(C7-B7)*J7"
    sheet["D8"] = "=(E5-D5)*100+(E6-D6)*100-(E7-D7)*J7"
    sheet["F8"] = "=(G5-F5)*100+(G6-F6)*100-(G7-F7)*J7"
    sheet["H8"] = "=(I5-H5)*100+(I6-H6)*100-(I7-H7)*J7"
    sheet["K8"] = "=K5+K6-K7"
    sheet["B9"] = "=K5+K6"
    sheet["B10"] = "=B9-K7"

    table_rows = _price_table_rows(values)
    if len(table_rows) == 4:
        middle_style = [_cell_snapshot(sheet.cell(5, column)) for column in range(13, 22)]
        bottom_style = [_cell_snapshot(sheet.cell(6, column)) for column in range(13, 22)]
        for column, snapshot in zip(range(13, 22), middle_style, strict=True):
            _apply_cell_snapshot(sheet.cell(6, column), snapshot)
        for column, snapshot in zip(range(13, 22), bottom_style, strict=True):
            _apply_cell_snapshot(sheet.cell(7, column), snapshot)

    market_basis_periods = PERIOD_KEYS if values.sharp_price_source == "bill_sharp" else PERIOD_KEYS[1:]
    market_formula = "+".join(
        _number(getattr(values.billed_period_energy, period)) for period in market_basis_periods
    )
    for row, period, label in table_rows:
        components = values.price_components[period]
        sheet.cell(row, 13, label)
        for column, key in zip(range(14, 19), ("energy", "transmission", "line_loss", "system", "fund"), strict=True):
            sheet.cell(row, column, float(components[key]))
        sheet.cell(row, 19, f"={_money(values.market_allocation_fee)}/({market_formula})")
        sheet.cell(row, 20, f"=SUM(N{row}:S{row})")
        sheet.cell(row, 21, f"=T{row}*{_number(SETTLEMENT_FACTOR)}")
    if values.sharp_price_source == "bill_sharp":
        settlement_rows = ((12, "B8", "=U4"), (13, "D8", "=U5"), (14, "F8", "=U6"), (15, "H8", "=U7"))
    else:
        settlement_rows = ((12, "B8", "=U6"), (13, "D8", "=U6"), (14, "F8", "=U5"), (15, "H8", "=U4"))
    for row, energy_cell, price_value in settlement_rows:
        sheet.cell(row, 4, f"={energy_cell}")
        sheet.cell(row, 6, price_value)
        sheet.cell(row, 9, f"=D{row}*F{row}")
    sheet["I16"] = "=SUM(I12:I15)"
    sheet["I17"] = "=I16/1.06*1.01"
    _enable_recalculation(workbook)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def write_yangjiang_summary(
    template_path: Path | str,
    output_path: Path | str,
    records: dict[str, YangjiangValues | YangjiangSummaryValues],
    *,
    materialize_calculations: bool = False,
) -> Path:
    workbook = load_workbook(template_path)
    sheet = workbook.active
    style_rows = [[_cell_snapshot(sheet.cell(row, column)) for column in range(1, 15)] for row in range(8, 12)]
    heights = [sheet.row_dimensions[row].height for row in range(8, 12)]
    data_start_row = 4 if materialize_calculations else 12
    if sheet.max_row >= data_start_row:
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row >= data_start_row:
                sheet.unmerge_cells(str(merged))
        sheet.delete_rows(data_start_row, sheet.max_row - data_start_row + 1)
    for index, month in enumerate(sorted(records)):
        row = data_start_row + index * 4
        sheet.insert_rows(row, 4)
        for offset in range(4):
            sheet.row_dimensions[row + offset].height = heights[offset]
            for column in range(1, 15):
                _apply_cell_snapshot(sheet.cell(row + offset, column), style_rows[offset][column - 1])
        for column in (1, 2, 3, 4, 5, 6, 9, 10, 11, 12, 13, 14):
            sheet.merge_cells(start_row=row, start_column=column, end_row=row + 3, end_column=column)
        values = records[month]
        year, month_number = map(int, month.split("-"))
        sheet.cell(row, 1, float(CAPACITY_KWP))
        sheet.cell(row, 2, date(year, month_number, 1))
        sheet.cell(row, 2).number_format = 'yyyy"年"m"月"'
        sheet.cell(row, 3, float(values.summary_total_generation))
        sheet.cell(row, 4, float(values.grid_meter_period_total_energy))
        sheet.cell(row, 5, float(values.grid_settlement_price))
        sheet.cell(row, 6, float(values.self_use_energy))
        for offset, (label, price) in enumerate(
            (
                ("尖", values.period_prices.sharp),
                ("峰", values.period_prices.peak),
                ("平", values.period_prices.flat),
                ("谷", values.period_prices.valley),
            )
        ):
            sheet.cell(row + offset, 7, label)
            sheet.cell(row + offset, 8, float(Decimal(_number(price, places=9))))
            sheet.cell(row + offset, 8).number_format = "0.000000000"
        sheet.cell(row, 9, float(Decimal(_money(values.self_use_fee_before_tax_adjustment))))
        grid_income = values.grid_meter_period_total_energy * values.grid_settlement_price
        total_income = values.self_use_fee_before_tax_adjustment + grid_income
        self_use_ratio = values.self_use_energy / values.summary_total_generation
        sheet.cell(row, 10, float(grid_income) if materialize_calculations else f"=D{row}*E{row}")
        sheet.cell(row, 11, float(total_income) if materialize_calculations else f"=I{row}+J{row}")
        sheet.cell(row, 12, float(self_use_ratio) if materialize_calculations else f"=F{row}/C{row}")
        days = calendar.monthrange(year, month_number)[1]
        effective_hours = values.summary_total_generation / EFFECTIVE_HOUR_CAPACITY_KW / Decimal(days)
        sheet.cell(row, 13, float(effective_hours) if materialize_calculations else f"=C{row}/600/{days}")
        sheet.cell(row, 14, getattr(values, "note", "") or "2026/3/18并网")
    _enable_recalculation(workbook)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def _combine_sources(runtime_root: Path | str, task: dict[str, Any]) -> YangjiangValues:
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    recognized = {
        role: sources.get(role, {}).get("recognized", {}) if isinstance(sources.get(role), dict) else {}
        for role in REQUIRED_ROLES
    }
    if any(not isinstance(recognized[role], dict) or not recognized[role] for role in REQUIRED_ROLES):
        raise YangjiangHongnengError("阳江宏能科技五份材料尚未收齐。")
    bill = recognized["electricity_bill"]
    settlement = recognized["grid_settlement"]
    meter_export = recognized["grid_meter_export"]
    current = {
        suffix: PeriodReadings.from_dict(recognized[f"generation_meter_{suffix}"].get("current_readings"), f"{suffix}电表")
        for suffix in METER_SUFFIXES
    }
    prior = _prior_generation_readings(runtime_root, task["month"])
    grid_prior_readings, grid_prior_total, prior_multiplier = _prior_grid_export(runtime_root, task["month"])
    component_raw = bill.get("price_components")
    if not isinstance(component_raw, dict):
        raise YangjiangHongnengError("电费账单缺少分时电价组成。")
    billed_period_energy = PeriodReadings.from_dict(bill.get("billed_period_energy"), "账单分时电量")
    components: dict[str, dict[str, Decimal]] = {}
    for key in PERIOD_KEYS:
        raw = component_raw.get(key)
        if not isinstance(raw, dict):
            if key == "sharp" and billed_period_energy.sharp <= 0:
                raw = component_raw.get("peak")
            if not isinstance(raw, dict):
                raise YangjiangHongnengError(f"电费账单缺少{key}电价组成。")
        components[key] = {name: _decimal(raw.get(name), f"{key}{name}") for name in ("energy", "transmission", "line_loss", "system", "fund", "market")}
    prices = _period_prices(task["month"], components, billed_period_energy)
    values = YangjiangValues(
        month=task["month"],
        prior_generation_readings=prior,
        current_generation_readings=current,
        grid_prior_readings=grid_prior_readings,
        grid_current_readings=PeriodReadings.from_dict(
            meter_export.get("reverse_period_readings"), "抄表文件反向有功本期示数"
        ),
        grid_prior_total_reading=grid_prior_total,
        grid_current_total_reading=_decimal(meter_export.get("reverse_total_reading"), "抄表文件反向有功总电能"),
        grid_meter_multiplier=_decimal(meter_export.get("grid_meter_multiplier"), "抄表文件综合倍率"),
        period_prices=prices,
        price_components=components,
        market_allocation_fee=_decimal(bill.get("market_allocation_fee"), "市场化分摊总费用"),
        billed_period_energy=billed_period_energy,
        grid_settlement_energy=_decimal(settlement.get("grid_settlement_energy"), "结算单上网电量"),
        grid_settlement_price=_decimal(settlement.get("grid_settlement_price"), "结算单上网电价"),
        grid_settlement_fee=_decimal(settlement.get("grid_settlement_fee"), "结算单上网电费"),
    )
    if values.grid_meter_multiplier <= 0:
        raise YangjiangHongnengError("上网电量表倍率必须大于 0。")
    if values.grid_meter_multiplier != prior_multiplier:
        raise YangjiangHongnengError("上网电量表倍率与上月历史记录不一致，请人工核对。")
    if any(
        current_value < prior_value
        for suffix in METER_SUFFIXES
        for prior_value, current_value in zip(prior[suffix].values(), current[suffix].values(), strict=True)
    ):
        raise YangjiangHongnengError("发电表本期累计示数小于上期示数，请核对月份或截图。")
    if any(current_value < prior_value for prior_value, current_value in zip(values.grid_prior_readings.values(), values.grid_current_readings.values(), strict=True)):
        raise YangjiangHongnengError("电表抄表文件的反向有功本期累计示数小于上期示数。")
    if values.grid_current_total_reading < values.grid_prior_total_reading:
        raise YangjiangHongnengError("电表抄表文件的反向有功总电能小于上期示数。")
    if values.grid_settlement_energy <= 0 or values.grid_settlement_price <= 0:
        raise YangjiangHongnengError("供电所结算上网电量和上网电价必须大于 0。")
    settlement_income = values.grid_settlement_energy * values.grid_settlement_price
    if abs(settlement_income - values.grid_settlement_fee) > Decimal("0.02"):
        raise YangjiangHongnengError("供电所结算单的上网电量、电价与电费不一致。")
    allowed_energy_difference = max(Decimal("50"), values.grid_settlement_energy * Decimal("0.001"))
    if abs(values.grid_meter_total_energy - values.grid_settlement_energy) > allowed_energy_difference:
        raise YangjiangHongnengError("抄表反向总电量与供电所结算电量差异过大，请人工复核。")
    return values


def _extract_generation_meter_readings(raw: OcrRawResult) -> PeriodReadings:
    values = []
    for period, key in zip(PERIODS, PERIOD_KEYS, strict=True):
        value = _row_value_for_labels(raw, (f"功率{PERIOD_KEYS.index(key) + 1}电能", period), prefer_last=True)
        if value is None:
            value = _row_value_for_labels(raw, ("正向有功", period), prefer_last=True)
        if value is None:
            raise YangjiangHongnengError(f"电表截图无法识别{period}时段累计示数。")
        values.append(value)
    return PeriodReadings(*values)


def _extract_electricity_bill(raw: OcrRawResult, month: str) -> dict[str, Any]:
    prior: list[Decimal] = []
    current: list[Decimal] = []
    multiplier: Decimal | None = None
    billed: list[Decimal] = []
    fallback_readings = _structured_grid_reading_rows(raw)
    fallback_prices = _structured_price_rows(raw)
    component_names = {
        "energy": "电能电费",
        "transmission": "输配电费",
        "line_loss": "线损电费",
        "system": "系统运行费用",
    }
    rates: dict[str, dict[str, Decimal]] = {key: {} for key in PERIOD_KEYS}
    for index, (period, key) in enumerate(zip(PERIODS, PERIOD_KEYS, strict=True)):
        reading_numbers = _numbers_after_matching_label(raw, ("正有功", period), minimum=3)
        if len(reading_numbers) < 3:
            reading_numbers = fallback_readings[index] if len(fallback_readings) == 4 else []
        if len(reading_numbers) < 3:
            raise YangjiangHongnengError(f"电费账单无法识别{period}时段上期、本期示数和倍率。")
        prior.append(reading_numbers[0])
        current.append(reading_numbers[1])
        if multiplier is None:
            multiplier = reading_numbers[2]
        bill_numbers = _numbers_after_matching_label(raw, ("电能电费", period), minimum=3)
        if len(bill_numbers) < 3:
            bill_numbers = fallback_prices[0][index] if fallback_prices else []
        if len(bill_numbers) < 3:
            raise YangjiangHongnengError(f"电费账单无法识别{period}时段结算电量。")
        billed.append(bill_numbers[0])
        rates[key]["energy"] = bill_numbers[1]
        for component_index, (component_key, label) in enumerate(component_names.items()):
            if component_key == "energy":
                continue
            numbers = _numbers_after_matching_label(raw, (label, period), minimum=2)
            if len(numbers) < 2:
                numbers = fallback_prices[component_index][index] if fallback_prices else []
            if len(numbers) < 2:
                raise YangjiangHongnengError(f"电费账单无法识别{period}时段{label}单价。")
            rates[key][component_key] = numbers[1]
    fund, fund_fee = _find_fund_values(raw)
    market_fee = _find_market_allocation_fee(raw)
    billed_total = sum(billed, Decimal("0"))
    if billed_total <= 0:
        raise YangjiangHongnengError("电费账单分时电量合计必须大于 0。")
    if abs(fund * billed_total - fund_fee) > Decimal("0.05"):
        raise YangjiangHongnengError("电费账单基金及附加费单价与总金额不一致。")
    billed_readings = PeriodReadings(*billed)
    sharp_active = _has_effective_sharp_price(month, billed_readings, rates)
    market_basis = sum(billed if sharp_active else billed[1:], Decimal("0"))
    if market_basis <= 0:
        raise YangjiangHongnengError("电费账单峰、平、谷结算电量合计必须大于 0。")
    market_rate = market_fee / market_basis
    for key in PERIOD_KEYS:
        rates[key]["fund"] = fund
        rates[key]["market"] = market_rate
    # The workbook has controlled price rows for valley, flat and peak. Sharp falls
    # back to peak whenever the bill has no effective sharp energy/energy price.
    return {
        "month": month,
        "utility_meter_number": GRID_METER_NUMBER,
        "purchased_prior_readings": PeriodReadings(*prior).to_dict(),
        "purchased_current_readings": PeriodReadings(*current).to_dict(),
        "utility_meter_multiplier": _number(multiplier or GRID_METER_MULTIPLIER),
        "billed_period_energy": PeriodReadings(*billed).to_dict(),
        "sharp_tariff_active": sharp_active,
        "fund_allocation_fee": _money(fund_fee),
        "market_allocation_fee": _money(market_fee),
        "market_allocation_rate": _number(market_rate),
        "price_components": {
            key: {name: _number(value) for name, value in components.items()}
            for key, components in rates.items()
        },
    }


def _extract_grid_settlement(raw: OcrRawResult, month: str) -> dict[str, Any]:
    numbers = _numbers_after_matching_label(raw, ("市场化电费",), minimum=3)
    if len(numbers) < 3:
        numbers = _numbers_after_matching_label(raw, ("抄见电量",), minimum=3)
    if len(numbers) < 3:
        raise YangjiangHongnengError("阳江宏能科技结算单无法识别上网电量、电价和电费。")
    energy, price, fee = numbers[:3]
    return {
        "month": month,
        "grid_settlement_energy": _number(energy),
        "grid_settlement_price": _number(price, places=6),
        "grid_settlement_fee": _money(fee),
    }


def _find_fund_values(raw: OcrRawResult) -> tuple[Decimal, Decimal]:
    rows = _ocr_rows(raw)
    for row in rows:
        if "基金及附加费" not in _compact("".join(block.text for block in row)):
            continue
        values = _fund_values_from_row(row)
        if values is not None:
            return values

    # The grid PDF may render Chinese labels as missing glyphs. Its fund row is
    # still structurally located after the 16 time-of-use component rows.
    structured = _structured_price_row_entries(raw)
    start_index = structured[15][0] if len(structured) >= 16 else -1
    for row_index, row in enumerate(rows):
        if row_index <= start_index:
            continue
        values = _fund_values_from_row(row)
        if values is not None:
            return values
    raise YangjiangHongnengError("电费账单无法识别基金及附加费单价和总金额。")


def _fund_values_from_row(row: list[OcrBlock]) -> tuple[Decimal, Decimal] | None:
    numeric = [
        (index, value)
        for index, block in enumerate(sorted(row, key=lambda item: item.bbox[0]))
        for value in [_numeric_block(block.text)]
        if value is not None
    ]
    rates = [
        (index, value)
        for index, value in numeric
        if Decimal("0") < value < Decimal("0.1")
        and "." in _number(value)
        and len(_number(value).split(".")[-1]) >= 6
    ]
    for rate_index, rate in reversed(rates):
        fees = [value for index, value in numeric if index > rate_index and Decimal("1") < value < Decimal("100000")]
        if fees:
            return rate, fees[0]
    return None


def _find_market_allocation_fee(raw: OcrRawResult) -> Decimal:
    numbers = _numbers_after_matching_label(raw, ("市场化分摊总费用",), minimum=1)
    positive = [value for value in numbers if value > 0]
    if positive:
        return positive[-1]
    structured = _structured_price_row_entries(raw)
    if len(structured) >= 16:
        last_component_index = structured[15][0]
        numeric_rows = _numeric_rows(raw)
        for row_index, values in numeric_rows:
            if row_index <= last_component_index:
                continue
            # Item (17) contains its index, two zero columns and the amount.
            # Taking the rightmost amount avoids drifting to later rows such as
            # the 8% discount field when Chinese labels are missing from OCR.
            if len(values) >= 3 and Decimal("1") < values[-1] < Decimal("100000"):
                return values[-1]
    raise YangjiangHongnengError("电费账单无法识别市场化分摊总费用。")


def _structured_grid_reading_rows(raw: OcrRawResult) -> list[list[Decimal]]:
    candidates: list[list[Decimal]] = []
    for _, numbers in _numeric_rows(raw):
        if len(numbers) < 4:
            continue
        prior, current, multiplier, energy = numbers[:4]
        if multiplier < Decimal("10") or current < prior:
            continue
        expected = (current - prior) * multiplier
        if abs(expected - energy) <= Decimal("1"):
            candidates.append(numbers)
    return candidates[:4] if len(candidates) >= 4 else []


def _structured_price_rows(raw: OcrRawResult) -> list[list[list[Decimal]]]:
    entries = _structured_price_row_entries(raw)
    if len(entries) != 16:
        return []
    values = [numbers for _, numbers in entries]
    return [values[offset : offset + 4] for offset in range(0, 16, 4)]


def _structured_price_row_entries(raw: OcrRawResult) -> list[tuple[int, list[Decimal]]]:
    start_index = _price_table_start_index(raw)
    numeric_rows = [(index, numbers) for index, numbers in _numeric_rows(raw) if index >= start_index]
    entries: list[tuple[int, list[Decimal]]] = []
    cursor = 0
    while cursor < len(numeric_rows) and len(entries) < 4:
        row_index, numbers = numeric_rows[cursor]
        cursor += 1
        triplet = _price_triplet(numbers)
        if triplet is not None:
            entries.append((row_index, triplet))
    if len(entries) != 4:
        return []
    billed_energy = [entry[1][0] for entry in entries]
    for _ in range(3):
        for expected_energy in billed_energy:
            matched: tuple[int, list[Decimal]] | None = None
            while cursor < len(numeric_rows):
                row_index, numbers = numeric_rows[cursor]
                cursor += 1
                triplet = _price_triplet(numbers, expected_energy=expected_energy)
                if triplet is not None:
                    matched = (row_index, triplet)
                    break
            if matched is None:
                return []
            entries.append(matched)
    return entries


def _price_table_start_index(raw: OcrRawResult) -> int:
    for row_index, row in enumerate(_ocr_rows(raw)):
        text = _compact("".join(block.text for block in row)).lower()
        if "electricitybillinformation" in text or "电费账单信息" in text:
            return row_index + 1
    return 0


def _price_triplet(
    numbers: list[Decimal],
    *,
    expected_energy: Decimal | None = None,
) -> list[Decimal] | None:
    matches: list[tuple[Decimal, int, list[Decimal]]] = []
    for start in range(max(0, len(numbers) - 2)):
        energy, rate, fee = numbers[start : start + 3]
        if energy < 0 or not Decimal("0") <= rate < Decimal("2") or fee < 0:
            continue
        if expected_energy is not None and abs(energy - expected_energy) > Decimal("0.05"):
            continue
        difference = abs(energy * rate - fee)
        if difference <= Decimal("0.05"):
            matches.append((difference, -start, [energy, rate, fee]))
    if not matches:
        return None
    return min(matches, key=lambda item: (item[0], item[1]))[2]


def _numeric_rows(raw: OcrRawResult) -> list[tuple[int, list[Decimal]]]:
    result: list[tuple[int, list[Decimal]]] = []
    for row_index, row in enumerate(_ocr_rows(raw)):
        numbers = [
            value
            for block in sorted(row, key=lambda item: item.bbox[0])
            for value in [_numeric_block(block.text)]
            if value is not None
        ]
        if numbers:
            result.append((row_index, numbers))
    return result


def _numbers_after_matching_label(raw: OcrRawResult, labels: tuple[str, ...], *, minimum: int) -> list[Decimal]:
    for row in _ocr_rows(raw):
        label_blocks = [block for block in row if all(label in _compact(block.text) for label in labels)]
        if not label_blocks:
            combined = _compact("".join(block.text for block in row))
            if not all(label in combined for label in labels):
                continue
            label_x = min(block.bbox[0] for block in row)
        else:
            label_x = max(block.bbox[2] for block in label_blocks)
        numbers = [
            value
            for block in sorted(row, key=lambda item: item.bbox[0])
            if block.bbox[0] >= label_x - 2
            for value in [_numeric_block(block.text)]
            if value is not None
        ]
        if len(numbers) >= minimum:
            return numbers
    return []


def _row_value_for_labels(raw: OcrRawResult, labels: tuple[str, ...], *, prefer_last: bool) -> Decimal | None:
    for row in _ocr_rows(raw, tolerance=14):
        combined = _compact("".join(block.text for block in row))
        if not all(label in combined for label in labels):
            continue
        numbers = [_numeric_block(block.text) for block in sorted(row, key=lambda item: item.bbox[0])]
        values = [value for value in numbers if value is not None]
        if values:
            return values[-1] if prefer_last else values[0]
    return None


def _numbers_near_label(raw: OcrRawResult, label: str, *, following_rows: int) -> list[Decimal]:
    rows = _ocr_rows(raw)
    for index, row in enumerate(rows):
        if label not in _compact("".join(block.text for block in row)):
            continue
        result = []
        for nearby in rows[index : index + following_rows + 1]:
            result.extend(value for value in (_numeric_block(block.text) for block in nearby) if value is not None)
        return result
    return []


def _ocr_rows(raw: OcrRawResult, *, tolerance: float = 10) -> list[list[OcrBlock]]:
    all_rows: list[list[OcrBlock]] = []
    for page in raw.pages:
        rows: list[list[OcrBlock]] = []
        for block in sorted(page.blocks, key=lambda item: ((item.bbox[1] + item.bbox[3]) / 2, item.bbox[0])):
            center = (block.bbox[1] + block.bbox[3]) / 2
            target = next(
                (
                    row
                    for row in rows
                    if abs(center - sum((item.bbox[1] + item.bbox[3]) / 2 for item in row) / len(row)) <= tolerance
                ),
                None,
            )
            if target is None:
                rows.append([block])
            else:
                target.append(block)
        all_rows.extend(sorted((sorted(row, key=lambda item: item.bbox[0]) for row in rows), key=lambda row: row[0].bbox[1]))
    return all_rows


def _values_from_dict(raw: Any, *, expected_month: str) -> YangjiangValues:
    if not isinstance(raw, dict):
        raise YangjiangHongnengError("阳江宏能科技任务缺少已确认数据。")
    month = _validate_month(str(raw.get("month") or expected_month))
    if month != expected_month:
        raise YangjiangHongnengError("阳江宏能科技任务月份不一致。")
    prior_raw = raw.get("prior_generation_readings")
    current_raw = raw.get("current_generation_readings")
    if not isinstance(prior_raw, dict) or not isinstance(current_raw, dict):
        raise YangjiangHongnengError("任务缺少发电表累计示数。")
    component_raw = raw.get("price_components")
    if not isinstance(component_raw, dict):
        raise YangjiangHongnengError("任务缺少电价组成。")
    components = {
        period: {key: _decimal(values.get(key), f"{period}{key}") for key in ("energy", "transmission", "line_loss", "system", "fund", "market")}
        for period, values in component_raw.items()
        if isinstance(values, dict)
    }
    grid_multiplier = _decimal(raw.get("grid_meter_multiplier"), "上网表倍率")
    settlement_energy = _decimal(raw.get("grid_settlement_energy"), "上网结算电量")
    grid_prior_total_raw = raw.get("grid_prior_total_reading")
    grid_current_total_raw = raw.get("grid_current_total_reading")
    if grid_prior_total_raw is None or grid_current_total_raw is None:
        # Read compatibility for pre-v2 effective records. Such a record may be
        # rendered in history, but it cannot seed the next month without a real
        # meter-export total reading.
        grid_prior_total = Decimal("0")
        grid_current_total = settlement_energy / grid_multiplier
    else:
        grid_prior_total = _decimal(grid_prior_total_raw, "上网表上期总电能")
        grid_current_total = _decimal(grid_current_total_raw, "上网表本期总电能")
    billed_period_energy = PeriodReadings.from_dict(raw.get("billed_period_energy"), "账单分时电量")
    stored_period_prices = PeriodReadings.from_dict(raw.get("period_prices"), "分时结算单价")
    if not _has_effective_sharp_price(month, billed_period_energy, components):
        period_prices = PeriodReadings(
            stored_period_prices.peak,
            stored_period_prices.peak,
            stored_period_prices.flat,
            stored_period_prices.valley,
        )
    elif "sharp" in components:
        period_prices = _period_prices(month, components, billed_period_energy)
    else:
        period_prices = stored_period_prices
    return YangjiangValues(
        month=month,
        prior_generation_readings={suffix: PeriodReadings.from_dict(prior_raw.get(suffix), f"{suffix}上期") for suffix in METER_SUFFIXES},
        current_generation_readings={suffix: PeriodReadings.from_dict(current_raw.get(suffix), f"{suffix}本期") for suffix in METER_SUFFIXES},
        grid_prior_readings=PeriodReadings.from_dict(raw.get("grid_prior_readings"), "上网表上期"),
        grid_current_readings=PeriodReadings.from_dict(raw.get("grid_current_readings"), "上网表本期"),
        grid_prior_total_reading=grid_prior_total,
        grid_current_total_reading=grid_current_total,
        grid_meter_multiplier=grid_multiplier,
        period_prices=period_prices,
        price_components=components,
        market_allocation_fee=_decimal(raw.get("market_allocation_fee"), "市场化分摊总费用"),
        billed_period_energy=billed_period_energy,
        grid_settlement_energy=settlement_energy,
        grid_settlement_price=_decimal(raw.get("grid_settlement_price"), "上网结算电价"),
        grid_settlement_fee=_decimal(raw.get("grid_settlement_fee"), "上网结算电费"),
    )


def _previous_month(month: str) -> str:
    year, month_number = map(int, month.split("-"))
    if month_number == 1:
        return f"{year - 1}-12"
    return f"{year}-{month_number - 1:02d}"


def _prior_generation_readings(runtime_root: Path | str, month: str) -> dict[str, PeriodReadings]:
    prior_month = _previous_month(month)
    if prior_month == SEED_MONTH:
        return {suffix: PeriodReadings.from_dict(readings, f"{suffix}四月基准") for suffix, readings in SEED_GENERATION_READINGS.items()}
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    record = (series.get("months") or {}).get(prior_month)
    values = record.get("values") if isinstance(record, dict) else None
    current = values.get("current_generation_readings") if isinstance(values, dict) else None
    if not isinstance(current, dict):
        raise YangjiangHongnengError(f"缺少 {prior_month} 已完成数据，不能跨月计算累计电表示数。")
    return {suffix: PeriodReadings.from_dict(current.get(suffix), f"{suffix}{prior_month}本期") for suffix in METER_SUFFIXES}


def _prior_grid_export(runtime_root: Path | str, month: str) -> tuple[PeriodReadings, Decimal, Decimal]:
    prior_month = _previous_month(month)
    if prior_month == SEED_MONTH:
        return (
            PeriodReadings.from_dict(SEED_GRID_PERIOD_READINGS, "四月反向有功历史基准"),
            SEED_GRID_TOTAL_READING,
            GRID_METER_MULTIPLIER,
        )
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    record = (series.get("months") or {}).get(prior_month)
    values = record.get("values") if isinstance(record, dict) else None
    if not isinstance(values, dict):
        raise YangjiangHongnengError(f"缺少 {prior_month} 已完成数据，不能跨月计算上网电表示数。")
    current = values.get("grid_current_readings")
    current_total = values.get("grid_current_total_reading")
    multiplier = values.get("grid_meter_multiplier")
    if not isinstance(current, dict) or current_total is None or multiplier is None:
        raise YangjiangHongnengError(
            f"{prior_month} 历史记录缺少电表抄表 XLSX 数据，请先按五份材料重新完成该月。"
        )
    return (
        PeriodReadings.from_dict(current, f"{prior_month}反向有功本期"),
        _decimal(current_total, f"{prior_month}反向有功总电能"),
        _decimal(multiplier, f"{prior_month}上网表倍率"),
    )


def _effective_records(runtime_root: Path | str) -> dict[str, YangjiangValues]:
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    result = {}
    for month, record in (series.get("months") or {}).items():
        values = record.get("values") if isinstance(record, dict) else None
        if isinstance(values, dict):
            result[month] = _values_from_dict(values, expected_month=month)
    return result


def effective_yangjiang_summary_records(
    runtime_root: Path | str,
) -> dict[str, YangjiangValues | YangjiangSummaryValues]:
    """Return summary-only history overlaid by verified operational records."""

    series = _read_json(_series_path(runtime_root), default={"months": {}, "historical_months": {}})
    result: dict[str, YangjiangValues | YangjiangSummaryValues] = {}
    for month, record in (series.get("historical_months") or {}).items():
        values = record.get("values") if isinstance(record, dict) else None
        if isinstance(values, dict):
            result[month] = YangjiangSummaryValues.from_dict(values, expected_month=month)
    result.update(_effective_records(runtime_root))
    return result


def _summary_materialization_enabled(runtime_root: Path | str) -> bool:
    series = _read_json(_series_path(runtime_root), default={})
    return series.get("summary_materialize_calculations") is True


def _promote_effective(runtime_root: Path | str, task: dict[str, Any], values: YangjiangValues, outputs: list[Path]) -> None:
    target = _series_path(runtime_root)
    series = _read_json(target, default={"site_key": SITE_KEY, "months": {}})
    months = series.setdefault("months", {})
    current = months.get(values.month) if isinstance(months.get(values.month), dict) else {}
    history = list(current.get("history") or [])
    if current.get("task_id"):
        history.append({key: current.get(key) for key in ("task_id", "revision", "values", "outputs")})
    months[values.month] = {
        "task_id": task["task_id"],
        "revision": task["revision"],
        "values": values.to_dict(),
        "outputs": [str(output) for output in outputs],
        "history": history,
        "updated_at": _utc_now(),
    }
    _write_json(target, series)


def _source_pages(source: Path, output_dir: Path) -> list[RenderedPage]:
    if source.suffix.lower() == ".pdf":
        return render_pdf_to_images(source, output_dir)
    return [RenderedPage(source_file=source, page=1, image_path=source)]


def _validate_source(path: Path) -> None:
    if not path.is_file() or path.is_symlink() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise YangjiangHongnengError("阳江宏能科技仅接收 JPG、PNG、PDF 和 XLSX 原始材料。")
    if path.stat().st_size <= 0 or path.stat().st_size > MAX_SOURCE_BYTES:
        raise YangjiangHongnengError("文件为空或超过 30 MB，本次文件未登记。")
    header = path.read_bytes()[:12]
    suffix = path.suffix.lower()
    valid = (suffix == ".pdf" and header.startswith(b"%PDF-")) or (
        suffix in {".jpg", ".jpeg"} and header.startswith(b"\xff\xd8\xff")
    ) or (suffix == ".png" and header.startswith(b"\x89PNG\r\n\x1a\n")) or (
        suffix == ".xlsx" and header.startswith(b"PK\x03\x04")
    )
    if not valid:
        raise YangjiangHongnengError("文件扩展名与实际内容不一致，本次文件未登记。")


def _month_from_text(text: str, expected_month: str) -> str:
    matches = {
        f"{year}-{int(month):02d}"
        for year, month in re.findall(r"(20\d{2})\s*(?:[-./年])\s*(0?[1-9]|1[0-2])", text)
    }
    compact_expected = expected_month.replace("-", "")
    if compact_expected in re.sub(r"\D", "", text):
        matches.add(expected_month)
    if expected_month in matches:
        return expected_month
    if len(matches) == 1:
        return next(iter(matches))
    raise YangjiangHongnengError("无法唯一识别文件账期。")


def _excel_reading_date(value: Any, *, source_label: str = "电表抄表 XLSX") -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise YangjiangHongnengError(f"{source_label}无法识别抄表时间。")


def _numeric_block(text: str) -> Decimal | None:
    compact = text.replace(",", "").strip()
    if not re.fullmatch(r"-?\d+(?:\.\d+)?(?:\s*(?:kwh|Kwh|KWh|元))?", compact):
        return None
    compact = re.sub(r"\s*(?:kwh|Kwh|KWh|元)$", "", compact)
    try:
        return Decimal(compact)
    except InvalidOperation:
        return None


def _settlement_price(components: dict[str, Decimal]) -> Decimal:
    return sum(components.values(), Decimal("0")) * SETTLEMENT_FACTOR


def _has_effective_sharp_price(
    month: str,
    billed_period_energy: PeriodReadings,
    components: dict[str, dict[str, Decimal]],
) -> bool:
    sharp_components = components.get("sharp")
    return (
        month not in SHARP_TARIFF_INACTIVE_MONTHS
        and billed_period_energy.sharp > 0
        and isinstance(sharp_components, dict)
        and sharp_components.get("energy", Decimal("0")) > 0
    )


def _period_prices(
    month: str,
    components: dict[str, dict[str, Decimal]],
    billed_period_energy: PeriodReadings,
) -> PeriodReadings:
    sharp_components = (
        components["sharp"]
        if _has_effective_sharp_price(month, billed_period_energy, components)
        else components["peak"]
    )
    return PeriodReadings(
        _settlement_price(sharp_components),
        _settlement_price(components["peak"]),
        _settlement_price(components["flat"]),
        _settlement_price(components["valley"]),
    )


def _price_table_rows(values: YangjiangValues) -> tuple[tuple[int, str, str], ...]:
    if values.sharp_price_source == "bill_sharp":
        return ((4, "sharp", "尖"), (5, "peak", "峰"), (6, "flat", "平"), (7, "valley", "谷"))
    return ((4, "valley", "谷"), (5, "flat", "平"), (6, "peak", "峰"))


def _write_reading_row(sheet: Any, row: int, prior: PeriodReadings, current: PeriodReadings, multiplier: Decimal) -> None:
    for column, value in zip((2, 4, 6, 8), prior.values(), strict=True):
        sheet.cell(row, column, float(value))
    for column, value in zip((3, 5, 7, 9), current.values(), strict=True):
        sheet.cell(row, column, float(value))
    sheet.cell(row, 10, float(multiplier))


def _cell_snapshot(cell: Any) -> dict[str, Any]:
    return {
        "style": copy(cell._style),
        "number_format": cell.number_format,
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "protection": copy(cell.protection),
    }


def _apply_cell_snapshot(cell: Any, snapshot: dict[str, Any]) -> None:
    cell._style = copy(snapshot["style"])
    cell.number_format = snapshot["number_format"]
    cell.font = copy(snapshot["font"])
    cell.fill = copy(snapshot["fill"])
    cell.border = copy(snapshot["border"])
    cell.alignment = copy(snapshot["alignment"])
    cell.protection = copy(snapshot["protection"])


def _enable_recalculation(workbook: Any) -> None:
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _chinese_month(month: int) -> str:
    names = "一二三四五六七八九十"
    if month <= 10:
        return names[month - 1]
    return "十一" if month == 11 else "十二"


def _clean_energy(value: Decimal) -> Decimal:
    integral = value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return integral if abs(value - integral) <= Decimal("0.01") else value


def _compact(value: str) -> str:
    return re.sub(r"[\s\u3000()（）]", "", unicodedata.normalize("NFKC", value))


def _role_label(role: str) -> str:
    return {
        "generation_meter_2065": "2065电表截图",
        "generation_meter_2066": "2066电表截图",
        "electricity_bill": "电费账单",
        "grid_settlement": "阳江宏能科技结算单",
        "grid_meter_export": "电表抄表 XLSX/图片/PDF",
    }.get(role, role)


def _task_root(runtime_root: Path | str, task_id: str) -> Path:
    return Path(runtime_root) / "yangjiang_hongneng_tasks" / task_id


def _series_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "yangjiang_hongneng" / "series.json"


def _load_task(runtime_root: Path | str, task_id: str) -> tuple[Path, dict[str, Any]]:
    clean = _safe_task_id(task_id)
    root = _task_root(runtime_root, clean)
    task = _read_json(root / "task.json")
    _assert_task_identity(task, clean, str(task.get("month") or ""))
    return root, task


def _assert_task_identity(task: dict[str, Any], task_id: str, month: str) -> None:
    if task.get("task_id") != task_id or task.get("business_profile") != BUSINESS_PROFILE:
        raise YangjiangHongnengError("阳江宏能科技任务身份不匹配。")
    if task.get("site_key") != SITE_KEY or task.get("month") != month:
        raise YangjiangHongnengError("阳江宏能科技任务跨站区或月份。")


def _next_revision(runtime_root: Path | str, month: str) -> int:
    current = (_read_json(_series_path(runtime_root), default={"months": {}}).get("months") or {}).get(month)
    return int(current.get("revision") or 0) + 1 if isinstance(current, dict) else 1


def _safe_task_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", value):
        raise YangjiangHongnengError("阳江宏能科技任务编号无效。")
    return value


def _validate_month(value: str) -> str:
    if not re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", value):
        raise YangjiangHongnengError("月份必须使用 YYYY-MM 格式。")
    return value


def _decimal(value: Any, field: str) -> Decimal:
    try:
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, AttributeError, TypeError, ValueError) as exc:
        raise YangjiangHongnengError(f"{field}不是有效数字。") from exc


def _optional_decimal(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _number(value: Decimal, *, places: int | None = None) -> str:
    if places is not None:
        value = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _money(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), ".2f")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_json(path: Path, *, default: Any = None) -> Any:
    if not path.is_file():
        if default is not None:
            return default
        raise YangjiangHongnengError(f"阳江宏能科技任务状态不存在：{path}")
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise YangjiangHongnengError(f"阳江宏能科技任务状态损坏：{path}") from exc


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
