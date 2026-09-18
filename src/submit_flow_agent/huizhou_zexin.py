"""Independent monthly settlement flow for Huizhou Zexin."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.properties import CalcProperties

from submit_flow_agent.ocr.base import OcrAdapter, OcrRawResult, write_ocr_raw_result
from submit_flow_agent.ocr.factory import create_ocr_adapter
from submit_flow_agent.pdf_renderer import render_pdf_to_images


BUSINESS_PROFILE = "huizhou_zexin_two_pdf_settlement"
SITE_KEY = "huizhou_zexin"
SITE_NAME = "惠州泽鑫"
PROJECT_NAME = "惠州泽鑫新能源科技有限公司"
SUPPLY_UNIT = "秋长供电所"
GRID_CUSTOMER_NO = "0313060362996094"
SELF_USE_UNIT_PRICE = Decimal("0.6")
OUTPUT_DOCX = "惠州泽鑫电费结算单.docx"
OUTPUT_SUMMARY = "惠州泽鑫汇总表.xlsx"
REQUIRED_ROLES = ("grid_bill", "manual_generation")
EDITABLE_FIELDS = {
    "total_generation",
    "grid_energy",
    "grid_tax_exclusive_fee",
    "grid_tax_inclusive_fee",
}


class HuizhouZexinError(RuntimeError):
    """Raised when the fixed Huizhou Zexin contract cannot advance."""


@dataclass(frozen=True)
class ZexinValues:
    month: str
    total_generation: Decimal
    grid_energy: Decimal
    grid_tax_exclusive_fee: Decimal
    grid_tax_inclusive_fee: Decimal

    @property
    def grid_tax(self) -> Decimal:
        return self.grid_tax_inclusive_fee - self.grid_tax_exclusive_fee

    @property
    def grid_unit_price(self) -> Decimal:
        return self.grid_tax_inclusive_fee / self.grid_energy

    @property
    def self_use_energy(self) -> Decimal:
        return self.total_generation - self.grid_energy

    @property
    def self_use_ratio(self) -> Decimal:
        return self.self_use_energy / self.total_generation

    @property
    def self_use_fee(self) -> Decimal:
        return self.self_use_energy * SELF_USE_UNIT_PRICE

    @property
    def total_income(self) -> Decimal:
        return self.grid_tax_inclusive_fee + self.self_use_fee

    def to_dict(self) -> dict[str, str]:
        return {
            "month": self.month,
            "project_name": PROJECT_NAME,
            "supply_unit": SUPPLY_UNIT,
            "grid_customer_no": GRID_CUSTOMER_NO,
            "total_generation": _number(self.total_generation),
            "grid_energy": _number(self.grid_energy),
            "grid_tax_exclusive_fee": _money(self.grid_tax_exclusive_fee),
            "grid_tax_inclusive_fee": _money(self.grid_tax_inclusive_fee),
            "grid_tax": _money(self.grid_tax),
            "grid_unit_price": _number(self.grid_unit_price, places=6),
            "self_use_energy": _number(self.self_use_energy),
            "self_use_ratio": _number(self.self_use_ratio, places=6),
            "self_use_unit_price": _number(SELF_USE_UNIT_PRICE),
            "self_use_fee": _money(self.self_use_fee),
            "total_income": _money(self.total_income),
        }


def prepare_huizhou_zexin(runtime_root: Path | str, *, task_id: str, month: str) -> dict[str, Any]:
    clean_task_id = _safe_task_id(task_id)
    clean_month = _validate_month(month)
    root = _task_root(runtime_root, clean_task_id)
    if root.exists():
        task = _read_json(root / "task.json")
        _assert_task_identity(task, clean_task_id, clean_month)
        return task
    root.mkdir(parents=True, exist_ok=False)
    task = {
        "task_id": clean_task_id,
        "business_profile": BUSINESS_PROFILE,
        "site_key": SITE_KEY,
        "site_name": SITE_NAME,
        "month": clean_month,
        "revision": _next_revision(runtime_root, clean_month),
        "status": "collecting_files",
        "stage": "collecting_files",
        "sources": {},
        "source_count": 0,
        "recognized_data": {},
        "confirmed_data": {},
        "missing_fields": [],
        "outputs": [],
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
    }
    _write_json(root / "task.json", task)
    return task


def collect_huizhou_zexin_source(
    runtime_root: Path | str,
    task_id: str,
    source: Path | str,
    *,
    ocr_adapter: OcrAdapter | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_files", "awaiting_confirmation", "need_review"}:
        raise HuizhouZexinError("当前惠州泽鑫任务不接收上网电费单 PDF。")
    source_path = Path(source)
    _validate_pdf(source_path)
    digest = _sha256_file(source_path)
    existing_sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    if any(isinstance(item, dict) and item.get("sha256") == digest for item in existing_sources.values()):
        return {**task, "duplicate": True}

    staged = root / "inputs" / f"source-{digest[:16]}.pdf"
    staged.parent.mkdir(parents=True, exist_ok=True)
    if not staged.exists():
        shutil.copy2(source_path, staged)
    pages = render_pdf_to_images(staged, root / "rendered" / digest[:16])
    raw = (ocr_adapter or create_ocr_adapter()).recognize(staged, pages)
    role, extracted = recognize_huizhou_zexin_source(raw, expected_month=task["month"])
    if role != "grid_bill":
        staged.unlink(missing_ok=True)
        raise HuizhouZexinError(
            "惠州泽鑫现在只需上传上网电费单 PDF；月总发电量请发送文字，例如：6月份总发电量168160千瓦时。"
        )
    source_month = extracted.get("month")
    if source_month != task["month"]:
        raise HuizhouZexinError(
            f"PDF 账期为 {source_month or '无法识别'}，与当前填报月份 {task['month']} 不一致；本次文件未登记。"
        )
    if role in existing_sources:
        raise HuizhouZexinError("上网电费单已登记，请勿重复上传另一份同类文件。")

    write_ocr_raw_result(raw, root / "ocr" / f"{role}.json", relative_to=root)
    existing_sources[role] = {
        "role": role,
        "name": source_path.name,
        "stored_path": str(staged.relative_to(root)).replace("\\", "/"),
        "sha256": digest,
        "recognized": extracted,
    }
    task["sources"] = existing_sources
    return _refresh_task_inputs(root, task)


def record_huizhou_zexin_generation(
    runtime_root: Path | str,
    task_id: str,
    *,
    month: str,
    total_generation: str | int | float | Decimal,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_files", "awaiting_confirmation", "need_review"}:
        raise HuizhouZexinError("当前惠州泽鑫任务不接收月总发电量。")
    clean_month = _validate_month(month)
    if clean_month != task["month"]:
        raise HuizhouZexinError(
            f"填写月份 {clean_month} 与当前填报月份 {task['month']} 不一致，本次数据未登记。"
        )
    generation = _decimal(total_generation, "月总发电量")
    if generation <= 0:
        raise HuizhouZexinError("月总发电量必须大于 0，本次数据未登记。")

    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    existing = sources.get("manual_generation")
    normalized_value = _number(generation)
    if (
        isinstance(existing, dict)
        and isinstance(existing.get("recognized"), dict)
        and existing["recognized"].get("total_generation") == normalized_value
    ):
        return {**task, "duplicate": True}
    sources["manual_generation"] = {
        "role": "manual_generation",
        "source_type": "operator_text",
        "recorded_at": _utc_now(),
        "recognized": {
            "month": clean_month,
            "total_generation": normalized_value,
        },
    }
    task["sources"] = sources
    return _refresh_task_inputs(root, task)


def _refresh_task_inputs(root: Path, task: dict[str, Any]) -> dict[str, Any]:
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    task["source_count"] = sum(role in sources for role in REQUIRED_ROLES)
    task["updated_at"] = _utc_now()
    if not all(role in sources for role in REQUIRED_ROLES):
        task["status"] = task["stage"] = "collecting_files"
        _write_json(root / "task.json", task)
        return task

    combined = _combine_sources(sources, expected_month=task["month"])
    missing = [field for field in EDITABLE_FIELDS if not combined.get(field)]
    task.update(
        {
            "recognized_data": combined,
            "missing_fields": sorted(missing),
            "status": "need_review" if missing else "awaiting_confirmation",
            "stage": "need_review" if missing else "awaiting_confirmation",
        }
    )
    _write_json(root / "task.json", task)
    return task


def update_huizhou_zexin(runtime_root: Path | str, task_id: str, values: dict[str, Any]) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"awaiting_confirmation", "need_review"}:
        raise HuizhouZexinError("当前惠州泽鑫任务不在数据核对阶段。")
    unknown = sorted(set(values) - EDITABLE_FIELDS)
    if unknown:
        raise HuizhouZexinError(f"不支持修改这些字段：{', '.join(unknown)}")
    normalized = _normalize_values(
        {**task.get("recognized_data", {}), **task.get("confirmed_data", {}), **values},
        expected_month=task["month"],
    )
    task.update(
        {
            "confirmed_data": normalized.to_dict(),
            "missing_fields": [],
            "status": "awaiting_confirmation",
            "stage": "awaiting_confirmation",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def confirm_huizhou_zexin(
    runtime_root: Path | str,
    task_id: str,
    values: dict[str, Any] | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"awaiting_confirmation", "need_review"}:
        raise HuizhouZexinError("当前惠州泽鑫任务不在数据确认阶段。")
    normalized = _normalize_values(
        {**task.get("recognized_data", {}), **task.get("confirmed_data", {}), **(values or {})},
        expected_month=task["month"],
    )
    task.update(
        {
            "confirmed_data": normalized.to_dict(),
            "missing_fields": [],
            "status": "ready_to_run",
            "stage": "ready_to_run",
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    return task


def run_huizhou_zexin(
    runtime_root: Path | str,
    task_id: str,
    *,
    template_dir: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] == "completed":
        return task
    if task["status"] != "ready_to_run":
        raise HuizhouZexinError("请先核对并确认识别数据，再发送“开始运行”。")
    values = _normalize_values(task["confirmed_data"], expected_month=task["month"])
    templates = Path(template_dir)
    docx_template = templates / "电费结算单.docx"
    summary_template = templates / OUTPUT_SUMMARY
    for template in (docx_template, summary_template):
        if not template.is_file():
            raise HuizhouZexinError(f"受控模板不存在：{template}")

    records = _effective_records(runtime_root)
    records[values.month] = values
    output_dir = root / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_output = output_dir / OUTPUT_DOCX
    summary_output = output_dir / OUTPUT_SUMMARY
    write_settlement_docx(docx_template, docx_output, values)
    write_summary_workbook(
        summary_template,
        summary_output,
        records,
        materialize_calculations=_summary_materialization_enabled(runtime_root),
    )
    outputs = [docx_output, summary_output]
    if any(not output.is_file() or output.stat().st_size <= 0 for output in outputs):
        raise HuizhouZexinError("惠州泽鑫成果生成不完整，任务未完成。")
    task.update(
        {
            "status": "completed",
            "stage": "completed",
            "outputs": [str(output.relative_to(root)).replace("\\", "/") for output in outputs],
            "completed_at": _utc_now(),
            "updated_at": _utc_now(),
        }
    )
    _write_json(root / "task.json", task)
    _promote_effective(runtime_root, task, values, outputs)
    return task


def get_huizhou_zexin(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    return _load_task(runtime_root, task_id)[1]


def list_huizhou_zexin_outputs(runtime_root: Path | str, task_id: str) -> list[str]:
    root, task = _load_task(runtime_root, task_id)
    return [str((root / item).resolve()) for item in task.get("outputs", []) if (root / item).is_file()]


def recognize_huizhou_zexin_source(
    raw: OcrRawResult,
    *,
    expected_month: str,
) -> tuple[str, dict[str, str]]:
    blocks = [block.text.strip() for page in raw.pages for block in page.blocks if block.text.strip()]
    source_month = _month_from_blocks(blocks, expected_month)
    tax_values = _find_tax_summary(blocks)
    if tax_values is not None:
        energy, tax_exclusive, tax_amount, tax_inclusive = tax_values
        return "grid_bill", {
            "month": source_month,
            "grid_energy": _number(energy),
            "grid_tax_exclusive_fee": _money(tax_exclusive),
            "grid_tax": _money(tax_amount),
            "grid_tax_inclusive_fee": _money(tax_inclusive),
        }
    if GRID_CUSTOMER_NO in blocks:
        grid_energy = _month_history_value(blocks, expected_month)
        result = {"month": source_month}
        if grid_energy is not None:
            result["grid_energy"] = _number(grid_energy)
        return "grid_bill", result
    generation = _month_history_value(blocks, expected_month)
    if generation is None:
        generation = _find_generation_measurement(blocks)
    if generation is None and "0313060357636990" not in blocks:
        raise HuizhouZexinError("无法识别上网电费单，请上传清晰的原始 PDF。")
    result = {"month": source_month}
    if generation is not None:
        result["total_generation"] = _number(generation)
    return "generation_bill", result


def write_settlement_docx(template_path: Path | str, output_path: Path | str, values: ZexinValues) -> Path:
    document = Document(template_path)
    if not document.tables or len(document.tables[0].rows) < 4 or len(document.tables[0].columns) < 8:
        raise HuizhouZexinError("电费结算单模板结构不符合受控版式。")
    table = document.tables[0]
    month_number = int(values.month.split("-")[1])
    row = table.rows[2].cells
    row[0].text = f"{month_number}月"
    row[1].text = _number(values.total_generation)
    row[2].text = _number(values.grid_energy)
    row[3].text = _number(values.self_use_energy)
    row[4].text = f"{_number(values.self_use_ratio * 100, places=2)}%"
    row[5].text = _number(SELF_USE_UNIT_PRICE)
    row[6].text = _integer(values.self_use_fee)
    row[7].text = "后附南网电费单"
    for table_row in table.rows[:3]:
        for cell in table_row.cells[:7]:
            _center_cell(cell)
    _set_centered_bold_cell(
        table.cell(3, 0),
        f"{month_number}月份消纳总电量(度)：{_number(values.self_use_energy)}",
    )
    _set_centered_bold_cell(
        table.cell(3, 4),
        f"购电方应付总电费：{_integer(values.self_use_fee)}",
    )
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def write_summary_workbook(
    template_path: Path | str,
    output_path: Path | str,
    records: dict[str, ZexinValues],
    *,
    materialize_calculations: bool = False,
) -> Path:
    workbook = load_workbook(template_path)
    sheet = workbook.active
    existing = _summary_template_records(sheet)
    existing.update(records)
    months = sorted(existing)
    style = _row_style_snapshot(sheet, 4, 9)
    height = sheet.row_dimensions[4].height
    if sheet.max_row >= 4:
        sheet.delete_rows(4, sheet.max_row - 3)
    for row_number, month in enumerate(months, start=4):
        _apply_row_style(sheet, row_number, style, height)
        values = existing[month]
        year, month_number = map(int, month.split("-"))
        sheet.cell(row_number, 1, date(year, month_number, 1))
        sheet.cell(row_number, 1).number_format = 'yyyy"年"m"月"'
        sheet.cell(row_number, 2, float(values.total_generation))
        sheet.cell(row_number, 3, float(values.grid_energy))
        sheet.cell(row_number, 4, float(values.grid_unit_price))
        sheet.cell(row_number, 5, float(values.grid_tax_inclusive_fee) if materialize_calculations else f"=C{row_number}*D{row_number}")
        sheet.cell(row_number, 6, float(values.self_use_energy) if materialize_calculations else f"=B{row_number}-C{row_number}")
        sheet.cell(row_number, 7, float(SELF_USE_UNIT_PRICE))
        sheet.cell(row_number, 8, float(values.self_use_fee) if materialize_calculations else f"=F{row_number}*G{row_number}")
        sheet.cell(row_number, 9, float(values.total_income) if materialize_calculations else f"=H{row_number}+E{row_number}")
    _enable_recalculation(workbook)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _combine_sources(sources: dict[str, Any], *, expected_month: str) -> dict[str, str]:
    values: dict[str, Any] = {"month": expected_month}
    for role in REQUIRED_ROLES:
        source = sources.get(role)
        if isinstance(source, dict) and isinstance(source.get("recognized"), dict):
            values.update(source["recognized"])
    missing = [field for field in EDITABLE_FIELDS if not values.get(field)]
    if missing:
        return {
            "month": expected_month,
            "project_name": PROJECT_NAME,
            "supply_unit": SUPPLY_UNIT,
            "grid_customer_no": GRID_CUSTOMER_NO,
            "self_use_unit_price": _number(SELF_USE_UNIT_PRICE),
            **{key: str(value) for key, value in values.items() if value is not None},
        }
    return _normalize_values(values, expected_month=expected_month).to_dict()


def _normalize_values(values: dict[str, Any], *, expected_month: str) -> ZexinValues:
    month = _validate_month(str(values.get("month") or expected_month))
    if month != expected_month:
        raise HuizhouZexinError(f"识别月份 {month} 与任务月份 {expected_month} 不一致。")
    normalized = ZexinValues(
        month=month,
        total_generation=_decimal(values.get("total_generation"), "月总发电量"),
        grid_energy=_decimal(values.get("grid_energy"), "上网电量"),
        grid_tax_exclusive_fee=_decimal(values.get("grid_tax_exclusive_fee"), "上网不含税电费"),
        grid_tax_inclusive_fee=_decimal(values.get("grid_tax_inclusive_fee"), "上网含税电费"),
    )
    if normalized.total_generation <= 0 or normalized.grid_energy < 0:
        raise HuizhouZexinError("月总发电量必须大于 0，上网电量不能小于 0。")
    if normalized.grid_energy > normalized.total_generation:
        raise HuizhouZexinError("上网电量不能大于月总发电量。")
    if normalized.grid_energy == 0:
        raise HuizhouZexinError("上网电量为 0 时无法计算上网电价，请人工复核。")
    if normalized.grid_tax_exclusive_fee < 0 or normalized.grid_tax_inclusive_fee < 0:
        raise HuizhouZexinError("电费不能小于 0。")
    if normalized.grid_tax_inclusive_fee < normalized.grid_tax_exclusive_fee:
        raise HuizhouZexinError("上网含税电费不能小于不含税电费。")
    return normalized


def _find_tax_summary(blocks: list[str]) -> tuple[Decimal, Decimal, Decimal, Decimal] | None:
    numbers = [_numeric(block) for block in blocks]
    for index in range(len(numbers) - 4):
        energy, exclusive, rate, tax, inclusive = numbers[index : index + 5]
        if None in {energy, exclusive, rate, tax, inclusive}:
            continue
        assert energy is not None and exclusive is not None and rate is not None and tax is not None and inclusive is not None
        if rate not in {Decimal("1"), Decimal("3"), Decimal("6"), Decimal("9"), Decimal("13")}:
            continue
        if energy <= 0 or inclusive <= 0:
            continue
        if abs((inclusive - exclusive) - tax) <= Decimal("0.05"):
            return energy, exclusive, tax, inclusive
    return None


def _find_generation_measurement(blocks: list[str]) -> Decimal | None:
    numbers = [_numeric(block) for block in blocks]
    for index, value in enumerate(numbers):
        if value == Decimal("4000"):
            candidates = [item for item in numbers[index + 1 : index + 5] if item is not None and item > 0]
            if candidates:
                return max(candidates)
    return None


def _month_history_value(blocks: list[str], expected_month: str) -> Decimal | None:
    compact_month = expected_month.replace("-", "")
    candidates: list[Decimal] = []
    for index, block in enumerate(blocks):
        if re.sub(r"\D", "", block) != compact_month:
            continue
        for follower in blocks[index + 1 : index + 4]:
            value = _numeric(follower)
            if value is not None and value > 0 and value != Decimal(compact_month):
                candidates.append(value)
                break
    return candidates[-1] if candidates else None


def _month_from_blocks(blocks: list[str], expected_month: str) -> str:
    expected_prefix = expected_month
    joined = "\n".join(blocks)
    matches = re.findall(r"(20\d{2})\s*[-./年]\s*(0?[1-9]|1[0-2])(?:\s*[-./月]\s*(?:0?[1-9]|[12]\d|3[01])日?)?", joined)
    months = {f"{year}-{int(month):02d}" for year, month in matches}
    compact = expected_month.replace("-", "")
    if compact in {re.sub(r"\D", "", block) for block in blocks}:
        months.add(expected_month)
    if expected_prefix in months:
        return expected_month
    if len(months) == 1:
        return next(iter(months))
    raise HuizhouZexinError("无法唯一识别 PDF 账期。")


def _summary_template_records(sheet: Any) -> dict[str, ZexinValues]:
    result: dict[str, ZexinValues] = {}
    for row in range(4, sheet.max_row + 1):
        month = _month_from_excel_value(sheet.cell(row, 1).value)
        if not month:
            continue
        total = _decimal(sheet.cell(row, 2).value, "月总发电量")
        grid = _decimal(sheet.cell(row, 3).value, "上网电量")
        inclusive = grid * _decimal(sheet.cell(row, 4).value, "上网电价")
        result[month] = ZexinValues(month, total, grid, inclusive, inclusive)
    return result


def _effective_records(runtime_root: Path | str) -> dict[str, ZexinValues]:
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    result: dict[str, ZexinValues] = {}
    for month, record in (series.get("months") or {}).items():
        if isinstance(record, dict) and isinstance(record.get("values"), dict):
            result[month] = _normalize_values(record["values"], expected_month=month)
    return result


def _summary_materialization_enabled(runtime_root: Path | str) -> bool:
    series = _read_json(_series_path(runtime_root), default={})
    return series.get("summary_materialize_calculations") is True


def _promote_effective(
    runtime_root: Path | str,
    task: dict[str, Any],
    values: ZexinValues,
    outputs: list[Path],
) -> None:
    target = _series_path(runtime_root)
    series = _read_json(target, default={"site_key": SITE_KEY, "months": {}})
    months = series.setdefault("months", {})
    current = months.get(values.month) if isinstance(months.get(values.month), dict) else {}
    history = list(current.get("history") or [])
    if current.get("task_id"):
        history.append({key: current.get(key) for key in ("task_id", "revision", "values", "outputs")})
    months[values.month] = {
        "task_id": task["task_id"],
        "revision": task["revision"],
        "values": values.to_dict(),
        "outputs": [str(output) for output in outputs],
        "history": history,
        "updated_at": _utc_now(),
    }
    _write_json(target, series)


def _row_style_snapshot(sheet: Any, row: int, columns: int) -> list[Any]:
    return [copy(sheet.cell(row, column)._style) for column in range(1, columns + 1)]


def _apply_row_style(sheet: Any, row: int, styles: list[Any], height: float | None) -> None:
    sheet.row_dimensions[row].height = height
    for column, style in enumerate(styles, start=1):
        sheet.cell(row, column)._style = copy(style)


def _set_centered_bold_cell(cell: Any, text: str) -> None:
    cell.text = text
    _center_cell(cell, bold=True)


def _center_cell(cell: Any, *, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if bold:
            for run in paragraph.runs:
                run.bold = True


def _enable_recalculation(workbook: Any) -> None:
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _month_from_excel_value(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.strftime("%Y-%m")
    if isinstance(value, (int, float)):
        text = str(int(value))
        if re.fullmatch(r"20\d{4}", text):
            return f"{text[:4]}-{text[4:]}"
        try:
            parsed = from_excel(value)
            return parsed.strftime("%Y-%m")
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})\D*(0?[1-9]|1[0-2])", text)
    return f"{match.group(1)}-{int(match.group(2)):02d}" if match else None


def _validate_pdf(path: Path) -> None:
    if not path.is_file() or path.is_symlink() or path.suffix.lower() != ".pdf":
        raise HuizhouZexinError("惠州泽鑫仅接收 PDF 原始凭证。")
    with path.open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise HuizhouZexinError("文件扩展名与 PDF 内容不一致，本次文件未登记。")


def _task_root(runtime_root: Path | str, task_id: str) -> Path:
    return Path(runtime_root) / "huizhou_zexin_tasks" / task_id


def _series_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "huizhou_zexin" / "series.json"


def _load_task(runtime_root: Path | str, task_id: str) -> tuple[Path, dict[str, Any]]:
    clean = _safe_task_id(task_id)
    root = _task_root(runtime_root, clean)
    task = _read_json(root / "task.json")
    _assert_task_identity(task, clean, str(task.get("month") or ""))
    return root, task


def _assert_task_identity(task: dict[str, Any], task_id: str, month: str) -> None:
    if task.get("task_id") != task_id or task.get("business_profile") != BUSINESS_PROFILE:
        raise HuizhouZexinError("惠州泽鑫任务身份不匹配。")
    if task.get("site_key") != SITE_KEY or task.get("month") != month:
        raise HuizhouZexinError("惠州泽鑫任务跨站区或月份。")


def _next_revision(runtime_root: Path | str, month: str) -> int:
    current = (_read_json(_series_path(runtime_root), default={"months": {}}).get("months") or {}).get(month)
    return int(current.get("revision") or 0) + 1 if isinstance(current, dict) else 1


def _safe_task_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", value):
        raise HuizhouZexinError("惠州泽鑫任务编号无效。")
    return value


def _validate_month(value: str) -> str:
    if not re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", value):
        raise HuizhouZexinError("月份必须使用 YYYY-MM 格式。")
    return value


def _numeric(value: str) -> Decimal | None:
    compact = value.replace(",", "").strip()
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", compact):
        return None
    try:
        return Decimal(compact)
    except InvalidOperation:
        return None


def _decimal(value: Any, field: str) -> Decimal:
    if isinstance(value, str) and value.startswith("="):
        raise HuizhouZexinError(f"{field}缺少已计算数值。")
    try:
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, AttributeError, TypeError, ValueError) as exc:
        raise HuizhouZexinError(f"{field}不是有效数字。") from exc


def _number(value: Decimal, *, places: int | None = None) -> str:
    if places is not None:
        value = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _money(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), ".2f")


def _integer(value: Decimal) -> str:
    return format(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP), "f")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_json(path: Path, *, default: Any = None) -> Any:
    if not path.is_file():
        if default is not None:
            return default
        raise HuizhouZexinError(f"惠州泽鑫任务状态不存在：{path}")
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise HuizhouZexinError(f"惠州泽鑫任务状态损坏：{path}") from exc


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
