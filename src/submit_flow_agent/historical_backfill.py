"""Deterministic, auditable historical backfill for school summary workbooks."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from submit_flow_agent.config import load_site_config
from submit_flow_agent.excel_writer import write_detail_excel, write_summary_excel
from submit_flow_agent.normalizer import normalize_recognized_data
from submit_flow_agent.pricing import _snapshot_hash, confirm_site_pricing, propose_site_pricing
from submit_flow_agent.task_series import record_historical_revision, resolve_history_workbooks
from submit_flow_agent.validator import validate_monthly_data


class HistoricalBackfillError(RuntimeError):
    """Raised when historical source data cannot be safely staged or applied."""


def backfill_school_history(
    source_path: Path | str,
    *,
    site_key: str,
    through_month: str,
    runtime_root: Path | str = "runtime",
    config_path: Path | str = "config/sites.json",
    sheet_name: str = "Sheet1",
    verified_site_key: str | None = None,
    route_path: Path | str | None = None,
    apply: bool = False,
    replace_existing: bool = False,
    rehearsal_root: Path | str | None = None,
) -> dict[str, Any]:
    """Stage a cumulative baseline and optionally publish it to site history."""

    source = Path(source_path)
    runtime = Path(runtime_root)
    config = Path(config_path)
    if not source.is_file() or source.is_symlink():
        raise HistoricalBackfillError(f"Historical source workbook is missing or unsafe: {source}")
    _validate_month(through_month)

    workspace = Path(rehearsal_root) if rehearsal_root else Path(tempfile.mkdtemp(prefix="school-history-backfill-"))
    workspace.mkdir(parents=True, exist_ok=True)
    staged_config = workspace / "sites.json"
    staged_route = workspace / "site-routes.json" if route_path else None
    _copy_file(config, staged_config)
    _absolutize_template_dirs(staged_config, config.parent.resolve())
    identity_changes = _repair_xinan_identity(staged_config, site_key=site_key, verified_site_key=verified_site_key)
    route_changes: list[dict[str, str]] = []
    if route_path:
        _copy_file(Path(route_path), staged_route)
        route_changes = _repair_xinan_routes(staged_route, site_key=site_key, legacy_site_key=verified_site_key)

    site = load_site_config(site_key, staged_config)
    if site.meter_topology_profile != "single_meter_school":
        raise HistoricalBackfillError("This backfill command currently accepts single-meter school sites only.")
    target_history = runtime / "site_history" / site.site_key / through_month
    revision = _next_history_revision(target_history) if apply and target_history.exists() else 1
    next_month = _next_month(through_month)
    source_records = _read_source_records(source, sheet_name=sheet_name, through_month=next_month)
    records = [record for record in source_records if record["month"] <= through_month]
    next_record = next((record for record in source_records if record["month"] == next_month), None)
    if not records or records[-1]["month"] != through_month:
        raise HistoricalBackfillError(f"Source does not contain the requested through_month {through_month}.")
    if any(record["meter_no"] != site.meter_no for record in records):
        raise HistoricalBackfillError("Source meter number does not match the target site's configured meter.")
    _apply_verified_history(
        source_records,
        runtime_root=runtime,
        verified_site_keys=tuple(dict.fromkeys(key for key in (verified_site_key, site.site_key) if key)),
        target_meter_no=site.meter_no,
    )

    history_dir = _stage_history(
        workspace,
        records=records,
        site_key=site.site_key,
        source_path=source,
        through_month=through_month,
        config_path=staged_config,
        revision=revision,
    )
    pricing_rehearsal = _stage_pricing_history(
        workspace,
        records=source_records,
        site_key=site.site_key,
        runtime_root=runtime,
        config_path=staged_config,
    )
    rehearsal = _verify_staged_baseline(
        workspace,
        records=records,
        site_key=site.site_key,
        through_month=through_month,
        config_path=staged_config,
        next_record=next_record,
    )
    result: dict[str, Any] = {
        "status": "ready_to_apply" if not apply else "applied",
        "site_key": site.site_key,
        "site_name": site.site_name,
        "through_month": through_month,
        "months": [record["month"] for record in records],
        "source": str(source.resolve()),
        "history_dir": str(history_dir),
        "identity_changes": identity_changes,
        "route_changes": route_changes,
        "verified_overrides": [
            {
                "month": record["month"],
                "fields": record.get("verified_fields", []),
                "sources": record.get("verified_sources", []),
                "differences": record.get("verified_differences", []),
            }
            for record in records
            if record.get("verified_fields")
        ],
        "rehearsal": rehearsal,
        "pricing_rehearsal": pricing_rehearsal,
        "production_write": False,
    }
    if not apply:
        _write_json_atomic(workspace / "backfill_plan.json", result)
        return result

    if target_history.exists() and not replace_existing:
        raise HistoricalBackfillError(f"Target history already exists; refusing to overwrite: {target_history}")
    backup_root = runtime / "backups" / f"historical_backfill_{site.site_key}_{_timestamp()}"
    _backup_scope(
        backup_root,
        config_path=config,
        runtime_root=runtime,
        site_key=site.site_key,
        route_path=Path(route_path) if route_path else None,
    )
    _repair_xinan_identity(config, site_key=site_key, verified_site_key=verified_site_key)
    if route_path and staged_route:
        _copy_file(staged_route, Path(route_path))
    _publish_history_revision(
        history_dir,
        target_history=target_history,
        backup_root=backup_root,
        replace_existing=replace_existing,
    )
    staged_pricing = workspace / "runtime" / "site_pricing" / site.site_key
    if staged_pricing.exists():
        shutil.copytree(staged_pricing, runtime / "site_pricing" / site.site_key, dirs_exist_ok=True)
    manifest = _read_json(target_history / "manifest.json")
    series_result = record_historical_revision(runtime, manifest)
    result.update(
        {
            "production_write": True,
            "history_dir": str(target_history.resolve()),
            "backup_root": str(backup_root.resolve()),
            "series_dir": str((runtime / "site_series" / site.site_key / through_month).resolve()),
            "revision": manifest["revision"],
            "series_result": series_result,
        }
    )
    _write_json_atomic(backup_root / "apply_receipt.json", result)
    return result


def _absolutize_template_dirs(config_path: Path, original_config_dir: Path) -> None:
    payload = _read_json(config_path)
    for site in payload.get("sites") or []:
        if not isinstance(site, dict):
            continue
        value = site.get("template_dir")
        if isinstance(value, str) and value and not Path(value).is_absolute():
            site["template_dir"] = str((original_config_dir / value).resolve())
    _write_json_atomic(config_path, payload)


def _repair_xinan_identity(config_path: Path, *, site_key: str, verified_site_key: str | None) -> list[dict[str, str]]:
    payload = _read_json(config_path)
    sites = payload.get("sites")
    if not isinstance(sites, list):
        raise HistoricalBackfillError("Site config must contain a sites list.")
    changes: list[dict[str, str]] = []
    target = next((item for item in sites if isinstance(item, dict) and item.get("site_key") == site_key), None)
    if target is None:
        raise HistoricalBackfillError(f"Target site_key is missing: {site_key}")
    if site_key == "site_5f14cc7ab0e2":
        if target.get("site_name") != "新安高中部":
            changes.append({"field": "site_name", "before": str(target.get("site_name")), "after": "新安高中部"})
            target["site_name"] = "新安高中部"
        target["aliases"] = list(
            dict.fromkeys([*(target.get("aliases") or []), "新安中学", "深圳新安中学高中部", "新安高中"])
        )
    if verified_site_key:
        legacy = next((item for item in sites if isinstance(item, dict) and item.get("site_key") == verified_site_key), None)
        if legacy is not None and legacy.get("retired") is not True:
            legacy["retired"] = True
            changes.append({"field": f"{verified_site_key}.retired", "before": "false", "after": "true"})
    _write_json_atomic(config_path, payload)
    return changes


def _repair_xinan_routes(route_path: Path, *, site_key: str, legacy_site_key: str | None) -> list[dict[str, str]]:
    payload = _read_json(route_path)
    changes: list[dict[str, str]] = []
    for section in ("activeDirectSites", "canonicalSiteRoutes"):
        records = payload.get(section)
        if not isinstance(records, dict):
            continue
        for route_id, record in records.items():
            if not isinstance(record, dict) or record.get("siteKey") not in {site_key, legacy_site_key}:
                continue
            before_key = str(record.get("siteKey") or "")
            before_label = str(record.get("label") or "")
            record["siteKey"] = site_key
            record["label"] = "光伏填报｜新安高中部"
            if before_key != site_key or before_label != record["label"]:
                changes.append(
                    {"route": str(route_id), "before": f"{before_key}|{before_label}", "after": f"{site_key}|{record['label']}"}
                )
    _write_json_atomic(route_path, payload)
    return changes


def _read_source_records(source: Path, *, sheet_name: str, through_month: str) -> list[dict[str, Any]]:
    # Some operator workbooks omit worksheet dimension metadata, which makes
    # openpyxl read-only mode report max_row=None even though rows are present.
    formulas = load_workbook(source, data_only=False, read_only=False)
    values = load_workbook(source, data_only=True, read_only=False)
    try:
        if sheet_name not in formulas.sheetnames or sheet_name not in values.sheetnames:
            raise HistoricalBackfillError(f"Source worksheet is missing: {sheet_name}")
        formula_sheet = formulas[sheet_name]
        value_sheet = values[sheet_name]
        records: list[dict[str, Any]] = []
        for row in range(4, formula_sheet.max_row + 1):
            month = _month_text(value_sheet.cell(row, 2).value)
            if month is None or month > through_month:
                continue
            meter_no = str(value_sheet.cell(row, 1).value or "").replace("表号", "").strip()
            record = {
                "month": month,
                "source_row": row,
                "meter_no": meter_no,
                "previous_reading": _decimal(value_sheet.cell(row, 3).value, f"C{row}"),
                "current_reading": _decimal(value_sheet.cell(row, 4).value, f"D{row}"),
                "multiplier": _decimal(value_sheet.cell(row, 5).value, f"E{row}"),
                "generation_kwh": _decimal(value_sheet.cell(row, 6).value, f"F{row}"),
                "grid_export_kwh": _decimal(value_sheet.cell(row, 7).value, f"G{row}"),
                "grid_price": _decimal(value_sheet.cell(row, 8).value, f"H{row}"),
                "grid_fee": _decimal(value_sheet.cell(row, 9).value, f"I{row}"),
                "effective_unit_price": _decimal(value_sheet.cell(row, 11).value, f"K{row}"),
                "generation_formula": formula_sheet.cell(row, 6).value,
                "source_type": "historical_source",
            }
            calculated = (record["current_reading"] - record["previous_reading"]) * record["multiplier"]
            rounded = calculated.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
            if min(abs(calculated - record["generation_kwh"]), abs(rounded - record["generation_kwh"])) > Decimal("0.01"):
                raise HistoricalBackfillError(f"Generation value conflicts with readings at source row {row}.")
            if not record["meter_no"]:
                raise HistoricalBackfillError(f"Meter number is blank at source row {row}.")
            records.append(record)
        records.sort(key=lambda item: item["month"])
        return records
    finally:
        formulas.close()
        values.close()


def _apply_verified_history(
    records: list[dict[str, Any]],
    *,
    runtime_root: Path,
    verified_site_keys: tuple[str, ...],
    target_meter_no: str,
) -> None:
    if not verified_site_keys:
        return
    for verified_site_key in verified_site_keys:
        for record in records:
            month = record["month"]
            history_manifest = runtime_root / "site_history" / verified_site_key / month / "manifest.json"
            if not history_manifest.is_file():
                continue
            manifest = _read_json(history_manifest)
            task_id = manifest.get("effective_task_id") or manifest.get("source_task_id")
            task_path = runtime_root / "tasks" / str(task_id) / "task.json"
            if not task_path.is_file():
                continue
            task = _read_json(task_path)
            report_rel = next(
                (item for item in task.get("outputs") or [] if str(item).endswith("validation_report.json")),
                None,
            )
            report_path = task_path.parent / str(report_rel) if report_rel else None
            if report_path is None or not report_path.is_file():
                continue
            normalized = _read_json(report_path).get("normalized_data")
            if not isinstance(normalized, dict):
                continue
            if Decimal(str(normalized.get("previous_reading"))) != record["previous_reading"]:
                raise HistoricalBackfillError(f"Verified history reading conflicts with source for {month}.")
            if Decimal(str(normalized.get("current_reading"))) != record["current_reading"]:
                raise HistoricalBackfillError(f"Verified history reading conflicts with source for {month}.")
            fields = list(record.get("verified_fields") or [])
            differences = list(record.get("verified_differences") or [])
            for source_field, target_field in (
                ("total_generation_kwh", "generation_kwh"),
                ("grid_export_kwh", "grid_export_kwh"),
                ("grid_price", "grid_price"),
                ("grid_fee", "grid_fee"),
                ("effective_unit_price", "effective_unit_price"),
            ):
                if source_field not in normalized:
                    continue
                verified_value = Decimal(str(normalized[source_field]))
                previous_value = record[target_field]
                if previous_value != verified_value:
                    differences.append(
                        {
                            "field": target_field,
                            "source_value": _decimal_text(previous_value),
                            "verified_value": _decimal_text(verified_value),
                            "verified_site_key": verified_site_key,
                            "task_id": str(task_id),
                        }
                    )
                record[target_field] = verified_value
                if target_field not in fields:
                    fields.append(target_field)
            record["meter_no"] = target_meter_no
            record["source_type"] = "verified_runtime_preferred"
            record["verified_fields"] = fields
            record["verified_differences"] = differences
            record["verified_sources"] = [
                *list(record.get("verified_sources") or []),
                {"site_key": verified_site_key, "task_id": str(task_id)},
            ]


def _stage_history(
    workspace: Path,
    *,
    records: list[dict[str, Any]],
    site_key: str,
    source_path: Path,
    through_month: str,
    config_path: Path,
    revision: int = 1,
) -> Path:
    site = load_site_config(site_key, config_path)
    project = workspace / "project"
    project.mkdir(parents=True, exist_ok=True)
    for template in Path(site.template_dir).glob("*.xlsx"):
        shutil.copy2(template, project / template.name)
    report_dir = workspace / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    summary_output: Path | None = None
    target_report: Path | None = None
    for record in records:
        report = validate_monthly_data(normalize_recognized_data(_recognized_payload(record, site))).to_dict()
        if not report["can_write_excel"]:
            raise HistoricalBackfillError(f"Validation failed for {record['month']}: {report['errors']}")
        report_path = report_dir / f"{record['month']}.json"
        _write_json_atomic(report_path, report)
        summary_output = workspace / f"summary-{record['month']}.xlsx"
        write_summary_excel(
            report_path,
            project,
            summary_output,
            site_key=site.site_key,
            config_path=config_path,
            pricing_snapshot=_pricing_snapshot(site.site_key, record),
        )
        shutil.copy2(summary_output, project / site.summary_workbook)
        if record["month"] == through_month:
            target_report = report_path
    if summary_output is None or target_report is None:
        raise HistoricalBackfillError("No source records were staged.")
    revision_name = f"r{revision:04d}"
    revision_dir = workspace / "runtime" / "site_history" / site.site_key / through_month / "revisions" / revision_name
    revision_dir.mkdir(parents=True, exist_ok=True)
    summary_name = site.output_workbook_name("summary", through_month)
    detail_name = site.output_workbook_name("detail", through_month)
    summary_path = revision_dir / summary_name
    detail_path = revision_dir / detail_name
    _ensure_meter_column_visible(summary_output, site.summary_sheet)
    shutil.copy2(summary_output, summary_path)
    write_detail_excel(
        target_report,
        project,
        detail_path,
        site_key=site.site_key,
        config_path=config_path,
        pricing_snapshot=_pricing_snapshot(site.site_key, records[-1]),
    )
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    source_task_id = f"historical_backfill_{site.site_key}_{through_month.replace('-', '_')}_{revision_name}"
    manifest = {
        "site_key": site.site_key,
        "site_name": site.site_name,
        "month": through_month,
        "revision": revision,
        "revision_name": revision_name,
        "source_task_id": source_task_id,
        "effective_revision": revision,
        "effective_task_id": source_task_id,
        "created_at": now,
        "updated_at": now,
        "migration": {
            "kind": "historical_school_backfill",
            "source": str(source_path.resolve()),
            "source_sha256": _sha256(source_path),
            "months": [record["month"] for record in records],
        },
        "baseline_dependency": None,
        "staleness": {
            "is_stale": False,
            "upstream_changed_month": None,
            "upstream_changed_revision": None,
            "upstream_changed_task_id": None,
            "marked_at": None,
            "cleared_at": now,
        },
        "workbooks": {
            "detail_workbook": {"path": f"revisions/{revision_name}/{detail_name}", "sha256": _sha256(detail_path)},
            "summary_workbook": {"path": f"revisions/{revision_name}/{summary_name}", "sha256": _sha256(summary_path)},
        },
    }
    history_dir = revision_dir.parent.parent
    _write_json_atomic(
        revision_dir / "manifest.json",
        {key: value for key, value in manifest.items() if key not in {"effective_revision", "effective_task_id"}},
    )
    _write_json_atomic(history_dir / "manifest.json", manifest)
    return history_dir


def _next_history_revision(history_dir: Path) -> int:
    manifest_path = history_dir / "manifest.json"
    if not manifest_path.is_file() or manifest_path.is_symlink():
        return 1
    manifest = _read_json(manifest_path)
    try:
        return int(manifest.get("effective_revision") or manifest.get("revision") or 0) + 1
    except (TypeError, ValueError) as exc:
        raise HistoricalBackfillError(f"Existing history revision is invalid: {manifest_path}") from exc


def _publish_history_revision(
    staged_history: Path,
    *,
    target_history: Path,
    backup_root: Path,
    replace_existing: bool,
) -> None:
    target_history.parent.mkdir(parents=True, exist_ok=True)
    publish_source = staged_history
    if target_history.exists():
        if not replace_existing:
            raise HistoricalBackfillError(f"Target history already exists; refusing to overwrite: {target_history}")
        staged_manifest = _read_json(staged_history / "manifest.json")
        revision_name = str(staged_manifest["revision_name"])
        merged_history = staged_history.parent / f"{staged_history.name}-merged"
        if merged_history.exists():
            shutil.rmtree(merged_history)
        shutil.copytree(target_history, merged_history)
        merged_revision = merged_history / "revisions" / revision_name
        if merged_revision.exists():
            raise HistoricalBackfillError(f"History revision already exists: {merged_revision}")
        shutil.copytree(staged_history / "revisions" / revision_name, merged_revision)
        _write_json_atomic(merged_history / "manifest.json", staged_manifest)
        publish_source = merged_history

    publish_temp = target_history.parent / f".{target_history.name}.historical-backfill-{_timestamp()}"
    if publish_temp.exists():
        raise HistoricalBackfillError(f"Temporary publish directory already exists: {publish_temp}")
    shutil.copytree(publish_source, publish_temp)
    displaced: Path | None = None
    try:
        if target_history.exists():
            displaced = backup_root / "replaced_target" / target_history.name
            displaced.parent.mkdir(parents=True, exist_ok=True)
            os.replace(target_history, displaced)
        os.replace(publish_temp, target_history)
    except Exception:
        if displaced is not None and displaced.exists() and not target_history.exists():
            os.replace(displaced, target_history)
        raise
    finally:
        if publish_temp.exists():
            shutil.rmtree(publish_temp)


def _ensure_meter_column_visible(workbook_path: Path, sheet_name: str) -> None:
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook[sheet_name]
        current = sheet.column_dimensions["A"].width or 0
        sheet.column_dimensions["A"].width = max(current, 36)
        workbook.save(workbook_path)
    finally:
        workbook.close()


def _verify_staged_baseline(
    workspace: Path,
    *,
    records: list[dict[str, Any]],
    site_key: str,
    through_month: str,
    config_path: Path,
    next_record: dict[str, Any] | None,
) -> dict[str, Any]:
    site = load_site_config(site_key, config_path)
    baseline = resolve_history_workbooks(
        workspace / "runtime",
        site_key=site.site_key,
        month=through_month,
        workbook_names=(site.detail_workbook, site.summary_workbook),
    )
    if baseline is None:
        raise HistoricalBackfillError("Staged history cannot be resolved as a baseline.")
    expected_months = [record["month"] for record in records]
    workbook = load_workbook(baseline["summary_workbook"], data_only=False, read_only=True)
    try:
        sheet = workbook[site.summary_sheet]
        actual_months = [
            month
            for row in range(site.summary_start_row, sheet.max_row + 1)
            if (month := _month_text(sheet.cell(row, 2).value))
        ]
    finally:
        workbook.close()
    if actual_months != expected_months:
        raise HistoricalBackfillError(f"Staged summary month sequence mismatch: {actual_months}")
    result = {
        "baseline_resolved": True,
        "summary_months": actual_months,
        "next_month": _next_month(through_month),
        "next_month_will_use_baseline": True,
    }
    if next_record is None:
        return result
    project = workspace / "next-month-project"
    project.mkdir(parents=True, exist_ok=True)
    shutil.copy2(baseline["detail_workbook"], project / site.detail_workbook)
    shutil.copy2(baseline["summary_workbook"], project / site.summary_workbook)
    report = validate_monthly_data(normalize_recognized_data(_recognized_payload(next_record, site))).to_dict()
    if not report["can_write_excel"]:
        raise HistoricalBackfillError(f"Next-month rehearsal validation failed: {report['errors']}")
    report_path = workspace / "reports" / f"rehearsal-{next_record['month']}.json"
    _write_json_atomic(report_path, report)
    output = workspace / f"rehearsal-{next_record['month']}.xlsx"
    write_summary_excel(
        report_path,
        project,
        output,
        site_key=site.site_key,
        config_path=config_path,
        pricing_snapshot=_pricing_snapshot(site.site_key, next_record),
    )
    workbook = load_workbook(output, data_only=False, read_only=False)
    try:
        sheet = workbook[site.summary_sheet]
        generated_months = [
            month
            for row in range(site.summary_start_row, sheet.max_row + 1)
            if (month := _month_text(sheet.cell(row, 2).value))
        ]
    finally:
        workbook.close()
    expected_generated = [*expected_months, next_record["month"]]
    if generated_months != expected_generated:
        raise HistoricalBackfillError(f"Next-month rehearsal sequence mismatch: {generated_months}")
    result.update(
        {
            "next_month_generated": True,
            "generated_summary_months": generated_months,
            "generated_summary": str(output.resolve()),
        }
    )
    return result


def _stage_pricing_history(
    workspace: Path,
    *,
    records: list[dict[str, Any]],
    site_key: str,
    runtime_root: Path,
    config_path: Path,
) -> dict[str, Any]:
    staged_runtime = workspace / "runtime"
    source_pricing = runtime_root / "site_pricing" / site_key
    target_pricing = staged_runtime / "site_pricing" / site_key
    if source_pricing.exists() and not target_pricing.exists():
        shutil.copytree(source_pricing, target_pricing)
    migrated: list[dict[str, Any]] = []
    for record in records:
        if record.get("source_type") != "verified_runtime_preferred":
            continue
        proposal = propose_site_pricing(
            staged_runtime,
            site_key=site_key,
            month=record["month"],
            proposal={"discount_rate": "0.95", "effective_unit_price": _decimal_text(record["effective_unit_price"])},
            config_path=config_path,
        )
        if not proposal.ready_for_confirmation or not proposal.proposal_id:
            raise HistoricalBackfillError(f"Pricing rehearsal is incomplete for {record['month']}.")
        snapshot = confirm_site_pricing(
            staged_runtime,
            proposal_id=proposal.proposal_id,
            config_path=config_path,
            pricing_source="historical_backfill_verified_runtime",
        )
        migrated.append(
            {
                "month": record["month"],
                "effective_unit_price": snapshot["effective_unit_price"],
                "pricing_version": snapshot["pricing_version"],
            }
        )
    return {"migrated": migrated, "staged_pricing_dir": str(target_pricing)}


def _recognized_payload(record: dict[str, Any], site: Any) -> dict[str, Any]:
    month = record["month"]
    school_consumption = record["generation_kwh"] - record["grid_export_kwh"]
    values = {
        "site_key": site.site_key,
        "site_name": site.site_name,
        "project_name": site.project_name,
        "month": month,
        "meter_no": site.meter_no,
        "period_start": f"{month}-01",
        "period_end": f"{_next_month(month)}-01",
        "previous_reading": record["previous_reading"],
        "current_reading": record["current_reading"],
        "multiplier": record["multiplier"],
        "total_generation_kwh": record["generation_kwh"],
        "grid_export_kwh": record["grid_export_kwh"],
        "grid_price": record["grid_price"],
        "grid_fee": record["grid_fee"],
        "school_base_price": record["effective_unit_price"],
        "school_discount_rate": Decimal("0.95"),
        "effective_unit_price": record["effective_unit_price"],
        "company_name": site.company_name,
        "statement_date": date.today().isoformat(),
    }

    def field(value: Any) -> dict[str, Any]:
        return {"value": _json_value(value), "source_type": record["source_type"], "source_file": "historical backfill"}

    return {
        "site_key": site.site_key,
        "site_name": site.site_name,
        "project_name": site.project_name,
        "month": month,
        "mapping_mode": "historical_backfill",
        "fields": {key: field(value) for key, value in values.items()},
        "meter_rows": [
            {
                "meter_no": field(site.meter_no),
                "previous_reading": field(record["previous_reading"]),
                "current_reading": field(record["current_reading"]),
                "multiplier": field(record["multiplier"]),
                "generation_kwh": field(record["generation_kwh"]),
                "grid_export_kwh": field(record["grid_export_kwh"]),
                "school_consumption_kwh": field(school_consumption),
                "detail_enabled": field(True),
            }
        ],
        "validation_not_applicable_fields": [],
    }


def _pricing_snapshot(site_key: str, record: dict[str, Any]) -> dict[str, Any]:
    month = record["month"]
    snapshot = {
        "site_key": site_key,
        "month": month,
        "discount_rate": "0.95",
        "effective_unit_price": _decimal_text(record["effective_unit_price"]),
        "base_unit_price": None,
        "pricing_version": 0,
        "pricing_source": "historical_backfill",
        "confirmed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "effective_month": month,
        "proposal_id": f"historical-backfill-{site_key}-{month}",
        "proposal_hash": _sha256_text(f"{site_key}|{month}|{record['effective_unit_price']}"),
    }
    snapshot["pricing_hash"] = _snapshot_hash(snapshot)
    return snapshot


def _backup_scope(backup_root: Path, *, config_path: Path, runtime_root: Path, site_key: str, route_path: Path | None) -> None:
    backup_root.mkdir(parents=True, exist_ok=False)
    _copy_file(config_path, backup_root / "config" / "sites.json")
    if route_path:
        _copy_file(route_path, backup_root / "openclaw" / "site-routes.json")
    for name in ("site_history", "site_series", "site_pricing"):
        source = runtime_root / name / site_key
        if source.exists():
            shutil.copytree(source, backup_root / "runtime" / name / site_key)


def _copy_file(source: Path, target: Path) -> None:
    if not source.is_file() or source.is_symlink():
        raise HistoricalBackfillError(f"Required file is missing or unsafe: {source}")
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        shutil.copy2(source, temporary)
        os.replace(temporary, target)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_json(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise HistoricalBackfillError(f"Invalid JSON file: {path}") from exc
    if not isinstance(payload, dict):
        raise HistoricalBackfillError(f"JSON file must contain an object: {path}")
    return payload


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2, default=_json_value)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _month_text(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, str):
        text = value.strip().replace("年", "-").replace("月", "")
        parts = text.split("-")
        if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
            return f"{int(parts[0]):04d}-{int(parts[1]):02d}"
    return None


def _next_month(month: str) -> str:
    year, number = (int(part) for part in month.split("-", 1))
    return f"{year + 1:04d}-01" if number == 12 else f"{year:04d}-{number + 1:02d}"


def _validate_month(month: str) -> None:
    try:
        year, number = (int(part) for part in month.split("-", 1))
        date(year, number, 1)
    except (TypeError, ValueError) as exc:
        raise HistoricalBackfillError(f"Month must be YYYY-MM, got {month!r}.") from exc


def _decimal(value: Any, cell: str) -> Decimal:
    if value is None or value == "":
        raise HistoricalBackfillError(f"Required source cell is blank: {cell}")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise HistoricalBackfillError(f"Source cell must be numeric: {cell}={value!r}") from exc


def _decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S_%f")
