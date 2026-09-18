"""Dedicated monthly reporting workflow for Southern University of Science and Technology."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

from submit_flow_agent.ocr.base import OcrAdapter, OcrBlock, OcrRawResult, write_ocr_raw_result
from submit_flow_agent.ocr.factory import create_ocr_adapter
from submit_flow_agent.pdf_renderer import RenderedPage, render_pdf_to_images


BUSINESS_PROFILE = "nanke_base_station_monthly_reporting_v1"
SITE_KEY = "nanke"
SITE_NAME = "南科大"
METER_NUMBER = "09001SF00000042509389672"
DISCOUNT_RATE = Decimal("0.88")
OUTPUT_DETAIL = "南科大结算单.xlsx"
OUTPUT_SUMMARY = "南科大汇总表.xlsx"
REQUIRED_ROLES = (
    "generation_statement",
    "grid_fee_settlement",
    "grid_energy_settlement",
    "base_station_notice",
)
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".pdf"}
MAX_SOURCE_BYTES = 30 * 1024 * 1024


class NankeError(RuntimeError):
    """Raised when the dedicated Nanke workflow cannot advance safely."""


@dataclass(frozen=True)
class NankeValues:
    month: str
    meter_number: str
    prior_reading: Decimal
    current_reading: Decimal
    multiplier: Decimal
    total_generation: Decimal
    grid_energy: Decimal
    grid_price: Decimal
    grid_fee: Decimal
    base_station_energy: Decimal
    base_station_price: Decimal
    base_station_fee: Decimal

    @property
    def calculated_generation(self) -> Decimal:
        return ((self.current_reading - self.prior_reading) * self.multiplier).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )

    @property
    def school_energy(self) -> Decimal:
        return self.total_generation - self.grid_energy - self.base_station_energy

    @property
    def school_price(self) -> Decimal:
        return self.base_station_price * DISCOUNT_RATE

    @property
    def school_fee(self) -> Decimal:
        return self.school_energy * self.school_price

    @property
    def grid_formula_fee(self) -> Decimal:
        return self.grid_energy * self.grid_price

    @property
    def total_income(self) -> Decimal:
        # The controlled workbook excludes the base-station payment from PV income.
        return self.school_fee + self.grid_formula_fee

    def to_dict(self) -> dict[str, str]:
        return {
            "month": self.month,
            "meter_number": self.meter_number,
            "prior_reading": _number(self.prior_reading),
            "current_reading": _number(self.current_reading),
            "multiplier": _number(self.multiplier),
            "total_generation": _number(self.total_generation),
            "calculated_generation": _number(self.calculated_generation),
            "grid_energy": _number(self.grid_energy),
            "grid_price": _number(self.grid_price, places=6),
            "grid_fee": _money(self.grid_fee),
            "grid_formula_fee": _money(self.grid_formula_fee),
            "base_station_energy": _number(self.base_station_energy),
            "base_station_price": _number(self.base_station_price),
            "base_station_fee": _money(self.base_station_fee),
            "school_energy": _number(self.school_energy),
            "discount_rate": _number(DISCOUNT_RATE),
            "school_price": _number(self.school_price, places=5),
            "school_fee": _money(self.school_fee),
            "total_income": _money(self.total_income),
        }


def prepare_nanke(runtime_root: Path | str, *, task_id: str, month: str) -> dict[str, Any]:
    task_id = _safe_task_id(task_id)
    month = _validate_month(month)
    root = _task_root(runtime_root, task_id)
    if root.exists():
        task = _read_json(root / "task.json")
        if task.get("task_id") != task_id or task.get("month") != month:
            raise NankeError("任务编号已被其他南科大月份使用。")
        return task
    root.mkdir(parents=True)
    task = {
        "task_id": task_id,
        "business_profile": BUSINESS_PROFILE,
        "site_key": SITE_KEY,
        "site_name": SITE_NAME,
        "month": month,
        "revision": _next_revision(runtime_root, month),
        "status": "collecting_files",
        "stage": "collecting_files",
        "sources": {},
        "source_count": 0,
        "recognized_data": {},
        "confirmed_data": {},
        "missing_fields": list(REQUIRED_ROLES),
        "warnings": [],
        "outputs": [],
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
    }
    _write_json(root / "task.json", task)
    return task


def collect_nanke_source(
    runtime_root: Path | str,
    task_id: str,
    source: Path | str,
    *,
    ocr_adapter: OcrAdapter | None = None,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] not in {"collecting_files", "awaiting_confirmation", "need_review"}:
        raise NankeError("当前南科大任务不接收文件。")
    source_path = Path(source)
    _validate_source(source_path)
    digest = _sha256_file(source_path)
    sources = task.get("sources") if isinstance(task.get("sources"), dict) else {}
    if any(isinstance(item, dict) and item.get("sha256") == digest for item in sources.values()):
        return {**task, "duplicate": True}

    staged = root / "inputs" / f"source-{digest[:16]}{source_path.suffix.lower()}"
    staged.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, staged)
    pages = _source_pages(staged, root / "rendered" / digest[:16])
    raw = (ocr_adapter or create_ocr_adapter()).recognize(staged, pages)
    role, extracted = recognize_nanke_source(raw, expected_month=task["month"])
    if role in sources:
        staged.unlink(missing_ok=True)
        raise NankeError(f"{_role_label(role)}已登记，请勿重复上传同类文件。")
    source_month = extracted.get("month")
    if source_month != task["month"]:
        staged.unlink(missing_ok=True)
        raise NankeError(
            f"文件账期为 {source_month or '无法识别'}，与当前填报月份 {task['month']} 不一致；本次文件未登记。"
        )
    write_ocr_raw_result(raw, root / "ocr" / f"{role}.json", relative_to=root)
    sources[role] = {
        "role": role,
        "name": source_path.name,
        "stored_path": str(staged.relative_to(root)).replace("\\", "/"),
        "sha256": digest,
        "recognized": extracted,
    }
    task["sources"] = sources
    task["source_count"] = sum(role in sources for role in REQUIRED_ROLES)
    task["missing_fields"] = [role for role in REQUIRED_ROLES if role not in sources]
    task["updated_at"] = _utc_now()
    if task["missing_fields"]:
        task["status"] = task["stage"] = "collecting_files"
        _write_json(root / "task.json", task)
        return task

    values, warnings = _combine_sources(task)
    task.update(
        recognized_data=values.to_dict(),
        warnings=warnings,
        status="awaiting_confirmation",
        stage="awaiting_confirmation",
        updated_at=_utc_now(),
    )
    _write_json(root / "task.json", task)
    return task


def confirm_nanke(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] != "awaiting_confirmation":
        raise NankeError("当前南科大任务不在数据确认阶段。")
    values = _values_from_dict(task.get("recognized_data"), expected_month=task["month"])
    task.update(
        confirmed_data=values.to_dict(),
        status="ready_to_run",
        stage="ready_to_run",
        updated_at=_utc_now(),
    )
    _write_json(root / "task.json", task)
    return task


def run_nanke(
    runtime_root: Path | str,
    task_id: str,
    *,
    template_dir: Path | str,
) -> dict[str, Any]:
    root, task = _load_task(runtime_root, task_id)
    if task["status"] == "completed":
        return task
    if task["status"] != "ready_to_run":
        raise NankeError("请先核对并确认识别数据，再发送“开始运行”。")
    values = _values_from_dict(task.get("confirmed_data"), expected_month=task["month"])
    templates = Path(template_dir)
    detail_template = templates / OUTPUT_DETAIL
    summary_template = templates / OUTPUT_SUMMARY
    if not detail_template.is_file() or not summary_template.is_file():
        raise NankeError("南科大受控模板不完整。")
    output_dir = root / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    detail_output = output_dir / OUTPUT_DETAIL
    summary_output = output_dir / OUTPUT_SUMMARY
    write_nanke_detail(detail_template, detail_output, values)
    records = _effective_records(runtime_root)
    records[values.month] = values
    write_nanke_summary(
        summary_template,
        summary_output,
        records,
        materialize_calculations=_summary_materialization_enabled(runtime_root),
    )
    outputs = [detail_output, summary_output]
    if any(not output.is_file() or output.stat().st_size <= 0 for output in outputs):
        raise NankeError("南科大成果生成不完整，任务未完成。")
    task.update(
        status="completed",
        stage="completed",
        outputs=[str(output.relative_to(root)).replace("\\", "/") for output in outputs],
        completed_at=_utc_now(),
        updated_at=_utc_now(),
    )
    _write_json(root / "task.json", task)
    _promote_effective(runtime_root, task, values, outputs)
    return task


def get_nanke(runtime_root: Path | str, task_id: str) -> dict[str, Any]:
    return _load_task(runtime_root, task_id)[1]


def list_nanke_outputs(runtime_root: Path | str, task_id: str) -> list[str]:
    root, task = _load_task(runtime_root, task_id)
    return [str((root / item).resolve()) for item in task.get("outputs", []) if (root / item).is_file()]


def recognize_nanke_source(raw: OcrRawResult, *, expected_month: str) -> tuple[str, dict[str, Any]]:
    text = "\n".join(block.text for page in raw.pages for block in page.blocks)
    compact = _compact(text)
    month = _month_from_text(text, expected_month)
    if METER_NUMBER in compact and ("发电量" in compact or "补贴核算单" in compact):
        values = _generation_values(raw)
        return "generation_statement", {"month": month, "meter_number": METER_NUMBER, **values}
    if "电费结算单" in compact and ("市场化电费" in compact or "结算含税" in compact):
        energy, price, fee = _market_fee_values(raw)
        return "grid_fee_settlement", {
            "month": month,
            "grid_energy": _number(energy),
            "grid_price": _number(price, places=6),
            "grid_fee": _money(fee),
        }
    if "电量结算单" in compact and ("日电量" in compact or "抄表基本信息" in compact):
        return "grid_energy_settlement", {
            "month": month,
            "grid_energy": _number(_energy_total(raw)),
        }
    if "缴费通知书" in compact and ("基站" in compact or "CMGD" in compact.upper()):
        energy, price, fee = _base_station_values(text)
        return "base_station_notice", {
            "month": month,
            "base_station_energy": _number(energy),
            "base_station_price": _number(price),
            "base_station_fee": _money(fee),
        }
    raise NankeError("无法识别材料类型，请上传发电单、电费结算单、电量结算单和基站缴费通知书。")


def write_nanke_detail(template: Path | str, output: Path | str, values: NankeValues) -> Path:
    workbook = load_workbook(template)
    try:
        sheet = workbook["Sheet1"]
        year, month = map(int, values.month.split("-"))
        title = str(sheet["A1"].value or "")
        sheet["A1"] = re.sub(r"[一二三四五六七八九十]+月份", f"{_chinese_month(month)}月份", title)
        sheet["E2"] = date(year, month, 1)
        sheet["I2"] = _next_month_date(year, month)
        sheet["A5"] = f"电表{values.meter_number}"
        for cell, value in {
            "B5": values.prior_reading,
            "C5": values.current_reading,
            "D5": values.multiplier,
            "E5": values.total_generation,
            "F5": values.grid_energy,
            "G5": values.base_station_energy,
        }.items():
            sheet[cell] = float(value)
        sheet["H5"] = "=E5-F5-G5"
        sheet["I5"] = f"={_number(DISCOUNT_RATE)}*{_number(values.base_station_price)}"
        sheet["J5"] = "=H5*I5"
        sheet["K5"] = "=J5"
        workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        Path(output).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()
    return Path(output)


def write_nanke_summary(
    template: Path | str,
    output: Path | str,
    records: dict[str, NankeValues],
    *,
    materialize_calculations: bool = False,
) -> Path:
    workbook = load_workbook(template)
    try:
        sheet = workbook["Sheet1"]
        template_row = 4
        ordered = sorted(records.items())
        for index, (month, values) in enumerate(ordered, start=4):
            if index > sheet.max_row:
                _copy_row(sheet, template_row, index, 15)
            _write_summary_row(sheet, index, month, values, materialize_calculations=materialize_calculations)
        workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        Path(output).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()
    return Path(output)


def _write_summary_row(
    sheet: Any,
    row: int,
    month: str,
    values: NankeValues,
    *,
    materialize_calculations: bool = False,
) -> None:
    year, month_number = map(int, month.split("-"))
    raw_values: list[Any] = [
        f"电表{values.meter_number}",
        f"{year}年{month_number}月",
        float(values.prior_reading),
        float(values.current_reading),
        float(values.multiplier),
        float(values.total_generation),
        float(values.grid_energy),
        float(values.base_station_energy),
        float(values.base_station_price),
        float(values.grid_price),
    ]
    for column, value in enumerate(raw_values, start=1):
        sheet.cell(row, column).value = value
    sheet.cell(row, 11).value = float(values.grid_formula_fee) if materialize_calculations else f"=J{row}*G{row}"
    sheet.cell(row, 12).value = float(values.school_energy) if materialize_calculations else f"=F{row}-G{row}-H{row}"
    sheet.cell(row, 13).value = float(values.school_price) if materialize_calculations else f"={_number(DISCOUNT_RATE)}*I{row}"
    sheet.cell(row, 14).value = float(values.school_fee) if materialize_calculations else f"=L{row}*M{row}"
    sheet.cell(row, 15).value = float(values.total_income) if materialize_calculations else f"=N{row}+K{row}"


def _copy_row(sheet: Any, source_row: int, target_row: int, max_column: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _combine_sources(task: dict[str, Any]) -> tuple[NankeValues, list[str]]:
    sources = task["sources"]
    generation = sources["generation_statement"]["recognized"]
    fee = sources["grid_fee_settlement"]["recognized"]
    energy = sources["grid_energy_settlement"]["recognized"]
    base = sources["base_station_notice"]["recognized"]
    if _decimal(fee["grid_energy"], "电费结算电量") != _decimal(energy["grid_energy"], "电量结算合计"):
        raise NankeError("电费结算单与电量结算单的上网电量不一致，请核对材料。")
    values = NankeValues(
        month=task["month"],
        meter_number=str(generation["meter_number"]),
        prior_reading=_decimal(generation["prior_reading"], "上期示数"),
        current_reading=_decimal(generation["current_reading"], "本期示数"),
        multiplier=_decimal(generation["multiplier"], "倍率"),
        total_generation=_decimal(generation["total_generation"], "总发电量"),
        grid_energy=_decimal(fee["grid_energy"], "上网电量"),
        grid_price=_decimal(fee["grid_price"], "上网电价"),
        grid_fee=_decimal(fee["grid_fee"], "上网电费"),
        base_station_energy=_decimal(base["base_station_energy"], "基站用电量"),
        base_station_price=_decimal(base["base_station_price"], "基站单价"),
        base_station_fee=_decimal(base["base_station_fee"], "基站电费"),
    )
    if values.meter_number != METER_NUMBER:
        raise NankeError(f"发电单表号不是南科大受控表号 {METER_NUMBER}。")
    if values.calculated_generation != values.total_generation:
        raise NankeError("发电单总发电量与（本期示数-上期示数）×倍率四舍五入结果不一致。")
    if values.school_energy < 0:
        raise NankeError("总发电量不足以扣除上网电量和基站用电量，请核对材料。")
    warnings: list[str] = []
    if abs(values.grid_formula_fee - values.grid_fee) > Decimal("0.01"):
        warnings.append("上网电量×电价与结算单含税金额存在超过0.01元差异。")
    if abs(values.base_station_energy * values.base_station_price - values.base_station_fee) > Decimal("0.01"):
        warnings.append("基站用电量×单价与通知书请款金额存在超过0.01元差异。")
    return values, warnings


def _generation_values(raw: OcrRawResult) -> dict[str, str]:
    for row in _ocr_rows(raw):
        texts = [block.text for block in row]
        if not any(METER_NUMBER in _compact(text) for text in texts):
            continue
        numbers: list[Decimal] = []
        seen_meter = False
        for text in texts:
            if METER_NUMBER in _compact(text):
                seen_meter = True
                continue
            if seen_meter:
                numbers.extend(_numbers(text))
        if len(numbers) >= 4:
            return {
                "prior_reading": _number(numbers[0]),
                "current_reading": _number(numbers[1]),
                "multiplier": _number(numbers[2]),
                "total_generation": _number(numbers[3]),
            }
    # OCR may split one table row into nearby lines; use the known meter as anchor.
    blocks = [block.text for page in raw.pages for block in page.blocks]
    anchor = next((index for index, text in enumerate(blocks) if METER_NUMBER in _compact(text)), -1)
    values = [number for text in blocks[anchor + 1 : anchor + 12] for number in _numbers(text)] if anchor >= 0 else []
    if len(values) < 4:
        raise NankeError("发电单未能完整识别上期示数、本期示数、倍率和发电量。")
    return {
        "prior_reading": _number(values[0]),
        "current_reading": _number(values[1]),
        "multiplier": _number(values[2]),
        "total_generation": _number(values[3]),
    }


def _market_fee_values(raw: OcrRawResult) -> tuple[Decimal, Decimal, Decimal]:
    rows = _ocr_rows(raw)
    for index, row in enumerate(rows):
        if "市场化电费" not in _compact("".join(block.text for block in row)):
            continue
        numbers = [value for nearby in rows[index : index + 2] for block in nearby for value in _numbers(block.text)]
        candidates = [value for value in numbers if value > 0]
        for position in range(len(candidates) - 2):
            energy, price, fee = candidates[position : position + 3]
            if energy >= 1 and Decimal("0") < price < Decimal("2"):
                if abs(energy * price - fee) <= Decimal("0.05"):
                    return energy, price, fee
    raise NankeError("电费结算单未能完整识别上网电量、电价和电费。")


def _energy_total(raw: OcrRawResult) -> Decimal:
    rows = _ocr_rows(raw)
    for row in reversed(rows):
        text = _compact("".join(block.text for block in row))
        if "合计" in text:
            values = [value for block in row for value in _numbers(block.text)]
            if values:
                return values[-1]
    raise NankeError("电量结算单未能识别月度合计电量。")


def _base_station_values(text: str) -> tuple[Decimal, Decimal, Decimal]:
    compact = _compact(text).replace(",", "")
    energy = _search_decimal(compact, r"(?:用电度数|用电量)(?:为|[:：])?(\d+(?:\.\d+)?)")
    price = _search_decimal(compact, r"(?:电费)?单价(?:为|[:：])?(\d+(?:\.\d+)?)")
    fee = _search_decimal(compact, r"(?:请款金额|金额)(?:为|[:：])?(\d+(?:\.\d+)?)")
    if energy is None or price is None or fee is None:
        raise NankeError("基站缴费通知书未能完整识别用电量、单价和请款金额。")
    return energy, price, fee


def _values_from_dict(raw: Any, *, expected_month: str) -> NankeValues:
    if not isinstance(raw, dict):
        raise NankeError("南科大任务缺少已确认数据。")
    month = _validate_month(str(raw.get("month") or expected_month))
    if month != expected_month:
        raise NankeError("南科大任务月份不一致。")
    return NankeValues(
        month=month,
        meter_number=str(raw.get("meter_number") or ""),
        prior_reading=_decimal(raw.get("prior_reading"), "上期示数"),
        current_reading=_decimal(raw.get("current_reading"), "本期示数"),
        multiplier=_decimal(raw.get("multiplier"), "倍率"),
        total_generation=_decimal(raw.get("total_generation"), "总发电量"),
        grid_energy=_decimal(raw.get("grid_energy"), "上网电量"),
        grid_price=_decimal(raw.get("grid_price"), "上网电价"),
        grid_fee=_decimal(raw.get("grid_fee"), "上网电费"),
        base_station_energy=_decimal(raw.get("base_station_energy"), "基站用电量"),
        base_station_price=_decimal(raw.get("base_station_price"), "基站单价"),
        base_station_fee=_decimal(raw.get("base_station_fee"), "基站电费"),
    )


def _effective_records(runtime_root: Path | str) -> dict[str, NankeValues]:
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    result: dict[str, NankeValues] = {}
    for month, record in (series.get("months") or {}).items():
        values = record.get("values") if isinstance(record, dict) else None
        if isinstance(values, dict):
            result[month] = _values_from_dict(values, expected_month=month)
    return result


def _summary_materialization_enabled(runtime_root: Path | str) -> bool:
    series = _read_json(_series_path(runtime_root), default={})
    return series.get("summary_materialize_calculations") is True


def _promote_effective(
    runtime_root: Path | str,
    task: dict[str, Any],
    values: NankeValues,
    outputs: list[Path],
) -> None:
    target = _series_path(runtime_root)
    series = _read_json(target, default={"site_key": SITE_KEY, "months": {}})
    months = series.setdefault("months", {})
    current = months.get(values.month) if isinstance(months.get(values.month), dict) else {}
    history = list(current.get("history") or [])
    if current.get("task_id"):
        history.append({key: current.get(key) for key in ("task_id", "revision", "values", "outputs")})
    months[values.month] = {
        "task_id": task["task_id"],
        "revision": task["revision"],
        "values": values.to_dict(),
        "outputs": [str(output) for output in outputs],
        "history": history,
        "updated_at": _utc_now(),
    }
    _write_json(target, series)


def _month_from_text(text: str, expected_month: str) -> str:
    digits = re.sub(r"\D", "", text)
    if expected_month.replace("-", "") in digits:
        return expected_month
    matches = {
        f"{year}-{int(month):02d}"
        for year, month in re.findall(r"(20\d{2})\s*(?:[-./年])\s*(0?[1-9]|1[0-2])", text)
    }
    if expected_month in matches:
        return expected_month
    if len(matches) == 1:
        return next(iter(matches))
    raise NankeError("无法唯一识别文件账期。")


def _ocr_rows(raw: OcrRawResult, *, tolerance: float = 10) -> list[list[OcrBlock]]:
    result: list[list[OcrBlock]] = []
    for page in raw.pages:
        rows: list[list[OcrBlock]] = []
        for block in sorted(page.blocks, key=lambda item: ((item.bbox[1] + item.bbox[3]) / 2, item.bbox[0])):
            center = (block.bbox[1] + block.bbox[3]) / 2
            row = next(
                (
                    candidate
                    for candidate in rows
                    if abs(center - sum((item.bbox[1] + item.bbox[3]) / 2 for item in candidate) / len(candidate)) <= tolerance
                ),
                None,
            )
            if row is None:
                rows.append([block])
            else:
                row.append(block)
        result.extend(sorted((sorted(row, key=lambda item: item.bbox[0]) for row in rows), key=lambda row: row[0].bbox[1]))
    return result


def _numbers(text: str) -> list[Decimal]:
    values: list[Decimal] = []
    for item in re.findall(r"(?<!\d)-?\d+(?:\.\d+)?(?!\d)", text.replace(",", "")):
        try:
            values.append(Decimal(item))
        except InvalidOperation:
            continue
    return values


def _search_decimal(text: str, pattern: str) -> Decimal | None:
    match = re.search(pattern, text, re.IGNORECASE)
    return Decimal(match.group(1)) if match else None


def _source_pages(source: Path, output_dir: Path) -> list[RenderedPage]:
    if source.suffix.lower() == ".pdf":
        return render_pdf_to_images(source, output_dir)
    return [RenderedPage(source_file=source, page=1, image_path=source)]


def _validate_source(path: Path) -> None:
    if not path.is_file() or path.is_symlink() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise NankeError("南科大仅接收 JPG、PNG 和 PDF 原始材料。")
    if path.stat().st_size <= 0 or path.stat().st_size > MAX_SOURCE_BYTES:
        raise NankeError("文件为空或超过 30 MB，本次文件未登记。")
    header = path.read_bytes()[:12]
    suffix = path.suffix.lower()
    valid = (suffix == ".pdf" and header.startswith(b"%PDF-")) or (
        suffix in {".jpg", ".jpeg"} and header.startswith(b"\xff\xd8\xff")
    ) or (suffix == ".png" and header.startswith(b"\x89PNG\r\n\x1a\n"))
    if not valid:
        raise NankeError("文件扩展名与实际内容不一致，本次文件未登记。")


def _role_label(role: str) -> str:
    return {
        "generation_statement": "发电单",
        "grid_fee_settlement": "电费结算单",
        "grid_energy_settlement": "电量结算单",
        "base_station_notice": "基站缴费通知书",
    }.get(role, role)


def _load_task(runtime_root: Path | str, task_id: str) -> tuple[Path, dict[str, Any]]:
    root = _task_root(runtime_root, _safe_task_id(task_id))
    task_file = root / "task.json"
    if not task_file.is_file():
        raise NankeError(f"南科大任务不存在：{task_id}")
    task = _read_json(task_file)
    if task.get("business_profile") != BUSINESS_PROFILE:
        raise NankeError("任务不属于南科大专用流程。")
    return root, task


def _task_root(runtime_root: Path | str, task_id: str) -> Path:
    return Path(runtime_root) / "nanke_tasks" / task_id


def _series_path(runtime_root: Path | str) -> Path:
    return Path(runtime_root) / "nanke" / "series.json"


def _next_revision(runtime_root: Path | str, month: str) -> int:
    series = _read_json(_series_path(runtime_root), default={"months": {}})
    current = (series.get("months") or {}).get(month)
    return int(current.get("revision") or 0) + 1 if isinstance(current, dict) else 1


def _next_month_date(year: int, month: int) -> date:
    return date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)


def _chinese_month(month: int) -> str:
    return ("一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一", "十二")[month - 1]


def _validate_month(month: str) -> str:
    match = re.fullmatch(r"(20\d{2})-(0[1-9]|1[0-2])", str(month).strip())
    if not match:
        raise NankeError("月份必须使用 YYYY-MM 格式。")
    return match.group(0)


def _safe_task_id(task_id: str) -> str:
    value = str(task_id).strip()
    if not re.fullmatch(r"[A-Za-z0-9_.-]{1,128}", value):
        raise NankeError("任务编号格式不合法。")
    return value


def _decimal(value: Any, field: str) -> Decimal:
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise NankeError(f"{field}不是有效数字。") from exc
    if not result.is_finite():
        raise NankeError(f"{field}不是有限数字。")
    return result


def _number(value: Decimal, *, places: int | None = None) -> str:
    if places is not None:
        value = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _money(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), "f")


def _compact(text: str) -> str:
    return re.sub(r"\s+", "", str(text)).replace("（", "(").replace("）", ")")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_json(path: Path, *, default: Any = None) -> Any:
    if not path.is_file():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
