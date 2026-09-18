"""Template profile registry and dynamic site registration."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import unicodedata
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterator

from submit_flow_agent.config import ConfigError, DEFAULT_CONFIG_PATH, load_site_config
from submit_flow_agent.task_store import utc_now_iso


DEFAULT_PROFILE_PATH = Path("config/template_profiles.json")
RESERVED_METADATA_KEYS = {
    "site_key",
    "template_dir",
    "detail_workbook",
    "summary_workbook",
    "template_hashes",
    "command",
    "shell",
    "cmd",
    "script",
    "python",
    "sql",
    "path",
}
PRICING_METADATA_KEYS = {
    "base_unit_price",
    "discount_rate",
    "effective_unit_price",
    "effective_school_price",
    "school_base_price",
    "school_discount_rate",
}
PRICING_CONFIRMATION_FIELDS = (
    {
        "key": "discount_rate",
        "type": "decimal",
        "label_zh": "折扣率",
        "required": True,
        "source_strategy": "pricing_confirmation_required",
        "semantic": "dynamic_header_only",
    },
    {
        "key": "effective_unit_price",
        "type": "decimal",
        "label_zh": "结算单价",
        "required": True,
        "source_strategy": "pricing_confirmation_required",
        "semantic": "direct_calculation_input",
    },
)
METER_TOPOLOGY_PROFILES = {"single_meter_school", "dual_meter_school"}



class TemplateProfileError(RuntimeError):
    """Raised when template profile registration is unsafe or invalid."""


def list_template_profiles(profiles_path: Path | str = DEFAULT_PROFILE_PATH) -> dict[str, list[dict[str, Any]]]:
    profiles = _load_profiles(profiles_path)
    return {
        "template_profiles": [
            {
                "profile_key": profile["profile_key"],
                "display_name": profile.get("display_name") or profile["profile_key"],
                "default": bool(profile.get("default")),
                "legacy_read_only": bool(profile.get("legacy_read_only")),
                "detail_profile_key": profile.get("detail_profile_key"),
                "summary_profile_key": profile.get("summary_profile_key"),
                "detail_profile": profile["detail_profile"],
                "summary_profile": profile["summary_profile"],
                "detail_start_row": profile.get("detail_start_row"),
                "summary_start_row": profile.get("summary_start_row"),
                "summary_end_column": profile.get("summary_end_column"),
                "summary_merge_columns": profile.get("summary_merge_columns") or [],
                "summary_site_name_cell": profile.get("summary_site_name_cell"),
                "summary_group_spacer_rows": profile.get("summary_group_spacer_rows", 0),
                "summary_month_cell_alignment": profile.get("summary_month_cell_alignment"),
                "validation_not_applicable_fields": profile.get("validation_not_applicable_fields") or [],
                "meter_topology_profile": profile.get("meter_topology_profile"),
                "meter_topology": profile.get("meter_topology") or {},
                "metadata_fields": _metadata_field_specs(profile),
                "pricing_confirmation_fields": list(PRICING_CONFIRMATION_FIELDS),
                "defaults": _non_pricing_defaults(profile),
                "output_rules": dict(profile.get("output_rules") or {}),
                "workbooks": _public_workbook_contract(profile),
            }
            for profile in profiles
        ]
    }


def auto_register_site_from_profile(
    *,
    site_name: str,
    profile_key: str | None = None,
    meter_topology_profile: str | None = None,
    aliases: list[str] | tuple[str, ...] = (),
    metadata: dict[str, Any] | None = None,
    config_path: Path | str = DEFAULT_CONFIG_PATH,
    profiles_path: Path | str = DEFAULT_PROFILE_PATH,
) -> dict[str, Any]:
    clean_metadata = _metadata_object(metadata)
    clean_site_name = _clean_text(site_name, "site_name")
    clean_aliases = _clean_aliases(aliases)
    existing_site = _existing_registered_site(config_path, clean_site_name, clean_aliases)
    if existing_site is not None:
        return {
            "status": "reused",
            "registered": False,
            "idempotent": True,
            "site_key": existing_site.site_key,
            "site_name": existing_site.site_name,
            "profile_key": existing_site.template_profile_key,
            "template_profile_key": existing_site.template_profile_key,
            "profile_selection": "existing_site",
            "aliases": list(dict.fromkeys((existing_site.site_name, *existing_site.aliases))),
        }

    metadata_topology = clean_metadata.pop("meter_topology_profile", None)
    selected_topology = meter_topology_profile or (str(metadata_topology).strip() if metadata_topology is not None else None)
    if profile_key is None and selected_topology:
        profile_key = _profile_key_for_meter_topology(selected_topology, profiles_path)
    if profile_key is None:
        return {
            "status": "waiting_for_meter_topology",
            "registered": False,
            "idempotent": False,
            "site_name": clean_site_name,
            "profile_key": None,
            "profile_selection": "pending_meter_topology",
            "error_code": "METER_TOPOLOGY_REQUIRED",
            "message": "Select the site meter topology before registration.",
            "meter_topology_profiles": sorted(METER_TOPOLOGY_PROFILES),
        }
    profile = _profile_by_key(profile_key, profiles_path)
    _reject_disallowed_metadata(profile, clean_metadata)

    missing = _missing_required_metadata(profile, clean_metadata)
    if missing:
        return {
            "status": "waiting_for_required_metadata",
            "registered": False,
            "idempotent": False,
            "site_name": clean_site_name,
            "profile_key": profile["profile_key"],
            "profile_selection": "meter_topology" if selected_topology else "explicit",
            "missing_metadata": missing,
        }

    try:
        result = register_site_from_profile(
            site_name=clean_site_name,
            profile_key=profile["profile_key"],
            aliases=list(clean_aliases),
            metadata=clean_metadata,
            config_path=config_path,
            profiles_path=profiles_path,
        )
    except TemplateProfileError as exc:
        message = str(exc)
        if _is_identity_conflict(message):
            return {
                "status": "conflict",
                "registered": False,
                "idempotent": False,
                "site_name": clean_site_name,
                "profile_key": profile["profile_key"],
                "profile_selection": "meter_topology" if selected_topology else "explicit",
                "error_code": "SITE_IDENTITY_CONFLICT",
                "message": message,
            }
        if _is_metadata_need_information(message):
            return {
                "status": "need_information",
                "registered": False,
                "idempotent": False,
                "site_name": clean_site_name,
                "profile_key": profile["profile_key"],
                "profile_selection": "meter_topology" if selected_topology else "explicit",
                "error_code": "METADATA_NEEDS_INFORMATION",
                "message": message,
                "metadata_fields": _metadata_field_specs(profile),
            }
        raise

    return {
        **result,
        "status": "registered" if result.get("registered") else "reused",
        "profile_selection": "meter_topology" if selected_topology else "explicit",
    }

def register_site_from_profile(
    *,
    site_name: str,
    profile_key: str,
    aliases: list[str] | tuple[str, ...] = (),
    metadata: dict[str, Any] | None = None,
    config_path: Path | str = DEFAULT_CONFIG_PATH,
    profiles_path: Path | str = DEFAULT_PROFILE_PATH,
) -> dict[str, Any]:
    profile = _profile_by_key(profile_key, profiles_path)
    clean_site_name = _clean_text(site_name, "site_name")
    clean_aliases = _clean_aliases(aliases)
    clean_metadata = _validate_metadata(profile, metadata or {})
    site_key = _site_key_for(clean_site_name)
    site = _site_from_profile(site_key, clean_site_name, clean_aliases, profile, clean_metadata)
    signature = _site_signature(site)
    site["dynamic_profile_key"] = profile["profile_key"]
    site["dynamic_site_signature"] = signature

    config_file = Path(config_path)
    with _config_lock(config_file):
        payload = _read_config_payload(config_file)
        sites = payload.setdefault("sites", [])
        if not isinstance(sites, list):
            raise TemplateProfileError(f"{config_file} must contain a sites list.")
        existing = _find_existing_dynamic_site(sites, site_key, clean_site_name, clean_aliases)
        if existing is not None:
            if existing.get("dynamic_site_signature") == signature:
                _ensure_registration_audit(config_file.parent, existing, profile)
            existing_profile_key = existing.get("template_profile_key") or profile["profile_key"]
            return {
                "registered": False,
                "idempotent": True,
                "site_key": existing["site_key"],
                "site_name": existing["site_name"],
                "profile_key": existing_profile_key,
                "template_profile_key": existing_profile_key,
                "reused_existing_site": True,
            }

        _validate_new_site_identity(site, sites)
        _validate_parseable_site(site, config_file)
        templates_root = config_file.parent / "templates"
        _ensure_safe_template_root(templates_root)
        target_template_dir = templates_root / site_key
        if target_template_dir.exists() or target_template_dir.is_symlink():
            raise TemplateProfileError(f"Template target already exists for generated site_key: {site_key}")

        payload["sites"] = [*sites, site]
        created_template_dir = False
        try:
            with tempfile.TemporaryDirectory(dir=config_file.parent) as tmpdir:
                tmp_root = Path(tmpdir)
                tmp_config = tmp_root / config_file.name
                tmp_template_dir = tmp_root / "templates" / site_key
                tmp_template_dir.mkdir(parents=True)
                _write_profile_workbooks(profile, site, clean_metadata, tmp_template_dir, Path(profiles_path))
                tmp_config.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
                os.replace(tmp_template_dir, target_template_dir)
                created_template_dir = True
                os.replace(tmp_config, config_file)
        except Exception as exc:
            if created_template_dir:
                _rollback_created_template_dir(target_template_dir, templates_root)
            if isinstance(exc, TemplateProfileError):
                raise
            raise TemplateProfileError(
                "Dynamic site registration failed before publish; no config entry was committed."
            ) from exc

        _ensure_registration_audit(config_file.parent, site, profile)
        return {
            "registered": True,
            "idempotent": False,
            "site_key": site_key,
            "site_name": clean_site_name,
            "profile_key": profile["profile_key"],
            "config_path": str(config_file),
        }


def _load_profiles(profiles_path: Path | str) -> list[dict[str, Any]]:
    path = Path(profiles_path)
    if not path.is_file():
        raise TemplateProfileError(f"Template profile registry does not exist: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise TemplateProfileError(f"Template profile registry is not valid JSON: {path}") from exc
    profiles = payload.get("template_profiles") if isinstance(payload, dict) else None
    role_profiles = _load_role_profiles(payload) if isinstance(payload, dict) else {"detail": {}, "summary": {}}
    if not isinstance(profiles, list) or not profiles:
        raise TemplateProfileError(f"{path} must contain a non-empty template_profiles list.")
    validated: list[dict[str, Any]] = []
    seen: set[str] = set()
    default_count = 0
    for raw_profile in profiles:
        if not isinstance(raw_profile, dict):
            raise TemplateProfileError("Each template profile must be an object.")
        profile = _compose_template_profile(raw_profile, role_profiles)
        key = _clean_key(profile.get("profile_key"), "profile_key")
        if key in seen:
            raise TemplateProfileError(f"Duplicate template profile key: {key}")
        seen.add(key)
        if bool(profile.get("default")):
            default_count += 1
        for field in ("detail_profile", "summary_profile", "detail_workbook", "summary_workbook", "detail_sheet", "summary_sheet"):
            _clean_text(profile.get(field), field)
        fields = _metadata_field_specs(profile)
        field_keys = [field["key"] for field in fields]
        if len(field_keys) != len(set(field_keys)):
            raise TemplateProfileError(f"Duplicate metadata field in profile {key}.")
        _validate_profile_workbook_specs(profile)
        _validate_meter_topology_profile(profile)
        validated.append({**profile, "profile_key": key, "metadata_fields": fields})
    if default_count > 1:
        raise TemplateProfileError("Template profile registry must declare at most one default profile.")
    return validated


def _load_role_profiles(payload: dict[str, Any]) -> dict[str, dict[str, dict[str, Any]]]:
    raw_roles = payload.get("role_profiles") or {}
    if not isinstance(raw_roles, dict):
        raise TemplateProfileError("role_profiles must be an object when declared.")
    registry: dict[str, dict[str, dict[str, Any]]] = {"detail": {}, "summary": {}}
    for role in ("detail", "summary"):
        entries = raw_roles.get(role) or []
        if not isinstance(entries, list):
            raise TemplateProfileError(f"role_profiles.{role} must be a list.")
        for entry in entries:
            if not isinstance(entry, dict):
                raise TemplateProfileError(f"Each role_profiles.{role} entry must be an object.")
            key = _clean_key(entry.get("profile_key"), f"{role}_profile_key")
            if key in registry[role]:
                raise TemplateProfileError(f"Duplicate {role} role profile key: {key}")
            workbook = entry.get("workbook")
            if not isinstance(workbook, dict):
                raise TemplateProfileError(f"Role profile {key} must declare workbook.")
            _clean_text(workbook.get("workbook_name"), f"{key}.workbook_name")
            _clean_text(workbook.get("sheet"), f"{key}.sheet")
            _clean_text(workbook.get("path"), f"{key}.path")
            registry[role][key] = {**entry, "profile_key": key}
    return registry


def _compose_template_profile(
    profile: dict[str, Any],
    role_profiles: dict[str, dict[str, dict[str, Any]]],
) -> dict[str, Any]:
    detail_key = profile.get("detail_profile_key")
    summary_key = profile.get("summary_profile_key")
    if detail_key is None and summary_key is None:
        return dict(profile)
    if detail_key is None or summary_key is None:
        raise TemplateProfileError("Composable template profiles must declare both detail_profile_key and summary_profile_key.")
    detail = _role_profile_by_key(role_profiles, "detail", detail_key)
    summary = _role_profile_by_key(role_profiles, "summary", summary_key)
    composed = dict(profile)
    composed["detail_profile_key"] = detail["profile_key"]
    composed["summary_profile_key"] = summary["profile_key"]
    composed["detail_profile"] = detail["writer_profile"]
    composed["summary_profile"] = summary["writer_profile"]
    composed["detail_workbook"] = detail["workbook"]["workbook_name"]
    composed["detail_sheet"] = detail["workbook"]["sheet"]
    composed["summary_workbook"] = summary["workbook"]["workbook_name"]
    composed["summary_sheet"] = summary["workbook"]["sheet"]
    composed["detail_start_row"] = detail.get("start_row", composed.get("detail_start_row", 5))
    composed["summary_start_row"] = summary.get("start_row", composed.get("summary_start_row", 4))
    composed["detail_company_cell"] = detail.get("company_cell", composed.get("detail_company_cell"))
    composed["detail_statement_date_cell"] = detail.get("statement_date_cell", composed.get("detail_statement_date_cell"))
    composed["summary_end_column"] = summary.get("end_column", composed.get("summary_end_column", 13))
    composed["summary_merge_columns"] = summary.get("merge_columns", composed.get("summary_merge_columns", ["I", "M"]))
    composed["summary_site_name_cell"] = summary.get("site_name_cell", composed.get("summary_site_name_cell"))
    layout = summary.get("layout") if isinstance(summary.get("layout"), dict) else {}
    composed["summary_group_spacer_rows"] = summary.get(
        "group_spacer_rows",
        layout.get("group_spacer_rows", composed.get("summary_group_spacer_rows", 0)),
    )
    composed["summary_month_cell_alignment"] = summary.get(
        "month_cell_alignment",
        layout.get("month_cell_alignment", composed.get("summary_month_cell_alignment")),
    )
    composed["validation_not_applicable_fields"] = list(
        dict.fromkeys(
            [
                *list(summary.get("validation_not_applicable_fields") or []),
                *list(profile.get("validation_not_applicable_fields") or []),
            ]
        )
    )
    composed["output_rules"] = {
        **dict(summary.get("output_rules") or {}),
        **dict(detail.get("output_rules") or {}),
        **dict(profile.get("output_rules") or {}),
    }
    composed["meter_topology"] = {
        **dict(detail.get("meter_topology") or {}),
        **dict(summary.get("meter_topology") or {}),
        **dict(profile.get("meter_topology") or {}),
    }
    composed["workbook_masters"] = {
        "detail": _workbook_master_from_role(detail),
        "summary": _workbook_master_from_role(summary),
    }
    return composed


def _role_profile_by_key(
    role_profiles: dict[str, dict[str, dict[str, Any]]],
    role: str,
    key_value: Any,
) -> dict[str, Any]:
    key = _clean_key(key_value, f"{role}_profile_key")
    profile = role_profiles.get(role, {}).get(key)
    if profile is None:
        raise TemplateProfileError(f"UNSUPPORTED_{role.upper()}_PROFILE profile_key={key}")
    return profile


def _workbook_master_from_role(role_profile: dict[str, Any]) -> dict[str, Any]:
    workbook = dict(role_profile["workbook"])
    return {
        "path": workbook["path"],
        "sha256": workbook["sha256"],
        "workbook_name": workbook["workbook_name"],
        "sheet": workbook["sheet"],
        "layout": dict(role_profile.get("layout") or {}),
        "allowed_parameter_cells": dict(role_profile.get("allowed_parameter_cells") or {}),
    }


def _profile_by_key(profile_key: str, profiles_path: Path | str) -> dict[str, Any]:
    key = _clean_key(profile_key, "profile_key")
    for profile in _load_profiles(profiles_path):
        if profile["profile_key"] == key:
            return profile
    raise TemplateProfileError(f"UNSUPPORTED_PROFILE profile_key={key}")


def _profile_key_for_meter_topology(meter_topology_profile: str, profiles_path: Path | str) -> str:
    topology = _clean_key(meter_topology_profile, "meter_topology_profile")
    if topology not in METER_TOPOLOGY_PROFILES:
        raise TemplateProfileError(f"UNSUPPORTED_METER_TOPOLOGY meter_topology_profile={topology}")
    matches = [
        profile["profile_key"]
        for profile in _load_profiles(profiles_path)
        if profile.get("meter_topology_profile") == topology and not profile.get("legacy_read_only")
    ]
    if len(matches) != 1:
        raise TemplateProfileError(f"meter_topology_profile={topology} must map to exactly one active template profile.")
    return matches[0]


def _default_profile(profiles_path: Path | str) -> dict[str, Any]:
    defaults = [profile for profile in _load_profiles(profiles_path) if profile.get("default") is True]
    if len(defaults) != 1:
        raise TemplateProfileError("Template profile registry must declare exactly one default profile.")
    return defaults[0]


def _validate_metadata(profile: dict[str, Any], metadata: dict[str, Any]) -> dict[str, Any]:
    metadata = _metadata_object(metadata)
    fields = _metadata_field_specs(profile)
    _reject_disallowed_metadata(profile, metadata)
    required = {field["key"] for field in fields if field.get("required") is True}
    merged = {**_non_pricing_defaults(profile), **metadata}
    missing = [key for key in sorted(required) if _blank(merged.get(key))]
    if missing:
        raise TemplateProfileError(f"Missing required metadata: {', '.join(missing)}")
    meters = merged.get("meters")
    if not isinstance(meters, list) or not meters:
        raise TemplateProfileError("metadata.meters must be a non-empty list.")
    meters = _normalize_metadata_meters(profile, meters)
    merged["meters"] = meters
    topology = profile.get("meter_topology") or {}
    if isinstance(topology, dict):
        detail_meters = [meter for meter in meters if bool(meter.get("detail_enabled"))]
        if topology.get("supports_multiple_meters") is False and (len(detail_meters) != 1 or len(meters) != 1):
            raise TemplateProfileError("Profile requires exactly one meter for this layout.")
        expected_detail_count = topology.get("expected_detail_meter_count")
        if expected_detail_count is not None and len(detail_meters) != int(expected_detail_count):
            raise TemplateProfileError(
                f"Profile requires exactly {expected_detail_count} detail-enabled meter rows for this layout."
            )
        expected_summary_count = topology.get("expected_summary_meter_count")
        if expected_summary_count is not None and len(meters) != int(expected_summary_count):
            raise TemplateProfileError(f"Profile requires exactly {expected_summary_count} summary meter rows for this layout.")
    return merged


def _existing_registered_site(config_path: Path | str, site_name: str, aliases: tuple[str, ...]) -> Any | None:
    for label in (site_name, *aliases):
        try:
            return load_site_config(site_key=label, config_path=config_path)
        except ConfigError as exc:
            if "was not found" in str(exc):
                continue
            raise
    return None


def _normalize_metadata_meters(profile: dict[str, Any], meters: list[Any]) -> list[dict[str, Any]]:
    topology = profile.get("meter_topology") if isinstance(profile.get("meter_topology"), dict) else {}
    output_rules = profile.get("output_rules") if isinstance(profile.get("output_rules"), dict) else {}
    expected_detail_count = _optional_positive_int(topology.get("expected_detail_meter_count"), "expected_detail_meter_count")
    recipient_meter_no = str(output_rules.get("grid_export_recipient_meter_no") or "").strip()
    normalized: list[dict[str, Any]] = []
    for index, meter in enumerate(meters, start=1):
        if not isinstance(meter, dict):
            raise TemplateProfileError("Each meter metadata entry must be an object.")
        for field in ("meter_no", "multiplier"):
            if _blank(meter.get(field)):
                raise TemplateProfileError(f"meter metadata missing {field}.")
        meter_no = _clean_text(meter.get("meter_no"), "meter_no")
        normalized.append(
            {
                "meter_no": meter_no,
                "multiplier": str(meter.get("multiplier")).strip(),
                "role": _metadata_meter_role(meter, index),
                "detail_enabled": _metadata_meter_detail_enabled(meter, index, len(meters), expected_detail_count, topology),
                "grid_export_allocation": _metadata_meter_grid_allocation(
                    meter,
                    index,
                    len(meters),
                    meter_no,
                    recipient_meter_no,
                    expected_detail_count,
                ),
            }
        )
    return normalized


def _metadata_meter_role(meter: dict[str, Any], index: int) -> str:
    if not _blank(meter.get("role")):
        return str(meter["role"]).strip()
    if index == 1:
        return "primary"
    if index == 2:
        return "secondary"
    return f"meter_{index}"


def _metadata_meter_detail_enabled(
    meter: dict[str, Any],
    index: int,
    meter_count: int,
    expected_detail_count: int | None,
    topology: dict[str, Any],
) -> bool:
    if "detail_enabled" in meter:
        return _metadata_bool(meter.get("detail_enabled"), "detail_enabled")
    if topology.get("supports_multiple_meters") is False or meter_count == 1:
        return True
    if expected_detail_count is not None:
        return index <= expected_detail_count
    return index == 1


def _metadata_meter_grid_allocation(
    meter: dict[str, Any],
    index: int,
    meter_count: int,
    meter_no: str,
    recipient_meter_no: str,
    expected_detail_count: int | None,
) -> str:
    if not _blank(meter.get("grid_export_allocation")):
        allocation = str(meter["grid_export_allocation"]).strip()
        if allocation not in {"site_total", "all", "none"}:
            raise TemplateProfileError(f"Unsupported meter grid_export_allocation: {allocation}")
        return allocation
    if recipient_meter_no:
        return "site_total" if meter_no == recipient_meter_no else "none"
    if meter_count == 1:
        return "site_total"
    if expected_detail_count is not None and expected_detail_count < meter_count:
        return "site_total" if index == 1 else "none"
    return "site_total" if index == meter_count else "none"


def _metadata_bool(value: Any, field: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        text = value.strip().lower()
        if text in {"true", "1", "yes"}:
            return True
        if text in {"false", "0", "no"}:
            return False
    raise TemplateProfileError(f"meter metadata {field} must be boolean.")


def _optional_positive_int(value: Any, field: str) -> int | None:
    if value is None:
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise TemplateProfileError(f"Profile meter_topology.{field} must be an integer.") from exc
    if parsed <= 0:
        raise TemplateProfileError(f"Profile meter_topology.{field} must be positive.")
    return parsed

def _metadata_object(metadata: dict[str, Any] | None) -> dict[str, Any]:
    if metadata is None:
        return {}
    if not isinstance(metadata, dict):
        raise TemplateProfileError("metadata must be a JSON object.")
    return dict(metadata)


def _reject_disallowed_metadata(profile: dict[str, Any], metadata: dict[str, Any]) -> None:
    allowed = {field["key"] for field in _metadata_field_specs(profile)}
    for key in metadata:
        if key in RESERVED_METADATA_KEYS or key not in allowed:
            raise TemplateProfileError(f"Metadata field is not allowed by profile: {key}")


def _missing_required_metadata(profile: dict[str, Any], metadata: dict[str, Any]) -> list[dict[str, Any]]:
    merged = {**_non_pricing_defaults(profile), **metadata}
    missing: list[dict[str, Any]] = []
    for field in _metadata_field_specs(profile):
        if field.get("required") is True and not _metadata_field_has_value(field, merged.get(field["key"])):
            missing.append(field)
    return missing


def _metadata_field_has_value(field: dict[str, Any], value: Any) -> bool:
    if str(field.get("type") or "") == "meter_list":
        if not isinstance(value, list) or not value:
            return False
        for meter in value:
            if not isinstance(meter, dict):
                return False
            if _blank(meter.get("meter_no")) or _blank(meter.get("multiplier")):
                return False
        return True
    if isinstance(value, (list, dict)):
        return bool(value)
    return not _blank(value)


def _is_identity_conflict(message: str) -> bool:
    return "identity conflict" in message.casefold() or "already exists" in message.casefold()


def _is_metadata_need_information(message: str) -> bool:
    markers = (
        "missing required metadata",
        "metadata.meters must be a non-empty list",
        "meter metadata missing",
        "requires exactly one meter",
        "requires exactly",
    )
    folded = message.casefold()
    return any(marker in folded for marker in markers)


def _site_from_profile(
    site_key: str,
    site_name: str,
    aliases: tuple[str, ...],
    profile: dict[str, Any],
    metadata: dict[str, Any],
) -> dict[str, Any]:
    detail_workbook = _site_workbook_name(site_name, "用电统计表")
    summary_workbook = _site_workbook_name(site_name, "汇总表")
    return {
        "site_key": site_key,
        "site_name": site_name,
        "template_profile_key": profile["profile_key"],
        "detail_profile_key": profile.get("detail_profile_key"),
        "summary_profile_key": profile.get("summary_profile_key"),
        "meter_topology_profile": profile.get("meter_topology_profile"),
        "project_name": _clean_text(metadata.get("project_name"), "project_name"),
        "aliases": list(aliases),
        "meter_no": str(metadata["meters"][0]["meter_no"]),
        "company_name": _clean_text(metadata.get("company_name"), "company_name"),
        "multiplier": str(metadata["meters"][0]["multiplier"]),
        "meters": metadata["meters"],
        "school_base_price": "1",
        "school_discount_rate": "1",
        "effective_school_price": None,
        "pricing_required": True,
        "template_dir": f"templates/{site_key}",
        "detail_workbook": detail_workbook,
        "detail_sheet": profile["detail_sheet"],
        "summary_workbook": summary_workbook,
        "summary_sheet": profile["summary_sheet"],
        "detail_profile": profile["detail_profile"],
        "summary_profile": profile["summary_profile"],
        "detail_start_row": int(profile.get("detail_start_row") or 5),
        "summary_start_row": int(profile.get("summary_start_row") or 4),
        "detail_company_cell": profile.get("detail_company_cell"),
        "detail_statement_date_cell": profile.get("detail_statement_date_cell"),
        "summary_end_column": profile.get("summary_end_column") or 13,
        "summary_merge_columns": profile.get("summary_merge_columns") or ["I", "M"],
        "summary_site_name_cell": profile.get("summary_site_name_cell"),
        "summary_group_spacer_rows": profile.get("summary_group_spacer_rows", 0),
        "summary_month_cell_alignment": profile.get("summary_month_cell_alignment"),
        "validation_not_applicable_fields": profile.get("validation_not_applicable_fields") or [],
        "issuer_aliases": [],
    }


def _validate_parseable_site(site: dict[str, Any], config_file: Path) -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "sites.json"
        path.write_text(json.dumps({"sites": [site]}, ensure_ascii=False), encoding="utf-8")
        try:
            load_site_config(str(site["site_key"]), config_path=path)
        except (ConfigError, KeyError) as exc:
            raise TemplateProfileError(f"Generated site config is invalid: {exc}") from exc


def _validate_new_site_identity(site: dict[str, Any], existing_sites: list[Any]) -> None:
    new_values = _identity_values(site)
    new_meters = _meter_numbers(site)
    for raw in existing_sites:
        if not isinstance(raw, dict):
            continue
        for new_field, new_value in new_values:
            for old_field, old_value in _identity_values(raw):
                if new_value.casefold() == old_value.casefold():
                    raise TemplateProfileError(f"Site identity conflict: {new_field} conflicts with {old_field}.")
        duplicate_meters = new_meters & _meter_numbers(raw)
        if duplicate_meters:
            existing_name = str(raw.get("site_name") or raw.get("site_key") or "existing site").strip()
            raise TemplateProfileError(
                "Site identity conflict: meter_no "
                f"{sorted(duplicate_meters)[0]} is already assigned to {existing_name}."
            )


def _find_existing_dynamic_site(
    sites: list[Any], site_key: str, site_name: str, aliases: tuple[str, ...]
) -> dict[str, Any] | None:
    wanted = {site_key.casefold(), _collapse(site_name).casefold(), *(_collapse(alias).casefold() for alias in aliases)}
    for raw in sites:
        if not isinstance(raw, dict):
            continue
        current = {_collapse(value).casefold() for _, value in _identity_values(raw)}
        if wanted & current:
            return raw
    return None


def _identity_values(site: dict[str, Any]) -> list[tuple[str, str]]:
    values = [("site_key", str(site.get("site_key") or "")), ("site_name", str(site.get("site_name") or ""))]
    values.extend(("alias", str(alias)) for alias in (site.get("aliases") or []))
    return [(field, _collapse(value)) for field, value in values if _collapse(value)]


def _meter_numbers(site: dict[str, Any]) -> set[str]:
    values: set[str] = set()
    primary = str(site.get("meter_no") or "").strip()
    if primary:
        values.add(primary.casefold())
    meters = site.get("meters")
    if isinstance(meters, list):
        for meter in meters:
            if not isinstance(meter, dict):
                continue
            meter_no = str(meter.get("meter_no") or "").strip()
            if meter_no:
                values.add(meter_no.casefold())
    return values


def _read_config_payload(config_file: Path) -> dict[str, Any]:
    try:
        payload = json.loads(config_file.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise TemplateProfileError(f"Site config is not valid JSON: {config_file}") from exc
    if not isinstance(payload, dict):
        raise TemplateProfileError(f"{config_file} must contain a JSON object.")
    return payload


def _write_profile_workbooks(
    profile: dict[str, Any],
    site: dict[str, Any],
    metadata: dict[str, Any],
    target_dir: Path,
    profiles_path: Path,
) -> None:
    masters = profile.get("workbook_masters")
    if not isinstance(masters, dict):
        raise TemplateProfileError(f"Profile {profile['profile_key']} does not declare workbook_masters.")
    for role in ("detail", "summary"):
        spec = masters.get(role)
        if not isinstance(spec, dict):
            raise TemplateProfileError(f"Profile {profile['profile_key']} missing {role} workbook master.")
        target = target_dir / str(site[f"{role}_workbook"])
        source = _validated_master_path(spec, profiles_path.parent)
        shutil.copy2(source, target, follow_symlinks=False)
        _apply_parameter_cells(target, spec, site, metadata)


def _validated_master_path(spec: dict[str, Any], config_root: Path) -> Path:
    relative = spec.get("path")
    if not isinstance(relative, str) or not relative.strip():
        raise TemplateProfileError("Workbook master must declare an internal config path.")
    candidate = Path(relative)
    if candidate.is_absolute() or ".." in candidate.parts:
        raise TemplateProfileError("Workbook master path must stay inside config.")
    path = (config_root / candidate).resolve()
    try:
        path.relative_to(config_root.resolve())
    except ValueError as exc:
        raise TemplateProfileError("Workbook master path escapes config root.") from exc
    if path.is_symlink() or not path.is_file():
        raise TemplateProfileError(f"Workbook master is missing or unsafe: {relative}")
    expected = str(spec.get("sha256") or "").casefold()
    if len(expected) != 64:
        raise TemplateProfileError("Workbook master must declare a SHA-256 hash.")
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual.casefold() != expected:
        raise TemplateProfileError(f"Workbook master hash mismatch: {relative}")
    return path


def _apply_parameter_cells(target: Path, spec: dict[str, Any], site: dict[str, Any], metadata: dict[str, Any]) -> None:
    cells = spec.get("allowed_parameter_cells") or {}
    if not cells:
        return
    if not isinstance(cells, dict):
        raise TemplateProfileError("allowed_parameter_cells must be an object.")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TemplateProfileError("openpyxl is required to parameterize workbook masters.") from exc

    sheet_name = str(spec.get("sheet") or "")
    workbook = load_workbook(target)
    try:
        if sheet_name not in workbook.sheetnames:
            raise TemplateProfileError(f"Workbook master sheet is missing: {sheet_name}")
        worksheet = workbook[sheet_name]
        for cell, cell_spec in cells.items():
            if not isinstance(cell, str) or not cell.strip():
                raise TemplateProfileError("allowed_parameter_cells contains an invalid cell reference.")
            worksheet[cell] = _parameter_cell_value(cell_spec, site, metadata)
        workbook.save(target)
    finally:
        workbook.close()


def _parameter_cell_value(cell_spec: Any, site: dict[str, Any], metadata: dict[str, Any]) -> Any:
    if isinstance(cell_spec, str):
        source = cell_spec
        template = None
    elif isinstance(cell_spec, dict):
        source = cell_spec.get("source")
        template = cell_spec.get("format")
    else:
        raise TemplateProfileError("Parameter cell declaration must be a string or object.")
    if not isinstance(source, str) or not source:
        raise TemplateProfileError("Parameter cell declaration missing source.")
    if source.startswith("metadata."):
        key = source.split(".", 1)[1]
        value = metadata.get(key)
    elif source.startswith("site."):
        key = source.split(".", 1)[1]
        value = site.get(key)
    else:
        raise TemplateProfileError(f"Unsupported parameter cell source: {source}")
    if _blank(value):
        raise TemplateProfileError(f"Parameter cell source is blank: {source}")
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        raise TemplateProfileError(f"Parameter cell source looks like a formula: {source}")
    if template is None:
        return value
    if not isinstance(template, str) or "{value}" not in template:
        raise TemplateProfileError("Parameter cell format must contain {value}.")
    return template.replace("{value}", str(value))


def _append_registration_audit(config_root: Path, site: dict[str, Any], profile: dict[str, Any]) -> None:
    record = _registration_audit_record(site, profile)
    audit_path = config_root / "site_registration_audit.jsonl"
    with audit_path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def _ensure_registration_audit(config_root: Path, site: dict[str, Any], profile: dict[str, Any]) -> None:
    record = _registration_audit_record(site, profile)
    audit_path = config_root / "site_registration_audit.jsonl"
    if _audit_event_exists(audit_path, str(record["event_id"])):
        return
    try:
        _append_registration_audit(config_root, site, profile)
    except OSError as exc:
        raise TemplateProfileError(
            "Dynamic site registration was published but audit was not durably recorded; retry the same request to recover."
        ) from exc


def _registration_audit_record(site: dict[str, Any], profile: dict[str, Any]) -> dict[str, Any]:
    signature = str(site.get("dynamic_site_signature") or _site_signature(site))
    event_id = hashlib.sha256(f"dynamic_site_registered:{site['site_key']}:{signature}".encode("utf-8")).hexdigest()
    return {
        "event_id": event_id,
        "ts": utc_now_iso(),
        "event": "dynamic_site_registered",
        "site_key": site["site_key"],
        "site_name": site["site_name"],
        "profile_key": profile["profile_key"],
        "site_signature": signature,
    }


def _audit_event_exists(audit_path: Path, event_id: str) -> bool:
    if not audit_path.is_file() or audit_path.is_symlink():
        return False
    for line in audit_path.read_text(encoding="utf-8-sig").splitlines():
        if not line.strip():
            continue
        try:
            payload = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(payload, dict) and payload.get("event_id") == event_id:
            return True
    return False


def _non_pricing_defaults(profile: dict[str, Any]) -> dict[str, Any]:
    defaults = dict(profile.get("defaults") or {})
    return {key: value for key, value in defaults.items() if key not in PRICING_METADATA_KEYS}


def _metadata_field_specs(profile: dict[str, Any]) -> list[dict[str, Any]]:
    raw_fields = profile.get("metadata_fields")
    if raw_fields is None:
        raw_fields = [
            {"key": key, "type": "string", "label_zh": key, "required": True, "source_strategy": "operator_supplied"}
            for key in profile.get("required_metadata") or []
        ]
        raw_fields.extend(
            {"key": key, "type": "string", "label_zh": key, "required": False, "source_strategy": "operator_supplied"}
            for key in profile.get("optional_metadata") or []
        )
    if not isinstance(raw_fields, list) or not raw_fields:
        raise TemplateProfileError("metadata_fields must be a non-empty list.")
    defaults = _non_pricing_defaults(profile)
    fields: list[dict[str, Any]] = []
    for raw in raw_fields:
        if not isinstance(raw, dict):
            raise TemplateProfileError("Each metadata field descriptor must be an object.")
        key = _clean_key(raw.get("key"), "metadata field key")
        if key in PRICING_METADATA_KEYS:
            continue
        if key in RESERVED_METADATA_KEYS:
            raise TemplateProfileError(f"Reserved metadata field is not allowed: {key}")
        required = bool(raw.get("required"))
        strategy = str(raw.get("source_strategy") or "operator_supplied")
        if strategy not in {"operator_supplied", "operator_supplied_or_document_derived", "document_derived"}:
            raise TemplateProfileError(f"Unsupported metadata source_strategy for {key}.")
        default_source = "profile_default" if key in defaults else None
        fields.append(
            {
                "key": key,
                "type": str(raw.get("type") or "string"),
                "label_zh": str(raw.get("label_zh") or key),
                "required": required,
                "source_strategy": strategy,
                "default_source": default_source,
                "default": defaults.get(key),
            }
        )
    return fields


def _public_workbook_contract(profile: dict[str, Any]) -> dict[str, Any]:
    masters = profile["workbook_masters"]
    public: dict[str, Any] = {}
    for role in ("detail", "summary"):
        spec = masters[role]
        public[role] = {
            "workbook_name": spec["workbook_name"],
            "sheet": spec["sheet"],
            "layout": spec.get("layout") or {},
            "allowed_parameter_cells": spec.get("allowed_parameter_cells") or {},
        }
    return public


def _validate_profile_workbook_specs(profile: dict[str, Any]) -> None:
    masters = profile.get("workbook_masters")
    if not isinstance(masters, dict):
        raise TemplateProfileError(f"Profile {profile['profile_key']} must declare workbook_masters.")
    for role in ("detail", "summary"):
        spec = masters.get(role)
        if not isinstance(spec, dict):
            raise TemplateProfileError(f"Profile {profile['profile_key']} missing {role} workbook master.")
        workbook_name = _clean_text(spec.get("workbook_name"), f"{role}.workbook_name")
        if workbook_name != profile[f"{role}_workbook"]:
            raise TemplateProfileError(f"Profile {profile['profile_key']} {role} workbook name mismatch.")
        if _clean_text(spec.get("sheet"), f"{role}.sheet") != profile[f"{role}_sheet"]:
            raise TemplateProfileError(f"Profile {profile['profile_key']} {role} sheet mismatch.")
        _clean_text(spec.get("path"), f"{role}.path")
        sha = str(spec.get("sha256") or "")
        if len(sha) != 64 or any(char not in "0123456789abcdefABCDEF" for char in sha):
            raise TemplateProfileError(f"Profile {profile['profile_key']} {role} master hash must be SHA-256.")


def _validate_meter_topology_profile(profile: dict[str, Any]) -> None:
    value = profile.get("meter_topology_profile")
    if value is None:
        return
    topology = _clean_key(value, "meter_topology_profile")
    if topology not in METER_TOPOLOGY_PROFILES:
        raise TemplateProfileError(f"Unsupported meter_topology_profile: {topology}")
    profile["meter_topology_profile"] = topology


def _ensure_safe_template_root(templates_root: Path) -> None:
    if templates_root.is_symlink():
        raise TemplateProfileError("Template root must not be a symbolic link.")
    if templates_root.exists() and not templates_root.is_dir():
        raise TemplateProfileError("Template root must be a directory.")
    templates_root.mkdir(exist_ok=True)


def _rollback_created_template_dir(target_template_dir: Path, templates_root: Path) -> None:
    if not target_template_dir.exists() and not target_template_dir.is_symlink():
        return
    try:
        target_template_dir.resolve().relative_to(templates_root.resolve())
    except ValueError as exc:
        raise TemplateProfileError("Rollback target escapes template root.") from exc
    if target_template_dir.is_dir() and not target_template_dir.is_symlink():
        shutil.rmtree(target_template_dir)
    else:
        target_template_dir.unlink()


@contextmanager
def _config_lock(config_file: Path) -> Iterator[None]:
    config_file.parent.mkdir(parents=True, exist_ok=True)
    lock_path = config_file.with_suffix(config_file.suffix + ".lock")
    with lock_path.open("a+b") as handle:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"0")
            handle.flush()
        _lock_file(handle)
        try:
            yield
        finally:
            _unlock_file(handle)


def _lock_file(handle: Any) -> None:
    if os.name == "nt":
        import msvcrt

        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
        return
    import fcntl

    fcntl.flock(handle.fileno(), fcntl.LOCK_EX)


def _unlock_file(handle: Any) -> None:
    if os.name == "nt":
        import msvcrt

        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        return
    import fcntl

    fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _site_key_for(site_name: str) -> str:
    digest = hashlib.sha256(_collapse(site_name).casefold().encode("utf-8")).hexdigest()[:12]
    return f"site_{digest}"


def _site_workbook_name(site_name: str, role_label: str) -> str:
    clean_name = _clean_text(site_name, "site_name")
    if any(ord(character) < 32 or character in '<>:"/\\|?*' for character in clean_name):
        raise TemplateProfileError("site_name contains characters that are unsafe in workbook file names.")
    clean_name = clean_name.rstrip(" .")
    if not clean_name:
        raise TemplateProfileError("site_name does not produce a safe workbook file name.")
    return f"{clean_name}{role_label}.xlsx"


def _site_signature(site: dict[str, Any]) -> str:
    payload = {key: value for key, value in site.items() if key not in {"dynamic_site_signature"}}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _clean_aliases(aliases: list[str] | tuple[str, ...]) -> tuple[str, ...]:
    if not isinstance(aliases, (list, tuple)):
        raise TemplateProfileError("aliases must be a list of strings.")
    return tuple(dict.fromkeys(_clean_text(alias, "alias") for alias in aliases))


def _clean_key(value: Any, field: str) -> str:
    text = _clean_text(value, field)
    path = Path(text)
    if path.is_absolute() or ".." in path.parts or len(path.parts) != 1:
        raise TemplateProfileError(f"{field} must be a single safe key segment.")
    return text


def _clean_text(value: Any, field: str) -> str:
    if value is None or not str(value).strip():
        raise TemplateProfileError(f"{field} is required.")
    return str(value).strip()


def _collapse(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())











