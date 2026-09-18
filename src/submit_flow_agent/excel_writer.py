"""Excel writing for monthly output workbooks."""

from __future__ import annotations

import json
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from submit_flow_agent.config import load_site_config
from submit_flow_agent.normalizer import uses_formula_generation
from submit_flow_agent.pricing import canonicalize_discount_rate, pricing_snapshot_for_recognition, validate_task_pricing_snapshot

MONTH_NUMBER_FORMAT = 'yyyy"年"m"月"'


class ExcelWriteError(RuntimeError):
    """Raised when an Excel output workbook cannot be written."""


@dataclass(frozen=True)
class DetailExcelWriteResult:
    month: str
    output: Path
    workbook: str
    sheet: str
    written_cells: list[str]

    def to_dict(self) -> dict[str, object]:
        return {
            "month": self.month,
            "output": str(self.output),
            "workbook": self.workbook,
            "sheet": self.sheet,
            "written_cells": self.written_cells,
        }


@dataclass(frozen=True)
class SummaryExcelWriteResult:
    month: str
    output: Path
    workbook: str
    sheet: str
    target_row: int
    written_cells: list[str]
    mode: str

    def to_dict(self) -> dict[str, object]:
        return {
            "month": self.month,
            "output": str(self.output),
            "workbook": self.workbook,
            "sheet": self.sheet,
            "target_row": self.target_row,
            "written_cells": self.written_cells,
            "mode": self.mode,
        }


def write_detail_excel(
    validation_report_path: Path | str,
    project_dir: Path | str,
    output_path: Path | str,
    *,
    site_key: str = "xinan_high_school",
    config_path: Path | str = "config/sites.json",
    pricing_snapshot: dict[str, Any] | None = None,
) -> DetailExcelWriteResult:
    """Write the school detail workbook copy from a validation report."""

    report = _load_writable_validation_report(Path(validation_report_path))
    normalized = _normalized_data(report)
    site = load_site_config(site_key=site_key, config_path=config_path)
    pricing = _pricing_context(site, report, normalized, pricing_snapshot)
    template_path = Path(project_dir) / site.detail_workbook
    if not template_path.exists():
        raise ExcelWriteError(f"Detail workbook template does not exist: {template_path}")

    output = _copy_template(template_path, output_path, "detail workbook template")

    workbook = load_workbook(output)
    if site.detail_sheet not in workbook.sheetnames:
        raise ExcelWriteError(f"Detail worksheet '{site.detail_sheet}' does not exist in {output}.")
    worksheet = workbook[site.detail_sheet]

    month = _required_str(report, "month")
    formula_generation = uses_formula_generation(month)
    month_number = int(month.split("-", 1)[1])
    title = f"{site.project_name}{month_number}月学校用电统计表"
    period_start = _date_value(_required_normalized(normalized, "period_start"))
    period_end = _date_value(_required_normalized(normalized, "period_end"))
    statement_date = _date_value(_required_normalized(normalized, "statement_date"))

    row = site.detail_start_row
    writes: dict[str, Any] = {
        "A1": title,
        "E2": period_start,
        "H2": period_end,
    }
    if site.detail_profile == "multi_detail_rows":
        detail_rows = _detail_meter_rows(normalized)
        expected_count = len(site.detail_meters)
        if len(detail_rows) != expected_count:
            raise ExcelWriteError(
                f"Detail meter row count {len(detail_rows)} does not match configured detail meter count {expected_count}."
            )
        price_value = _school_price_cell_value(pricing["base_unit_price"], pricing["discount_rate"], pricing["effective_unit_price"])
        total_terms: list[str] = []
        for offset, detail_row in enumerate(detail_rows):
            current_row = row + offset
            writes.update(
                {
                    f"A{current_row}": f"表号{_required_meter_value(detail_row, 'meter_no')}",
                    f"B{current_row}": _required_meter_value(detail_row, "previous_reading"),
                    f"C{current_row}": _required_meter_value(detail_row, "current_reading"),
                    f"D{current_row}": _required_meter_value(detail_row, "multiplier"),
                    f"E{current_row}": (
                        f"=ROUND((C{current_row}-B{current_row})*D{current_row},0)"
                        if formula_generation else _required_meter_value(detail_row, "generation_kwh")
                    ),
                    f"F{current_row}": _required_meter_value(detail_row, "grid_export_kwh"),
                    f"G{current_row}": f"=E{current_row}-F{current_row}",
                    f"H{current_row}": price_value,
                    f"I{current_row}": f"=G{current_row}*H{current_row}",
                }
            )
            total_terms.append(f"I{current_row}")
        if len(detail_rows) > 1:
            _ensure_vertical_merge(worksheet, "J", row, len(detail_rows))
        writes[f"J{row}"] = "=ROUND(" + "+".join(total_terms) + ",2)"
        compaction_state = None
    else:
        detail_row = _detail_meter_row(normalized)
        writes.update(
            {
                f"A{row}": f"表号{_required_meter_value(detail_row, 'meter_no')}",
                f"B{row}": _required_meter_value(detail_row, "previous_reading"),
                f"C{row}": _required_meter_value(detail_row, "current_reading"),
                f"D{row}": _required_meter_value(detail_row, "multiplier"),
                f"E{row}": (
                    f"=ROUND((C{row}-B{row})*D{row},0)"
                    if formula_generation else _required_meter_value(detail_row, "generation_kwh")
                ),
                f"F{row}": _required_meter_value(detail_row, "grid_export_kwh"),
                f"G{row}": f"=E{row}-F{row}",
                f"H{row}": _school_price_cell_value(pricing["base_unit_price"], pricing["discount_rate"], pricing["effective_unit_price"]),
                f"I{row}": f"=G{row}*H{row}",
                f"J{row}": f"=I{row}",
            }
        )
        company_value = _required_normalized(normalized, "company_name") if site.detail_company_cell else None
        compaction_state = None
        if _uses_compacted_single_detail_layout(site):
            compaction_state = _compact_single_detail_row(worksheet, row, company_value=company_value)

    writes[_detail_price_header_cell(site)] = _price_header_text(pricing["discount_rate"])
    for cell, value in writes.items():
        if _is_merged_non_anchor(worksheet, cell):
            continue
        _write_worksheet_cell(worksheet, cell, value)

    company_value = _required_normalized(normalized, "company_name") if site.detail_company_cell else None
    footer_writes: dict[str, Any] = {}
    if site.detail_company_cell:
        footer_writes[_cell_after_detail_compaction(site.detail_company_cell, row, compaction_state)] = company_value
    if site.detail_statement_date_cell:
        footer_writes[_cell_after_detail_compaction(site.detail_statement_date_cell, row, compaction_state)] = statement_date
    for cell, value in footer_writes.items():
        _write_worksheet_cell(worksheet, cell, value)
    writes.update(footer_writes)

    _assert_pricing_writes(writes, pricing)
    workbook.save(output)
    workbook.close()

    return DetailExcelWriteResult(
        month=month,
        output=output,
        workbook=site.detail_workbook,
        sheet=site.detail_sheet,
        written_cells=list(writes),
    )


def write_summary_excel(
    validation_report_path: Path | str,
    project_dir: Path | str,
    output_path: Path | str,
    *,
    site_key: str = "xinan_high_school",
    config_path: Path | str = "config/sites.json",
    pricing_snapshot: dict[str, Any] | None = None,
) -> SummaryExcelWriteResult:
    """Write or update the summary workbook copy from a validation report."""

    report = _load_writable_validation_report(Path(validation_report_path))
    normalized = _normalized_data(report)
    site = load_site_config(site_key=site_key, config_path=config_path)
    pricing = _pricing_context(site, report, normalized, pricing_snapshot)
    template_path = Path(project_dir) / site.summary_workbook
    if not template_path.exists():
        raise ExcelWriteError(f"Summary workbook template does not exist: {template_path}")

    output = _copy_template(template_path, output_path, "summary workbook template")

    workbook = load_workbook(output)
    if site.summary_sheet not in workbook.sheetnames:
        raise ExcelWriteError(f"Summary worksheet '{site.summary_sheet}' does not exist in {output}.")
    worksheet = workbook[site.summary_sheet]
    header_writes = _write_summary_site_name(worksheet, site)

    month = _required_str(report, "month")
    formula_generation = uses_formula_generation(month)
    month_date = _month_date(month)
    meter_rows = _summary_meter_rows(normalized)
    summary_group_row_count = _summary_group_row_count(site, meter_rows)
    summary_group_spacer_rows = _summary_group_spacer_rows(site)
    target_row = _find_month_row(worksheet, month_date)
    final_group_before_write = _latest_summary_month_group(
        worksheet,
        summary_group_row_count,
        start_row=site.summary_start_row,
        spacer_rows=summary_group_spacer_rows,
    )
    mode = "update"
    if target_row is None:
        if final_group_before_write is not None:
            target_row = final_group_before_write + summary_group_row_count + summary_group_spacer_rows
            mode = "append"
        else:
            target_row = _find_blank_summary_row_group(
                worksheet,
                summary_group_row_count,
                start_row=site.summary_start_row,
                end_column=site.summary_end_column,
                spacer_rows=summary_group_spacer_rows,
            )
            mode = "blank_row"
    if target_row is None:
        target_row = site.summary_start_row
        mode = "append"

    if mode in {"blank_row", "append"}:
        previous_group_start = _find_previous_summary_month_group(
            worksheet,
            target_row,
            summary_group_row_count,
            start_row=site.summary_start_row,
            spacer_rows=summary_group_spacer_rows,
        )
        if previous_group_start is not None:
            _copy_row_style_group(worksheet, previous_group_start, target_row, summary_group_row_count)
        elif mode == "append":
            _copy_row_style_group(
                worksheet,
                max(site.summary_start_row, target_row - summary_group_row_count),
                target_row,
                summary_group_row_count,
            )

    if summary_group_row_count > 1:
        for column in site.summary_merge_columns:
            _ensure_vertical_merge(worksheet, column, target_row, summary_group_row_count)

    if site.summary_profile == "two_meter_ak_summary":
        writes = _two_meter_ak_summary_writes(
            site, normalized, meter_rows, month_date, target_row, pricing,
            formula_generation=formula_generation,
        )
    else:
        writes = _meter_rows_summary_writes(
            site, normalized, meter_rows, month_date, target_row, pricing,
            formula_generation=formula_generation,
            materialize_calculations=site.summary_materialize_calculations,
        )

    for cell, value in writes.items():
        if _is_merged_non_anchor(worksheet, cell):
            continue
        worksheet[cell] = value
    # Existing month groups are historical records. Do not rewrite them when
    # appending a later month; only the target month is generated here.
    repaired_total_cells: list[str] = []
    for offset in range(summary_group_row_count):
        month_cell = f"B{target_row + offset}"
        if not _is_merged_non_anchor(worksheet, month_cell):
            worksheet[month_cell].number_format = MONTH_NUMBER_FORMAT
    _apply_summary_month_cell_style(worksheet, target_row, summary_group_row_count, site)
    _ensure_summary_spacer_rows(
        worksheet,
        target_row,
        summary_group_row_count,
        summary_group_spacer_rows,
        end_column=site.summary_end_column,
    )
    _clear_unused_summary_body_borders(
        worksheet,
        summary_group_row_count,
        start_row=site.summary_start_row,
        end_column=site.summary_end_column,
        spacer_rows=summary_group_spacer_rows,
    )

    _assert_pricing_writes(writes, pricing)
    _reject_summary_placeholders(worksheet)
    workbook.save(output)
    workbook.close()

    return SummaryExcelWriteResult(
        month=month,
        output=output,
        workbook=site.summary_workbook,
        sheet=site.summary_sheet,
        target_row=target_row,
        written_cells=[*header_writes, *writes, *repaired_total_cells],
        mode=mode,
    )




def _write_summary_site_name(worksheet: Any, site: Any) -> list[str]:
    cell_ref = getattr(site, "summary_site_name_cell", None) or "A1"
    _write_worksheet_cell(worksheet, str(cell_ref), site.site_name)
    return [str(cell_ref)]


def _reject_summary_placeholders(worksheet: Any) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and "{site_name}" in value:
                raise ExcelWriteError(f"Summary workbook still contains unresolved site placeholder at {cell.coordinate}.")


def _apply_summary_month_cell_style(worksheet: Any, target_row: int, row_count: int, site: Any) -> None:
    alignment = getattr(site, "summary_month_cell_alignment", None)
    if not alignment:
        return
    if str(alignment).lower() != "center":
        raise ExcelWriteError(f"Unsupported summary_month_cell_alignment={alignment!r}.")
    source = worksheet[f"B{target_row}"]
    source.number_format = MONTH_NUMBER_FORMAT
    source.alignment = _centered_alignment(source.alignment)
    for offset in range(row_count):
        cell = worksheet[f"B{target_row + offset}"]
        if cell.coordinate != source.coordinate:
            _copy_cell_visual_style(source, cell)
        cell.number_format = MONTH_NUMBER_FORMAT
        cell.alignment = _centered_alignment(cell.alignment)


def _copy_cell_visual_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def _centered_alignment(alignment: Any) -> Alignment:
    return Alignment(
        horizontal="center",
        vertical="center",
        text_rotation=alignment.textRotation,
        wrap_text=alignment.wrap_text,
        shrink_to_fit=alignment.shrink_to_fit,
        indent=alignment.indent,
    )


def _ensure_summary_spacer_rows(
    worksheet: Any,
    target_row: int,
    row_count: int,
    spacer_rows: int,
    *,
    end_column: int,
) -> None:
    if spacer_rows <= 0:
        return
    first_spacer_row = target_row + row_count
    last_spacer_row = first_spacer_row + spacer_rows - 1
    _unmerge_ranges_intersecting_rows(worksheet, first_spacer_row, last_spacer_row)
    for row in range(first_spacer_row, last_spacer_row + 1):
        for column in range(1, end_column + 1):
            worksheet.cell(row=row, column=column).value = None


def _unmerge_ranges_intersecting_rows(worksheet: Any, first_row: int, last_row: int) -> None:
    for merged_range in list(worksheet.merged_cells.ranges):
        _, min_row, _, max_row = range_boundaries(str(merged_range))
        if min_row <= last_row and max_row >= first_row:
            worksheet.unmerge_cells(str(merged_range))


def _clear_unused_summary_body_borders(
    worksheet: Any,
    row_count: int,
    *,
    start_row: int,
    end_column: int,
    spacer_rows: int,
) -> None:
    final_group_start = _latest_summary_month_group(
        worksheet,
        row_count,
        start_row=start_row,
        spacer_rows=spacer_rows,
    )
    if final_group_start is None:
        return
    first_unused_row = final_group_start + row_count
    for row in range(first_unused_row, worksheet.max_row + 1):
        for column in range(1, end_column + 1):
            worksheet.cell(row=row, column=column).border = Border()


def _meter_rows_summary_writes(
    site: Any,
    normalized: dict[str, Any],
    meter_rows: list[dict[str, Any]],
    month_date: date,
    target_row: int,
    pricing: dict[str, Decimal],
    *,
    formula_generation: bool = True,
    materialize_calculations: bool = False,
) -> dict[str, Any]:
    writes: dict[str, Any] = {}
    price_value = _school_price_cell_value(pricing["base_unit_price"], pricing["discount_rate"], pricing["effective_unit_price"])
    price_decimal = Decimal(str(price_value))
    calculated_total_fee = (
        Decimal(str(_required_normalized(normalized, "calculated_total_fee")))
        if materialize_calculations
        else None
    )
    writes[f"K{site.summary_start_row - 1}"] = _price_header_text(pricing["discount_rate"])
    total_terms: list[str] = [f"I{target_row}"] if site.meter_topology_profile == "dual_meter_school" else []
    for offset, meter_row in enumerate(meter_rows):
        row = target_row + offset
        school_consumption = (
            Decimal(str(_required_meter_value(meter_row, "school_consumption_kwh")))
            if materialize_calculations
            else None
        )
        school_fee = school_consumption * price_decimal if school_consumption is not None else None
        meter_grid_energy = Decimal(str(_required_meter_value(meter_row, "grid_export_kwh")))
        grid_price = Decimal(str(_required_normalized(normalized, "grid_price")))
        calculated_grid_fee = meter_grid_energy * grid_price if materialize_calculations else None
        writes.update(
            {
                f"A{row}": f"表号{_required_meter_value(meter_row, 'meter_no')}",
                f"B{row}": month_date,
                f"C{row}": _required_meter_value(meter_row, "previous_reading"),
                f"D{row}": _required_meter_value(meter_row, "current_reading"),
                f"E{row}": _required_meter_value(meter_row, "multiplier"),
                f"F{row}": (
                    _required_meter_value(meter_row, "generation_kwh")
                    if materialize_calculations or not formula_generation
                    else f"=ROUND((D{row}-C{row})*E{row},0)"
                ),
                f"G{row}": _required_meter_value(meter_row, "grid_export_kwh"),
                f"H{row}": _required_normalized(normalized, "grid_price"),
                f"I{row}": _json_number(calculated_grid_fee) if calculated_grid_fee is not None else f"=G{row}*H{row}",
                f"J{row}": _json_number(school_consumption) if school_consumption is not None else f"=F{row}-G{row}",
                f"K{row}": price_value,
                f"L{row}": _json_number(school_fee) if school_fee is not None else f"=J{row}*K{row}",
            }
        )
        total_terms.append(f"L{row}")
        if site.meter_topology_profile == "single_meter_school":
            total_terms.append(f"I{row}")
    if materialize_calculations:
        writes[f"M{target_row}"] = _json_number(calculated_total_fee)
    elif site.meter_topology_profile == "dual_meter_school":
        school_fee_terms = "+".join(f"L{target_row + offset}" for offset in range(len(meter_rows)))
        writes[f"M{target_row}"] = f"=I{target_row}+ROUND({school_fee_terms},2)"
    else:
        writes[f"M{target_row}"] = "=" + "+".join(total_terms)
    return writes


def _repair_dual_meter_summary_totals(worksheet: Any, site: Any, row_count: int) -> list[str]:
    if site.meter_topology_profile != "dual_meter_school" or site.summary_profile != "meter_rows_summary":
        return []
    repaired: list[str] = []
    for row in range(site.summary_start_row, worksheet.max_row + 1):
        if _summary_month_group_key(worksheet, row, row_count) is None:
            continue
        school_fee_terms = "+".join(f"L{row + offset}" for offset in range(row_count))
        cell = f"M{row}"
        worksheet[cell] = f"=I{row}+ROUND({school_fee_terms},2)"
        repaired.append(cell)
    return repaired


def _two_meter_ak_summary_writes(
    site: Any,
    normalized: dict[str, Any],
    meter_rows: list[dict[str, Any]],
    month_date: date,
    target_row: int,
    pricing: dict[str, Decimal],
    *,
    formula_generation: bool = True,
) -> dict[str, Any]:
    writes: dict[str, Any] = {}
    price_value = _school_price_cell_value(pricing["base_unit_price"], pricing["discount_rate"], pricing["effective_unit_price"])
    writes[f"I{site.summary_start_row - 1}"] = _price_header_text(pricing["discount_rate"])
    total_terms: list[str] = []
    for offset, meter_row in enumerate(meter_rows):
        row = target_row + offset
        writes.update(
            {
                f"A{row}": f"表号{_required_meter_value(meter_row, 'meter_no')}",
                f"B{row}": month_date,
                f"C{row}": _required_meter_value(meter_row, "previous_reading"),
                f"D{row}": _required_meter_value(meter_row, "current_reading"),
                f"E{row}": _required_meter_value(meter_row, "multiplier"),
                f"F{row}": (
                    f"=ROUND((D{row}-C{row})*E{row},0)"
                    if formula_generation else _required_meter_value(meter_row, "generation_kwh")
                ),
                f"G{row}": _required_meter_value(meter_row, "grid_export_kwh"),
                f"H{row}": f"=F{row}-G{row}",
                f"I{row}": price_value,
                f"J{row}": f"=H{row}*I{row}",
            }
        )
        total_terms.append(f"J{row}")
    writes[f"K{target_row}"] = "=" + "+".join(total_terms)
    return writes

def _load_writable_validation_report(path: Path) -> dict[str, Any]:
    report = _load_validation_report(path)
    if not report.get("can_write_excel"):
        status = report.get("status", "unknown")
        errors = report.get("errors") or []
        detail = "; ".join(str(error) for error in errors) if errors else "can_write_excel=false"
        raise ExcelWriteError(f"Validation report blocks Excel writing: status={status}; {detail}")
    return report


def _load_validation_report(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise ExcelWriteError(f"validation_report.json does not exist: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise ExcelWriteError(f"validation_report.json is not valid JSON: {path}") from exc
    if not isinstance(payload, dict):
        raise ExcelWriteError("validation_report.json must contain a JSON object.")
    return payload


def _normalized_data(report: dict[str, Any]) -> dict[str, Any]:
    normalized = report.get("normalized_data")
    if not isinstance(normalized, dict):
        raise ExcelWriteError("Validation report must contain normalized_data.")
    return normalized


def _copy_template(template_path: Path, output_path: Path | str, label: str) -> Path:
    output = Path(output_path)
    if template_path.resolve() == output.resolve():
        raise ExcelWriteError(f"Refusing to overwrite the original {label}.")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output)
    return output


def _required_str(payload: dict[str, Any], key: str) -> str:
    value = payload.get(key)
    if value is None or str(value).strip() == "":
        raise ExcelWriteError(f"Validation report is missing '{key}'.")
    return str(value)


def _required_normalized(normalized: dict[str, Any], key: str) -> Any:
    if key not in normalized:
        raise ExcelWriteError(f"Validation report normalized_data is missing '{key}'.")
    return normalized[key]


def _detail_meter_row(normalized: dict[str, Any]) -> dict[str, Any]:
    return _detail_meter_rows(normalized)[0]


def _detail_meter_rows(normalized: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [row for row in _summary_meter_rows(normalized) if row.get("detail_enabled") is True]
    return rows or _summary_meter_rows(normalized)


def _summary_meter_rows(normalized: dict[str, Any]) -> list[dict[str, Any]]:
    rows = normalized.get("meter_rows")
    if not isinstance(rows, list) or not rows:
        return [
            {
                "meter_no": _required_normalized(normalized, "meter_no"),
                "previous_reading": _required_normalized(normalized, "previous_reading"),
                "current_reading": _required_normalized(normalized, "current_reading"),
                "multiplier": _required_normalized(normalized, "multiplier"),
                "generation_kwh": _required_normalized(normalized, "total_generation_kwh"),
                "grid_export_kwh": _required_normalized(normalized, "grid_export_kwh"),
                "detail_enabled": True,
            }
        ]
    if not all(isinstance(row, dict) for row in rows):
        raise ExcelWriteError("Validation report normalized_data.meter_rows must contain objects.")
    return rows



def _summary_group_row_count(site: Any, meter_rows: list[dict[str, Any]]) -> int:
    configured_count = len(getattr(site, "meters", ()) or ())
    row_count = configured_count or len(meter_rows)
    if row_count <= 0:
        raise ExcelWriteError("Summary meter row group must contain at least one row.")
    if len(meter_rows) != row_count:
        raise ExcelWriteError(
            f"Summary meter row count {len(meter_rows)} does not match configured row count {row_count}."
        )
    return row_count


def _summary_group_spacer_rows(site: Any) -> int:
    value = getattr(site, "summary_group_spacer_rows", 0)
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ExcelWriteError(f"summary_group_spacer_rows must be an integer, got {value!r}.") from exc
    if parsed < 0:
        raise ExcelWriteError("summary_group_spacer_rows must be non-negative.")
    return parsed


def _required_meter_value(row: dict[str, Any], key: str) -> Any:
    if key not in row:
        raise ExcelWriteError(f"Validation report meter row is missing '{key}'.")
    return row[key]


def _pricing_context(
    site: Any,
    report: dict[str, Any],
    normalized: dict[str, Any],
    pricing_snapshot: dict[str, Any] | None,
) -> dict[str, Decimal]:
    month = _required_str(report, "month")
    site_key = _required_str(report, "site_key")
    if pricing_snapshot is None:
        raise ExcelWriteError("Excel writing requires a frozen task pricing_snapshot.")
    snapshot = validate_task_pricing_snapshot(pricing_snapshot, site_key=site_key, month=month)
    values = pricing_snapshot_for_recognition(snapshot)
    context = {
        "base_unit_price": values["school_base_price"],
        "discount_rate": values["school_discount_rate"],
        "effective_unit_price": values["effective_unit_price"],
    }
    if site_key != site.site_key:
        raise ExcelWriteError("Validation report site_key does not match requested site.")
    _assert_normalized_pricing(normalized, context)
    return context


def _assert_normalized_pricing(normalized: dict[str, Any], pricing: dict[str, Decimal]) -> None:
    base = Decimal(str(_required_normalized(normalized, "school_base_price")))
    discount = Decimal(str(_required_normalized(normalized, "school_discount_rate")))
    if base != pricing["base_unit_price"] or discount != pricing["discount_rate"]:
        raise ExcelWriteError("Validation report pricing fields do not match task pricing snapshot.")


def _detail_price_header_cell(site: Any) -> str:
    return f"H{int(site.detail_start_row) - 1}"


def _price_header_text(discount_rate: Decimal) -> str:
    discount_text = _decimal_text(canonicalize_discount_rate(discount_rate) * Decimal("10"))
    return f"{discount_text}折单价（元/kWh）"


def _assert_pricing_writes(writes: dict[str, Any], pricing: dict[str, Decimal]) -> None:
    expected_header = _price_header_text(pricing["discount_rate"])
    expected_price = _school_price_cell_value(
        pricing["base_unit_price"],
        pricing["discount_rate"],
        pricing["effective_unit_price"],
    )
    header_values = [value for value in writes.values() if isinstance(value, str) and "单价" in value]
    if not header_values or any(value != expected_header for value in header_values):
        raise ExcelWriteError("Written price header does not match task pricing snapshot.")
    numeric_values = [value for value in writes.values() if isinstance(value, (int, float))]
    if expected_price not in numeric_values:
        raise ExcelWriteError("Written price value does not match task pricing snapshot.")


def _school_price_cell_value(
    base_price: Decimal,
    discount_rate: Decimal,
    effective_price: Decimal | None = None,
) -> int | float | str:
    if effective_price is not None:
        return _json_number(effective_price)
    raise ExcelWriteError("Excel writing requires pricing_snapshot.effective_unit_price; settlement price is never derived.")


def _date_value(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value)
    try:
        year, month, day = (int(part) for part in text.split("-"))
    except ValueError as exc:
        raise ExcelWriteError(f"Expected date value as YYYY-MM-DD, got {value!r}.") from exc
    return date(year, month, day)


def _month_date(month: str) -> date:
    try:
        year, month_number = (int(part) for part in month.split("-", 1))
    except ValueError as exc:
        raise ExcelWriteError(f"Expected month value as YYYY-MM, got {month!r}.") from exc
    return date(year, month_number, 1)


def _find_month_row(worksheet: Any, month_date: date) -> int | None:
    for row in range(1, worksheet.max_row + 1):
        value = _worksheet_cell_value(worksheet, row, 2)
        if _same_month(value, month_date):
            return row
    return None


def _find_blank_summary_row(worksheet: Any, *, start_row: int = 4, end_column: int = 13) -> int | None:
    for row in range(start_row, worksheet.max_row + 1):
        cells = [worksheet.cell(row=row, column=column) for column in range(2, end_column + 1)]
        if all(cell.value in (None, "") for cell in cells) and any(cell.has_style for cell in cells):
            return row
    return None


def _find_blank_summary_row_group(
    worksheet: Any,
    row_count: int,
    *,
    start_row: int = 4,
    end_column: int = 13,
    spacer_rows: int = 0,
) -> int | None:
    if row_count <= 1:
        return _find_blank_summary_row(worksheet, start_row=start_row, end_column=end_column)
    step = row_count + spacer_rows if spacer_rows else 1
    for row in range(start_row, worksheet.max_row + 1, step):
        if row + row_count - 1 > worksheet.max_row:
            break
        group_cells = [
            worksheet.cell(row=group_row, column=column)
            for group_row in range(row, row + row_count)
            for column in range(2, end_column + 1)
        ]
        spacer_cells = [
            worksheet.cell(row=spacer_row, column=column)
            for spacer_row in range(row + row_count, row + row_count + spacer_rows)
            if spacer_row <= worksheet.max_row
            for column in range(1, end_column + 1)
        ]
        if (
            all(cell.value in (None, "") for cell in group_cells)
            and all(cell.value in (None, "") for cell in spacer_cells)
            and any(cell.has_style for cell in group_cells)
        ):
            return row
    return None


def _same_month(value: Any, month_date: date) -> bool:
    if isinstance(value, datetime):
        return value.year == month_date.year and value.month == month_date.month
    if isinstance(value, date):
        return value.year == month_date.year and value.month == month_date.month
    if isinstance(value, (int, float)):
        # openpyxl uses 1899-12-30 as day zero for Windows Excel serials.
        serial_date = datetime.fromordinal(date(1899, 12, 30).toordinal() + int(value))
        return serial_date.year == month_date.year and serial_date.month == month_date.month
    if isinstance(value, str):
        return value.startswith(month_date.strftime("%Y-%m"))
    return False



def _copy_row_style_group(worksheet: Any, source_start_row: int, target_start_row: int, row_count: int) -> None:
    for offset in range(row_count):
        _copy_row_style(worksheet, source_start_row + offset, target_start_row + offset)


def _find_previous_summary_month_group(
    worksheet: Any,
    target_start_row: int,
    row_count: int,
    *,
    start_row: int,
    spacer_rows: int = 0,
) -> int | None:
    candidate = target_start_row - row_count
    while candidate >= start_row:
        if _summary_month_group_key(worksheet, candidate, row_count) is not None:
            return candidate
        candidate -= 1
    return None


def _latest_summary_month_group(worksheet: Any, row_count: int, *, start_row: int, spacer_rows: int) -> int | None:
    latest: int | None = None
    for row in range(start_row, worksheet.max_row + 1):
        if _summary_month_group_key(worksheet, row, row_count) is not None:
            latest = row
    return latest


def _summary_month_group_key(worksheet: Any, start_row: int, row_count: int) -> str | None:
    keys = [_month_key(_worksheet_cell_value(worksheet, start_row + offset, 2)) for offset in range(row_count)]
    if any(key is None for key in keys):
        return None
    first = keys[0]
    return first if all(key == first for key in keys) else None


def _worksheet_cell_value(worksheet: Any, row: int, column: int) -> Any:
    cell = worksheet.cell(row=row, column=column)
    if cell.value is not None:
        return cell.value
    coordinate = cell.coordinate
    for merged_range in worksheet.merged_cells.ranges:
        if coordinate in merged_range:
            return merged_range.start_cell.value
    return cell.value


def _month_key(value: Any) -> str | None:
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, date):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, (int, float)):
        serial_date = datetime.fromordinal(date(1899, 12, 30).toordinal() + int(value))
        return f"{serial_date.year:04d}-{serial_date.month:02d}"
    if isinstance(value, str):
        text = value.strip()
        if len(text) >= 7 and text[:4].isdigit() and text[4] in {"-", "/"} and text[5:7].isdigit():
            return f"{text[:4]}-{text[5:7]}"
        if "年" in text and "月" in text:
            year, rest = text.split("年", 1)
            month = rest.split("月", 1)[0]
            if year.isdigit() and month.isdigit():
                return f"{int(year):04d}-{int(month):02d}"
    return None

def _copy_row_style(worksheet: Any, source_row: int, target_row: int) -> None:
    if source_row < 1:
        return
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(row=source_row, column=column)
        target = worksheet.cell(row=target_row, column=column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)


def _uses_compacted_single_detail_layout(site: Any) -> bool:
    blank_row = int(site.detail_start_row) + 1
    company_row = _cell_row(site.detail_company_cell)
    return company_row is not None and company_row > blank_row


def _compact_single_detail_row(worksheet: Any, data_row: int, *, company_value: Any | None) -> str:
    """Remove the unused second detail row while preserving footer rows below it."""

    blank_row = data_row + 1
    if not _row_is_blank(worksheet, blank_row):
        if _row_is_already_compacted_footer(worksheet, blank_row, company_value=company_value):
            return "already_compacted"
        raise ExcelWriteError("Single detail row profile found non-empty second detail row.")

    merged_ranges = [str(item) for item in worksheet.merged_cells.ranges]
    for merged_range in merged_ranges:
        worksheet.unmerge_cells(merged_range)

    worksheet.delete_rows(blank_row, 1)
    for merged_range in merged_ranges:
        shifted = _shift_merged_range_after_row_delete(merged_range, blank_row)
        if shifted is not None:
            worksheet.merge_cells(shifted)
    return "deleted"


def _row_is_blank(worksheet: Any, row: int) -> bool:
    return all(worksheet.cell(row=row, column=column).value in (None, "") for column in range(1, worksheet.max_column + 1))


def _row_is_already_compacted_footer(worksheet: Any, row: int, *, company_value: Any | None) -> bool:
    if not _row_has_wide_merge(worksheet, row):
        return False
    values = [worksheet.cell(row=row, column=column).value for column in range(1, worksheet.max_column + 1)]
    nonblank = [str(value).strip() for value in values if value not in (None, "")]
    if len(nonblank) != 1:
        return False
    markers = {"供电公司名称"}
    if company_value not in (None, ""):
        markers.add(str(company_value).strip())
    return nonblank[0] in markers


def _row_has_wide_merge(worksheet: Any, row: int) -> bool:
    from openpyxl.utils.cell import range_boundaries

    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_row <= row <= max_row and min_col == 1 and max_col >= 10:
            return True
    return False


def _cell_after_detail_compaction(cell_ref: str, data_row: int, compaction_state: str | None) -> str:
    if compaction_state not in {"deleted", "already_compacted"}:
        return cell_ref
    from openpyxl.utils.cell import coordinate_from_string

    column, row = coordinate_from_string(cell_ref)
    blank_row = data_row + 1
    if row > blank_row:
        return f"{column}{row - 1}"
    return cell_ref


def _cell_row(cell_ref: str | None) -> int | None:
    if not cell_ref:
        return None
    from openpyxl.utils.cell import coordinate_from_string

    return int(coordinate_from_string(cell_ref)[1])


def _write_worksheet_cell(worksheet: Any, cell_ref: str, value: Any) -> None:
    worksheet[_merged_anchor_cell(worksheet, cell_ref)] = value


def _merged_anchor_cell(worksheet: Any, cell_ref: str) -> str:
    for merged_range in worksheet.merged_cells.ranges:
        if cell_ref in merged_range:
            return str(merged_range.start_cell.coordinate)
    return cell_ref


def _shift_merged_range_after_row_delete(range_text: str, deleted_row: int) -> str | None:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(range_text)
    if min_row == deleted_row and max_row == deleted_row:
        return None
    if min_row > deleted_row:
        min_row -= 1
        max_row -= 1
    elif min_row <= deleted_row <= max_row:
        max_row -= 1
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _is_merged_non_anchor(worksheet: Any, cell_ref: str) -> bool:
    for merged_range in worksheet.merged_cells.ranges:
        if cell_ref not in merged_range:
            continue
        return cell_ref != str(merged_range.start_cell.coordinate)
    return False


def _ensure_vertical_merge(worksheet: Any, column: str, start_row: int, row_count: int) -> None:
    range_text = f"{column}{start_row}:{column}{start_row + row_count - 1}"
    if any(str(merged_range) == range_text for merged_range in worksheet.merged_cells.ranges):
        return
    column_index = column_index_from_string(column)
    end_row = start_row + row_count - 1
    for merged_range in list(worksheet.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_col <= column_index <= max_col and min_row <= end_row and max_row >= start_row:
            worksheet.unmerge_cells(str(merged_range))
    worksheet.merge_cells(range_text)


def _json_number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")
