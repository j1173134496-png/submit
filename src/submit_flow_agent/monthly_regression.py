"""Manifest-driven monthly regression runner for MVP-022A."""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from submit_flow_agent.file_scanner import REQUIRED_FILE_TYPES, FileScanError, scan_input_files
from submit_flow_agent.pipeline import RunMonthResult, run_month


BLOCKED_BY_SAMPLE = "BLOCKED_BY_SAMPLE"
PASSED = "PASS"
FAILED = "FAIL"
NOT_RUN = "NOT_RUN"
DEFAULT_MANIFEST_PATH = Path("config/monthly_regression_manifest.json")
DEFAULT_MATRIX_PATH = Path("spec/reports/monthly_regression_matrix.md")
MONTH_PATTERN = re.compile(r"^20\d{2}-(0[1-9]|1[0-2])$")
PipelineRunner = Callable[..., RunMonthResult]


class MonthlyRegressionError(RuntimeError):
    """Raised when the monthly regression manifest or run is invalid."""


@dataclass(frozen=True)
class MonthlyRegressionCase:
    case_id: str
    site_key: str
    month: str
    project_dir: Path
    sample_status: str = "available"
    description: str = ""


@dataclass(frozen=True)
class RequiredSampleSet:
    label: str
    site_key: str
    exclude_months: tuple[str, ...] = ()
    required_file_types: tuple[str, ...] = REQUIRED_FILE_TYPES


@dataclass(frozen=True)
class MonthlyRegressionManifest:
    version: int
    cases: tuple[MonthlyRegressionCase, ...]
    required_sample_sets: tuple[RequiredSampleSet, ...] = ()
    source_path: Path | None = None


@dataclass
class MonthlyRegressionRow:
    case_id: str
    site_key: str
    month: str
    sample_status: str
    input_complete: bool
    detected_month: str | None = None
    result: str = NOT_RUN
    pipeline_status: str | None = None
    review_fields: list[str] = field(default_factory=list)
    validation_status: str | None = None
    can_write_excel: bool | None = None
    detail_excel: str | None = None
    summary_excel: str | None = None
    key_cells: dict[str, Any] = field(default_factory=dict)
    evidence_paths: dict[str, str] = field(default_factory=dict)
    hashes: dict[str, str] = field(default_factory=dict)
    message: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "case_id": self.case_id,
            "site_key": self.site_key,
            "month": self.month,
            "sample_status": self.sample_status,
            "input_complete": self.input_complete,
            "detected_month": self.detected_month,
            "result": self.result,
            "pipeline_status": self.pipeline_status,
            "review_fields": self.review_fields,
            "validation_status": self.validation_status,
            "can_write_excel": self.can_write_excel,
            "detail_excel": self.detail_excel,
            "summary_excel": self.summary_excel,
            "key_cells": self.key_cells,
            "evidence_paths": self.evidence_paths,
            "hashes": self.hashes,
            "message": self.message,
        }


@dataclass(frozen=True)
class MonthlyRegressionResult:
    rows: tuple[MonthlyRegressionRow, ...]
    matrix_path: Path

    @property
    def has_blocked_samples(self) -> bool:
        return any(row.result == BLOCKED_BY_SAMPLE for row in self.rows)

    @property
    def executed_case_count(self) -> int:
        return sum(row.pipeline_status is not None for row in self.rows)

    def to_dict(self) -> dict[str, Any]:
        return {
            "matrix_path": str(self.matrix_path),
            "executed_case_count": self.executed_case_count,
            "has_blocked_samples": self.has_blocked_samples,
            "rows": [row.to_dict() for row in self.rows],
        }


def load_monthly_regression_manifest(path: Path | str = DEFAULT_MANIFEST_PATH) -> MonthlyRegressionManifest:
    manifest_path = Path(path)
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError as exc:
        raise MonthlyRegressionError(f"Regression manifest does not exist: {manifest_path}") from exc
    except json.JSONDecodeError as exc:
        raise MonthlyRegressionError(f"Regression manifest is not valid JSON: {manifest_path}") from exc
    if not isinstance(payload, dict):
        raise MonthlyRegressionError("Regression manifest must contain a JSON object.")

    version = payload.get("version")
    if not isinstance(version, int) or version < 1:
        raise MonthlyRegressionError("Regression manifest must contain integer version >= 1.")

    base_dir = manifest_path.parent.parent if manifest_path.parent.name == "config" else manifest_path.parent
    cases = tuple(_parse_case(item, base_dir, manifest_path) for item in _list_field(payload, "cases"))
    required = tuple(
        _parse_required_sample_set(item, manifest_path)
        for item in payload.get("required_sample_sets", [])
    )
    return MonthlyRegressionManifest(version=version, cases=cases, required_sample_sets=required, source_path=manifest_path)


def run_monthly_regression(
    manifest_path: Path | str = DEFAULT_MANIFEST_PATH,
    *,
    output_base: Path | str = "tmp/mvp022a_multi_month_regression/outputs",
    matrix_path: Path | str = DEFAULT_MATRIX_PATH,
    config_path: Path | str = "config/sites.json",
    execute: bool = True,
    pipeline_runner: PipelineRunner = run_month,
) -> MonthlyRegressionResult:
    manifest = load_monthly_regression_manifest(manifest_path)
    rows: list[MonthlyRegressionRow] = []
    output_root = Path(output_base)
    for case in manifest.cases:
        rows.append(
            _run_case(
                case,
                output_root=output_root,
                config_path=Path(config_path),
                execute=execute,
                pipeline_runner=pipeline_runner,
            )
        )
    rows.extend(_missing_required_sample_rows(manifest, rows))
    target = write_monthly_regression_matrix(rows, matrix_path, manifest=manifest)
    return MonthlyRegressionResult(rows=tuple(rows), matrix_path=target)


def write_monthly_regression_matrix(
    rows: list[MonthlyRegressionRow] | tuple[MonthlyRegressionRow, ...],
    matrix_path: Path | str = DEFAULT_MATRIX_PATH,
    *,
    manifest: MonthlyRegressionManifest | None = None,
) -> Path:
    target = Path(matrix_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Monthly Regression Matrix",
        "",
        "Status: READY_FOR_MVP022_REAL_MONTHS",
        "",
        "MVP-022A validates the manifest/config driven regression framework using the available 2026-05 sample only. It does not claim real multi-month acceptance; MVP-022B remains blocked until at least two additional real month sample sets are provided.",
        "",
    ]
    if manifest and manifest.source_path:
        lines.extend([f"Manifest: `{manifest.source_path}`", ""])
    lines.extend(
        [
            "| Case | Site | Month | Input | Detected | Result | Pipeline | Review Fields | Validation | Excel | Evidence |",
            "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
        ]
    )
    for row in rows:
        excel = _excel_summary(row)
        evidence = _evidence_summary(row)
        lines.append(
            "| "
            + " | ".join(
                _md_cell(value)
                for value in (
                    row.case_id,
                    row.site_key,
                    row.month,
                    "complete" if row.input_complete else "missing",
                    row.detected_month or "-",
                    row.result,
                    row.pipeline_status or "-",
                    ", ".join(row.review_fields) if row.review_fields else "-",
                    _validation_summary(row),
                    excel,
                    evidence,
                )
            )
            + " |"
        )
    blocked = [row for row in rows if row.result == BLOCKED_BY_SAMPLE]
    lines.extend(["", "## Missing Real Month Samples", ""])
    if blocked:
        for row in blocked:
            lines.append(f"- `{row.case_id}`: {row.message}")
    else:
        lines.append("- None recorded by this manifest.")
    lines.extend(["", "## Row Details", ""])
    for row in rows:
        lines.extend([
            f"### {row.case_id}",
            "",
            "```json",
            json.dumps(row.to_dict(), ensure_ascii=False, indent=2, sort_keys=True),
            "```",
            "",
        ])
    target.write_text("\n".join(lines), encoding="utf-8")
    return target


def _run_case(
    case: MonthlyRegressionCase,
    *,
    output_root: Path,
    config_path: Path,
    execute: bool,
    pipeline_runner: PipelineRunner,
) -> MonthlyRegressionRow:
    row = MonthlyRegressionRow(
        case_id=case.case_id,
        site_key=case.site_key,
        month=case.month,
        sample_status=case.sample_status,
        input_complete=False,
    )
    if case.sample_status != "available":
        row.result = BLOCKED_BY_SAMPLE
        row.message = f"Sample set is marked {case.sample_status}; no pipeline run was attempted."
        return row

    try:
        scan = scan_input_files(case.project_dir)
    except FileScanError as exc:
        row.result = BLOCKED_BY_SAMPLE
        row.message = str(exc)
        return row

    row.input_complete = set(scan.files) == set(REQUIRED_FILE_TYPES)
    row.detected_month = scan.month
    row.hashes.update({f"input.{file_type}": _sha256(path) for file_type, path in scan.files.items()})
    row.evidence_paths.update({f"input.{file_type}": str(path) for file_type, path in scan.files.items()})
    if not row.input_complete:
        row.result = BLOCKED_BY_SAMPLE
        row.message = "Source directory does not contain all three required PDF types."
        return row
    if scan.month != case.month:
        row.result = BLOCKED_BY_SAMPLE
        row.message = f"Manifest month {case.month} does not match detected month {scan.month}."
        return row
    if not execute:
        row.result = NOT_RUN
        row.message = "Sample discovery passed; execute=false skipped pipeline run."
        return row

    case_output_base = output_root / case.case_id
    try:
        result = pipeline_runner(
            case.project_dir,
            site_key=case.site_key,
            output_base=case_output_base,
            config_path=config_path,
        )
    except Exception as exc:
        row.result = FAILED
        row.pipeline_status = "failed"
        row.message = f"{type(exc).__name__}: {exc}"
        return row

    _populate_from_result(row, result)
    row.result = PASSED if result.status == "pass" else result.status
    return row


def _populate_from_result(row: MonthlyRegressionRow, result: RunMonthResult) -> None:
    row.pipeline_status = result.status
    row.evidence_paths["output_dir"] = str(result.output_dir)
    row.evidence_paths["process_log"] = str(result.process_log)
    row.evidence_paths["recognized_data"] = str(result.recognized_data)
    row.evidence_paths["validation_report"] = str(result.validation_report)
    if result.review_report:
        row.evidence_paths["review_report"] = str(result.review_report)
    if result.detail_workbook:
        row.detail_excel = str(result.detail_workbook)
        row.evidence_paths["detail_excel"] = str(result.detail_workbook)
    if result.summary_workbook:
        row.summary_excel = str(result.summary_workbook)
        row.evidence_paths["summary_excel"] = str(result.summary_workbook)

    for label, path_text in row.evidence_paths.items():
        path = Path(path_text)
        if path.is_file():
            row.hashes[label] = _sha256(path)

    validation = _read_json_if_exists(result.validation_report)
    if validation:
        row.validation_status = str(validation.get("status") or "unknown")
        can_write = validation.get("can_write_excel")
        row.can_write_excel = bool(can_write) if isinstance(can_write, bool) else None
    if result.review_report:
        review = _read_json_if_exists(result.review_report)
        if review:
            issues = review.get("issues")
            if isinstance(issues, list):
                row.review_fields = [str(issue.get("field")) for issue in issues if isinstance(issue, dict) and issue.get("field")]
    row.key_cells.update(_excel_key_cells(result.detail_workbook, result.summary_workbook))


def _missing_required_sample_rows(
    manifest: MonthlyRegressionManifest,
    existing_rows: list[MonthlyRegressionRow],
) -> list[MonthlyRegressionRow]:
    rows: list[MonthlyRegressionRow] = []
    available_months_by_site: dict[str, set[str]] = {}
    for row in existing_rows:
        if row.sample_status == "available" and row.input_complete and row.detected_month:
            available_months_by_site.setdefault(row.site_key, set()).add(row.detected_month)
    consumed_months_by_site: dict[str, set[str]] = {}
    for required in manifest.required_sample_sets:
        consumed = consumed_months_by_site.setdefault(required.site_key, set())
        available = sorted(available_months_by_site.get(required.site_key, set()) - set(required.exclude_months) - consumed)
        if available:
            consumed.add(available[0])
            continue
        rows.append(
            MonthlyRegressionRow(
                case_id=required.label,
                site_key=required.site_key,
                month="TBD",
                sample_status="missing",
                input_complete=False,
                result=BLOCKED_BY_SAMPLE,
                message=(
                    f"Missing one confirmed real month sample set for site {required.site_key}; "
                    f"must include {', '.join(required.required_file_types)} and must not reuse excluded months "
                    f"{', '.join(required.exclude_months) or '-'}."
                ),
            )
        )
    return rows


def _parse_case(item: object, base_dir: Path, manifest_path: Path) -> MonthlyRegressionCase:
    if not isinstance(item, dict):
        raise MonthlyRegressionError(f"{manifest_path} cases entries must be objects.")
    case_id = _required_string(item, "case_id", manifest_path)
    site_key = _required_string(item, "site_key", manifest_path)
    month = _required_string(item, "month", manifest_path)
    if not MONTH_PATTERN.match(month):
        raise MonthlyRegressionError(f"Regression case {case_id} has invalid month {month!r}.")
    project_dir = Path(_required_string(item, "project_dir", manifest_path))
    if not project_dir.is_absolute():
        project_dir = (base_dir / project_dir).resolve()
    return MonthlyRegressionCase(
        case_id=case_id,
        site_key=site_key,
        month=month,
        project_dir=project_dir,
        sample_status=str(item.get("sample_status") or "available"),
        description=str(item.get("description") or ""),
    )


def _parse_required_sample_set(item: object, manifest_path: Path) -> RequiredSampleSet:
    if not isinstance(item, dict):
        raise MonthlyRegressionError(f"{manifest_path} required_sample_sets entries must be objects.")
    label = _required_string(item, "label", manifest_path)
    site_key = _required_string(item, "site_key", manifest_path)
    exclude = item.get("exclude_months", [])
    if not isinstance(exclude, list) or not all(isinstance(value, str) for value in exclude):
        raise MonthlyRegressionError(f"Required sample set {label} exclude_months must be a string list.")
    file_types = item.get("required_file_types", list(REQUIRED_FILE_TYPES))
    if not isinstance(file_types, list) or not all(isinstance(value, str) for value in file_types):
        raise MonthlyRegressionError(f"Required sample set {label} required_file_types must be a string list.")
    return RequiredSampleSet(
        label=label,
        site_key=site_key,
        exclude_months=tuple(exclude),
        required_file_types=tuple(file_types),
    )


def _list_field(payload: dict[str, Any], key: str) -> list[Any]:
    value = payload.get(key)
    if not isinstance(value, list):
        raise MonthlyRegressionError(f"Regression manifest must contain a '{key}' list.")
    return value


def _required_string(payload: dict[str, Any], key: str, path: Path) -> str:
    value = payload.get(key)
    if not isinstance(value, str) or not value.strip():
        raise MonthlyRegressionError(f"{path} entry is missing string field {key!r}.")
    return value


def _read_json_if_exists(path: Path | None) -> dict[str, Any] | None:
    if path is None or not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError:
        return None
    return payload if isinstance(payload, dict) else None


def _excel_key_cells(detail_path: Path | None, summary_path: Path | None) -> dict[str, Any]:
    cells: dict[str, Any] = {}
    if detail_path and detail_path.exists():
        workbook = load_workbook(detail_path, data_only=False)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            for cell in ("A1", "E2", "H2", "A7"):
                cells[f"detail.{cell}"] = _cell_value(sheet[cell].value)
        finally:
            workbook.close()
    if summary_path and summary_path.exists():
        workbook = load_workbook(summary_path, data_only=False)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            for cell in ("B6", "F6", "M6"):
                cells[f"summary.{cell}"] = _cell_value(sheet[cell].value)
        finally:
            workbook.close()
    return cells


def _cell_value(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _validation_summary(row: MonthlyRegressionRow) -> str:
    if row.validation_status is None:
        return "-"
    if row.can_write_excel is None:
        return row.validation_status
    return f"{row.validation_status}; can_write_excel={str(row.can_write_excel).lower()}"


def _excel_summary(row: MonthlyRegressionRow) -> str:
    detail = "yes" if row.detail_excel else "no"
    summary = "yes" if row.summary_excel else "no"
    return f"detail={detail}; summary={summary}"


def _evidence_summary(row: MonthlyRegressionRow) -> str:
    if not row.evidence_paths:
        return "-"
    return ", ".join(sorted(row.evidence_paths))


def _md_cell(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


