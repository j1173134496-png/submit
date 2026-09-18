from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from submit_flow_agent.monthly_summary import (
    MonthlySummaryError,
    MonthlySummaryValues,
    collect_monthly_summary_source,
    confirm_monthly_summary,
    prepare_monthly_summary,
    run_monthly_summary,
    update_monthly_summary,
    write_monthly_summary_workbook,
    recognize_monthly_summary_values,
)
from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult
from submit_flow_agent.pdf_renderer import RenderedPage


class StaticOcrAdapter:
    def __init__(self, lines: list[str]) -> None:
        self.lines = lines

    def recognize(self, source_file: Path, pages: list[RenderedPage]) -> OcrRawResult:
        blocks = [
            OcrBlock(text=line, bbox=(0.0, float(index * 20), 300.0, float(index * 20 + 18)), confidence=0.99)
            for index, line in enumerate(self.lines)
        ]
        return OcrRawResult(source_file=source_file, pages=[OcrPage(1, pages[0].image_path, blocks)])


class MonthlySummaryTest(unittest.TestCase):
    def test_coordinate_table_maps_headers_to_the_value_row(self) -> None:
        headers = [
            ("\u8d2d\u7535\u6708\u4efd", 760, 112, 843, 144),
            ("\u603b\u7535\u91cf", 935, 116, 1000, 145),
            ("\u542b\u7a0e\u7535\u8d39", 1115, 118, 1194, 146),
        ]
        values = [
            ("202604", 766, 178, 836, 206),
            ("6621.6", 932, 179, 1001, 207),
            ("2075.55", 1117, 182, 1191, 206),
        ]
        blocks = [OcrBlock(text, bbox, 0.99) for text, *bbox in headers + values]
        raw = OcrRawResult(Path("voucher.jpg"), [OcrPage(1, Path("voucher.jpg"), blocks)])
        recognized, confidence = recognize_monthly_summary_values(raw)
        self.assertEqual(
            recognized,
            {"month": "2026-04", "total_energy": "6621.6", "tax_inclusive_fee": "2075.55"},
        )
        self.assertEqual(confidence["month"], 0.99)

    def test_writer_uses_actual_days_for_april_and_leap_february(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = _template(root / "template.xlsx")
            output = root / "output.xlsx"
            write_monthly_summary_workbook(
                template,
                output,
                {
                    "2026-04": MonthlySummaryValues("2026-04", Decimal("6621.6"), Decimal("2075.55")),
                    "2028-02": MonthlySummaryValues("2028-02", Decimal("2900"), Decimal("900")),
                },
            )
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["Sheet1"]
            rows = {sheet.cell(row, 4).value.strftime("%Y-%m"): row for row in range(3, sheet.max_row + 1)}
            self.assertEqual(sheet["A1"].value, "\u6c64\u897f\u656c\u8001\u9662")
            self.assertIn("A1:L1", {str(item) for item in sheet.merged_cells.ranges})
            self.assertEqual(sheet.cell(rows["2026-04"], 12).value, f"=IFERROR(G{rows['2026-04']}/86.14/30,0)")
            self.assertEqual(sheet.cell(rows["2028-02"], 12).value, f"=IFERROR(G{rows['2028-02']}/86.14/29,0)")
            self.assertIn("2026-03", rows)

    def test_writer_handles_a_template_without_calculation_properties(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = _template(root / "template.xlsx")
            workbook = load_workbook(template)
            workbook.calculation = None
            workbook.save(template)
            output = root / "output.xlsx"
            write_monthly_summary_workbook(
                template,
                output,
                {"2026-05": MonthlySummaryValues("2026-05", Decimal("100"), Decimal("30"))},
            )
            written = load_workbook(output, data_only=False)
            self.assertIsNotNone(written.calculation)
            self.assertTrue(written.calculation.fullCalcOnLoad)

    def test_same_month_revision_replaces_effective_row_and_preserves_history(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = _template(root / "template.xlsx")
            first = prepare_monthly_summary(root, task_id="first", month="2026-05")
            self.assertEqual(first["revision"], 1)
            _collect(root, "first", "202605", "100", "30")
            update_monthly_summary(
                root,
                "first",
                {"month": "2026-05", "total_energy": "100", "tax_inclusive_fee": "30"},
            )
            confirm_monthly_summary(root, "first")
            run_monthly_summary(root, "first", template_path=template)

            second = prepare_monthly_summary(root, task_id="second", month="2026-05")
            self.assertEqual(second["revision"], 2)
            _collect(root, "second", "202605", "120", "36")
            update_monthly_summary(
                root,
                "second",
                {"month": "2026-05", "total_energy": "120", "tax_inclusive_fee": "36"},
            )
            confirm_monthly_summary(root, "second")
            completed = run_monthly_summary(root, "second", template_path=template)
            output = root / "monthly_summary_tasks" / "second" / completed["outputs"][0]
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["Sheet1"]
            may_rows = [row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 4).value.strftime("%Y-%m") == "2026-05"]
            self.assertEqual(may_rows, [5])
            self.assertEqual(sheet.cell(5, 7).value, 120)
            series = (root / "monthly_summaries" / "tangxi_nursing_home" / "series.json").read_text(encoding="utf-8")
            self.assertIn('"task_id": "first"', series)

    def test_month_mismatch_is_rejected_without_changing_expected_month(self) -> None:
        from submit_flow_agent.monthly_summary import get_monthly_summary

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "voucher.png"
            source.write_bytes(b"fake-image")
            prepare_monthly_summary(root, task_id="mismatch", month="2026-05")
            with self.assertRaisesRegex(MonthlySummaryError, "2026-04"):
                collect_monthly_summary_source(
                    root,
                    "mismatch",
                    source,
                    ocr_adapter=StaticOcrAdapter(
                        [
                            "\u8d2d\u7535\u6708\u4efd\uff1a202604",
                            "\u603b\u7535\u91cf\uff1a6621.6",
                            "\u542b\u7a0e\u7535\u8d39\uff1a2075.55",
                        ]
                    ),
                )
            task = get_monthly_summary(root, "mismatch")
            self.assertEqual(task["status"], "collecting_file")
            self.assertIsNone(task["source_file"])

    def test_tax_components_are_only_an_auxiliary_warning(self) -> None:
        from submit_flow_agent.monthly_summary import _recognition_warnings

        self.assertEqual(
            _recognition_warnings(
                {"tax_inclusive_fee": "2075.55", "tax_exclusive_fee": "1836.77", "tax_amount": "238.78"}
            ),
            [],
        )
        self.assertEqual(len(_recognition_warnings(
            {"tax_inclusive_fee": "2075.55", "tax_exclusive_fee": "1800", "tax_amount": "238.78"}
        )), 1)


def _template(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.merge_cells("A1:L1")
    sheet["A1"] = "\u6c64\u897f\u656c\u8001\u9662"
    sheet.append(
        [
            "\u7535\u7ad9\u540d\u79f0",
            "\u552e\u7535\u7c7b\u578b",
            "\u7535\u4ef7\u8ba1\u8d39\u65b9\u5f0f",
            "\u6708\u4efd",
            "\u6d88\u7eb3\u7535\u91cf",
            "\u4e0a\u7f51\u7535\u91cf",
            "\u603b\u53d1\u7535\u91cf",
            "\u6d88\u7eb3\u6536\u76ca",
            "\u4e0a\u7f51\u7535\u4ef7",
            "\u4e0a\u7f51\u6536\u76ca",
            "\u603b\u6536\u76ca",
            "\u65e5\u5747\u6709\u6548\u5c0f\u65f6\u6570",
        ]
    )
    for month, energy, fee in (("2026-03", 7084.8, 3650.81), ("2026-04", 6621.6, 2075.55)):
        row = sheet.max_row + 1
        year, number = map(int, month.split("-"))
        sheet.append(
            [
                "\u656c\u8001\u9662\u5149\u4f0f\u9879\u76ee",
                "\u5168\u989d\u4e0a\u7f51",
                "\u56fa\u5b9a\u7535\u4ef7",
                __import__("datetime").date(year, number, 1),
                0,
                energy,
                energy,
                0,
                f"=J{row}/G{row}",
                fee,
                fee,
                f"=G{row}/86.14/31",
            ]
        )
    workbook.save(path)
    return path


def _collect(root: Path, task_id: str, month: str, energy: str, fee: str) -> None:
    source = root / f"{task_id}.png"
    source.write_bytes(b"fake-image")
    collect_monthly_summary_source(
        root,
        task_id,
        source,
        ocr_adapter=StaticOcrAdapter(
            [
                f"\u8d2d\u7535\u6708\u4efd\uff1a{month}",
                f"\u603b\u7535\u91cf\uff1a{energy}",
                f"\u542b\u7a0e\u7535\u8d39\uff1a{fee}",
            ]
        ),
    )


if __name__ == "__main__":
    unittest.main()
