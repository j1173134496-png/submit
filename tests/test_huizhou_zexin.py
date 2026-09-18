from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook

from submit_flow_agent.huizhou_zexin import (
    HuizhouZexinError,
    SELF_USE_UNIT_PRICE,
    ZexinValues,
    collect_huizhou_zexin_source,
    confirm_huizhou_zexin,
    list_huizhou_zexin_outputs,
    prepare_huizhou_zexin,
    record_huizhou_zexin_generation,
    recognize_huizhou_zexin_source,
    run_huizhou_zexin,
    write_settlement_docx,
)
from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult


TEMPLATES = Path("config/templates/huizhou_zexin")


class FakeAdapter:
    def __init__(self, blocks: list[str]) -> None:
        self.blocks = blocks

    def recognize(self, source_file: Path, pages: list[object]) -> OcrRawResult:
        return _raw(source_file, self.blocks)


def _raw(source: Path, texts: list[str]) -> OcrRawResult:
    blocks = [OcrBlock(text, (0.0, float(index * 10), 100.0, float(index * 10 + 8)), 0.99) for index, text in enumerate(texts)]
    return OcrRawResult(source, [OcrPage(1, source.with_suffix(".png"), blocks)])


def _pdf(path: Path, marker: bytes) -> Path:
    path.write_bytes(b"%PDF-1.4\n" + marker)
    return path


def _grid_blocks() -> list[str]:
    return [
        "2026-04-01",
        "0313060362996094",
        "57480.00",
        "32033.30",
        "13",
        "4164.33",
        "36197.63",
        "202604",
        "57480",
    ]


def _generation_blocks() -> list[str]:
    return ["2026-04-01", "0313060357636990", "4000.00", "157600.00", "0.00", "157600.00", "202604", "157600"]


def test_recognizes_both_pdf_roles_from_content() -> None:
    role, grid = recognize_huizhou_zexin_source(_raw(Path("grid.pdf"), _grid_blocks()), expected_month="2026-04")
    assert role == "grid_bill"
    assert grid["grid_energy"] == "57480"
    assert grid["grid_tax_exclusive_fee"] == "32033.30"
    assert grid["grid_tax_inclusive_fee"] == "36197.63"

    role, generation = recognize_huizhou_zexin_source(
        _raw(Path("generation.pdf"), _generation_blocks()),
        expected_month="2026-04",
    )
    assert role == "generation_bill"
    assert generation["total_generation"] == "157600"


def test_pdf_and_manual_generation_flow_is_order_independent_and_generates_two_outputs(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("submit_flow_agent.huizhou_zexin.render_pdf_to_images", lambda source, output: [])
    runtime = tmp_path / "runtime"
    task_id = "zexin_2026_04_test"
    prepare_huizhou_zexin(runtime, task_id=task_id, month="2026-04")
    grid = _pdf(tmp_path / "anything-b.pdf", b"grid")

    first = record_huizhou_zexin_generation(
        runtime,
        task_id,
        month="2026-04",
        total_generation="157600",
    )
    assert first["source_count"] == 1
    assert first["stage"] == "collecting_files"
    duplicate = record_huizhou_zexin_generation(
        runtime,
        task_id,
        month="2026-04",
        total_generation="157600",
    )
    assert duplicate["duplicate"] is True
    assert duplicate["source_count"] == 1

    reviewed = collect_huizhou_zexin_source(runtime, task_id, grid, ocr_adapter=FakeAdapter(_grid_blocks()))
    assert reviewed["stage"] == "awaiting_confirmation"
    data = reviewed["recognized_data"]
    assert data["self_use_energy"] == "100120"
    assert data["self_use_ratio"] == "0.635279"
    assert data["self_use_unit_price"] == str(SELF_USE_UNIT_PRICE)
    assert data["self_use_fee"] == "60072.00"
    assert data["total_income"] == "96269.63"

    confirm_huizhou_zexin(runtime, task_id)
    completed = run_huizhou_zexin(runtime, task_id, template_dir=TEMPLATES)
    assert completed["stage"] == "completed"
    outputs = [Path(item) for item in list_huizhou_zexin_outputs(runtime, task_id)]
    assert [output.name for output in outputs] == ["惠州泽鑫电费结算单.docx", "惠州泽鑫汇总表.xlsx"]
    assert all(output.is_file() and output.stat().st_size > 0 for output in outputs)
    assert not (outputs[0].parent / "惠阳供电局秋长供电所电量电费结算明细表.xlsx").exists()

    document = Document(outputs[0])
    assert document.tables[0].cell(2, 1).text == "157600"
    assert document.tables[0].cell(2, 3).text == "100120"
    assert document.tables[0].cell(2, 6).text == "60072"
    for row_number in range(3):
        for column_number in range(7):
            cell = document.tables[0].cell(row_number, column_number)
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER for paragraph in cell.paragraphs)
    for cell in (document.tables[0].cell(3, 0), document.tables[0].cell(3, 4)):
        assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
        assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert all(run.bold is True for run in cell.paragraphs[0].runs)

    summary = load_workbook(outputs[1], data_only=False).active
    assert summary["A4"].number_format == 'yyyy"年"m"月"'
    assert summary["B4"].value == 157600
    assert summary["C4"].value == 57480
    assert summary["E4"].value == "=C4*D4"
    assert summary["F4"].value == "=B4-C4"
    assert summary["H4"].value == "=F4*G4"
    assert summary["I4"].value == "=H4+E4"

    revision = prepare_huizhou_zexin(runtime, task_id="zexin_2026_04_revision_2", month="2026-04")
    assert revision["revision"] == 2


def test_manual_generation_rejects_wrong_month_and_generation_pdf(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("submit_flow_agent.huizhou_zexin.render_pdf_to_images", lambda source, output: [])
    runtime = tmp_path / "runtime"
    task_id = "zexin_manual_contract"
    prepare_huizhou_zexin(runtime, task_id=task_id, month="2026-04")

    with pytest.raises(HuizhouZexinError, match="与当前填报月份 2026-04 不一致"):
        record_huizhou_zexin_generation(
            runtime,
            task_id,
            month="2026-05",
            total_generation="157600",
        )

    generation = _pdf(tmp_path / "generation.pdf", b"generation")
    with pytest.raises(HuizhouZexinError, match="只需上传上网电费单 PDF"):
        collect_huizhou_zexin_source(
            runtime,
            task_id,
            generation,
            ocr_adapter=FakeAdapter(_generation_blocks()),
        )


def test_grid_pdf_can_arrive_before_manual_generation(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("submit_flow_agent.huizhou_zexin.render_pdf_to_images", lambda source, output: [])
    runtime = tmp_path / "runtime"
    task_id = "zexin_grid_first"
    prepare_huizhou_zexin(runtime, task_id=task_id, month="2026-04")
    grid = _pdf(tmp_path / "grid-first.pdf", b"grid")

    collected = collect_huizhou_zexin_source(
        runtime,
        task_id,
        grid,
        ocr_adapter=FakeAdapter(_grid_blocks()),
    )
    assert collected["source_count"] == 1
    assert collected["stage"] == "collecting_files"

    reviewed = record_huizhou_zexin_generation(
        runtime,
        task_id,
        month="2026-04",
        total_generation="157600",
    )
    assert reviewed["source_count"] == 2
    assert reviewed["stage"] == "awaiting_confirmation"
    assert reviewed["recognized_data"]["self_use_energy"] == "100120"


def test_docx_rounds_non_integer_fee_for_display_only(tmp_path: Path) -> None:
    values = ZexinValues(
        "2026-05",
        Decimal("157601"),
        Decimal("57480"),
        Decimal("32033.30"),
        Decimal("36197.63"),
    )
    output = tmp_path / "settlement.docx"
    write_settlement_docx(TEMPLATES / "电费结算单.docx", output, values)
    document = Document(output)
    assert values.self_use_fee == Decimal("60072.6")
    assert document.tables[0].cell(2, 6).text == "60073"
    assert document.tables[0].cell(3, 4).text == "购电方应付总电费：60073"


def test_incomplete_grid_bill_enters_review_instead_of_becoming_generation(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("submit_flow_agent.huizhou_zexin.render_pdf_to_images", lambda source, output: [])
    runtime = tmp_path / "runtime"
    task_id = "zexin_incomplete_grid"
    prepare_huizhou_zexin(runtime, task_id=task_id, month="2026-04")
    grid = _pdf(tmp_path / "grid.pdf", b"grid")
    record_huizhou_zexin_generation(
        runtime,
        task_id,
        month="2026-04",
        total_generation="157600",
    )
    incomplete = ["2026-04-01", "0313060362996094", "202604", "57480"]
    reviewed = collect_huizhou_zexin_source(runtime, task_id, grid, ocr_adapter=FakeAdapter(incomplete))
    assert reviewed["stage"] == "need_review"
    assert reviewed["sources"]["grid_bill"]["role"] == "grid_bill"
    assert reviewed["recognized_data"]["grid_energy"] == "57480"
    assert reviewed["missing_fields"] == ["grid_tax_exclusive_fee", "grid_tax_inclusive_fee"]
