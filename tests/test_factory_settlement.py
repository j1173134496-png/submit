from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from submit_flow_agent.factory_settlement import (
    FactorySettlementError,
    collect_factory_settlement_source,
    confirm_factory_settlement,
    confirm_factory_settlement_price,
    list_factory_settlement_outputs,
    prepare_factory_settlement,
    recognize_factory_workbook,
    run_factory_settlement,
    update_factory_settlement_price,
)


FIXTURES = Path("tests/fixtures/factory_meizhou_v1")
MULTI_FIXTURES = Path("tests/fixtures/factory_meizhou_multi_v1")
TEMPLATE = Path("config/templates/factory_meizhou_v1/梅州3号厂房汇总表.xlsx")
FEE_SOURCE = FIXTURES / "2026.5电费结算单（一期）.xlsx"
ENERGY_SOURCE = FIXTURES / "2026.5电量结算单(一期).xlsx"
JUNE_FEE_SOURCE = FIXTURES / "2026.6电费结算单（一期）.xlsx"
JUNE_ENERGY_SOURCE = FIXTURES / "2026.6电量结算单(一期).xlsx"
MULTI_FEE_SOURCE = MULTI_FIXTURES / "2026.5电费结算单（二期）.xlsx"
MULTI_ENERGY_SOURCE = MULTI_FIXTURES / "2026.5电量结算单（二期）.xlsx"


def test_recognizes_factory_workbooks_and_reconciles_source_values() -> None:
    fee_role, fee = recognize_factory_workbook(FEE_SOURCE, expected_month="2026-05")
    assert fee_role == "fee_settlement"
    assert fee["period_energy"] == {"尖": "2952", "峰": "50502", "平": "41342", "谷": "6660"}
    assert fee["self_use_energy"] == "101456"
    assert fee["self_use_fee"] == "31438.41"

    energy_role, energy = recognize_factory_workbook(ENERGY_SOURCE, expected_month="2026-05")
    assert energy_role == "energy_settlement"
    assert energy["generation_meters"] == ["97844", "98092"]
    assert energy["meter_total_generation"] == "195936"
    assert energy["total_generation"] == "101456"
    assert energy["grid_energy"] == "94480"
    assert energy["self_use_energy"] == "101456"


def test_recognizes_formula_backed_june_workbooks_with_excel_float_tolerance() -> None:
    formulas = load_workbook(JUNE_FEE_SOURCE, data_only=False)
    try:
        assert formulas.worksheets[0]["G10"].value == "=Sheet2!B9"
    finally:
        formulas.close()

    fee_role, fee = recognize_factory_workbook(JUNE_FEE_SOURCE, expected_month="2026-06")
    assert fee_role == "fee_settlement"
    assert fee["period_energy"] == {"尖": "6442", "峰": "46460", "平": "44426", "谷": "7404"}
    assert fee["self_use_energy"] == "104732"
    assert fee["self_use_fee"] == "32431.47"

    energy_role, energy = recognize_factory_workbook(JUNE_ENERGY_SOURCE, expected_month="2026-06")
    assert energy_role == "energy_settlement"
    assert energy["generation_meters"] == ["94512", "94780"]
    assert energy["meter_total_generation"] == "189292"
    assert energy["total_generation"] == "104732"
    assert energy["grid_energy"] == "84560"
    assert energy["self_use_energy"] == "104732"


def test_seed_june_summary_applies_display_only_period_override(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "factory_meizhou_2026_06"
    prepare_factory_settlement(runtime, task_id=task_id, site_name="梅州3号厂房", month="2026-06")
    confirm_factory_settlement_price(runtime, task_id)
    collect_factory_settlement_source(runtime, task_id, JUNE_FEE_SOURCE)
    reviewed = collect_factory_settlement_source(runtime, task_id, JUNE_ENERGY_SOURCE)
    assert reviewed["recognized_data"]["period_energy"] == {
        "尖": "6442",
        "峰": "46460",
        "平": "44426",
        "谷": "7404",
    }
    assert reviewed["recognized_data"]["self_use_fee"] == "32431.47"
    assert reviewed["recognized_data"]["total_fee"] == "70737.15"

    confirm_factory_settlement(runtime, task_id)
    run_factory_settlement(runtime, task_id, template_path=TEMPLATE)
    output = Path(list_factory_settlement_outputs(runtime, task_id)[0])
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook.active
        assert [sheet.cell(9, col).value for col in range(6, 10)] == [27354, 25548, 44426, 7404]
        assert sheet["L9"].value == pytest.approx(32431.47)
        assert sheet["M9"].value == "=E9+L9"
    finally:
        workbook.close()


def test_multi_meter_layout_creates_an_isolated_new_factory(tmp_path: Path) -> None:
    fee_role, fee = recognize_factory_workbook(MULTI_FEE_SOURCE, expected_month="2026-05")
    assert fee_role == "fee_settlement"
    assert fee["period_energy"] == {"尖": "889", "峰": "18741", "平": "14143", "谷": "3013"}
    assert fee["self_use_energy"] == "36786"
    assert fee["self_use_fee"] == "11334.89"

    energy_role, energy = recognize_factory_workbook(MULTI_ENERGY_SOURCE, expected_month="2026-05")
    assert energy_role == "energy_settlement"
    assert energy["generation_meters"] == ["78873", "157656", "77709", "80037", "77652", "78480", "77739"]
    assert energy["meter_total_generation"] == "628146"
    assert energy["total_generation"] == "36786"
    assert energy["grid_energy"] == "591360"

    runtime = tmp_path / "runtime"
    task_id = "factory_meizhou_multi_2026_05"
    prepare_factory_settlement(
        runtime,
        task_id=task_id,
        site_name="梅州2、4、5厂房",
        month="2026-05",
        allow_create=True,
    )
    update_factory_settlement_price(runtime, task_id, grid_unit_price="0.453")
    confirm_factory_settlement_price(runtime, task_id)
    collect_factory_settlement_source(runtime, task_id, MULTI_ENERGY_SOURCE)
    collect_factory_settlement_source(runtime, task_id, MULTI_FEE_SOURCE)
    confirm_factory_settlement(runtime, task_id)
    run_factory_settlement(runtime, task_id, template_path=TEMPLATE)
    output = Path(list_factory_settlement_outputs(runtime, task_id)[0])
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook.active
        assert sheet["A1"].value == "梅州2、4、5厂房"
        assert sheet["A5"].value.strftime("%Y-%m") == "2026-05"
        assert sheet["B5"].value == 36786
        assert sheet["C5"].value == 591360
        assert [sheet.cell(5, col).value for col in range(6, 10)] == [889, 18741, 14143, 3013]
        assert sheet["L5"].value == pytest.approx(11334.89)
        assert sheet["M5"].value == "=E5+L5"
        assert sheet["A9"].value is None
    finally:
        workbook.close()


def test_multi_meter_layout_reports_the_real_source_month_on_mismatch() -> None:
    with pytest.raises(FactorySettlementError, match="月份为 2026-05，.*2026-06"):
        recognize_factory_workbook(MULTI_ENERGY_SOURCE, expected_month="2026-06")
    with pytest.raises(FactorySettlementError, match="月份为 2026-05，.*2026-06"):
        recognize_factory_workbook(MULTI_FEE_SOURCE, expected_month="2026-06")


def test_seed_site_confirms_price_collects_both_orders_and_appends_may(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "factory_meizhou_2026_05"
    prepared = prepare_factory_settlement(
        runtime,
        task_id=task_id,
        site_name="梅州三号厂房",
        month="2026-05",
    )
    assert prepared["site_name"] == "梅州3号厂房"
    assert prepared["pricing"]["grid_unit_price"] == "0.453"
    assert prepared["pricing"]["source"] == "site_baseline"

    collecting = confirm_factory_settlement_price(runtime, task_id)
    assert collecting["stage"] == "collecting_files"
    first = collect_factory_settlement_source(runtime, task_id, FEE_SOURCE)
    assert first["source_count"] == 1
    assert first["stage"] == "collecting_files"
    duplicate = collect_factory_settlement_source(runtime, task_id, FEE_SOURCE)
    assert duplicate["duplicate"] is True
    reviewed = collect_factory_settlement_source(runtime, task_id, ENERGY_SOURCE)
    assert reviewed["source_count"] == 2
    assert reviewed["stage"] == "awaiting_confirmation"
    assert reviewed["recognized_data"]["grid_fee"] == "42799.44"
    assert reviewed["recognized_data"]["total_fee"] == "74237.85"
    assert all(check["status"] == "pass" for check in reviewed["validation_checks"])

    confirm_factory_settlement(runtime, task_id)
    completed = run_factory_settlement(runtime, task_id, template_path=TEMPLATE)
    assert completed["stage"] == "completed"
    outputs = [Path(item) for item in list_factory_settlement_outputs(runtime, task_id)]
    assert [item.name for item in outputs] == ["梅州3号厂房汇总表.xlsx"]

    workbook = load_workbook(outputs[0], data_only=False)
    sheet = workbook.active
    assert sheet["A1"].value == "梅州3号厂房"
    assert sheet["A5"].value.strftime("%Y-%m") == "2026-04"
    assert sheet["A9"].value.strftime("%Y-%m") == "2026-05"
    assert sheet["A5"].number_format == 'yyyy"年"m"月"'
    assert sheet["A9"].number_format == 'yyyy"年"m"月"'
    assert sheet["B5"].value == 171660
    assert sheet["C5"].value == 137920
    assert sheet["D5"].value == pytest.approx(0.453)
    assert sheet["E5"].value == "=C5*D5"
    assert [sheet.cell(5, col).value for col in range(6, 10)] == [10276, 9124, 11132, 3208]
    assert [sheet.cell(row, 11).value for row in range(5, 9)] == pytest.approx(
        [0.3984558, 0.3984558, 0.2343894, 0.0890712]
    )
    assert sheet["L5"].value == pytest.approx(10625.01)
    assert sheet["M5"].value == "=E5+L5"
    assert sheet["B9"].value == 101456
    assert sheet["C9"].value == 94480
    assert sheet["D9"].value == pytest.approx(0.453)
    assert sheet["E9"].value == "=C9*D9"
    assert [sheet.cell(9, col).value for col in range(6, 10)] == [2952, 50502, 41342, 6660]
    assert [sheet.cell(row, 10).value for row in range(9, 13)] == ["尖", "峰", "平", "谷"]
    assert sheet["L9"].value == pytest.approx(31438.41)
    assert sheet["M9"].value == "=E9+L9"


def test_new_factory_does_not_inherit_meizhou_price_or_april_history(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "factory_new_2026_05"
    prepared = prepare_factory_settlement(
        runtime,
        task_id=task_id,
        site_name="测试工业园厂房",
        month="2026-05",
        allow_create=True,
    )
    assert prepared["is_new_site"] is True
    assert prepared["pricing"]["grid_unit_price"] is None
    with pytest.raises(FactorySettlementError, match="没有历史上网电价"):
        confirm_factory_settlement_price(runtime, task_id)

    update_factory_settlement_price(runtime, task_id, grid_unit_price="0.5")
    confirm_factory_settlement_price(runtime, task_id)
    collect_factory_settlement_source(runtime, task_id, ENERGY_SOURCE)
    collect_factory_settlement_source(runtime, task_id, FEE_SOURCE)
    confirm_factory_settlement(runtime, task_id)
    run_factory_settlement(runtime, task_id, template_path=TEMPLATE)
    output = Path(list_factory_settlement_outputs(runtime, task_id)[0])
    workbook = load_workbook(output, data_only=False)
    sheet = workbook.active
    assert sheet["A1"].value == "测试工业园厂房"
    assert sheet["A5"].value.strftime("%Y-%m") == "2026-05"
    assert sheet["A5"].number_format == 'yyyy"年"m"月"'
    assert sheet["D5"].value == pytest.approx(0.5)
    assert sheet["A9"].value is None


def test_source_month_mismatch_is_rejected_without_registration(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "factory_wrong_month"
    prepare_factory_settlement(runtime, task_id=task_id, site_name="梅州3号厂房", month="2026-06")
    confirm_factory_settlement_price(runtime, task_id)
    with pytest.raises(FactorySettlementError, match="与当前填报月份 2026-06 不一致"):
        collect_factory_settlement_source(runtime, task_id, FEE_SOURCE)
