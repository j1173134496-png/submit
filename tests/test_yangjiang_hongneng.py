from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult
from submit_flow_agent.yangjiang_hongneng import (
    OUTPUT_DETAIL,
    OUTPUT_SUMMARY,
    PeriodReadings,
    _period_prices,
    _values_from_dict,
    collect_yangjiang_hongneng_source,
    confirm_yangjiang_hongneng,
    list_yangjiang_hongneng_outputs,
    prepare_yangjiang_hongneng,
    recognize_yangjiang_hongneng_meter_export,
    recognize_yangjiang_hongneng_source,
    run_yangjiang_hongneng,
)


TEMPLATE_DIR = Path("config/templates/yangjiang_hongneng")


class QueueOcr:
    def __init__(self, results: list[OcrRawResult]) -> None:
        self.results = iter(results)

    def recognize(self, source_file: Path, pages: list[object]) -> OcrRawResult:
        result = next(self.results)
        return OcrRawResult(source_file, [OcrPage(page.page, page.image_path, result.pages[0].blocks) for page in pages[:1]])


def _raw(name: str, rows: list[list[str]]) -> OcrRawResult:
    blocks: list[OcrBlock] = []
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            x = float(column_index * 220)
            y = float(row_index * 24)
            blocks.append(OcrBlock(value, (x, y, x + 200, y + 18), 0.99))
    source = Path(name)
    return OcrRawResult(source, [OcrPage(1, source.with_suffix(".png"), blocks)])


def _meter_raw(suffix: str, readings: tuple[str, str, str, str]) -> OcrRawResult:
    return _raw(
        f"{suffix}.png",
        [
            [f"阳江1.2逆变器{suffix}"],
            ["(上月)正向有功功率1电能（尖）", readings[0]],
            ["(上月)正向有功功率2电能（峰）", readings[1]],
            ["(上月)正向有功功率3电能（平）", readings[2]],
            ["(上月)正向有功功率4电能（谷）", readings[3]],
        ],
    )


def _bill_raw() -> OcrRawResult:
    rows = [["中国南方电网公司 广东电网公司 电费通知单"], ["用电开始时间：20260501"]]
    reading_values = {
        "尖": ("165.44", "167.70", "1000", "2260.00"),
        "峰": ("207.79", "212.44", "1000", "4650.00"),
        "平": ("375.80", "388.69", "1000", "12890.00"),
        "谷": ("131.44", "137.22", "1000", "5780.00"),
    }
    billed_values = {
        "尖": ("0.00", "0", "0.00"),
        "峰": ("6910.35", "0.64211700", "4437.25"),
        "平": ("12890.64", "0.37771600", "4869.00"),
        "谷": ("5780.29", "0.14353200", "829.66"),
    }
    component_rates = {
        "输配电费": {"尖": "0.08610000", "峰": "0.06890000", "平": "0.04050000", "谷": "0.01540000"},
        "线损电费": {"尖": "0.03360000", "峰": "0.02690000", "平": "0.01580000", "谷": "0.00600000"},
        "系统运行费用": {"尖": "0.12350000", "峰": "0.09880000", "平": "0.05810000", "谷": "0.02210000"},
    }
    for period in ("尖", "峰", "平", "谷"):
        rows.append([f"正有功{period}", *reading_values[period]])
        rows.append([f"电能电费({period})", *billed_values[period]])
        for label, rates in component_rates.items():
            rows.append([f"{label}({period})", billed_values[period][0], rates[period], "1.00"])
    rows.extend(
        [
            ["市场化分摊总费用", "0", "0", "952.35"],
            ["基金及附加费", "0.02766875", "707.81"],
        ]
    )
    return _raw("electricity.pdf", rows)


def _glyph_broken_bill_raw() -> OcrRawResult:
    rows = [["000317100301492144"], ["20260501"], ["Electricity Bill Information"]]
    readings = (
        ("165.44", "167.70", "1000", "2260.00"),
        ("207.79", "212.44", "1000", "4650.00"),
        ("375.80", "388.69", "1000", "12890.00"),
        ("131.44", "137.22", "1000", "5780.00"),
    )
    # Keep the real PDF's extra OCR tokens before the business columns.
    rows[2:2] = [["□", *values] for values in readings]
    groups = (
        (("0", "0", "0"), ("6910.35", "0.642117", "4437.25"), ("12890.64", "0.377716", "4869"), ("5780.29", "0.143532", "829.66")),
        (("10", "0", "0.0861", "0"), ("6910.35", "0.0689", "476.12"), ("0", "12890.64", "0.0405", "522.07"), ("5780.29", "0.0154", "89.02")),
        (("0", "0", "0", "10", "0", "0.0336", "0"), ("10", "6910.35", "0.0269", "185.89"), ("0", "0", "12890.64", "0.0158", "203.67"), ("5780.29", "0.006", "34.68")),
        (("10", "0", "0.1235", "0"), ("100", "0", "0", "10", "6910.35", "0.0988", "682.74"), ("0", "10", "12890.64", "0.0581", "748.95"), ("0", "0", "0", "5780.29", "0.0221", "127.74")),
    )
    rows.extend([list(values) for group in groups for values in group])
    rows.extend([["0", "0", "952.35"], ["0.79", "0.055", "1561.84", "0.02766875", "707.81"]])
    return _raw("0317100301492144.pdf", rows)


def _june_bill_raw() -> OcrRawResult:
    rows = [["000317100301492144"], ["20260601"], ["Electricity Bill Information"]]
    readings = (
        ("167.70", "170.42", "1000", "2720.00"),
        ("212.44", "217.78", "1000", "5340.00"),
        ("388.69", "401.72", "1000", "13030.00"),
        ("137.22", "143.63", "1000", "6410.00"),
    )
    rows[2:2] = [["□", *values] for values in readings]
    groups = (
        (("140.01", "0.804993", "112.71"), ("7920.40", "0.644202", "5102.33"), ("13030.65", "0.378942", "4937.86"), ("6410.32", "0.143998", "923.07")),
        (("140.01", "0.0861", "12.05"), ("7920.40", "0.0689", "545.72"), ("13030.65", "0.0405", "527.74"), ("6410.32", "0.0154", "98.72")),
        (("140.01", "0.0336", "4.70"), ("7920.40", "0.0269", "213.06"), ("13030.65", "0.0158", "205.88"), ("6410.32", "0.006", "38.46")),
        (("140.01", "0.0646", "9.04"), ("7920.40", "0.0517", "409.48"), ("13030.65", "0.0304", "396.13"), ("6410.32", "0.0116", "74.36")),
    )
    rows.extend([list(values) for group in groups for values in group])
    rows.extend([["17", "0", "0", "971.38"], ["0.9", "0.8", "0.05", "1441.03", "0.02766875", "760.94", "0.00"], ["630", "22.60", "0", "8"]])
    return _raw("0317100301492144-202606.pdf", rows)


def _settlement_raw() -> OcrRawResult:
    return _raw(
        "settlement.pdf",
        [
            ["电费结算单"],
            ["购电月份", "2026年05月"],
            ["抄见电量", "电价", "电费"],
            ["市场化电费", "42700", "0.216491", "9244.16"],
        ],
    )


def _image(path: Path) -> Path:
    color = (sum(path.name.encode("utf-8")) % 255, 120, 200)
    Image.new("RGB", (8, 8), color).save(path)
    return path


def _meter_export(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet["D2"] = "03591SG00000002100073001"
    sheet["G2"] = "2026-06-01"
    sheet["N2"] = 81.6
    for cell, value in zip(("O2", "P2", "Q2", "R2"), (23.69, 21.34, 34.38, 2.17), strict=True):
        sheet[cell] = value
    sheet["W2"] = 1000
    workbook.save(path)
    workbook.close()
    return path


def _legacy_meter_export(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet["D2"] = "03591SG00000002100073001"
    sheet["G2"] = "2026-05-01"
    sheet["J2"] = 38.9
    sheet["O2"] = 1000
    workbook.save(path)
    workbook.close()
    return path


def _meter_export_image_raw() -> OcrRawResult:
    return _raw(
        "电表抄表.png",
        [
            [
                "序号",
                "用户名称",
                "用户编号",
                "表计资产编号",
                "测量点号",
                "用户类别",
                "数据时间",
                "正向有功总(kWh)",
                "正向有功尖(kWh)",
                "正向有功峰(kWh)",
                "正向有功平(kWh)",
                "正向有功谷(kWh)",
                "正向无功总(kVarh)",
                "反向有功总(kWh)",
                "反向有功尖(kWh)",
                "反向有功峰(kWh)",
                "反向有功平(kWh)",
                "反向有功谷(kWh)",
                "反向无功总(kVarh)",
                "采集时间",
                "综合倍率",
            ],
            [
                "1",
                "阳江宏能科技有限公司",
                "031710038369919",
                "03171003836991903591SG0000000210007301",
                "1",
                "地方电厂户",
                "2026-07-01",
                "933.56",
                "170.42",
                "217.78",
                "401.72",
                "143.63",
                "525.66",
                "117.44",
                "33.81",
                "30.16",
                "49.6",
                "3.85",
                "52.71",
                "2026-07-01 01:34:33",
                "1000",
            ],
        ],
    )


def test_recognizes_all_five_roles(tmp_path: Path) -> None:
    role, meter = recognize_yangjiang_hongneng_source(
        _meter_raw("2065", ("219.03", "204.01", "308.94", "23.68")),
        expected_month="2026-05",
    )
    assert role == "generation_meter_2065"
    assert meter["current_readings"]["flat"] == "308.94"

    role, bill = recognize_yangjiang_hongneng_source(_bill_raw(), expected_month="2026-05")
    assert role == "electricity_bill"
    assert bill["purchased_current_readings"]["valley"] == "137.22"
    assert bill["billed_period_energy"]["peak"] == "6910.35"
    assert bill["fund_allocation_fee"] == "707.81"
    assert bill["market_allocation_fee"] == "952.35"
    assert bill["market_allocation_rate"] == "0.03722839513894535378995890745"

    role, settlement = recognize_yangjiang_hongneng_source(_settlement_raw(), expected_month="2026-05")
    assert role == "grid_settlement"
    assert settlement["grid_settlement_energy"] == "42700"
    assert settlement["grid_settlement_price"] == "0.216491"

    role, meter_export = recognize_yangjiang_hongneng_meter_export(
        _meter_export(tmp_path / "宏能.xlsx"),
        expected_month="2026-05",
    )
    assert role == "grid_meter_export"
    assert meter_export["reverse_total_reading"] == "81.6"
    assert meter_export["reverse_period_readings"] == {
        "sharp": "23.69",
        "peak": "21.34",
        "flat": "34.38",
        "valley": "2.17",
    }


def test_recognizes_grid_meter_export_from_image_for_june() -> None:
    role, meter_export = recognize_yangjiang_hongneng_source(
        _meter_export_image_raw(),
        expected_month="2026-06",
    )
    assert role == "grid_meter_export"
    assert meter_export == {
        "month": "2026-06",
        "grid_meter_number": "03591SG00000002100073001",
        "reading_date": "2026-07-01",
        "reverse_total_reading": "117.44",
        "reverse_period_readings": {
            "sharp": "33.81",
            "peak": "30.16",
            "flat": "49.6",
            "valley": "3.85",
        },
        "grid_meter_multiplier": "1000",
        "schema": "time_of_use_image_v1",
    }


def test_recognizes_bill_when_chinese_glyphs_are_missing() -> None:
    role, bill = recognize_yangjiang_hongneng_source(
        _glyph_broken_bill_raw(),
        expected_month="2026-05",
    )
    assert role == "electricity_bill"
    assert bill["price_components"]["peak"]["system"] == "0.0988"
    assert bill["price_components"]["valley"]["system"] == "0.0221"
    assert bill["fund_allocation_fee"] == "707.81"
    assert bill["market_allocation_fee"] == "952.35"


def test_june_bill_uses_item_17_and_excludes_sharp_from_market_basis() -> None:
    role, bill = recognize_yangjiang_hongneng_source(
        _june_bill_raw(),
        expected_month="2026-06",
    )
    assert role == "electricity_bill"
    assert bill["fund_allocation_fee"] == "760.94"
    assert bill["market_allocation_fee"] == "971.38"
    assert bill["market_allocation_rate"] == "0.03550187728172967947145921421"
    assert bill["sharp_tariff_active"] is False
    assert bill["price_components"]["sharp"]["energy"] == "0.804993"
    assert bill["price_components"]["valley"]["market"] == bill["market_allocation_rate"]


def test_april_legacy_export_uses_the_controlled_historical_period_seed(tmp_path: Path) -> None:
    role, meter_export = recognize_yangjiang_hongneng_meter_export(
        _legacy_meter_export(tmp_path / "宏能(1).xlsx"),
        expected_month="2026-04",
    )
    assert role == "grid_meter_export"
    assert meter_export["reverse_total_reading"] == "38.9"
    assert meter_export["reverse_period_readings"] == {
        "sharp": "11.69",
        "peak": "10.24",
        "flat": "16.31",
        "valley": "0.65",
    }


def test_full_may_flow_generates_only_two_workbooks(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "yangjiang_2026_05_test"
    prepared = prepare_yangjiang_hongneng(runtime, task_id=task_id, month="2026-05")
    assert prepared["settlement_factor"] == "0.67"

    raws = [
        _meter_raw("2065", ("219.03", "204.01", "308.94", "23.68")),
        _meter_raw("2066", ("219.69", "204.32", "309.29", "23.55")),
        _bill_raw(),
        _settlement_raw(),
    ]
    adapter = QueueOcr(raws)
    for index in range(4):
        task = collect_yangjiang_hongneng_source(
            runtime,
            task_id,
            _image(tmp_path / f"source-{index}.png"),
            ocr_adapter=adapter,
        )
    task = collect_yangjiang_hongneng_source(
        runtime,
        task_id,
        _meter_export(tmp_path / "宏能.xlsx"),
    )
    assert task["status"] == "awaiting_confirmation"
    assert task["source_count"] == 5
    recognized = task["recognized_data"]
    assert recognized["generation_period_energy"] == {
        "sharp": "20459",
        "peak": "19755",
        "flat": "30018",
        "valley": "2755",
    }
    assert recognized["grid_meter_period_energy"] == {
        "sharp": "12000",
        "peak": "11100",
        "flat": "18070",
        "valley": "1520",
    }
    assert recognized["grid_meter_period_total_energy"] == "42690"
    assert recognized["grid_meter_total_energy"] == "42700"
    assert recognized["self_use_period_energy"] == {
        "sharp": "8459",
        "peak": "8655",
        "flat": "11948",
        "valley": "1235",
    }
    assert recognized["self_use_energy"] == "30297"
    assert recognized["summary_total_generation"] == "72987"
    assert recognized["period_prices"]["sharp"] == recognized["period_prices"]["peak"]
    assert recognized["sharp_price_source"] == "peak_fallback"
    assert recognized["self_use_fee_before_tax_adjustment"] == "15005.69"
    assert recognized["grid_income"] == "9242.00"
    assert recognized["total_income"] == "24247.69"

    legacy = json.loads(json.dumps(recognized))
    legacy["price_components"].pop("sharp")
    legacy["period_prices"]["sharp"] = legacy["period_prices"]["peak"]
    restored = _values_from_dict(legacy, expected_month="2026-05")
    assert restored.period_prices.sharp == restored.period_prices.peak
    assert restored.self_use_fee_before_tax_adjustment.quantize(Decimal("0.01")) == Decimal("15005.69")

    confirmed = confirm_yangjiang_hongneng(runtime, task_id)
    assert confirmed["status"] == "ready_to_run"
    completed = run_yangjiang_hongneng(runtime, task_id, template_dir=TEMPLATE_DIR)
    assert completed["status"] == "completed"
    outputs = [Path(path) for path in list_yangjiang_hongneng_outputs(runtime, task_id)]
    assert [path.name for path in outputs] == [OUTPUT_DETAIL, OUTPUT_SUMMARY]

    detail = load_workbook(outputs[0], data_only=False)
    try:
        sheet = detail.active
        assert "五月份电量对账单" in sheet["A1"].value
        assert sheet["C5"].value == 219.69
        assert sheet["C6"].value == 219.03
        assert sheet["I17"].value == "=I16/1.06*1.01"
        assert sheet["F12"].value == "=U6"
        assert sheet["S4"].value == "=952.35/(6910.35+12890.64+5780.29)"
        assert sheet["T4"].value == "=SUM(N4:S4)"
        assert sheet["U6"].value == "=T6*0.67"
    finally:
        detail.close()

    summary = load_workbook(outputs[1], data_only=False)
    try:
        sheet = summary.active
        assert sheet["B12"].value.strftime("%Y-%m") == "2026-05"
        assert sheet["C12"].value == 72987
        assert sheet["D12"].value == 42690
        assert sheet["F12"].value == 30297
        assert sheet["G12"].value == "尖"
        assert sheet["G13"].value == "峰"
        assert sheet["G14"].value == "平"
        assert sheet["G15"].value == "谷"
        assert sheet["H12"].value == sheet["H13"].value
        assert sheet["H12"].number_format == "0.000000000"
        assert sheet["I12"].value == 15005.69
        assert sheet["I12"].number_format == "0.00_);[Red]\\(0.00\\)"
        assert sheet["M12"].value == "=C12/600/31"
        assert "A12:A15" in {str(merged) for merged in sheet.merged_cells.ranges}
        assert sheet["G15"].border.bottom.style == "thin"
    finally:
        summary.close()


def test_uses_actual_sharp_price_when_bill_has_effective_sharp_energy() -> None:
    components = {
        key: {
            "energy": Decimal(energy),
            "transmission": Decimal("0.01"),
            "line_loss": Decimal("0.01"),
            "system": Decimal("0.01"),
            "fund": Decimal("0.01"),
            "market": Decimal("0.01"),
        }
        for key, energy in {
            "sharp": "0.90",
            "peak": "0.60",
            "flat": "0.40",
            "valley": "0.20",
        }.items()
    }
    billed = PeriodReadings(Decimal("1"), Decimal("2"), Decimal("3"), Decimal("4"))
    june_prices = _period_prices("2026-06", components, billed)
    assert june_prices.sharp == june_prices.peak

    prices = _period_prices("2026-07", components, billed)
    assert prices.sharp != prices.valley
    assert prices.sharp == Decimal("0.6365")


def test_february_effective_hours_use_calendar_days(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    try:
        prepare_yangjiang_hongneng(runtime, task_id="leap", month="2028-02")
    except Exception as exc:
        assert "缺少 2028-01 已完成数据" in str(exc)
    else:
        raise AssertionError("cross-month cumulative readings must require the prior month")
