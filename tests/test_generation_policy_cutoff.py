from copy import deepcopy
from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import shutil
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from submit_flow_agent.excel_writer import (
    _meter_rows_summary_writes,
    _two_meter_ak_summary_writes,
    write_detail_excel,
    write_summary_excel,
)
from submit_flow_agent.field_mapper import build_manual_baseline_recognized_data
from submit_flow_agent.normalizer import NormalizationError, normalize_recognized_data, uses_formula_generation
from submit_flow_agent.validator import validate_monthly_data
from tests.pricing_helpers import static_pricing_snapshot


PROJECT = Path('深圳新安中学高中部650KW分布式光伏项目')


@pytest.mark.parametrize('month,expected', [('2026-07',False),('2026-08',True),('2026-09',True),('2027-01',True)])
def test_policy_month_boundary(month,expected):
    assert uses_formula_generation(month) is expected


@pytest.mark.parametrize('month',['2026-00','2026-13','not-a-month'])
def test_invalid_policy_month(month):
    with pytest.raises(NormalizationError): uses_formula_generation(month)


@pytest.mark.parametrize('month,expected', [('2026-07',Decimal('49567.5')),('2026-08',Decimal('49568'))])
def test_rounding_does_not_rewrite_july_source(month,expected):
    payload=build_manual_baseline_recognized_data(PROJECT).to_dict()
    payload['month']=month
    payload['meter_rows'][0]['generation_kwh']['value']=49567.5
    result=normalize_recognized_data(payload)
    assert result.total_generation_kwh==expected
    assert result.meter_rows[0].reported_generation_kwh==Decimal('49567.5')


@pytest.mark.parametrize('month', ['2026-07','2026-08'])
def test_detail_and_summary_use_same_policy_and_preserve_earlier_rows(tmp_path,month):
    payload=build_manual_baseline_recognized_data(PROJECT).to_dict()
    payload['month']=month
    payload['fields']['period_start']['value']=month+'-01'
    payload['fields']['period_end']['value']='2026-08-01' if month=='2026-07' else '2026-09-01'
    report=validate_monthly_data(normalize_recognized_data(payload)).to_dict()
    assert report['can_write_excel']
    report_path=tmp_path/'validation.json'
    report_path.write_text(json.dumps(report,ensure_ascii=False),encoding='utf8')
    project=tmp_path/'project';project.mkdir()
    for file in PROJECT.glob('*.xlsx'):shutil.copy2(file,project/file.name)
    baseline=project/'汇总表格.xlsx'
    w=load_workbook(baseline);s=w['Sheet1']
    before=[[(c.value,deepcopy(c._style)) for c in row] for row in s.iter_rows(min_row=4,max_row=5,max_col=13)]
    w.close()
    price=static_pricing_snapshot(month=month)
    detail=tmp_path/'detail.xlsx';summary=tmp_path/'summary.xlsx'
    write_detail_excel(report_path,project,detail,pricing_snapshot=price)
    result=write_summary_excel(report_path,project,summary,pricing_snapshot=price)
    w=load_workbook(summary);s=w['Sheet1'];r=result.target_row
    assert [[(c.value,c._style) for c in row] for row in s.iter_rows(min_row=4,max_row=5,max_col=13)]==before
    assert s[f'F{r}'].value==(f'=ROUND((D{r}-C{r})*E{r},0)' if month>='2026-08' else 49568)
    w.close()
    w=load_workbook(detail);s=w.worksheets[0]
    assert s['E5'].value==('=ROUND((C5-B5)*D5,0)' if month>='2026-08' else 49568)
    w.close()


def test_both_dual_summary_profiles_preserve_legacy_generation():
    row={'meter_no':'m1','previous_reading':1,'current_reading':2.5,'multiplier':3,'generation_kwh':4.5,'grid_export_kwh':1}
    site=SimpleNamespace(summary_start_row=4,meter_topology_profile='dual_meter_school')
    price={'base_unit_price':Decimal('1'),'discount_rate':Decimal('.9'),'effective_unit_price':Decimal('.6')}
    for writer in (_meter_rows_summary_writes,_two_meter_ak_summary_writes):
        values=writer(site,{'grid_price':.3},[row,{**row,'meter_no':'m2'}],date(2026,7,1),4,price,formula_generation=False)
        assert values['F4']==4.5
        assert values['F5']==4.5
