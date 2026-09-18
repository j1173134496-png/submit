from __future__ import annotations

import json
import tempfile
import unittest
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from submit_flow_agent.excel_writer import ExcelWriteError, _price_header_text, write_detail_excel
from submit_flow_agent.field_mapper import build_manual_baseline_recognized_data
from submit_flow_agent.normalizer import normalize_recognized_data
from submit_flow_agent.validator import validate_monthly_data
from tests.pricing_helpers import static_pricing_snapshot


PROJECT_DIR = Path("深圳新安中学高中部650KW分布式光伏项目")
TEMPLATE = PROJECT_DIR / "新安中学（高中部）光伏发电项目学校用电统计表.xlsx"


class DetailExcelWriterTests(unittest.TestCase):
    def test_price_header_hides_binary_float_artifact(self) -> None:
        self.assertEqual(_price_header_text(Decimal("0.8800000000000001")), "8.8折单价（元/kWh）")
        self.assertEqual(_price_header_text(Decimal("0.8699999999999999")), "8.7折单价（元/kWh）")

    def test_write_detail_excel_copy(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            validation_path = root / "validation_report.json"
            output_path = root / "新安中学（高中部）光伏发电项目学校用电统计表_2026-05.xlsx"
            validation_path.write_text(json.dumps(_validation_report(), ensure_ascii=False), encoding="utf-8")
            original_size = TEMPLATE.stat().st_size

            result = write_detail_excel(validation_path, PROJECT_DIR, output_path, pricing_snapshot=static_pricing_snapshot())

            self.assertTrue(output_path.exists())
            self.assertEqual(result.month, "2026-05")
            self.assertEqual(result.sheet, "1")
            self.assertIn("A1", result.written_cells)
            self.assertEqual(TEMPLATE.stat().st_size, original_size)

            workbook = load_workbook(output_path, data_only=False)
            self.assertIn("1", workbook.sheetnames)
            worksheet = workbook["1"]
            self.assertEqual(worksheet["A1"].value, "新安中学（高中部）光伏发电项目5月学校用电统计表")
            self.assertEqual(worksheet["E2"].value.strftime("%Y-%m-%d"), "2026-05-01")
            self.assertEqual(worksheet["H2"].value.strftime("%Y-%m-%d"), "2026-06-01")
            self.assertEqual(worksheet["A5"].value, "表号09001SF00000042508945216")
            self.assertEqual(worksheet["B5"].value, 816.82)
            self.assertEqual(worksheet["C5"].value, 1147.27)
            self.assertEqual(worksheet["D5"].value, 150)
            self.assertEqual(worksheet["E5"].value, 49568)
            self.assertEqual(worksheet["F5"].value, 1680)
            self.assertEqual(worksheet["G5"].value, "=E5-F5")
            self.assertEqual(worksheet["H5"].value, 0.6998)
            self.assertEqual(worksheet["I5"].value, "=G5*H5")
            self.assertEqual(worksheet["J5"].value, "=I5")
            self.assertEqual(worksheet["A6"].value, "福域(深圳)新能源有限公司")
            self.assertEqual(worksheet["A7"].value.strftime("%Y-%m-%d"), "2026-06-21")
            workbook.close()

    def test_write_detail_excel_blocks_failed_validation(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            validation_path = root / "validation_report.json"
            output_path = root / "blocked.xlsx"
            report = _validation_report()
            report["status"] = "fail"
            report["can_write_excel"] = False
            report["errors"] = ["发电量公式校验失败"]
            validation_path.write_text(json.dumps(report, ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(ExcelWriteError, "blocks Excel writing"):
                write_detail_excel(validation_path, PROJECT_DIR, output_path)

            self.assertFalse(output_path.exists())

    def test_write_detail_excel_refuses_template_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            validation_path = Path(tmpdir) / "validation_report.json"
            validation_path.write_text(json.dumps(_validation_report(), ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(ExcelWriteError, "Refusing to overwrite"):
                write_detail_excel(validation_path, PROJECT_DIR, TEMPLATE, pricing_snapshot=static_pricing_snapshot())


def _validation_report() -> dict[str, object]:
    recognized = build_manual_baseline_recognized_data(PROJECT_DIR).to_dict()
    report = validate_monthly_data(normalize_recognized_data(recognized)).to_dict()
    return deepcopy(report)


if __name__ == "__main__":
    unittest.main()
