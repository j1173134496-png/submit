from __future__ import annotations

import json
import shutil
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from submit_flow_agent.config import load_site_config
from submit_flow_agent.excel_writer import write_summary_excel
from submit_flow_agent.historical_backfill import _pricing_snapshot, backfill_school_history


def test_rehearsal_builds_cumulative_july_baseline_without_production_write(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _source_workbook(source)
    config = tmp_path / "sites.json"
    payload = json.loads(Path("config/sites.json").read_text(encoding="utf-8-sig"))
    target = next(item for item in payload["sites"] if item["site_key"] == "site_5f14cc7ab0e2")
    template_dir = tmp_path / "templates"
    shutil.copytree(Path("config") / target["template_dir"], template_dir)
    target["template_dir"] = str(template_dir)
    config.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    runtime = tmp_path / "production-runtime"
    route = tmp_path / "site-routes.json"
    route.write_text(
        json.dumps(
            {"canonicalSiteRoutes": {"route": {"siteKey": "xinan_high_school", "label": "光伏填报｜新安高中部"}}},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    result = backfill_school_history(
        source,
        site_key="site_5f14cc7ab0e2",
        through_month="2026-07",
        runtime_root=runtime,
        config_path=config,
        verified_site_key="xinan_high_school",
        route_path=route,
        rehearsal_root=tmp_path / "rehearsal",
    )

    assert result["production_write"] is False
    assert result["rehearsal"]["summary_months"] == [
        "2026-01",
        "2026-02",
        "2026-03",
        "2026-04",
        "2026-05",
        "2026-06",
        "2026-07",
    ]
    assert not (runtime / "site_history").exists()
    history = Path(result["history_dir"])
    manifest = json.loads((history / "manifest.json").read_text(encoding="utf-8"))
    summary = history / manifest["workbooks"]["summary_workbook"]["path"]
    workbook = load_workbook(summary, data_only=False)
    sheet = workbook["Sheet1"]
    assert [sheet[f"F{row}"].value for row in range(4, 11)] == [6, 25851, 49623, 47043, 49567.5, 53966, 47552]
    assert all("390320" in sheet[f"A{row}"].value for row in range(4, 11))
    assert [sheet[f"M{row}"].value for row in range(4, 9)] == [
        3.98886,
        12875.24915,
        30945.68939,
        30689.26747,
        32243.379675,
    ]
    workbook.close()

    values_workbook = load_workbook(summary, data_only=True)
    values_sheet = values_workbook["Sheet1"]
    assert all(values_sheet[f"I{row}"].value is not None for row in range(4, 11))
    assert all(values_sheet[f"J{row}"].value is not None for row in range(4, 11))
    assert all(values_sheet[f"L{row}"].value is not None for row in range(4, 11))
    assert all(values_sheet[f"M{row}"].value is not None for row in range(4, 11))
    july_values = [values_sheet[f"{column}10"].value for column in "ABCDEFGHIJKLM"]
    values_workbook.close()

    site = load_site_config("site_5f14cc7ab0e2", config)
    retroactive_project = tmp_path / "retroactive-project"
    retroactive_project.mkdir()
    shutil.copy2(summary, retroactive_project / site.summary_workbook)
    replacement = write_summary_excel(
        tmp_path / "rehearsal" / "reports" / "2026-06.json",
        retroactive_project,
        tmp_path / "retroactive-june.xlsx",
        site_key=site.site_key,
        config_path=config,
        pricing_snapshot=_pricing_snapshot(
            site.site_key,
            {"month": "2026-06", "effective_unit_price": Decimal("0.66481")},
        ),
    )
    assert replacement.mode == "update"
    assert replacement.target_row == 9
    replaced = load_workbook(replacement.output, data_only=True)
    replaced_sheet = replaced[site.summary_sheet]
    assert [
        f"{replaced_sheet[f'B{row}'].value.year:04d}-{replaced_sheet[f'B{row}'].value.month:02d}"
        for row in range(4, 11)
    ] == result["months"]
    assert replaced_sheet["F9"].value == 53966
    assert [replaced_sheet[f"{column}10"].value for column in "ABCDEFGHIJKLM"] == july_values
    replaced.close()


def test_apply_replace_preserves_prior_revision_and_updates_series(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _source_workbook(source)
    config = tmp_path / "sites.json"
    payload = json.loads(Path("config/sites.json").read_text(encoding="utf-8-sig"))
    target = next(item for item in payload["sites"] if item["site_key"] == "site_5f14cc7ab0e2")
    template_dir = tmp_path / "templates"
    shutil.copytree(Path("config") / target["template_dir"], template_dir)
    target["template_dir"] = str(template_dir)
    config.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    runtime = tmp_path / "runtime"

    first = backfill_school_history(
        source,
        site_key="site_5f14cc7ab0e2",
        through_month="2026-07",
        runtime_root=runtime,
        config_path=config,
        apply=True,
        rehearsal_root=tmp_path / "rehearsal-1",
    )
    second = backfill_school_history(
        source,
        site_key="site_5f14cc7ab0e2",
        through_month="2026-07",
        runtime_root=runtime,
        config_path=config,
        apply=True,
        replace_existing=True,
        rehearsal_root=tmp_path / "rehearsal-2",
    )

    history = runtime / "site_history" / "site_5f14cc7ab0e2" / "2026-07"
    manifest = json.loads((history / "manifest.json").read_text(encoding="utf-8"))
    series = json.loads(
        (runtime / "site_series" / "site_5f14cc7ab0e2" / "2026-07" / "series.json").read_text(encoding="utf-8")
    )
    assert first["revision"] == 1
    assert second["revision"] == 2
    assert (history / "revisions" / "r0001").is_dir()
    assert (history / "revisions" / "r0002").is_dir()
    assert manifest["effective_revision"] == 2
    assert series["effective_revision"] == 2
    assert Path(second["backup_root"]).is_dir()


def _source_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    rows = [
        ("2026-01", 0, 0.04, 6, 0, 0),
        ("2026-02", 0.04, 172.38, 25851, 7120, 0.059367),
        ("2026-03", 172.38, 503.2, 49623, 4760, 0.235361),
        ("2026-04", 503.2, 816.82, 47043, 2320, 0.412487),
        ("2026-05", 816.82, 1147.27, 49567.5, 1680, 0.242435),
        ("2026-06", 1147.27, 1507.04, 53966, 3520, 0.53608),
        ("2026-07", 1507.04, 1824.05, 47552, 10840, 0.414),
    ]
    for row, (month, previous, current, generation, grid, price) in enumerate(rows, start=4):
        sheet.cell(row, 1, "表号09001SF00000042509390320")
        year, number = (int(part) for part in month.split("-"))
        sheet.cell(row, 2, date(year, number, 1))
        sheet.cell(row, 3, previous)
        sheet.cell(row, 4, current)
        sheet.cell(row, 5, 150)
        sheet.cell(row, 6, generation)
        sheet.cell(row, 7, grid)
        sheet.cell(row, 8, price)
        sheet.cell(row, 9, grid * price)
        sheet.cell(row, 11, 0.66481)
    workbook.save(path)
    workbook.close()
