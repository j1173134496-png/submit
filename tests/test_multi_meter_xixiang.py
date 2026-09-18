from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from submit_flow_agent.excel_writer import write_detail_excel, write_summary_excel
from submit_flow_agent.field_mapper import build_ocr_recognized_data, write_recognized_data
from submit_flow_agent.normalizer import normalize_recognized_data
from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult, write_ocr_raw_result
from submit_flow_agent.validator import validate_data_to_file, validate_monthly_data
from tests.pricing_helpers import static_pricing_snapshot


class XixiangMultiMeterTests(unittest.TestCase):
    def test_electronic_statement_uses_header_columns_for_meter_generation(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            project_dir = root / "project"
            project_dir.mkdir()
            for name in ("华盛腾发电单202605.pdf", "华盛腾电费结算单202605.pdf", "华盛腾电量结算单202605.pdf"):
                (project_dir / name).write_bytes(b"%PDF-1.4\n")
            raw_dir = _write_xixiang_ocr_raw(root, project_dir, generation_blocks=_electronic_generation_blocks())

            payload = build_ocr_recognized_data(
                project_dir,
                raw_dir,
                site_key="xixiang_middle_school",
            ).to_dict()

            self.assertEqual(payload["meter_rows"][0]["previous_reading"]["value"], 1655.88)
            self.assertEqual(payload["meter_rows"][0]["current_reading"]["value"], 1926.96)
            self.assertEqual(payload["meter_rows"][0]["multiplier"]["value"], 200)
            self.assertEqual(payload["meter_rows"][0]["generation_kwh"]["value"], 54216)
            self.assertEqual(payload["meter_rows"][1]["previous_reading"]["value"], 1882.44)
            self.assertEqual(payload["meter_rows"][1]["current_reading"]["value"], 2176.25)
            self.assertEqual(payload["meter_rows"][1]["multiplier"]["value"], 150)
            self.assertEqual(payload["meter_rows"][1]["generation_kwh"]["value"], 44072)
            self.assertEqual(payload["fields"]["total_generation_kwh"]["value"], 98288)
            self.assertEqual(payload["fields"]["grid_export_kwh"]["value"], 10140)

    def test_split_meter_number_keeps_each_meter_bound_to_its_visual_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            project_dir = root / "project"
            project_dir.mkdir()
            for name in ("华盛腾发电单202605.pdf", "华盛腾电费结算单202605.pdf", "华盛腾电量结算单202605.pdf"):
                (project_dir / name).write_bytes(b"%PDF-1.4\n")
            blocks = _electronic_generation_blocks()
            # Match the real Shajing layout: the suffix sits below the row,
            # while the first fragment and numeric cells identify the row.
            blocks[12] = OcrBlock("08945216", (130, 305, 250, 335), 0.99)
            blocks[21] = OcrBlock("08945217", (130, 405, 250, 435), 0.99)
            raw_dir = _write_xixiang_ocr_raw(root, project_dir, generation_blocks=blocks)

            payload = build_ocr_recognized_data(
                project_dir,
                raw_dir,
                site_key="xixiang_middle_school",
            ).to_dict()

            self.assertEqual(payload["meter_rows"][0]["previous_reading"]["value"], 1655.88)
            self.assertEqual(payload["meter_rows"][0]["generation_kwh"]["value"], 54216)
            self.assertEqual(payload["meter_rows"][1]["previous_reading"]["value"], 1882.44)
            self.assertEqual(payload["meter_rows"][1]["generation_kwh"]["value"], 44072)

    def test_xixiang_two_meter_rows_validate_and_write_excel(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            project_dir = root / "project"
            project_dir.mkdir()
            for name in ("华盛腾发电单202605.pdf", "华盛腾电费结算单202605.pdf", "华盛腾电量结算单202605.pdf"):
                (project_dir / name).write_bytes(b"%PDF-1.4\n")
            for template in Path("config/templates/xixiang_middle_school").glob("*.xlsx"):
                shutil.copy2(template, project_dir / template.name)
            raw_dir = _write_xixiang_ocr_raw(root, project_dir)
            recognized_path = root / "recognized_data.json"
            validation_path = root / "validation_report.json"

            recognized = build_ocr_recognized_data(
                project_dir,
                raw_dir,
                site_key="xixiang_middle_school",
            )
            write_recognized_data(recognized, recognized_path)
            payload = recognized.to_dict()
            monthly = normalize_recognized_data(payload)
            report = validate_monthly_data(monthly)

            self.assertEqual(len(payload["meter_rows"]), 2)
            self.assertEqual(payload["fields"]["total_generation_kwh"]["value"], 120767)
            self.assertEqual(payload["fields"]["grid_export_kwh"]["value"], 10140)
            self.assertEqual(payload["meter_rows"][0]["school_consumption_kwh"]["value"], 53045)
            self.assertEqual(monthly.calculated_school_fee, Decimal("77416.77"))
            self.assertEqual(monthly.calculated_total_fee, Decimal("81148.340700"))
            self.assertEqual(report.status, "pass")
            self.assertTrue(report.can_write_excel)
            total_check = next(check for check in report.checks if check.name == "total_fee")
            self.assertEqual(total_check.message, "双电表汇总合计已按上网电费和各电表学校消纳电费相加计算。")

            validate_data_to_file(recognized_path, validation_path)
            detail_path = root / "detail.xlsx"
            summary_path = root / "summary.xlsx"
            write_detail_excel(validation_path, project_dir, detail_path, site_key="xixiang_middle_school", pricing_snapshot=static_pricing_snapshot("xixiang_middle_school"))
            write_summary_excel(validation_path, project_dir, summary_path, site_key="xixiang_middle_school", pricing_snapshot=static_pricing_snapshot("xixiang_middle_school"))

            detail = load_workbook(detail_path, data_only=False)
            detail_ws = detail["1"]
            self.assertEqual(detail_ws["A5"].value, "表号09001SF00000042508945216")
            self.assertEqual(detail_ws["E5"].value, 63185)
            self.assertEqual(detail_ws["F5"].value, 10140)
            self.assertEqual(detail_ws["G5"].value, "=E5-F5")
            self.assertEqual(detail_ws["H5"].value, 0.6998)
            self.assertEqual(detail_ws["A6"].value, "表号09001SF00000042508945217")
            self.assertEqual(detail_ws["E6"].value, 57582)
            self.assertEqual(detail_ws["F6"].value, 0)
            self.assertEqual(detail_ws["G6"].value, "=E6-F6")
            self.assertEqual(detail_ws["H6"].value, 0.6998)
            self.assertEqual(detail_ws["J5"].value, "=ROUND(I5+I6,2)")
            self.assertEqual(detail_ws["A7"].value, "深圳市华盛腾新能源科技有限公司")
            self.assertIsNotNone(detail_ws["A8"].value)
            self.assertIn("J5:J6", [str(item) for item in detail_ws.merged_cells.ranges])
            detail.close()

            summary = load_workbook(summary_path, data_only=False)
            summary_ws = summary["Sheet1"]
            self.assertEqual(summary_ws["A1"].value, "深圳市西乡中学")
            self.assertEqual(summary_ws["A4"].value, "表号09001SF00000042508945216")
            self.assertEqual(summary_ws["A5"].value, "表号09001SF00000042508945217")
            self.assertEqual(summary_ws["A6"].value, "表号09001SF00000042508945216")
            self.assertEqual(summary_ws["A7"].value, "表号09001SF00000042508945217")
            self.assertEqual(summary_ws["F6"].value, 63185)
            self.assertEqual(summary_ws["F7"].value, 57582)
            self.assertEqual(summary_ws["G6"].value, 10140)
            self.assertEqual(summary_ws["G7"].value, 0)
            self.assertEqual(summary_ws["K6"].value, 0.6998)
            self.assertEqual(summary_ws["J6"].value, 53045)
            self.assertEqual(summary_ws["M4"].value, "=L4+L5")
            self.assertEqual(summary_ws["M6"].value, 81148.3407)
            self.assertIn("M6:M7", [str(item) for item in summary_ws.merged_cells.ranges])
            summary.close()


def _write_xixiang_ocr_raw(
    root: Path,
    project_dir: Path,
    *,
    generation_blocks: list[OcrBlock] | None = None,
) -> Path:
    raw_dir = root / "raw_ocr"
    raw_dir.mkdir()
    samples = {
        "华盛腾发电单202605.pdf": generation_blocks or [
            OcrBlock("上期抄表时间", (1214, 258, 1304, 276), 0.99),
            OcrBlock("本期抄表时间", (1017, 258, 1108, 275), 0.99),
            OcrBlock("2026.05.01", (1214, 300, 1304, 318), 0.99),
            OcrBlock("2026.06.01", (1017, 300, 1108, 318), 0.99),
            OcrBlock("09001SF00000042508945216", (88, 365, 324, 379), 0.99),
            OcrBlock("2095.11", (386, 361, 455, 381), 0.99),
            OcrBlock("2516.34", (515, 362, 584, 380), 0.99),
            OcrBlock("150", (668, 362, 703, 382), 0.99),
            OcrBlock("63185", (827, 363, 881, 381), 0.99),
            OcrBlock("09001SF00000042508945217", (87, 403, 323, 417), 0.99),
            OcrBlock("1446.67", (388, 401, 456, 419), 0.99),
            OcrBlock("1734.58", (517, 401, 584, 419), 0.99),
            OcrBlock("200", (668, 400, 704, 420), 0.99),
            OcrBlock("57582", (826, 400, 882, 421), 0.99),
        ],
        "华盛腾电量结算单202605.pdf": [
            OcrBlock("合计", (241, 328, 282, 352), 0.99),
            OcrBlock("10140", (1290, 330, 1359, 351), 0.99),
        ],
        "华盛腾电费结算单202605.pdf": [
            OcrBlock("结算小计", (333, 455, 398, 477), 0.99),
            OcrBlock("10140", (818, 456, 863, 479), 0.99),
            OcrBlock("0.368005", (1055, 457, 1123, 478), 0.99),
            OcrBlock("3731.57", (1282, 455, 1343, 479), 0.99),
        ],
    }
    for filename, blocks in samples.items():
        result = OcrRawResult(
            source_file=project_dir / filename,
            pages=[OcrPage(page=1, image_path=Path("images") / f"{Path(filename).stem}-1.png", blocks=blocks)],
        )
        write_ocr_raw_result(result, raw_dir / f"{Path(filename).stem}.json")
    return raw_dir


def _electronic_generation_blocks() -> list[OcrBlock]:
    return [
        OcrBlock("上次抄表日期", (100, 100, 260, 130), 0.99),
        OcrBlock("2026-05-01", (280, 100, 430, 130), 0.99),
        OcrBlock("本次抄表日期", (500, 100, 660, 130), 0.99),
        OcrBlock("2026-06-01", (680, 100, 830, 130), 0.99),
        OcrBlock("上月行码", (400, 200, 480, 230), 0.99),
        OcrBlock("本月行码", (520, 200, 600, 230), 0.99),
        OcrBlock("度差", (640, 200, 700, 230), 0.99),
        OcrBlock("倍率", (760, 200, 820, 230), 0.99),
        OcrBlock("电表电量", (880, 200, 980, 230), 0.99),
        OcrBlock("换表电量", (1000, 200, 1100, 230), 0.99),
        OcrBlock("合计", (1140, 200, 1200, 230), 0.99),
        OcrBlock("09001SF000000425", (80, 250, 300, 280), 0.99),
        OcrBlock("08945216", (130, 285, 250, 315), 0.99),
        OcrBlock("1655.88", (400, 270, 480, 300), 0.99),
        OcrBlock("1926.96", (520, 270, 600, 300), 0.99),
        OcrBlock("271.08", (640, 270, 700, 300), 0.99),
        OcrBlock("200", (760, 270, 820, 300), 0.99),
        OcrBlock("54216", (880, 270, 980, 300), 0.99),
        OcrBlock("0", (1000, 270, 1100, 300), 0.99),
        OcrBlock("54216", (1140, 270, 1200, 300), 0.99),
        OcrBlock("09001SF000000425", (80, 350, 300, 380), 0.99),
        OcrBlock("08945217", (130, 385, 250, 415), 0.99),
        OcrBlock("1882.44", (400, 370, 480, 400), 0.99),
        OcrBlock("2176.25", (520, 370, 600, 400), 0.99),
        OcrBlock("293.81", (640, 370, 700, 400), 0.99),
        OcrBlock("150", (760, 370, 820, 400), 0.99),
        OcrBlock("44072", (880, 370, 980, 400), 0.99),
        OcrBlock("0", (1000, 370, 1100, 400), 0.99),
        OcrBlock("44072", (1140, 370, 1200, 400), 0.99),
    ]


if __name__ == "__main__":
    unittest.main()



