from __future__ import annotations

import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from submit_flow_agent.file_scanner import scan_input_files
from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult
from submit_flow_agent.ocr.local_stub import LocalSampleOcrAdapter
from submit_flow_agent.pdf_renderer import RenderedPage
from submit_flow_agent.pipeline import run_month
from submit_flow_agent.validator import ValidationReport
from tests.pricing_helpers import static_pricing_snapshot


PROJECT_DIR = Path("深圳新安中学高中部650KW分布式光伏项目")


class RunMonthPipelineTests(unittest.TestCase):
    def test_run_month_threads_controlled_month_into_field_mapping_scan(self) -> None:
        def controlled_scan(project_dir: Path, *, expected_month: str | None = None):
            self.assertEqual(expected_month, "2026-05")
            return scan_input_files(project_dir, expected_month=expected_month)

        with tempfile.TemporaryDirectory() as tmpdir:
            with patch("submit_flow_agent.field_mapper.scan_input_files", side_effect=controlled_scan):
                result = run_month(
                    PROJECT_DIR,
                    output_base=Path(tmpdir),
                    ocr_adapter=LocalSampleOcrAdapter(),
                    runtime_clock=_fixed_runtime_clock,
                    pricing_snapshot=static_pricing_snapshot(),
                    expected_month="2026-05",
                )

        self.assertEqual(result.status, "pass")
        self.assertEqual(result.month, "2026-05")

    def test_run_month_generates_all_outputs_and_log(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output_base = Path(tmpdir)

            result = run_month(
                PROJECT_DIR,
                output_base=output_base,
                ocr_adapter=LocalSampleOcrAdapter(),
                runtime_clock=_fixed_runtime_clock,
                pricing_snapshot=static_pricing_snapshot(),
            )
            stale_raw = result.output_dir / "raw_ocr" / "stale-retry.json"
            stale_image = result.output_dir / "raw_ocr" / "images" / "stale-retry.png"
            stale_raw.write_text("{}", encoding="utf-8")
            stale_image.write_bytes(b"stale")
            second = run_month(
                PROJECT_DIR,
                output_base=output_base,
                ocr_adapter=LocalSampleOcrAdapter(),
                runtime_clock=_fixed_runtime_clock,
                pricing_snapshot=static_pricing_snapshot(),
            )

            self.assertEqual(result.status, "pass")
            self.assertEqual(result.month, "2026-05")
            self.assertEqual(result.output_dir, output_base / "新安高中部" / "2026-05")
            self.assertEqual(second.output_dir, result.output_dir)
            self.assertFalse(stale_raw.exists())
            self.assertFalse(stale_image.exists())
            self.assertTrue(result.recognized_data.exists())
            self.assertTrue(result.validation_report.exists())
            self.assertTrue(result.detail_workbook and result.detail_workbook.exists())
            self.assertTrue(result.summary_workbook and result.summary_workbook.exists())
            self.assertTrue(result.process_log.exists())
            self.assertIsNone(result.review_report)

            recognized = json.loads(result.recognized_data.read_text(encoding="utf-8"))
            self.assertEqual(recognized["fields"]["statement_date"]["value"], "2026-07-17")
            self.assertEqual(recognized["fields"]["statement_date"]["source_type"], "system_runtime")

            log_text = result.process_log.read_text(encoding="utf-8")
            self.assertIn("run-month started", log_text)
            self.assertIn("project_dir=", log_text)
            self.assertIn("site_key=xinan_high_school", log_text)
            self.assertIn("statement_date=2026-07-17", log_text)
            self.assertIn("site_name=新安高中部", log_text)
            self.assertIn("source_file[generation_statement]=福域发电单202605-2.pdf", log_text)
            self.assertIn("raw_ocr_dir=", log_text)
            self.assertIn("mapping_mode=ocr_raw", log_text)
            self.assertTrue((result.output_dir / "raw_ocr").exists())
            self.assertIn("validation_status=pass", log_text)
            self.assertIn("can_write_excel=True", log_text)
            self.assertIn("detail_workbook=", log_text)
            self.assertIn("summary_workbook=", log_text)
            self.assertIn("run-month completed", log_text)

            detail_workbook = load_workbook(result.detail_workbook, data_only=False)
            self.assertIn("1", detail_workbook.sheetnames)
            self.assertEqual(detail_workbook["1"]["G5"].value, "=E5-F5")
            self.assertEqual(detail_workbook["1"]["A7"].value.strftime("%Y-%m-%d"), "2026-07-17")
            detail_workbook.close()

            summary_workbook = load_workbook(result.summary_workbook, data_only=False)
            self.assertIn("Sheet1", summary_workbook.sheetnames)
            self.assertEqual(summary_workbook["Sheet1"]["M6"].value, "=L6+I6")
            summary_workbook.close()

    def test_run_month_validation_failure_writes_review_report_and_skips_excel(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output_base = Path(tmpdir)
            failed_report = ValidationReport(
                site_key="xinan_high_school",
                site_name="新安高中部",
                month="2026-05",
                status="fail",
                can_write_excel=False,
                normalized_data={},
                checks=[],
                errors=["forced validation failure"],
                warnings=[],
            )

            with patch("submit_flow_agent.pipeline.validate_data_to_file", return_value=failed_report):
                result = run_month(
                    PROJECT_DIR,
                    output_base=output_base,
                    ocr_adapter=LocalSampleOcrAdapter(),
                    runtime_clock=_fixed_runtime_clock,
                pricing_snapshot=static_pricing_snapshot(),
            )

            output_dir = output_base / "新安高中部" / "2026-05"
            self.assertEqual(result.status, "need_review")
            self.assertTrue(result.review_report and result.review_report.exists())
            self.assertTrue((output_dir / "recognized_data.json").exists())
            self.assertFalse((output_dir / "新安高中部统计表_2026-05.xlsx").exists())
            self.assertFalse((output_dir / "新安高中部汇总表格_2026-05.xlsx").exists())
            review = json.loads((output_dir / "review_report.json").read_text(encoding="utf-8"))
            self.assertEqual(review["task_status"], "need_review")
            self.assertTrue(any(issue["message"] == "forced validation failure" for issue in review["issues"]))
            log_text = (output_dir / "process_log.txt").read_text(encoding="utf-8")
            self.assertIn("failed_reason=forced validation failure", log_text)
            self.assertIn("run-month need_review", log_text)

    def test_run_month_low_confidence_writes_review_report_and_skips_excel(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            with patch.object(LocalSampleOcrAdapter, "recognize", _low_confidence_recognize):
                result = run_month(
                    PROJECT_DIR,
                    output_base=Path(tmpdir),
                    ocr_adapter=LocalSampleOcrAdapter(),
                    runtime_clock=_fixed_runtime_clock,
                pricing_snapshot=static_pricing_snapshot(),
            )

            self.assertEqual(result.status, "need_review")
            self.assertTrue(result.review_report and result.review_report.exists())
            review = json.loads(result.review_report.read_text(encoding="utf-8"))
            self.assertTrue(any(issue["field"] == "grid_fee" and issue["reason"] == "low_confidence" for issue in review["issues"]))
            self.assertIsNone(result.detail_workbook)
            self.assertIsNone(result.summary_workbook)

    def test_run_month_manual_confirmed_data_revalidates_then_writes_excel(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            confirmed_path = root / "confirmed_data.json"
            confirmed_path.write_text(
                json.dumps(
                    {
                        "confirmed_by": "operator",
                        "confirmed_at": "2026-07-08T10:00:00+08:00",
                        "fields": {"grid_fee": {"value": 407.29, "reason": "人工核对电费结算单后确认"}},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.object(LocalSampleOcrAdapter, "recognize", _low_confidence_recognize):
                result = run_month(
                    PROJECT_DIR,
                    output_base=root,
                    confirmed_data=confirmed_path,
                    ocr_adapter=LocalSampleOcrAdapter(),
                    runtime_clock=_fixed_runtime_clock,
                pricing_snapshot=static_pricing_snapshot(),
            )

            recognized = json.loads(result.recognized_data.read_text(encoding="utf-8"))
            validation = json.loads(result.validation_report.read_text(encoding="utf-8"))
            self.assertEqual(result.status, "pass")
            self.assertTrue(result.detail_workbook and result.detail_workbook.exists())
            self.assertTrue(result.summary_workbook and result.summary_workbook.exists())
            self.assertEqual(recognized["mapping_mode"], "manual_confirmed")
            self.assertEqual(recognized["fields"]["grid_fee"]["source_type"], "manual_confirmed")
            self.assertTrue(validation["can_write_excel"])


def _fixed_runtime_clock() -> datetime:
    return datetime(2026, 7, 17, 9, 0, tzinfo=timezone.utc)


def _low_confidence_recognize(self: LocalSampleOcrAdapter, source_file: Path, pages: list[RenderedPage]) -> OcrRawResult:
    if "发电单" in source_file.name:
        texts = [
            "电表编号 09001SF00000042508945216",
            "本期起始日期 2026-05-01",
            "本期结束日期 2026-06-01",
            "上期示数 816.82",
            "本期示数 1147.27",
            "倍率 150",
            "月总发电量 49568 kWh",
            "填表日期 2026-06-21",
        ]
    elif "电量" in source_file.name:
        texts = ["结算月份 2026年05月", "上网电量 1680 kWh"]
    elif "电费" in source_file.name:
        texts = ["上网电价 0.242434 元/kWh", "上网电费 407.29 元"]
    else:
        texts = []
    blocks = [
        OcrBlock(
            text=text,
            bbox=(0.0, float(index * 24), 240.0, float(index * 24 + 20)),
            confidence=0.80 if "上网电费" in text else 0.98,
        )
        for index, text in enumerate(texts)
    ]
    return OcrRawResult(
        source_file=source_file,
        pages=[OcrPage(page=page.page, image_path=page.image_path, blocks=blocks if page.page == 1 else []) for page in pages],
    )


if __name__ == "__main__":
    unittest.main()

