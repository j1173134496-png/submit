from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from submit_flow_agent.nanke import (
    METER_NUMBER,
    OUTPUT_DETAIL,
    OUTPUT_SUMMARY,
    collect_nanke_source,
    confirm_nanke,
    list_nanke_outputs,
    prepare_nanke,
    recognize_nanke_source,
    run_nanke,
)
from submit_flow_agent.ocr.base import OcrBlock, OcrPage, OcrRawResult


TEMPLATE_DIR = Path("config/templates/nanke")


class QueueOcr:
    def __init__(self, results: list[OcrRawResult]) -> None:
        self.results = iter(results)

    def recognize(self, source_file: Path, pages: list[object]) -> OcrRawResult:
        result = next(self.results)
        return OcrRawResult(source_file, [OcrPage(1, Path("page.png"), result.pages[0].blocks)])


def _raw(name: str, rows: list[list[str]]) -> OcrRawResult:
    blocks: list[OcrBlock] = []
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            x = float(column_index * 220)
            y = float(row_index * 25)
            blocks.append(OcrBlock(value, (x, y, x + 200, y + 18), 0.99))
    source = Path(name)
    return OcrRawResult(source, [OcrPage(1, source.with_suffix(".png"), blocks)])


def _generation() -> OcrRawResult:
    return _raw(
        "generation.pdf",
        [
            ["分布式光伏发电补贴核算单", "发电月份", "202606"],
            ["表号", "上期示数", "本期示数", "倍率", "发电量(kWh)"],
            [METER_NUMBER, "1437.2", "1733.59", "150", "44459"],
        ],
    )


def _fee() -> OcrRawResult:
    return _raw(
        "fee.pdf",
        [
            ["电费结算单", "购电月份", "2026年06月"],
            ["2", "市场化电费", "2960", "0.480007", "1420.82"],
        ],
    )


def _energy() -> OcrRawResult:
    return _raw(
        "energy.pdf",
        [
            ["电量结算单", "购电月份", "2026年06月", "日电量"],
            ["合计", "2960"],
        ],
    )


def _base() -> OcrRawResult:
    return _raw(
        "base.png",
        [
            ["缴费通知书", "基站合同", "请款期间为：2026年6月1日至2026年6月30日"],
            ["用电度数为1477.35度。电费单价为0.698元/度，请款金额1031.19元。"],
        ],
    )


def _image(path: Path) -> Path:
    color = (sum(path.name.encode("utf-8")) % 255, 180, 220)
    Image.new("RGB", (12, 12), color).save(path)
    return path


def test_recognizes_four_roles() -> None:
    role, generation = recognize_nanke_source(_generation(), expected_month="2026-06")
    assert role == "generation_statement"
    assert generation["total_generation"] == "44459"

    role, fee = recognize_nanke_source(_fee(), expected_month="2026-06")
    assert role == "grid_fee_settlement"
    assert fee["grid_price"] == "0.480007"

    role, energy = recognize_nanke_source(_energy(), expected_month="2026-06")
    assert role == "grid_energy_settlement"
    assert energy["grid_energy"] == "2960"

    role, base = recognize_nanke_source(_base(), expected_month="2026-06")
    assert role == "base_station_notice"
    assert base["base_station_energy"] == "1477.35"
    assert base["base_station_price"] == "0.698"


def test_full_june_flow_generates_two_controlled_workbooks(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    task_id = "nanke_2026_06_test"
    prepare_nanke(runtime, task_id=task_id, month="2026-06")
    adapter = QueueOcr([_generation(), _fee(), _energy(), _base()])
    for index in range(4):
        task = collect_nanke_source(
            runtime,
            task_id,
            _image(tmp_path / f"source-{index}.png"),
            ocr_adapter=adapter,
        )
    assert task["status"] == "awaiting_confirmation"
    assert task["source_count"] == 4
    assert task["recognized_data"]["school_energy"] == "40021.65"
    assert task["recognized_data"]["school_price"] == "0.61424"
    assert task["recognized_data"]["grid_formula_fee"] == "1420.82"

    confirm_nanke(runtime, task_id)
    completed = run_nanke(runtime, task_id, template_dir=TEMPLATE_DIR)
    assert completed["status"] == "completed"
    outputs = [Path(path) for path in list_nanke_outputs(runtime, task_id)]
    assert [path.name for path in outputs] == [OUTPUT_DETAIL, OUTPUT_SUMMARY]

    detail = load_workbook(outputs[0], data_only=False)
    try:
        sheet = detail["Sheet1"]
        assert "六月份" in sheet["A1"].value
        assert sheet["A5"].value == f"电表{METER_NUMBER}"
        assert sheet["E5"].value == 44459
        assert sheet["G5"].value == 1477.35
        assert sheet["H5"].value == "=E5-F5-G5"
        assert sheet["I5"].value == "=0.88*0.698"
        assert sheet["J5"].value == "=H5*I5"
    finally:
        detail.close()

    summary = load_workbook(outputs[1], data_only=False)
    try:
        sheet = summary["Sheet1"]
        assert sheet["A4"].value == f"电表{METER_NUMBER}"
        assert sheet["B4"].value == "2026年6月"
        assert sheet["K4"].value == "=J4*G4"
        assert sheet["L4"].value == "=F4-G4-H4"
        assert sheet["N4"].value == "=L4*M4"
        assert sheet["O4"].value == "=N4+K4"
    finally:
        summary.close()
