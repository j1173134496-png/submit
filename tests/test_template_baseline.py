from __future__ import annotations

import hashlib
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from submit_flow_agent.config import load_site_config


PROJECT_DIR = next(Path(".").glob("*650KW*"))
CONTROLLED_TEMPLATE_DIR = Path("config/templates/xinan_high_school_blank")
SOURCE_DETAIL_HASH = "c74d062c9e093b0abf908a59c825e13a983797d9b7bc0255ece4ac974eacdd5a"
SOURCE_SUMMARY_HASH = "8a55e7869cfa9a3eb262ce2d509faa8a2883093eb78092a0d20bbecac8f0ac4b"
BLANK_DETAIL_HASH = "4ebff0cbf553fe9cbed60e717044af75c505a2d388c8d71ca74bd280f8221d9b"
BLANK_SUMMARY_HASH = "2fa1b5c407b49858bcb3ad9ee73f27e0d6d87ecafa10f4e10b49deb3a1d63037"


class TemplateBaselineTests(unittest.TestCase):
    def test_reusable_single_and_dual_profiles_allow_document_derived_project_name(self) -> None:
        payload = json.loads(Path("config/template_profiles.json").read_text(encoding="utf-8-sig"))
        profiles = {profile["profile_key"]: profile for profile in payload["template_profiles"]}

        for profile_key in ("xinan_single_meter_v1", "xixiang_meter_rows_v1"):
            fields = {
                field["key"]: field
                for field in profiles[profile_key]["metadata_fields"]
            }
            self.assertEqual(
                fields["project_name"]["source_strategy"],
                "operator_supplied_or_document_derived",
            )

        summary_profiles = {
            profile["profile_key"]: profile
            for profile in payload["role_profiles"]["summary"]
        }
        self.assertEqual(
            summary_profiles["summary_meter_rows_am_v1"]["meter_topology"]["total_due_formula"],
            "school_fee_rows_only",
        )
        for profile in payload["role_profiles"]["detail"]:
            self.assertEqual(
                profile["meter_topology"]["generation_value_source"],
                "rounded_meter_reading_formula",
            )
            self.assertEqual(profile["layout"]["generation_formula_column"], "E")
        for profile in payload["role_profiles"]["summary"]:
            self.assertEqual(
                profile["meter_topology"]["generation_value_source"],
                "rounded_meter_reading_formula",
            )
            self.assertEqual(profile["layout"]["generation_formula_column"], "F")

    def test_cleaned_source_and_controlled_template_hashes_match(self) -> None:
        site = load_site_config()
        pairs = [
            (PROJECT_DIR / site.detail_workbook, SOURCE_DETAIL_HASH, CONTROLLED_TEMPLATE_DIR / site.detail_workbook, BLANK_DETAIL_HASH),
            (PROJECT_DIR / site.summary_workbook, SOURCE_SUMMARY_HASH, CONTROLLED_TEMPLATE_DIR / site.summary_workbook, BLANK_SUMMARY_HASH),
        ]

        for source, source_hash, controlled, blank_hash in pairs:
            self.assertEqual(_sha256(source), source_hash)
            self.assertEqual(_sha256(controlled), blank_hash)
            self.assertNotEqual(_sha256(source), _sha256(controlled))

    def test_controlled_detail_template_dynamic_cells_are_blank(self) -> None:
        site = load_site_config()
        workbook = load_workbook(CONTROLLED_TEMPLATE_DIR / site.detail_workbook, data_only=False)
        worksheet = workbook[site.detail_sheet]

        self.assertIsNone(worksheet["E2"].value)
        self.assertIsNone(worksheet["H2"].value)
        self.assertTrue(all(worksheet.cell(row=5, column=column).value is None for column in range(1, 11)))
        self.assertEqual(worksheet["A6"].value, site.company_name)
        self.assertIsNone(worksheet["A7"].value)
        workbook.close()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


if __name__ == "__main__":
    unittest.main()
